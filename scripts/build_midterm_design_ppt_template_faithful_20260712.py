#!/usr/bin/env python3
"""Build a template-faithful WalkSafe Assist midterm design presentation.

Unlike the first redesign, this builder keeps the official application-software
template's native title tab, corner decoration, logo, footer band, and each
deliverable's original body pattern. Only example content is replaced.
"""

from __future__ import annotations

import io
import os
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from build_midterm_design_ppt_20260712 import (
    add_picture_fill,
    add_picture_fit,
    add_table,
    code_excerpt,
    delete_slide,
    sha256,
    textbox,
)


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    "/home/ddobagi/Downloads/[서식2] 2026년 한이음 드림업 제작설계서(택1, 개인정보 기재x)/"
    "[서식2] 2026년_제작설계서_응용소프트웨어(개인정보 기재x).pptx"
)
OUTPUT = Path("/home/ddobagi/Downloads/제작설계서_완성본_템플릿준수본_수정본.pptx")
EXPECTED_TEMPLATE_SHA256 = "75d2579807ea758185556b6d4349b82e71eb84451940c592434a55e1dadb3d1b"

ASSETS = ROOT / "docs/submission/form_materials/assets"
WEB_MOBILE = ASSETS / "web_main_mobile.png"
WEB_DESKTOP = ASSETS / "web_main_desktop.png"
ADMIN_DESKTOP = ASSETS / "admin_desktop.png"
DAMAGED_TACTILE = ROOT / "ai_tasks/walksafe_v3_relabel_20260521/images/wsv3_img_0021_51453cd280.jpg"
MENU_REFERENCE = Path("/home/ddobagi/Downloads/메뉴구성도.png")

FONT = "Noto Sans CJK KR"
CODE_FONT = "DejaVu Sans Mono"
BLUE = RGBColor(59, 90, 168)
MID_BLUE = RGBColor(111, 145, 196)
LIGHT_BLUE = RGBColor(224, 233, 247)
TEAL = RGBColor(24, 151, 151)
LIGHT_TEAL = RGBColor(232, 246, 245)
GOLD = RGBColor(230, 166, 49)
LIGHT_GOLD = RGBColor(253, 247, 228)
CORAL = RGBColor(205, 87, 61)
LIGHT_CORAL = RGBColor(252, 238, 234)
INK = RGBColor(37, 48, 63)
MUTED = RGBColor(91, 101, 114)
LINE = RGBColor(192, 200, 212)
PALE = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(206, 48, 42)

SELECTED_TEMPLATE_SLIDES = [1, 2, 7, 8, 10, 11, 14, 15, 16, 17, 18, 19, 20, 21, 22, 26, 27, 28, 29, 35]


def remove_shape(slide, shape) -> None:
    slide.shapes._spTree.remove(shape._element)


def set_text(shape, value: str, *, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def replace_text_preserve_style(shape, value: str) -> None:
    """Replace visible text while keeping the official template run style."""
    frame = shape.text_frame
    if not frame.paragraphs:
        return
    paragraph = frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = value
    for extra in frame.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def ppt_rect(slide, x, y, w, h, *, fill=WHITE, line_color=LINE, radius=False, width=1.0):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line_color
    shape.line.width = Pt(width)
    return shape


def ppt_line(slide, x1, y1, x2, y2, *, color=LINE, width=1.2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def small_tag(slide, value="중간보고") -> None:
    # The official gray right-end motif remains visible; no custom label is added.
    return None


def body_title(slide, value: str) -> None:
    textbox(slide, value, 0.55, 1.27, 8.9, 0.42, size=17, color=INK, bold=True, valign=MSO_ANCHOR.MIDDLE)
    ppt_line(slide, 0.55, 1.72, 9.45, 1.72, color=LINE, width=0.9)


def template_footer(slide) -> None:
    textbox(
        slide,
        "WalkSafe Assist · 중간 제작설계서 · 기준일 2026-07-11 · 개인정보 미기재",
        0.55, 7.06, 8.7, 0.18,
        size=7.7, color=MUTED, align=PP_ALIGN.CENTER,
    )


def clean_to_template_chrome(slide, title: str, tag="중간보고") -> None:
    title_shape = None
    kept_footer = False
    for shape in list(slide.shapes):
        x = shape.left / 914400
        y = shape.top / 914400
        w = shape.width / 914400
        h = shape.height / 914400
        text = " ".join(getattr(shape, "text", "").split())
        keep = False

        # Official blue title tab, horizontal lines, and corner arc.
        if abs(x - 0.12) < 0.08 and y < 0.1 and 3.1 < w < 3.7 and 1.0 < h < 1.4:
            keep = True
        elif shape.shape_type == 9 and y < 0.8:
            keep = True
        elif x < 0 and y < 0 and w < 1.5 and h < 1.5:
            keep = True
        # Official Hanium logo.
        elif shape.shape_type == 13 and 8.0 < x < 8.5 and y < 0.4:
            keep = True
        elif shape.shape_type == 13 and x > 9.3 and 0.4 < y < 0.8:
            keep = True
        # Native title textbox inside the blue tab.
        elif 0.25 < x < 0.55 and 0.66 < y < 0.95 and 2.8 < w < 3.4 and hasattr(shape, "text_frame"):
            keep = True
            title_shape = shape
        # Keep the native transparent sample-caption object, but empty its text.
        elif 4.0 < x < 4.4 and 0.5 < y < 0.75 and w > 5.2 and hasattr(shape, "text_frame"):
            keep = True
            set_text(shape, "", size=8, color=RED)
        # Keep the native bottom instruction band and replace only its text.
        elif y > 6.95 and w > 8.5 and hasattr(shape, "text_frame"):
            keep = True
            replace_text_preserve_style(shape, "")
            kept_footer = True
        # Preserve the native footer placeholder only when it exists.
        elif y > 6.85 and "한이음" in text:
            keep = True
            kept_footer = True

        if not keep:
            remove_shape(slide, shape)

    if title_shape is None:
        raise RuntimeError(f"template title shape not found for {title}")
    replace_text_preserve_style(title_shape, f"| {title}")
    small_tag(slide, tag)


def panel(slide, x, y, w, h, title, *, accent=BLUE, fill=WHITE):
    ppt_rect(slide, x, y, w, h, fill=fill, line_color=LINE, radius=False)
    ppt_rect(slide, x, y, w, 0.42, fill=accent, line_color=accent)
    textbox(slide, title, x + 0.12, y + 0.08, w - 0.24, 0.24, size=12.2, color=WHITE, bold=True)


def node(slide, x, y, w, h, title, body="", *, fill=LIGHT_BLUE, border=BLUE, title_size=11.5, body_size=9.5):
    ppt_rect(slide, x, y, w, h, fill=fill, line_color=border, radius=True, width=1.2)
    title_color = WHITE if fill == border else border
    body_color = WHITE if fill == border else INK
    textbox(slide, title, x + 0.12, y + 0.12, w - 0.24, 0.28, size=title_size, color=title_color, bold=True, align=PP_ALIGN.CENTER)
    if body:
        textbox(slide, body, x + 0.12, y + 0.48, w - 0.24, h - 0.58, size=body_size, color=body_color, align=PP_ALIGN.CENTER)


def arrow(slide, x, y, *, direction="right", color=MUTED):
    value = "→" if direction == "right" else "↓"
    textbox(slide, value, x, y, 0.35, 0.35, size=18, color=color, bold=True, align=PP_ALIGN.CENTER)


def menu_node(slide, x, y, w, value, *, size=8.2, fill=WHITE, bold=True):
    """Draw a node matching the supplied monochrome menu-tree reference."""
    border = RGBColor(133, 133, 133)
    shape = ppt_rect(slide, x, y, w, 0.48, fill=fill, line_color=border, width=0.9)
    textbox(
        slide,
        value,
        x + 0.04,
        y + 0.08,
        w - 0.08,
        0.30,
        size=size,
        color=INK,
        bold=bold,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    return shape


def normalize_package(path: Path, slide_count: int) -> None:
    app_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    with tempfile.TemporaryDirectory(prefix="walksafe-template-ppt-") as directory:
        normalized = Path(directory) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/app.xml":
                    root = ET.fromstring(data)
                    slides = root.find(f"{{{app_ns}}}Slides")
                    notes = root.find(f"{{{app_ns}}}Notes")
                    if slides is not None:
                        slides.text = str(slide_count)
                    if notes is not None:
                        notes.text = "0"
                    for stale_tag in ("HeadingPairs", "TitlesOfParts"):
                        stale = root.find(f"{{{app_ns}}}{stale_tag}")
                        if stale is not None:
                            root.remove(stale)
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                info = zipfile.ZipInfo(name, date_time=(2026, 7, 12, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, data)
        shutil.move(normalized, path)


def crop_picture(slide, path: Path, x, y, w, h, *, top=0.0, bottom=0.0, left=0.0, right=0.0):
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    picture.crop_top = top
    picture.crop_bottom = bottom
    picture.crop_left = left
    picture.crop_right = right
    picture.line.color.rgb = LINE
    picture.line.width = Pt(1)
    return picture


def prepare_presentation() -> Presentation:
    if not TEMPLATE.is_file():
        raise SystemExit(f"template not found: {TEMPLATE}")
    if sha256(TEMPLATE) != EXPECTED_TEMPLATE_SHA256:
        raise SystemExit("official template hash changed")
    prs = Presentation(TEMPLATE)
    if len(prs.slides) != 35 or (prs.slide_width, prs.slide_height) != (9144000, 6858000):
        raise SystemExit("unexpected official template dimensions")

    selected = {number - 1 for number in SELECTED_TEMPLATE_SLIDES}
    for index in range(len(prs.slides) - 1, -1, -1):
        if index not in selected:
            delete_slide(prs, index)
    if len(prs.slides) != 20:
        raise SystemExit("selected template slide count mismatch")

    core = prs.core_properties
    core.title = "WalkSafe Assist 중간 제작설계서 — 공식 템플릿 준수본"
    core.subject = "2026 한이음 드림업 제작설계서"
    core.author = "WalkSafe Assist 프로젝트"
    core.last_modified_by = "WalkSafe Assist 프로젝트"
    core.comments = "공식 응용소프트웨어 템플릿 배치 보존본"
    core.created = datetime(2026, 7, 12, 0, 0, 0)
    core.modified = datetime(2026, 7, 12, 0, 0, 0)
    return prs


def build() -> None:
    for path in (WEB_MOBILE, WEB_DESKTOP, ADMIN_DESKTOP, DAMAGED_TACTILE, MENU_REFERENCE):
        if not path.is_file():
            raise SystemExit(f"missing visual asset: {path}")

    prs = prepare_presentation()
    slides = list(prs.slides)

    # 1. Official cover: keep the graphic, logo, gray information box, and typography.
    cover = slides[0]
    for shape in cover.shapes:
        text = " ".join(getattr(shape, "text", "").split())
        if "SW개발/HW제작 설계서" in text:
            set_text(shape, "SW 개발 설계서", size=50, color=RGBColor(141, 45, 45), bold=True)
        elif text.startswith("프로젝트 명"):
            set_text(shape, "프로젝트 명 : AI를 이용한 카메라 기반 시각장애인 보행 지원 시스템", size=16, color=INK, bold=True)
        elif "2026. 00. 00" in text:
            set_text(
                shape,
                "2026. 07. 11.\nWalkSafe Assist · 중간 제작설계서\n주제품 Web/PWA · Android 연구 보조",
                size=16,
                color=MUTED,
                align=PP_ALIGN.CENTER,
            )

    # 2. Preserve the official full deliverables matrix and highlight only the general column.
    matrix_slide = slides[1]
    for shape in list(matrix_slide.shapes):
        text = " ".join(getattr(shape, "text", "").split())
        x = shape.left / 914400
        y = shape.top / 914400
        if "샘플" in text and hasattr(shape, "text_frame"):
            replace_text_preserve_style(shape, "")
        elif hasattr(shape, "text_frame") and "수행 단계별 주요 산출물" in text:
            replace_text_preserve_style(shape, "수행 단계별 주요 산출물")
    table = next(shape.table for shape in matrix_slide.shapes if getattr(shape, "has_table", False))
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            if col_index == 2:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_GOLD
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = INK

    # 3. Requirements 1 — canonical REQ-001~007 and the official six-column schema.
    slide = slides[2]
    clean_to_template_chrome(slide, "요구사항 정의서(1)", "사용자·보행")
    rows = [
        ["요구사항ID", "요구사항명", "기능ID", "기능명", "세부사항", "예외사항"],
        ["REQ-001", "Web/PWA 카메라 기반 위험 감지", "F-001", "위험 감지", "새 frame을 450ms 간격으로 v2 탐지 API에 전달해 overlay·위험평가", "권한 거부·stale·malformed·중복 응답은 탐지 불가; 종료 시 센서 해제"],
        ["REQ-002", "실탐지와 fake/demo 결과의 명시적 분리", "F-001·F-009", "위험 감지·모델 관리", "UI·응답·현장 기록에서 real/fake와 config·model hash를 구분", "real 실패 시 fake 자동 전환 금지; 승인 모델·Release 근거 없음"],
        ["REQ-003", "Android ARCore·TFLite 연구 보조", "F-009", "모델 연구·후보 관리", "ARCore depth와 unified 13-class TFLite, bbox-depth 정합을 연구", "load/hash/tensor 실패 때만 legacy fallback; 실외 FPS·정합 검증 미확보"],
        ["REQ-004", "객체별 안정 추적과 경로 위험 선별", "F-002", "거리·접근/경로 보조", "IoU 기반 1:1 추적 후 3 frame·700ms, ROI·접근 신호로 위험 선별", "안정화 전·경로 근거 부족 시 경고 안 함; bbox TTC 단독 high/STOP 금지"],
        ["REQ-005", "손상 점자블록 전용 자동·음성 요청 신고", "F-004", "자동/음성 요청 신고", "damaged class만 conf≥0.70·GPS≤15m·3 frame·700ms·동일 snapshot으로 신고", "GPS 불량·jump>20m·gap>1.8s·10분 cooldown이면 미전송"],
        ["REQ-006", "짧은 TTS·진동 위험 피드백", "F-003", "위험 안내", "위험>상호작용>길안내 우선순위와 객체별 TTS·high 진동 cooldown 적용", "high만 진동; 반복 발화 억제; TalkBack·청취성 현장 검증 미확보"],
        ["REQ-007", "음성 신고·목적지·길안내 제어", "F-005", "음성 제어·신고", "strict intent로 신고·목적지·길안내 action을 Web/Android에 전달", "부정문·검색 중·범위 밖 선택은 no-op/재질문; 소음 환경 미검증"],
    ]
    req_table = add_table(
        slide, rows, 0.30, 1.36, 9.40, 5.42,
        widths=[0.82, 1.55, 0.78, 1.35, 2.75, 2.15], font_size=7.5, header_fill=BLUE,
    )
    req_table.table.rows[0].height = Inches(0.48)
    for row_index in range(1, len(req_table.table.rows)):
        req_table.table.rows[row_index].height = Inches(0.705)
    for cell in req_table.table.rows[0].cells:
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(7.5)

    # 4. Requirements 2 — canonical REQ-008~014.
    slide = slides[3]
    clean_to_template_chrome(slide, "요구사항 정의서(2)", "저장·운영·품질")
    rows = [
        ["요구사항ID", "요구사항명", "기능ID", "기능명", "세부사항", "예외사항"],
        ["REQ-008", "신고 이미지·metadata의 신뢰 경계와 원자 저장", "F-004", "자동/음성 요청 신고", "검증 뒤 durable journal→원자 rename→DB commit, 시작 시 미완료 write 조정", "잘못된 입력 거부; corrupt journal은 fail-closed; 운영 복구·보존 시행 증거 미확보"],
        ["REQ-009", "PostGIS 위치 저장과 검수 추적", "F-006·F-007", "운영자 검수·Export", "정확히 5개 application table: reports+status/export/read audits+60초 rate events", "3종 audit만 append-only; rate event는 non-audit; 좌표 불일치 거부"],
        ["REQ-010", "관리자 검수·상태관리·기관용 자료 준비", "F-006·F-007", "운영자 검수·Export", "목록/상세/image마다 actor·resource·time·purpose read audit와 기관용 반출", "purpose 누락 시 거부; 외부 제출은 수동"],
        ["REQ-011", "설치 가능한 PWA와 안전한 offline 경계", "F-011(신규)", "PWA 설치·offline 경계", "opt-in manifest·SW·offline app shell 제공; 안전 API·변경 요청은 network-only", "기본 OFF; offline 탐지·신고·경로·음성 불가; 실폰 HTTPS 미검증"],
        ["REQ-012", "목적지 검색과 보행 경로 안내", "F-008", "목적지 안내", "TMAP POI·STAIR_AVOID 경로, 이탈·재탐색·도착과 점자블록 local steering", "거리 미상 단일 후보 자동선택 금지; tactile gate 실패·손상 시 TMAP 복귀"],
        ["REQ-013", "모델·데이터·class order 추적", "F-009", "모델 연구·후보 관리", "13-class order, PT/TFLite hash, 후보 registry·배포 상태를 추적", "hash/class 불일치 시 배포 금지; 독립 test·재분할 근거 미확보"],
        ["REQ-014", "개인정보·인증·반출·보존 경계", "F-007·F-010", "Export·개인정보/현장 기록", "Web frame/report·Android report·telemetry 각각 기본 OFF, 철회 시 in-flight 취소", "역할 교차·미동의 전송 차단; IAM·사람 검수·복원 증거 미확보"],
    ]
    req_table = add_table(
        slide, rows, 0.30, 1.36, 9.40, 5.42,
        widths=[0.82, 1.55, 0.78, 1.35, 2.75, 2.15], font_size=7.5, header_fill=BLUE,
    )
    req_table.table.rows[0].height = Inches(0.48)
    for row_index in range(1, len(req_table.table.rows)):
        req_table.table.rows[row_index].height = Inches(0.705)
    for cell in req_table.table.rows[0].cells:
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(7.5)

    # 5. User-facing service flow.
    slide = slides[4]
    clean_to_template_chrome(slide, "서비스 흐름도", "이용 흐름")
    panel(slide, 0.24, 1.42, 4.69, 5.35, "사용자 관점 서비스 흐름", accent=BLUE)
    panel(slide, 5.0, 1.42, 4.69, 5.35, "서비스 역할과 책임", accent=BLUE)
    flow_nodes = [
        (0.55, 2.2, "1 시작", "권한·카메라\nGPS·마이크", BLUE),
        (1.95, 2.2, "2 분석", "13-class 탐지\n인스턴스 안정화", TEAL),
        (3.35, 2.2, "3 판단", "경로·접근\n신뢰조건", GOLD),
        (1.2, 4.3, "위험", "TTS·진동", CORAL),
        (2.95, 4.3, "손상 시설", "신고 자료", TEAL),
    ]
    for x, y, title, body, color in flow_nodes:
        node(slide, x, y, 1.2, 1.15, title, body, fill=WHITE, border=color, title_size=10.8, body_size=8.4)
    arrow(slide, 1.67, 2.55)
    arrow(slide, 3.07, 2.55)
    ppt_line(slide, 3.95, 3.35, 1.8, 4.3, color=CORAL)
    ppt_line(slide, 3.95, 3.35, 3.55, 4.3, color=TEAL)
    roles = [
        (5.35, 2.12, "보행 사용자", "보행 중 화면을 계속 보지 않아도\n짧은 음성·진동으로 행동 판단", BLUE),
        (5.35, 3.45, "WalkSafe Assist", "현재 위험 안내와 손상 시설 신고를\n서로 다른 정책으로 처리", TEAL),
        (5.35, 4.78, "운영 관리자", "근거 확인 → 상태 변경 → 기관용 자료\n외부 제출은 사람의 수동 행위", CORAL),
    ]
    for x, y, title, body, color in roles:
        node(slide, x, y, 4.0, 1.02, title, body, fill=PALE, border=color, title_size=11.3, body_size=8.8)
    textbox(slide, "주제품: Web/PWA · Android: ARCore depth/TFLite 연구 보조", 5.3, 6.13, 4.05, 0.34, size=9.2, color=MUTED, align=PP_ALIGN.CENTER)

    # 6. Service/system architecture — infographic left, implementation summary right.
    slide = slides[5]
    clean_to_template_chrome(slide, "서비스 구성도", "시스템")
    panel(slide, 0.24, 1.42, 5.15, 5.34, "시스템 구성도", accent=BLUE, fill=PALE)
    node(slide, 0.57, 1.98, 4.50, 0.68, "스마트폰 센서", "Camera · GPS/방향 · Mic", fill=WHITE, border=BLUE, title_size=10.8, body_size=8.0)
    arrow(slide, 2.64, 2.61, direction="down", color=BLUE)
    node(slide, 0.57, 2.91, 4.50, 0.82, "Next.js Web/PWA + Same-origin BFF", "용도별 기본 OFF 동의 · 단일 replica/process lock", fill=LIGHT_BLUE, border=BLUE, title_size=11.2, body_size=8.2)
    ppt_line(slide, 2.82, 3.73, 2.82, 3.94, color=MUTED, width=1.2)
    ppt_line(slide, 1.93, 3.94, 4.40, 3.94, color=MUTED, width=1.2)
    ppt_line(slide, 1.93, 3.94, 1.93, 4.08, color=MUTED, width=1.2)
    ppt_line(slide, 4.40, 3.94, 4.40, 4.08, color=MUTED, width=1.2)
    textbox(slide, "HTTPS API", 2.32, 3.77, 1.0, 0.18, size=8.0, color=MUTED, align=PP_ALIGN.CENTER)
    node(slide, 0.48, 4.08, 2.90, 0.92, "FastAPI Backend", "인증 OpenAPI · TMAP-only · nonblocking ready", fill=LIGHT_TEAL, border=TEAL, title_size=10.2, body_size=7.4)
    node(slide, 3.67, 4.08, 1.46, 0.92, "Voice API", "전용 token\npre auth/length/rate\ncopy byte · post duration/queue\n1 worker/replica+lock", fill=LIGHT_GOLD, border=GOLD, title_size=9.8, body_size=5.8)
    ppt_line(slide, 1.93, 5.00, 1.93, 5.29, color=TEAL, width=1.2)
    ppt_line(slide, 1.20, 5.29, 4.43, 5.29, color=TEAL, width=1.2)
    for x in (1.20, 2.82, 4.43):
        ppt_line(slide, x, 5.29, x, 5.48, color=TEAL, width=1.2)
    node(slide, 0.48, 5.44, 1.45, 0.76, "YOLO26", "13종 탐지", fill=WHITE, border=CORAL, title_size=9.2, body_size=8.0)
    node(slide, 2.10, 5.44, 1.45, 0.76, "PostgreSQL", "5 tables · shared limiter", fill=WHITE, border=TEAL, title_size=8.7, body_size=6.9)
    node(slide, 3.72, 5.44, 1.42, 0.76, "TMAP-only", "POI · 계단 회피", fill=WHITE, border=GOLD, title_size=8.9, body_size=8.0)
    ppt_rect(slide, 0.50, 6.26, 4.62, 0.28, fill=WHITE, line_color=LINE, radius=True, width=0.8)
    textbox(slide, "탐지·경로·신고 응답 → 위험 선별 → TTS·진동·화면 안내", 0.60, 6.30, 4.42, 0.17, size=8.0, color=INK, bold=True, align=PP_ALIGN.CENTER)

    panel(slide, 5.62, 1.42, 4.08, 2.38, "FRONTEND", accent=BLUE, fill=WHITE)
    textbox(slide, "Next.js · React · TypeScript · PWA", 5.86, 1.99, 3.60, 0.26, size=9.4, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, "• Web frame/report 각각 기본 OFF 동의\n• Android report·telemetry도 별도 동의\n• 철회 시 in-flight 취소\n• 관리자 검수·Heatmap·자료 반출", 5.86, 2.34, 3.58, 1.12, size=8.5, color=INK)
    textbox(slide, "Same-origin BFF · 단일 replica/process lock", 5.91, 3.54, 3.48, 0.18, size=7.3, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

    panel(slide, 5.62, 3.98, 4.08, 2.78, "BACKEND", accent=TEAL, fill=WHITE)
    textbox(slide, "FastAPI · SQLAlchemy · PostGIS", 5.86, 4.55, 3.60, 0.26, size=9.4, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, "API", 5.86, 4.88, 0.52, 0.20, size=8.2, color=TEAL, bold=True)
    textbox(slide, "• 인증된 OpenAPI — detect/reports/admin\n• /navigation/* — TMAP-only POI·경로\n• 공유 원자 DB limiter · nonblocking readiness", 5.86, 5.10, 3.58, 0.85, size=8.3, color=INK)
    textbox(slide, "DATA · AI", 5.86, 5.97, 0.92, 0.20, size=8.2, color=TEAL, bold=True)
    textbox(slide, "DB sslmode=verify-full+gssencmode=disable · Voice global/actor/IP rate · post-copy inference queue · disconnect", 5.86, 6.20, 3.58, 0.28, size=5.5, color=INK)

    # 7. Single menu tree, following Downloads/메뉴구성도.png.
    slide = slides[6]
    clean_to_template_chrome(slide, "메뉴 구성도", "통합 메뉴")
    menu_line = RGBColor(142, 142, 142)
    ppt_rect(slide, 0.28, 1.39, 9.44, 5.35, fill=RGBColor(245, 245, 245), line_color=RGBColor(232, 232, 232), width=0.7)
    centers = [1.94, 2.89, 3.84, 4.79]
    # Field flow and its four branches.
    ppt_line(slide, 1.48, centers[0], 1.72, centers[0], color=menu_line, width=0.9)
    ppt_line(slide, 2.88, centers[0], 3.12, centers[0], color=menu_line, width=0.9)
    ppt_line(slide, 4.27, centers[0], 4.55, centers[0], color=menu_line, width=0.9)
    ppt_line(slide, 4.40, centers[0], 4.40, centers[-1], color=menu_line, width=0.9)
    for center in centers[1:]:
        ppt_line(slide, 4.40, center, 4.55, center, color=menu_line, width=0.9)
    for center in centers:
        ppt_line(slide, 5.90, center, 6.25, center, color=menu_line, width=0.9)
        ppt_line(slide, 7.60, center, 8.00, center, color=menu_line, width=0.9)
    # Admin branch starts at app launch and remains a separate /admin route.
    ppt_line(slide, 1.60, centers[0], 1.60, 5.84, color=menu_line, width=0.9)
    ppt_line(slide, 1.60, 5.84, 1.72, 5.84, color=menu_line, width=0.9)
    for x1, x2 in ((2.88, 3.12), (4.27, 4.55), (5.90, 6.25), (7.60, 8.00)):
        ppt_line(slide, x1, 5.84, x2, 5.84, color=menu_line, width=0.9)

    menu_node(slide, 0.43, 1.70, 1.05, "앱 실행")
    menu_node(slide, 1.72, 1.70, 1.16, "현장 인증")
    menu_node(slide, 3.12, 1.70, 1.15, "보행 보조 홈")
    field_rows = [
        (1.70, "실시간 보행", "객체·경로 위험", "TTS·진동 안내"),
        (2.65, "목적지 길안내", "목적지 검색·선택", "TMAP 도보 안내"),
        (3.60, "손상 점자블록 신고", "자동·음성 요청", "중복 판정·저장"),
        (4.55, "설정·세션", "권한·연결·보폭", "PWA·현장 기록"),
    ]
    for y, first, second, third in field_rows:
        menu_node(slide, 4.55, y, 1.35, first, size=7.9)
        menu_node(slide, 6.25, y, 1.35, second, size=7.6)
        menu_node(slide, 8.00, y, 1.45, third, size=7.6)
    menu_node(slide, 1.72, 5.60, 1.16, "관리자 인증")
    menu_node(slide, 3.12, 5.60, 1.15, "관리자 화면")
    menu_node(slide, 4.55, 5.60, 1.35, "신고 관리", size=7.9)
    menu_node(slide, 6.25, 5.60, 1.35, "목록·상세 검수", size=7.6)
    menu_node(slide, 8.00, 5.60, 1.45, "상태·지도·내보내기", size=7.0)

    # 8. Screen design — state/permission definition.
    slide = slides[8]
    clean_to_template_chrome(slide, "화면 설계서", "상태·권한")
    body_title(slide, "권한과 연결 상태를 확인한 뒤 보행 보조를 시작")
    add_picture_fill(slide, WEB_MOBILE, 0.58, 1.98, 2.55, 4.72)
    rows = [
        ["화면 상태", "표시 정보", "주요 행동", "오류·복구"],
        ["고지", "Web frame/report·Android report·telemetry", "각각 기본 OFF 동의", "철회 시 in-flight 취소"],
        ["권한 확인", "camera·GPS·mic", "브라우저 허용", "재요청 버튼"],
        ["보행 보조", "탐지 출처·현재 위험", "음성·길안내", "stale/unavailable 표시"],
        ["백그라운드", "세션 중지", "재개 후 재초기화", "이전 탐지 폐기"],
    ]
    add_table(slide, rows, 3.38, 2.02, 6.0, 3.6, widths=[1.3, 2.1, 1.7, 2.2], font_size=9.4, header_fill=BLUE)
    textbox(slide, "개발용 fixture UI · 실폰 권한·TalkBack·브라우저별 동작은 현장 검증 전", 3.55, 5.95, 5.65, 0.45, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)

    # 9. Screen design — main Web/PWA specification.
    slide = slides[9]
    clean_to_template_chrome(slide, "화면 설계서", "Web/PWA")
    body_title(slide, "현재 위험·위치·방향·길안내를 한 화면에서 확인")
    add_picture_fit(slide, WEB_DESKTOP, 0.55, 1.95, 5.25, 4.72)
    rows = [
        ["영역", "내용", "사용자 행동"],
        ["카메라", "미리보기·bbox·탐지 출처", "전방 상태 확인"],
        ["상태 카드", "현재 위험·GPS·방향·신고", "판단 불가 원인 확인"],
        ["음성", "안내 on/off·목적지·신고", "말로 실행·취소"],
        ["길안내", "목적지 후보·다음 안내·이탈", "번호 선택·재탐색"],
        ["복구", "권한·연결 오류", "재요청·중지"],
    ]
    add_table(slide, rows, 6.0, 1.95, 3.42, 3.92, widths=[1.0, 2.1, 1.5], font_size=8.7, header_fill=BLUE)
    textbox(slide, "개발용 fake-v2 합성 입력\n실모델·실폰 현장 성능 근거와 분리", 6.25, 6.05, 2.92, 0.48, size=9.2, color=RED, align=PP_ALIGN.CENTER)

    # 10. Screen design — Web/PWA UI views.
    slide = slides[10]
    clean_to_template_chrome(slide, "화면 설계서 - UI(SW)", "Web 화면")
    body_title(slide, "모바일과 데스크톱에서 같은 보행 상태를 유지")
    add_picture_fit(slide, WEB_MOBILE, 0.62, 1.92, 2.75, 4.78)
    add_picture_fit(slide, WEB_DESKTOP, 3.62, 1.92, 5.75, 3.65)
    add_table(slide, [
        ["화면", "확인 목적", "현재 근거"],
        ["모바일", "큰 상태·조작·한 손 사용", "viewport 렌더"],
        ["데스크톱", "카메라+상태 패널 배치", "개발 fixture"],
    ], 3.75, 5.78, 5.5, 0.83, widths=[1.2, 2.8, 1.5], font_size=8.7, header_fill=BLUE)

    # 11. Screen design — Admin UI.
    slide = slides[11]
    clean_to_template_chrome(slide, "화면 설계서 - UI(SW)", "Admin 화면")
    body_title(slide, "운영자는 근거를 확인한 뒤 상태와 반출을 결정")
    add_picture_fit(slide, ADMIN_DESKTOP, 0.52, 1.92, 7.0, 4.75)
    admin_steps = [
        (7.72, 2.05, "1 필터", "상태·유형·출처", BLUE),
        (7.72, 3.15, "2 검수", "이미지·좌표·bbox", TEAL),
        (7.72, 4.25, "3 상태", "담당자·사유·이력", GOLD),
        (7.72, 5.35, "4 반출", "검토 완료 건만 준비", CORAL),
    ]
    for x, y, title, body, color in admin_steps:
        node(slide, x, y, 1.62, 0.9, title, body, fill=WHITE, border=color, title_size=10, body_size=7.8)
    textbox(slide, "실제 Admin UI · 로컬 fixture 자료", 0.8, 6.38, 3.25, 0.22, size=8.8, color=RED, bold=True)

    # 12. ERD.
    slide = slides[12]
    clean_to_template_chrome(slide, "엔티티 관계도 - ERD", "데이터")
    body_title(slide, "신고·3종 감사·공유 제한을 정확히 5개 application table로 분리")
    node(slide, 0.65, 2.0, 4.25, 3.55, "reports", "id UUID · status\nclass·confidence·bbox·captured_at\nPoint(4326)·accuracy·heading\nimage_path·MIME·metadata JSONB", fill=LIGHT_TEAL, border=TEAL, title_size=16, body_size=11.3)
    node(slide, 5.55, 1.8, 3.8, 0.92, "report_status_audits", "상태 전이 · actor · time · APPEND-ONLY", fill=LIGHT_GOLD, border=GOLD, title_size=11.4, body_size=7.3)
    node(slide, 5.55, 2.88, 3.8, 0.92, "report_export_audits", "반출 목적·범위·hash · APPEND-ONLY", fill=LIGHT_CORAL, border=CORAL, title_size=11.4, body_size=7.3)
    node(slide, 5.55, 3.96, 3.8, 0.92, "report_read_audits", "list/detail/image · actor/resource/time/purpose", fill=LIGHT_BLUE, border=BLUE, title_size=11.4, body_size=7.0)
    node(slide, 5.55, 5.04, 3.8, 0.92, "actor_rate_limit_events", "공유 원자 60초 window · non-audit", fill=WHITE, border=TEAL, title_size=11.4, body_size=7.3)
    ppt_line(slide, 4.9, 2.25, 5.55, 2.25, color=GOLD, width=1.7)
    ppt_line(slide, 4.9, 4.42, 5.55, 4.42, color=BLUE, width=1.7)
    node(slide, 1.2, 5.78, 3.2, 0.88, "DB 밖 이미지 파일", "UUID 파일명 · 경로·MIME만 DB 저장", fill=LIGHT_BLUE, border=BLUE, title_size=10, body_size=7.1)
    textbox(slide, "status/export/read 3종 감사만 UPDATE·DELETE 금지", 5.55, 6.18, 3.8, 0.3, size=8.6, color=INK, bold=True, align=PP_ALIGN.CENTER)

    # 13. Risk function flow — preserve the template's metadata strip + large flow area.
    slide = slides[13]
    clean_to_template_chrome(slide, "기능 처리도(기능 흐름도)", "위험 안내")
    add_table(slide, [
        ["프로그램 ID", "기능명", "입력", "출력"],
        ["FLOW-01", "보행 위험 선별", "새 frame·detection·sensor", "TTS·진동 또는 표시만"],
    ], 0.45, 1.42, 9.1, 0.72, widths=[1.3, 2.4, 3.4, 2.3], font_size=8.6, header_fill=BLUE)
    flow_items = [
        (0.55, "새 frame", "기본 450ms\n완료 frame만", BLUE),
        (2.25, "인스턴스", "bbox·class\n동일 객체 추적", BLUE),
        (3.95, "안정화", "3 frame\n700ms", TEAL),
        (5.65, "위험 근거", "경로·접근\n거리·TTC", GOLD),
        (7.35, "피드백", "우선순위·cooldown\nTTS·고위험 진동", CORAL),
    ]
    for x, title, body, color in flow_items:
        node(slide, x, 2.65, 1.35, 1.45, title, body, fill=WHITE, border=color, title_size=10.7, body_size=8.5)
    for x in (1.92, 3.62, 5.32, 7.02):
        arrow(slide, x, 3.16)
    add_table(slide, [
        ["근거 부족", "처리"],
        ["stale·frozen·server unavailable", "판단 중지 · ‘위험 없음’으로 표시하지 않음"],
        ["경로 밖·접근 근거 없음", "display-only · TTS·진동 없음"],
        ["bbox 변화만 있는 stop 후보", "신뢰 거리 없으면 warning으로 제한"],
    ], 1.0, 4.65, 8.0, 1.65, widths=[3.5, 4.5], font_size=9, header_fill=BLUE)

    # 14. Report function flow.
    slide = slides[14]
    clean_to_template_chrome(slide, "기능 처리도(기능 흐름도)", "신고 처리")
    body_title(slide, "신고는 후보·서버·사람의 세 신뢰 단계를 통과")
    report_items = [
        (0.45, "후보 gate", "손상 class\nconf≥0.70\nGPS≤15m", GOLD),
        (2.22, "동일 촬영", "image·위치·시각\n3 frame·700ms", TEAL),
        (3.99, "서버 재검증", "class·model\nGPS·MIME", BLUE),
        (5.76, "원자 저장", "journal→rename→DB\nstartup 조정", TEAL),
        (7.53, "사람 검수", "근거·중복 확인\n기관용 자료", CORAL),
    ]
    for x, title, body, color in report_items:
        node(slide, x, 2.15, 1.5, 1.55, title, body, fill=WHITE, border=color, title_size=10.4, body_size=8.4)
    for x in (1.96, 3.73, 5.5, 7.27):
        arrow(slide, x, 2.73)
    node(slide, 0.85, 4.4, 2.5, 1.25, "자동 신고", "조건 충족 시 무음 저장\n완료 TTS 없음", fill=LIGHT_TEAL, border=TEAL, title_size=11.5, body_size=8.8)
    node(slide, 3.75, 4.4, 2.5, 1.25, "음성 요청 신고", "사용자가 명시하면\n성공·실패를 짧게 안내", fill=LIGHT_GOLD, border=GOLD, title_size=11.5, body_size=8.8)
    node(slide, 6.65, 4.4, 2.5, 1.25, "외부 제출", "기관 자동 API 없음\n관리자가 수동 제출", fill=LIGHT_CORAL, border=CORAL, title_size=11.5, body_size=8.8)
    textbox(slide, "중복 후보는 자동 병합하지 않고 관리자 검수 정보로만 사용", 2.0, 6.15, 6.0, 0.34, size=10.3, color=INK, bold=True, align=PP_ALIGN.CENTER)

    # 15. Program list.
    slide = slides[15]
    clean_to_template_chrome(slide, "프로그램 - 목록", "구현 목록")
    body_title(slide, "카메라 입력부터 관리자 검수까지 모듈 책임을 분리")
    rows = [
        ["영역", "대표 프로그램", "기능"],
        ["Capture", "page.tsx · useCamera.ts · useDetectionV2.ts", "카메라·센서·탐지 연결"],
        ["Risk", "risk-instance-tracker.ts · risk-evaluator.ts", "안정화·경로 위험 판단"],
        ["Route", "useNavigationGuidance.ts · tactile-route-policy.ts", "TMAP 진행·점자블록 보조"],
        ["Voice", "useVoiceCommands.ts · voice-intent-executor.ts", "목적지·신고·길안내 명령"],
        ["Report", "reports.py · report_storage.py", "journal 저장·시작 조정"],
        ["Admin", "admin/page · report_read_audit.py", "필터·상태·반출·조회 감사"],
        ["Model", "detect_v2.py · yolo_inference_adapter.py", "13-class 추론"],
        ["Android", "MainActivity.kt · depth/ · inference/", "ARCore·TFLite 연구 보조"],
    ]
    add_table(slide, rows, 0.42, 1.92, 9.16, 4.85, widths=[1.15, 5.2, 2.8], font_size=10, header_fill=BLUE)

    # 16. Table definition — preserve the template's two-column body pattern.
    slide = slides[16]
    clean_to_template_chrome(slide, "테이블 정의서 - ERD", "테이블")
    body_title(slide, "5개 application table로 신고 원천·3종 감사·공유 제한을 보존")
    panel(slide, 0.45, 1.9, 4.35, 4.85, "테이블 관계", accent=BLUE)
    panel(slide, 5.05, 1.9, 4.5, 4.85, "핵심 컬럼·제약", accent=BLUE)
    node(slide, 0.82, 2.65, 3.6, 1.25, "reports", "UUID · status · class · bbox · Point(4326) · image · metadata", fill=LIGHT_TEAL, border=TEAL, title_size=12, body_size=8.7)
    node(slide, 0.62, 4.35, 1.15, 1.35, "status", "상태 전이\nappend-only", fill=LIGHT_GOLD, border=GOLD, title_size=9.4, body_size=7.4)
    node(slide, 1.87, 4.35, 1.15, 1.35, "export", "반출 실행\nappend-only", fill=LIGHT_CORAL, border=CORAL, title_size=9.4, body_size=7.4)
    node(slide, 3.12, 4.35, 1.15, 1.35, "read", "목록/상세/image\nappend-only", fill=LIGHT_BLUE, border=BLUE, title_size=9.4, body_size=6.8)
    node(slide, 1.55, 5.9, 1.8, 0.65, "rate events", "공유 원자 60초 · non-audit", fill=WHITE, border=TEAL, title_size=8.5, body_size=6.3)
    ppt_line(slide, 2.62, 3.9, 1.65, 4.35, color=GOLD)
    ppt_line(slide, 2.62, 3.9, 3.6, 4.35, color=CORAL)
    add_table(slide, [
        ["테이블", "핵심 컬럼", "제약"],
        ["reports", "id·Point·image·metadata", "원천·CHECK/GiST"],
        ["status audit", "report·actor·전이·time", "append-only"],
        ["export audit", "audit IDs·profile·hash", "append-only"],
        ["read audit", "actor·resource·purpose·time", "append-only"],
        ["rate events", "actor digest·group·time", "60초·non-audit"],
    ], 5.28, 2.58, 4.02, 3.2, widths=[1.3, 2.2, 1.8], font_size=7.8, header_fill=BLUE)
    textbox(slide, "DB 밖 image: durable journal→원자 rename→DB commit · startup reconciliation", 5.22, 6.05, 4.15, 0.38, size=7.6, color=MUTED, align=PP_ALIGN.CENTER)

    # 17. Core source 1 — use the template's full bordered code area.
    slide = slides[17]
    clean_to_template_chrome(slide, "핵심 소스코드(1)", "Web 위험 정책")
    body_title(slide, "손상 점자블록은 경고가 아니라 신고 후보로 분기")
    ppt_rect(slide, 0.48, 1.93, 6.25, 4.72, fill=PALE, line_color=LINE, radius=False)
    source = code_excerpt("apps/web/app/_walksafe/risk-evaluator.ts", "export function evaluateTwoModelDetectionRisk", 14)
    textbox(slide, source, 0.68, 2.13, 5.85, 4.25, size=9.7, color=INK, font=CODE_FONT)
    add_table(slide, [
        ["입력", "판정", "출력"],
        ["detection", "손상 class", "경고 X"],
        ["context", "report-only", "신고 O"],
    ], 6.98, 2.12, 2.35, 1.95, widths=[1.2, 1.1, 1.1], font_size=7.8, header_fill=BLUE)
    node(slide, 7.02, 4.55, 2.28, 1.35, "정책 의미", "손상 시설은 기본 사용자 경고 없이 신고 흐름으로 전달", fill=LIGHT_TEAL, border=TEAL, title_size=11, body_size=8.3)
    textbox(slide, "apps/web/app/_walksafe/risk-evaluator.ts", 6.98, 6.18, 2.35, 0.28, size=7.5, color=MUTED, align=PP_ALIGN.CENTER)

    # 18. Core source 2.
    slide = slides[18]
    clean_to_template_chrome(slide, "핵심 소스코드(2)", "Backend 신고")
    body_title(slide, "대표 신고 ID를 유지하고 내부 중복 후보는 응답에서 제거")
    ppt_rect(slide, 0.48, 1.93, 6.25, 4.72, fill=PALE, line_color=LINE, radius=False)
    source = code_excerpt("backend/app/api/reports.py", "def _field_creation_response", 15)
    textbox(slide, source, 0.68, 2.13, 5.85, 4.25, size=9.2, color=INK, font=CODE_FONT)
    add_table(slide, [
        ["단계", "보호 규칙"],
        ["대표 신고", "primary id·duplicate_count 유지"],
        ["응답", "다른 중복 후보 ID 제거"],
        ["저장", "durable journal→rename→DB commit"],
        ["시작", "미완료 write reconciliation"],
    ], 6.98, 2.12, 2.35, 2.55, widths=[1.0, 2.0], font_size=7.8, header_fill=BLUE)
    node(slide, 7.02, 5.15, 2.28, 0.95, "결과", "사용자 응답 최소화 + 저장 일관성", fill=LIGHT_CORAL, border=CORAL, title_size=10.5, body_size=7.8)
    textbox(slide, "backend/app/api/reports.py · uploads.py", 6.98, 6.22, 2.35, 0.24, size=7.4, color=MUTED, align=PP_ALIGN.CENTER)

    # 19. Remove the second legacy menu slide and keep the official Thank-you slide.
    delete_slide(prs, 7)
    if len(prs.slides) != 19:
        raise RuntimeError("final slide count must be 19")

    for rel_id, rel in list(prs.part.rels.items()):
        reltype = rel.reltype.lower()
        if any(token in reltype for token in ("authors", "comment", "notesmaster", "person")):
            prs.part.drop_rel(rel_id)
    for slide in prs.slides:
        for rel_id, rel in list(slide.part.rels.items()):
            reltype = rel.reltype.lower()
            if any(token in reltype for token in ("notesslide", "comment", "person")):
                slide.part.drop_rel(rel_id)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix="walksafe-template-faithful-", suffix=".pptx", dir=OUTPUT.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        prs.save(temporary)
        normalize_package(temporary, len(prs.slides))
        os.replace(temporary, OUTPUT)
    finally:
        temporary.unlink(missing_ok=True)

    if sha256(TEMPLATE) != EXPECTED_TEMPLATE_SHA256:
        raise SystemExit("official source template changed while building")
    print(f"generated {OUTPUT} ({OUTPUT.stat().st_size} bytes, sha256={sha256(OUTPUT)})")


if __name__ == "__main__":
    build()

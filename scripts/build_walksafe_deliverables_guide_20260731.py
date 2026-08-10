#!/usr/bin/env python3
"""Build the non-canonical WalkSafe deliverables introduction DOCX.

This builder reads the current controlled JSON sources, creates accessible
charts, and writes one explanatory DOCX under docs/submission/drafts. It does
not modify DOC-01, DOC-05, or any existing deliverable.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import textwrap
import time
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence
from xml.etree import ElementTree
from zoneinfo import ZoneInfo

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager, patches
from matplotlib.colors import LinearSegmentedColormap

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
CHECKPOINT_PATH = ROOT / "docs/control/walksafe-project-continuation-checkpoint.json"
GAP_PATH = ROOT / "docs/control/audits/walksafe-implementation-gap-analysis-20260726-r021.json"
BACKLOG_PATH = ROOT / "docs/control/audits/walksafe-implementation-remediation-backlog-20260726-r021.json"
DOC01_PATH = ROOT / "docs/deliverables/00-control/artifact-register.json"
TEST_CASES_PATH = ROOT / "docs/deliverables/06-testing/registers/test-cases.json"
CLOSURE_LEDGER_PATH = (
    ROOT
    / "docs/control/execution/artifact-closure/run-20260727-001/packets/"
    "phase1-exact257-successor-r011/phase1-exact257-successor-ledger-r011.json"
)
DEFAULT_OUTPUT = ROOT / "docs/submission/drafts/워크세이프_산출물_소개서_20260731.docx"
FINALIZER = ROOT / "scripts/finalize_walksafe_deliverables_guide_20260731.py"

BUILD_DATE = "2026-07-31"
BUILD_TIME = datetime(2026, 7, 31, 0, 0, 0, tzinfo=ZoneInfo("Asia/Seoul"))
FONT_REGULAR = Path("/usr/share/fonts/truetype/nanum/NanumSquareR.ttf")
FONT_BOLD = Path("/usr/share/fonts/truetype/nanum/NanumSquareB.ttf")

FONT_KO = "NanumSquare"
FONT_CODE = "Liberation Mono"
NAVY = "17365D"
BLUE = "1F6FA8"
TEAL = "007A5E"
ORANGE = "B95300"
BURGUNDY = "963B4E"
PURPLE = "6F4C9B"
GRAY = "59636E"
LIGHT_BLUE = "E8F2F8"
LIGHT_TEAL = "E4F2ED"
LIGHT_ORANGE = "F9EDE4"
LIGHT_GRAY = "F2F4F6"
INK = "17212B"
MUTED = "52616D"

STATUS_LABELS = {
    "IMPLEMENTED": "공식 완료",
    "PARTIAL": "일부 구현",
    "CONFLICTING": "정한 기준과 다름",
    "MISSING": "핵심 구현 없음",
    "EVIDENCE_MISSING": "시험·증거 부족",
    "BLOCKED": "실측·검토·훈련 대기",
}
STATUS_COLORS = {
    "IMPLEMENTED": "#007A5E",
    "PARTIAL": "#1F6FA8",
    "CONFLICTING": "#B95300",
    "MISSING": "#963B4E",
    "EVIDENCE_MISSING": "#6F4C9B",
    "BLOCKED": "#59636E",
}
STATUS_HATCHES = {
    "IMPLEMENTED": "",
    "PARTIAL": "///",
    "CONFLICTING": "xx",
    "MISSING": "...",
    "EVIDENCE_MISSING": "\\\\",
    "BLOCKED": "++",
}
LIFECYCLE_LABELS = {
    "APPROVED_BASELINED": "승인 기준본",
    "ACTIVE": "계속 갱신",
    "DRAFT": "초안",
    "PLANNED": "계획·미실행",
}
LIFECYCLE_COLORS = {
    "APPROVED_BASELINED": "#1F6FA8",
    "ACTIVE": "#007A5E",
    "DRAFT": "#D89000",
    "PLANNED": "#7B8794",
}

CURRENT_OBSERVATION_KO = {
    "GAP-021": (
        "앱과 중간 서버는 기기별 로그인 상태를 따로 관리하고, 서버 원장은 한 계정에서 보행 하나만 "
        "활성화되게 합니다. 다른 기기의 시작 요청은 사용자가 확인하기 전까지 막고, 기기 전환을 "
        "승인하면 이전 보행을 먼저 끝냅니다. 같은 요청이 반복돼도 결과를 한 번만 만들며, 오프라인이던 "
        "이전 기기도 다시 연결되거나 사용기한이 지나면 종료됩니다."
    ),
    "GAP-023": (
        "앱은 권한마다 사용할 수 있는 기능을 나눕니다. 권한 거부·철회·만료 또는 설정 화면에서 돌아온 "
        "뒤 실제 권한을 다시 확인하고, 사용자가 명시적으로 재개하기 전에는 보행·자료 수집·전송을 막습니다."
    ),
    "GAP-024": (
        "앱과 중간 서버는 특정 자료 사용 동의 철회와 계정 삭제를 구분합니다. 새 처리는 즉시 막고, "
        "저장 위치별 삭제 진행 상태와 재시도·보류·문의 상태를 순서가 뒤로 가지 않게 기록합니다."
    ),
    "GAP-025": (
        "앱은 권한마다 사용할 수 있는 기능을 나눕니다. 권한 거부·철회·만료 또는 설정 화면에서 돌아온 "
        "뒤 실제 권한을 다시 확인하고, 사용자가 명시적으로 재개하기 전에는 보행·자료 수집·전송을 막습니다."
    ),
    "GAP-056": (
        "사용자 앱·중간 서버와 관리자 앱·본 서버의 권한 경계를 분리했습니다. 서버가 역할과 자료 소유자를 "
        "확인하고, 알 수 없는 관리자 요청은 거부합니다. 기기별 로그인은 교체·폐기할 수 있고, 관리자는 "
        "비밀번호와 일회용 인증번호를 함께 사용합니다. 위험한 작업은 다시 확인하며 허용·거부 결과를 모두 남깁니다."
    ),
}

BLOCKER_LABELS = {
    "ACTUAL_CI_NOT_RUN": "실제 자동 빌드·시험 미실행",
    "AIML18_NOT_RUN": "관련 AI·데이터 절차 미실행",
    "BASELINE_NOT_RUN": "기준본 확정 절차 미실행",
    "BASELINE_THRESHOLDS_NOT_MEASURED": "기준값 실측 미완료",
    "CHANNEL_UNSELECTED": "사용할 전달 채널 미선정",
    "CONFIG_CONFORMANCE_NOT_REVIEWED": "설정이 기준에 맞는지 미검토",
    "EVIDENCE_OR_EXTERNAL_VALUE_PENDING": "실행 증거 또는 외부 확인값 대기",
    "EXECUTION_NOT_RUN": "실제 실행 미완료",
    "FORMAL_REPRODUCTION_ZERO": "정식 재현 실행 기록 없음",
    "HUMAN_REVIEW_AFTER_EVIDENCE": "증거 확보 뒤 사람 검토 필요",
    "HUMAN_REVIEW_AFTER_EXECUTION": "실행 뒤 사람 검토 필요",
    "LEGAL_AND_QA_UNASSIGNED": "법률·품질 검토 담당 미지정",
    "LOAD_RESET_NOT_RUN": "부하 뒤 정상 복귀 확인 미실행",
    "NOT_ACTIVE_UNDER_CURRENT_BASELINE": "현재 기준본에서 사용 대상으로 활성화되지 않음",
    "NO_APPROVED_TARGET_ENVIRONMENT_OR_ACCOUNT": "승인된 실행 환경 또는 계정 없음",
    "NO_DESIGNATED_REVIEW_EVENT": "정해진 검토 일정 없음",
    "OPS_ROLES_NOT_ASSIGNED": "운영 담당 역할 미지정",
    "PENDING_ACTIVATION_EVALUATION": "사용 시작 가능 여부 평가 대기",
    "PLAN_ROLLBACK_NOT_RUN": "계획 되돌리기 훈련 미실행",
    "QA_UNASSIGNED": "품질 검토 담당 미지정",
    "RIGHTS_PRIVACY_SCHEMA_NOT_VERIFIED": "개인정보 권리 처리 자료 구조 미검증",
    "RIGHTS_PRIVACY_SPLIT_NOT_VERIFIED": "동의 철회와 계정 삭제 분리 미검증",
    "S25_EXACT_VERSION_UNKNOWN": "정확한 대상 버전 미확정",
    "SLA_NOT_APPROVED": "서비스 대응시간 기준 미승인",
    "TECH_APPROVAL_NOT_PERFORMED": "기술 승인 미실행",
    "TRIGGER_PARAMETERS_NOT_APPROVED": "실행을 시작할 기준값 미승인",
}

CATEGORY_ORDER = [
    "DOC",
    "MGT",
    "DSC",
    "REQ",
    "DES",
    "DEV",
    "TST",
    "SEC",
    "AIML",
    "REL",
    "OPS",
    "WS",
    "CLS",
]
CATEGORY_INFO = {
    "DOC": ("문서 관리", "문서 이름·버전·승인·위치를 관리", "docs/deliverables/00-control/"),
    "MGT": ("프로젝트 관리", "목표·일정·역할·위험·결정을 관리", "docs/deliverables/01-management/"),
    "DSC": ("사용자·문제 조사", "누구의 어떤 문제를 풀지 확인", "docs/deliverables/02-discovery/"),
    "REQ": ("무엇을 만들지", "제품이 해야 할 일과 합격 조건을 정의", "docs/deliverables/03-requirements/"),
    "DES": ("어떻게 만들지", "시스템·화면·데이터·보안 구조를 설계", "docs/deliverables/04-design/"),
    "DEV": ("실제 구현", "코드·빌드·개발 방법과 구현 근거를 관리", "docs/deliverables/05-implementation/"),
    "TST": ("시험", "시험 계획·결과·결함·인수 판단을 관리", "docs/deliverables/06-testing/"),
    "SEC": ("보안·개인정보", "보안과 개인정보 위험·검증 결과를 관리", "docs/deliverables/07-security/"),
    "AIML": ("데이터·AI 모델", "학습자료·AI 모델·평가·교체를 관리", "docs/deliverables/08-ai-ml-data/"),
    "REL": ("출시·전달", "배포할 버전·서명·배포·인수 기록을 관리", "docs/deliverables/09-release/"),
    "OPS": ("운영·복구", "관제·장애·백업·복원·비용을 관리", "docs/deliverables/10-operations/"),
    "WS": ("WalkSafe 안전·인수", "WalkSafe 고유 안전·실기기·현장 검증을 관리", "docs/deliverables/11-walksafe/"),
    "CLS": ("인계·종료", "최종 인수·운영 이관·서비스 종료를 관리", "docs/deliverables/12-closure/"),
}

EPIC_INFO = [
    ("EPIC-01", "제품 경계 정리", "Android 중심 제품을 정리하고 과거 Web 기능을 분리", "IMPLEMENTATION_READY"),
    ("EPIC-02", "안전한 보행 상태", "걷기 시작·중지와 권한 문제를 안전하게 처리", "IN_PROGRESS"),
    ("EPIC-03", "계정·관리자·보안", "일반 사용자와 관리자 계정·권한을 분리", "IN_PROGRESS"),
    ("EPIC-04", "길찾기·도착·이탈", "도착과 경로 이탈 때 사용자가 결정하게 개선", "PLANNED"),
    ("EPIC-05", "물체·거리·위험", "급한 위험부터 정확하고 짧게 안내", "PLANNED"),
    ("EPIC-06", "음성·진동·접근성", "손을 쓰지 않는 음성·진동·화면읽기 사용을 완성", "PLANNED"),
    ("EPIC-07", "원본자료 수명주기", "보행자료를 안전하게 모으고 기한에 맞춰 삭제", "PLANNED"),
    ("EPIC-08", "자동신고 대기함", "신고를 안전하게 보관했다가 조건에 맞춰 전송", "PLANNED"),
    ("EPIC-09", "서버 용량·장애", "저장공간·비용·장애를 관리하고 앱에 상태 전달", "PLANNED"),
    ("EPIC-10", "AI 모델 수명주기", "학습자료·평가·교체·이전 정상판 복구를 관리", "PLANNED"),
    ("EPIC-11", "출시·운영·복구", "서명·단계 배포·관제·백업복구를 준비", "PLANNED"),
    ("EPIC-12", "정식 검증·출시 관문", "실제 폰·현장·사용자 시험과 5개 관문을 통과", "PLANNED"),
]

EPIC_GAP_COUNTS = {
    "EPIC-01": {"PARTIAL": 4, "EVIDENCE_MISSING": 1},
    "EPIC-02": {"PARTIAL": 13},
    "EPIC-03": {"PARTIAL": 2, "MISSING": 3},
    "EPIC-04": {"CONFLICTING": 3, "EVIDENCE_MISSING": 1},
    "EPIC-05": {"CONFLICTING": 3},
    "EPIC-06": {"CONFLICTING": 2, "PARTIAL": 3, "MISSING": 1},
    "EPIC-07": {"MISSING": 1, "PARTIAL": 1},
    "EPIC-08": {"MISSING": 3, "CONFLICTING": 4, "PARTIAL": 1},
    "EPIC-09": {"MISSING": 3, "PARTIAL": 3, "CONFLICTING": 2},
    "EPIC-10": {"CONFLICTING": 2, "PARTIAL": 1},
    "EPIC-11": {"PARTIAL": 4},
    "EPIC-12": {"EVIDENCE_MISSING": 2, "BLOCKED": 5},
}

BUNDLE_PURPOSES = {
    "BND-DOC-MANUAL": "문서 이름·검토·승인·버전 규칙",
    "BND-DOC-REGISTER": "전체 산출물 목록과 변경 이력",
    "BND-MGT-CHARTER": "사업 이유·목표·성공 기준·범위",
    "BND-MGT-PMP": "일정·작업·자원·역할·품질 관리",
    "BND-MGT-CONTROL": "위험·결정·변경·진행률·회의 조치",
    "BND-DSC-DISCOVERY": "사용자 문제·조사·여정·대안·기술 가능성",
    "BND-DSC-PRODUCT": "제품 방향·측정값·최소 제품·로드맵",
    "BND-DSC-REGISTER": "우선 작업과 계속·중단 판단",
    "BND-REQ-BASELINE": "기능·품질·보안·접근성 요구사항",
    "BND-REQ-ACCEPTANCE": "사용자 상황과 합격 조건",
    "BND-REQ-TRACE": "요구→설계→코드→시험 연결과 용어",
    "BND-DES-ARCH": "전체 구조와 실행·배포·연결 흐름",
    "BND-DES-INTERFACE-DATA": "서버 연결·데이터베이스·보존·변경",
    "BND-DES-UX-ACCESS": "화면 흐름과 접근성 설계",
    "BND-DES-SECOPS": "로그인·권한·개인정보·오류·성능·복구",
    "BND-DEV-GUIDE": "개발환경 설치·실행·시험·협업 방법",
    "BND-DEV-CONFIGURATION": "코드·외부 부품·빌드·DB·시험자료 구성",
    "BND-DEV-QUALITY": "코드검토·자동검사·라이선스·빌드 이력",
    "BND-TST-PLAN": "무엇을 어디서 어떻게 시험할지 정의",
    "BND-TST-EVIDENCE": "단위시험부터 접근성·복구까지 실제 결과",
    "BND-TST-QUALITY": "결함·시험범위·종합 결과·출시 판단",
    "BND-SEC-PLAN": "보안·개인정보 계획과 위험 관리",
    "BND-SEC-VERIFICATION": "보안 자동검사와 침투시험 결과",
    "BND-SEC-RESPONSE": "취약점·예외·사고 대응과 감시",
    "BND-AIML-DATA": "데이터 출처·품질·표시·분할 관리",
    "BND-AIML-MODEL": "학습 방법·실험·모델 목록과 설명",
    "BND-AIML-EVALUATION": "정확도·실수·안전·휴대폰 성능 평가",
    "BND-AIML-OPS": "모델 배포·되돌리기·성능 저하 감시",
    "BND-REL-CONTROL": "배포 버전·파일·지문·서명·변경사항",
    "BND-REL-DEPLOYMENT": "배포·초기 확인·단계 배포·복구",
    "BND-REL-DELIVERY": "사용자·관리자 안내와 인수인계",
    "BND-OPS-GUIDE": "담당자·서비스 목표·감시·알림·장애 절차",
    "BND-OPS-RECOVERY": "백업·복원·허용 중단시간·재해복구",
    "BND-OPS-CONTROL": "권한·키·패치·장애·비용·삭제·외부 서비스",
    "BND-WS-SAFETY": "WalkSafe 안전·신고·권한·개인정보·제한 고지",
    "BND-WS-ACCEPTANCE": "GPS·경로·AI·음성·진동·접근성·실폰 시험",
    "BND-WS-FIELD-SAFETY": "현장 참여자 동의와 안전·중단 계획",
    "BND-CLS-CLOSURE": "완료·최종 인수·목표 결과·회고",
    "BND-CLS-HANDOVER": "최종 목록·소스 보관·남은 문제·운영 이관",
    "BND-CLS-DECOMMISSION": "계약·데이터·계정·키·서비스 종료",
}

GLOSSARY = [
    ("정본(canonical)", "판단과 승인에 공식적으로 사용하는 파일이다."),
    ("기준선(baseline)", "승인한 파일 묶음을 버전과 파일 지문으로 고정한 시점이다."),
    ("Gap·Backlog r021", "2026-07-26에 고정한 21번째 기능 차이 분석과 개선 작업 목록이다."),
    ("DOC-01", "257개 산출물 유형의 상태·위치·책임자를 관리하는 공식 등록부다."),
    ("폐쇄 원장 R011", "산출물 유형별로 기준상 닫힘에 해당하는지 투영한 11번째 원장이다. 프로젝트 완료율은 아니다."),
    ("체크포인트(checkpoint)", "현재 활성 작업 묶음·공식 자료 버전·다음 통제 상태를 가리키는 실행 포인터다."),
    ("산출물 유형", "만들어야 할 결과물의 종류다. 실제 파일 수와 같지 않다."),
    ("문서 묶음(bundle)", "여러 산출물 유형을 한 파일이나 문서군에 모은 단위다."),
    ("EPIC", "서로 관련된 기능을 크게 묶은 작업 단위다."),
    ("Gap", "정한 요구와 현재 코드·증거 사이의 차이다."),
    ("출시 관문(Release Gate)", "정식 배포 전에 반드시 통과해야 하는 실측·검토·훈련이다."),
    ("정식 시험(Formal Test)", "승인된 환경·절차·대상 버전으로 실행하고 원자료와 판정을 남기는 시험이다."),
    ("Android", "WalkSafe 사용자 앱이 실행되는 구글 스마트폰 운영체제다."),
    ("TMAP", "목적지 검색과 큰 보행 경로를 제공하는 외부 지도 서비스다."),
    ("TalkBack", "화면의 글과 버튼을 음성으로 읽고 손가락 동작으로 조작하게 하는 Android 기능이다."),
    ("TFLite", "휴대폰 안에서 AI를 실행하기 위한 가벼운 모델 파일 형식이다."),
    ("ARCore Depth", "휴대폰 카메라와 센서로 주변 공간의 대략적인 거리를 계산하는 Android 기술이다."),
    ("CameraX", "Android에서 카메라 영상을 안정적으로 읽게 돕는 기술이다. 이것만으로 믿을 만한 실제 거리를 알 수 있는 것은 아니다."),
    ("STT·TTS", "STT는 사용자의 음성을 글로 바꾸고, TTS는 글로 된 안내를 음성으로 읽는다."),
    ("TTC", "현재 움직임이 계속된다고 가정했을 때 충돌까지 남은 것으로 추정되는 시간이다."),
    ("GPS·IMU", "GPS는 위치를, IMU는 휴대폰의 회전과 움직임을 측정한다."),
    ("Gateway·API", "Gateway는 서버 요청을 검사하는 중간 출입구이고, API는 앱과 서버가 정보를 주고받는 약속이다."),
    ("OpenAPI", "앱과 서버가 주고받을 주소·입력·결과 형식을 기계가 읽을 수 있게 적은 통신 규격이다."),
    ("PostgreSQL·PostGIS", "계정·신고 정보를 저장하고 위치 검색도 처리하는 데이터베이스 기술이다."),
    ("TOTP", "비밀번호와 함께 입력하는 짧은 일회용 인증번호다."),
    ("CI/CD", "코드 변경 때 빌드·시험을 자동 실행하고 승인된 결과만 전달하는 절차다."),
    ("test fixture", "시험할 때 같은 조건을 다시 만들기 위해 고정해 둔 입력·환경·상태다."),
    ("SBOM", "배포물에 포함된 프로그램 부품과 버전을 적은 목록이다."),
    ("SLA", "장애 대응시간이나 복구시간처럼 서비스가 지키기로 승인한 운영 기준이다."),
    ("IoU·confidence·threshold", "IoU는 영역 겹침 정도, confidence는 AI의 확신값, threshold는 기능을 실행할 최소 기준값이다."),
    ("Keystore", "로그인 정보나 암호화 열쇠를 일반 파일보다 안전하게 보관하는 Android 저장공간이다."),
    ("HTTPS·TLS", "통신을 암호화하고 연결한 서버가 맞는지 확인하는 방식이다."),
    ("SHA-256", "파일이 바뀌었는지 확인하는 긴 디지털 지문이다."),
    ("manifest", "함께 시험·배포할 앱·서버·AI·설정의 정확한 목록이다."),
    ("rollback", "문제가 생겼을 때 검증된 이전 정상 버전으로 되돌리는 절차다."),
    ("queue", "지금 처리하지 못한 신고나 자료를 순서대로 기다리게 하는 대기함이다."),
    ("object storage·KMS", "큰 파일을 보관하는 저장소와 암호화 열쇠를 분리해 관리하는 서비스다."),
    ("fail-closed", "안전하다고 확인할 수 없으면 기능을 시작하지 않거나 멈추는 원칙이다."),
    ("무가림 원본", "얼굴·번호판·목소리를 가리지 않은 영상·음성·위치 등 민감한 원자료다."),
]


@dataclass(frozen=True)
class Footnote:
    marker: str
    note: str


@dataclass
class Snapshot:
    checkpoint: dict[str, Any]
    gaps: dict[str, Any]
    backlog: dict[str, Any]
    register: dict[str, Any]
    test_cases: dict[str, Any]
    closure: dict[str, Any]
    source_hashes: dict[str, str]
    generated_at: str
    feature_counts: Counter[str]
    lifecycle_counts: Counter[str]
    category_lifecycle: dict[str, Counter[str]]
    dependency_matrix: dict[tuple[str, str], int]
    dependency_count: int
    bundles: dict[str, list[dict[str, Any]]]


@dataclass(frozen=True)
class FigureSpec:
    path: Path
    title: str
    alt_text: str
    caption: str


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def pct(value: int, total: int) -> str:
    return f"{value / total * 100:.1f}%" if total else "—"


def shorten(text: str, limit: int) -> str:
    clean = re.sub(r"\s+", " ", text or "").strip()
    if len(clean) <= limit:
        return clean
    clipped = clean[: limit - 1].rsplit(" ", 1)[0]
    return (clipped or clean[: limit - 1]) + "…"


def simplify(text: str) -> str:
    clean = re.sub(r"\s+", " ", text or "").strip()
    replacements = (
        ("canonical", "정본"),
        ("Canonical", "정본"),
        ("artifact", "산출물"),
        ("schema", "자료 구조"),
        ("manifest", "파일 목록·검증 기록"),
        ("receipt", "실행 확인 기록"),
        ("hash", "파일 지문"),
        ("SHA-256", "SHA-256 파일 지문"),
        ("rollback", "이전 정상판 복구"),
        ("gateway", "중간 서버"),
        ("Gateway", "중간 서버"),
        ("backend", "본 서버"),
        ("Backend", "본 서버"),
        ("endpoint", "서버 접속 지점"),
        ("idempotency", "중복 생성 방지"),
        ("queue", "대기함"),
        ("object storage", "대용량 파일 저장소"),
        ("train·validation·test", "학습·조정·독립시험"),
        ("false positive·negative", "잘못 탐지함·놓침"),
        ("false positive", "잘못 탐지함"),
        ("false negative", "놓침"),
        ("hyperparameter", "학습 설정값"),
        ("source commit", "원본 코드 버전"),
        ("commit", "코드 버전"),
        ("lock 파일", "버전 고정 목록"),
        ("hash drift", "등록값과 현재 파일 지문 차이"),
        ("confidence", "확신값"),
        ("threshold", "최소 기준값"),
        ("protocol", "시험 절차"),
        ("runtime", "실행 중"),
        ("fixture", "고정 시험 입력"),
        ("provenance", "생성·출처 이력"),
        ("migration", "자료 구조 변경"),
        ("class", "분류 대상"),
        ("FPS", "초당 처리 화면 수"),
        ("IoU", "영역 겹침값"),
        ("SBOM", "프로그램 부품 목록"),
        ("IaC", "서버 환경 자동 구성"),
        ("SLA", "서비스 대응시간 기준"),
        ("dashboard-as-code", "자동 구성 대시보드"),
        ("Web/PWA", "과거 웹 앱"),
        ("README", "개발 안내문"),
        ("revision", "개정판"),
        ("GiB", "약 10억 바이트 단위"),
        ("CameraX", "Android 카메라 처리 기술"),
        ("epoch", "기준 시점 세대"),
        ("foreground", "앱이 화면 앞에서 실행 중인 상태"),
        ("refresh", "다시 불러오기"),
        ("tombstone", "삭제·폐기 표식"),
        ("root origin", "최상위 서버 주소"),
        ("ingress", "외부 접속 입구"),
        ("metadata", "파일 설명정보"),
        ("bbox", "물체 위치 상자"),
        ("sequence", "촬영 흐름"),
        ("source", "원본"),
        ("DB", "자료 저장소"),
        ("gate", "관문"),
        ("OS", "운영체제"),
        ("timeout", "기다림 제한시간"),
        ("telemetry", "운영 상태 신호"),
        ("register", "등록부"),
        ("instance", "실행 대상"),
        ("health", "정상상태"),
        ("generated", "자동 생성"),
        ("coverage", "검사 범위"),
        ("certificate", "인증서"),
        ("validation", "검증"),
        ("token", "임시 접속값"),
        ("snapshot", "시점 상태"),
        ("skip", "제외"),
        ("secret", "비밀값"),
        ("scan", "검사"),
        ("runbook", "운영 절차서"),
        ("payload", "전달 내용"),
        ("password", "비밀번호"),
        ("package", "묶음"),
        ("release", "출시"),
        ("header", "머리글"),
        ("split", "분리"),
        ("metric", "측정 지표"),
        ("live", "실제 연결"),
        ("export", "내보내기"),
        ("dependency", "앞뒤 의존관계"),
        ("credential", "접속 인증정보"),
        ("config", "설정"),
        ("checklist", "점검표"),
        ("cache", "임시 저장"),
        ("builder", "자동 생성 도구"),
        ("OpenAPI", "서버 통신 규격"),
        ("API", "통신 규격"),
        ("STT", "음성을 글로 바꾸기"),
        ("TTS", "글을 음성으로 읽기"),
        ("TLS", "암호화 통신"),
        ("WCAG", "웹 접근성 기준"),
        ("evidence", "증거"),
        ("formal", "정식"),
        ("NOT_RUN", "미실행"),
        ("NOT_ELIGIBLE", "출시 불가"),
    )
    for source, target in replacements:
        clean = clean.replace(source, target)
    return clean


def snapshot() -> Snapshot:
    sources = {
        "활성 체크포인트": CHECKPOINT_PATH,
        "기능 Gap r021": GAP_PATH,
        "개선 Backlog r021": BACKLOG_PATH,
        "산출물 등록부 DOC-01": DOC01_PATH,
        "정식 시험 목록": TEST_CASES_PATH,
        "산출물 폐쇄 원장 R011": CLOSURE_LEDGER_PATH,
    }
    checkpoint = load_json(CHECKPOINT_PATH)
    gaps = load_json(GAP_PATH)
    backlog = load_json(BACKLOG_PATH)
    register = load_json(DOC01_PATH)
    test_cases = load_json(TEST_CASES_PATH)
    closure = load_json(CLOSURE_LEDGER_PATH)
    source_hashes = {label: sha256(path) for label, path in sources.items()}

    assessments = gaps["assessments"]
    artifacts = register["artifacts"]
    feature_counts = Counter(row["status"] for row in assessments)
    lifecycle_counts = Counter(row["state"]["lifecycle_status"] for row in artifacts)
    category_lifecycle: dict[str, Counter[str]] = {
        category: Counter() for category in CATEGORY_ORDER
    }
    for row in artifacts:
        category_lifecycle[row["category"]][row["state"]["lifecycle_status"]] += 1

    code_to_category = {row["artifact_type_code"]: row["category"] for row in artifacts}
    dependency_edges: set[tuple[str, str]] = set()
    dependency_matrix: dict[tuple[str, str], int] = defaultdict(int)
    for row in artifacts:
        target = row["artifact_type_code"]
        for upstream in row["trace"]["upstream_types"]:
            edge = (upstream, target)
            dependency_edges.add(edge)
    for upstream, target in dependency_edges:
        dependency_matrix[(code_to_category[upstream], code_to_category[target])] += 1

    bundles: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in artifacts:
        bundles[row["bundle_id"]].append(row)

    assert len(assessments) == 68
    assert sum(feature_counts.values()) == 68
    assert feature_counts == Counter(
        {
            "PARTIAL": 32,
            "CONFLICTING": 16,
            "MISSING": 11,
            "EVIDENCE_MISSING": 4,
            "BLOCKED": 5,
        }
    )
    assert len([row for row in assessments if row["source_kind"] != "REMAINING_GATE"]) == 63
    assert len([row for row in assessments if row["source_kind"] == "REMAINING_GATE"]) == 5
    assert len(artifacts) == 257
    assert len(bundles) == 40
    assert lifecycle_counts == Counter(
        {"APPROVED_BASELINED": 102, "ACTIVE": 27, "DRAFT": 53, "PLANNED": 75}
    )
    assert len(dependency_edges) == 698
    assert closure["summaries"]["open_artifact_count"] == 131
    assert closure["summaries"]["global_artifact_completion_claim_count"] == 0

    generated_at = datetime.now(ZoneInfo("Asia/Seoul")).replace(microsecond=0).isoformat()
    return Snapshot(
        checkpoint=checkpoint,
        gaps=gaps,
        backlog=backlog,
        register=register,
        test_cases=test_cases,
        closure=closure,
        source_hashes=source_hashes,
        generated_at=generated_at,
        feature_counts=feature_counts,
        lifecycle_counts=lifecycle_counts,
        category_lifecycle=category_lifecycle,
        dependency_matrix=dependency_matrix,
        dependency_count=len(dependency_edges),
        bundles=dict(bundles),
    )


def configure_matplotlib() -> None:
    for font_path in (FONT_REGULAR, FONT_BOLD):
        if font_path.is_file():
            font_manager.fontManager.addfont(str(font_path))
    if FONT_REGULAR.is_file():
        family = font_manager.FontProperties(fname=str(FONT_REGULAR)).get_name()
    else:
        family = "DejaVu Sans"
    plt.rcParams.update(
        {
            "font.family": family,
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "text.color": "#17212B",
            "axes.labelcolor": "#17212B",
            "xtick.color": "#52616D",
            "ytick.color": "#52616D",
            "axes.titleweight": "bold",
        }
    )


def save_figure(fig: plt.Figure, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def add_figure_title(fig: plt.Figure, title: str, subtitle: str = "") -> None:
    fig.suptitle(title, x=0.04, y=0.98, ha="left", fontsize=16, fontweight="bold", color="#17365D")
    if subtitle:
        fig.text(0.04, 0.925, subtitle, ha="left", va="top", fontsize=9, color="#52616D")


def create_cover_art(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12, 5.8))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6)
    ax.axis("off")
    ax.add_patch(patches.Rectangle((0, 0), 12, 6, color="#102B4E"))
    ax.add_patch(patches.Circle((10.7, 5.2), 2.2, color="#1F6FA8", alpha=0.45))
    ax.add_patch(patches.Circle((1.0, 0.2), 2.7, color="#007A5E", alpha=0.34))
    path_x = [0.6, 2.1, 3.2, 4.6, 6.0, 7.4, 9.0, 11.5]
    path_y = [1.0, 1.7, 1.35, 2.25, 2.0, 3.1, 2.8, 4.1]
    ax.plot(path_x, path_y, color="#F4C95D", linewidth=6, solid_capstyle="round")
    for x, y in zip(path_x, path_y):
        ax.add_patch(patches.Circle((x, y), 0.11, color="#FFFFFF"))

    # Smartphone
    ax.add_patch(
        patches.FancyBboxPatch(
            (5.1, 1.0),
            2.0,
            3.7,
            boxstyle="round,pad=0.12,rounding_size=0.25",
            facecolor="#F8FAFC",
            edgecolor="#BFD3E2",
            linewidth=2,
        )
    )
    ax.add_patch(patches.Rectangle((5.35, 1.45), 1.5, 2.75, color="#DDEBF5"))
    ax.add_patch(patches.Circle((6.1, 1.2), 0.08, color="#52616D"))
    ax.add_patch(patches.Circle((6.1, 2.75), 0.55, color="#1F6FA8"))
    ax.text(6.1, 2.75, "!", ha="center", va="center", color="white", fontsize=24, fontweight="bold")

    # Voice and vibration waves
    for radius, alpha in ((0.35, 0.85), (0.65, 0.6), (0.95, 0.35)):
        ax.add_patch(
            patches.Arc(
                (7.5, 3.0),
                radius * 2,
                radius * 2,
                theta1=-55,
                theta2=55,
                color="#F4C95D",
                linewidth=3,
                alpha=alpha,
            )
        )
    # Tactile blocks
    for row in range(3):
        for col in range(5):
            ax.add_patch(patches.Circle((9.15 + col * 0.28, 1.0 + row * 0.28), 0.07, color="#F4C95D"))

    ax.text(0.7, 5.15, "WALKSAFE", fontsize=25, color="white", fontweight="bold", va="top")
    ax.text(
        0.7,
        4.55,
        "보행 중 가까운 위험과 큰 이동 방향을\n음성·진동으로 알려주는 보조 시스템",
        fontsize=14,
        color="#E8F2F8",
        va="top",
        linespacing=1.5,
    )
    save_figure(fig, path)


def chart_user_journey(path: Path) -> None:
    steps = [
        ("1", "앱 준비", "로그인·동의·권한 확인"),
        ("2", "목적지 말하기", "검색 결과를 듣고 선택"),
        ("3", "걷기 시작", "큰 이동 방향 안내"),
        ("4", "주변 살피기", "사람·차량·턱 등 후보 찾기"),
        ("5", "위험 알림", "급한 위험을 음성·진동으로 안내"),
        ("6", "손상 신고", "점자블록 신고 후보를 보관"),
        ("7", "안전 종료", "전송·삭제·상태 확인"),
    ]
    fig, ax = plt.subplots(figsize=(13, 5.2))
    add_figure_title(
        fig,
        "사용자가 경험하려는 전체 흐름",
        "파란색은 앱 안에서 이어지는 사용자 여정이며, 각 단계의 공식 시험은 아직 끝나지 않았습니다.",
    )
    ax.set_xlim(0, len(steps) * 2.1)
    ax.set_ylim(0, 5)
    ax.axis("off")
    for index, (number, title, detail) in enumerate(steps):
        x = index * 2.1 + 0.2
        y = 1.5 if index % 2 == 0 else 2.15
        if index:
            prev_y = 1.5 if (index - 1) % 2 == 0 else 2.15
            ax.annotate(
                "",
                xy=(x - 0.08, y + 0.68),
                xytext=(x - 0.44, prev_y + 0.68),
                arrowprops=dict(arrowstyle="->", color="#7B8794", lw=2),
            )
        box = patches.FancyBboxPatch(
            (x, y),
            1.65,
            1.35,
            boxstyle="round,pad=0.08,rounding_size=0.12",
            facecolor="#E8F2F8",
            edgecolor="#1F6FA8",
            linewidth=1.7,
        )
        ax.add_patch(box)
        ax.add_patch(patches.Circle((x + 0.22, y + 1.08), 0.17, color="#17365D"))
        ax.text(x + 0.22, y + 1.08, number, ha="center", va="center", color="white", fontsize=9, fontweight="bold")
        ax.text(x + 0.47, y + 1.08, title, ha="left", va="center", fontsize=10, fontweight="bold", color="#17365D")
        ax.text(
            x + 0.12,
            y + 0.60,
            textwrap.fill(detail, 13),
            ha="left",
            va="center",
            fontsize=8.2,
            color="#37474F",
            linespacing=1.3,
        )
    ax.text(
        0.2,
        0.45,
        "안전 경계: WalkSafe는 흰지팡이·안내견·보호자를 대신하거나 보행 안전을 보장하지 않습니다.",
        fontsize=9.5,
        color="#963B4E",
        fontweight="bold",
    )
    save_figure(fig, path)


def chart_architecture(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12.8, 7.2))
    add_figure_title(
        fig,
        "WalkSafe가 정보를 처리하는 구조",
        "쉬운 이름을 크게 쓰고, 바꿀 수 없는 기술명은 괄호에만 남겼습니다.",
    )
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 8)
    ax.axis("off")

    def box(x: float, y: float, w: float, h: float, title: str, lines: Sequence[str], color: str) -> None:
        ax.add_patch(
            patches.FancyBboxPatch(
                (x, y),
                w,
                h,
                boxstyle="round,pad=0.08,rounding_size=0.16",
                facecolor="white",
                edgecolor=color,
                linewidth=2,
            )
        )
        ax.add_patch(patches.Rectangle((x, y + h - 0.55), w, 0.55, color=color))
        ax.text(x + 0.18, y + h - 0.28, title, color="white", fontsize=10.5, fontweight="bold", va="center")
        for index, line in enumerate(lines):
            ax.text(x + 0.18, y + h - 0.92 - index * 0.42, "• " + line, fontsize=8.5, color="#37474F", va="top")

    box(
        0.4,
        3.4,
        3.2,
        3.4,
        "사용자 휴대폰 앱 (Android)",
        [
            "카메라로 주변 물체 후보 찾기",
            "지원 기기에서 거리·접근 정도 판단",
            "목적지·경로·이탈 안내",
            "음성·화면읽기·진동",
            "동의·신고·삭제 요청",
        ],
        "#1F6FA8",
    )
    box(
        4.25,
        4.05,
        2.7,
        2.5,
        "요청 검사 중간 서버",
        ["사용자 세션 확인", "동의·보행 상태 확인", "요청량·업로드 제한", "본 서버용 인증 추가"],
        "#007A5E",
    )
    box(
        7.6,
        3.4,
        2.8,
        3.4,
        "정보 처리 본 서버",
        ["길찾기·목적지 중계", "신고 저장·조회", "관리자 보안", "운영 준비상태 확인"],
        "#B95300",
    )
    box(
        10.95,
        4.35,
        1.65,
        2.15,
        "외부·저장",
        ["TMAP 지도", "신고 DB", "현재 로컬 파일"],
        "#6F4C9B",
    )
    box(
        4.25,
        0.8,
        2.7,
        2.35,
        "관리자 휴대폰 앱",
        ["로그인·추가 확인", "세션 조회·차단", "복구 흐름", "운영 업무 화면은 미완성"],
        "#59636E",
    )
    ax.annotate("", xy=(4.1, 5.25), xytext=(3.65, 5.25), arrowprops=dict(arrowstyle="->", lw=2.2, color="#52616D"))
    ax.annotate("", xy=(7.45, 5.25), xytext=(7.0, 5.25), arrowprops=dict(arrowstyle="->", lw=2.2, color="#52616D"))
    ax.annotate("", xy=(10.8, 5.25), xytext=(10.45, 5.25), arrowprops=dict(arrowstyle="->", lw=2.2, color="#52616D"))
    ax.annotate("", xy=(7.65, 3.25), xytext=(6.95, 2.45), arrowprops=dict(arrowstyle="->", lw=2.0, color="#59636E"))
    ax.text(3.78, 5.55, "HTTPS", fontsize=7.5, color="#52616D", ha="center")
    ax.text(7.2, 5.55, "내부 연결", fontsize=7.5, color="#52616D", ha="center")
    ax.text(7.35, 2.15, "관리자 전용 연결", fontsize=7.5, color="#52616D", ha="center")
    ax.text(
        0.45,
        0.25,
        "제품 경계: 사용자 Android 앱 + 별도 관리자 Android 앱 + 중간 서버 + 본 서버. Web/PWA와 Python 음성 서버는 현행 제품 경로가 아닙니다.",
        fontsize=8.7,
        color="#52616D",
    )
    save_figure(fig, path)


def chart_feature_status(snap: Snapshot, path: Path) -> None:
    policy_order = ["IMPLEMENTED", "PARTIAL", "CONFLICTING", "MISSING", "EVIDENCE_MISSING"]
    counts = {status: snap.feature_counts.get(status, 0) for status in policy_order}
    fig = plt.figure(figsize=(12.5, 7.2))
    grid = fig.add_gridspec(2, 1, height_ratios=[1.1, 1.2], hspace=0.5)
    add_figure_title(
        fig,
        "필요한 기능·정책 63개의 현재 충족 상태",
        "출시 전 확인 5개는 기능과 성격이 달라 아래쪽에 따로 표시했습니다.",
    )
    ax = fig.add_subplot(grid[0])
    left = 0
    for status in policy_order:
        value = counts[status]
        if value:
            ax.barh(
                [0],
                [value],
                left=left,
                height=0.48,
                color=STATUS_COLORS[status],
                edgecolor="white",
                hatch=STATUS_HATCHES[status],
                label=f"{STATUS_LABELS[status]} {value}",
            )
            ax.text(left + value / 2, 0, f"{value}\n{pct(value, 63)}", ha="center", va="center", fontsize=9, color="white", fontweight="bold")
        left += value
    ax.set_xlim(0, 63)
    ax.set_yticks([])
    ax.set_xlabel("기능·공통 규칙 63개")
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.grid(axis="x", color="#D9E1E8", linewidth=0.7)
    ax.legend(ncol=3, loc="upper center", bbox_to_anchor=(0.5, -0.32), frameon=False, fontsize=8.5)
    ax.text(
        0,
        -0.75,
        "중요: ‘일부 구현 32개’는 50.8% 완료가 아닙니다. 기능마다 남은 양이 다르고 정식 시험·승인이 끝나지 않았습니다.",
        fontsize=9,
        color="#963B4E",
        fontweight="bold",
        transform=ax.transData,
    )

    ax2 = fig.add_subplot(grid[1])
    ax2.set_xlim(0, 10)
    ax2.set_ylim(0, 3.2)
    ax2.axis("off")
    gate_titles = [
        "휴대폰\n대기자료 용량",
        "서버 용량상태\n전달 규칙",
        "민감 원본\n독립 검토",
        "실제 클라우드\n비용 측정",
        "관리자 폰 분실\n복구 훈련",
    ]
    for index, title in enumerate(gate_titles):
        x = 0.2 + index * 1.95
        ax2.add_patch(
            patches.FancyBboxPatch(
                (x, 0.85),
                1.65,
                1.55,
                boxstyle="round,pad=0.06,rounding_size=0.12",
                facecolor="#F2F4F6",
                edgecolor="#59636E",
                linewidth=1.5,
                hatch="++",
            )
        )
        ax2.text(x + 0.82, 1.78, title, ha="center", va="center", fontsize=8.7, fontweight="bold", color="#37474F")
        ax2.text(x + 0.82, 1.12, "미실행", ha="center", va="center", fontsize=8.5, color="#963B4E", fontweight="bold")
    ax2.text(0.2, 2.75, "출시 전 반드시 확인할 5개 관문: 통과 0 / 미실행 5", fontsize=11, color="#17365D", fontweight="bold")
    ax2.text(0.2, 0.3, "기준: Gap r021 · 기능·정책 분모 63, 출시 관문 분모 5 · 전체 프로젝트 완료율로 합산하지 않음", fontsize=8.7, color="#52616D")
    save_figure(fig, path)


def chart_epics(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(12.5, 8.4))
    add_figure_title(
        fig,
        "12개 큰 기능 묶음의 진행 위치",
        "‘구현 준비’는 출시 준비가 아니라 내부 구현 작업 묶음의 현재 위치를 뜻합니다.",
    )
    ax.set_xlim(0, 10)
    ax.set_ylim(-0.6, 12.4)
    ax.axis("off")
    status_color = {
        "IMPLEMENTATION_READY": "#007A5E",
        "IN_PROGRESS": "#1F6FA8",
        "PLANNED": "#7B8794",
    }
    status_label = {
        "IMPLEMENTATION_READY": "내부 구현 준비",
        "IN_PROGRESS": "진행 중",
        "PLANNED": "계획",
    }
    for row_index, (epic_id, title, description, status) in enumerate(reversed(EPIC_INFO)):
        y = row_index
        ax.add_patch(patches.Rectangle((0.2, y - 0.32), 9.4, 0.64, facecolor="#F7F9FA", edgecolor="#D7DEE5", linewidth=0.7))
        ax.add_patch(patches.Rectangle((0.2, y - 0.32), 0.15, 0.64, facecolor=status_color[status], edgecolor="none"))
        ax.text(0.48, y, epic_id, va="center", fontsize=8.5, color="#52616D", fontweight="bold")
        ax.text(1.5, y, title, va="center", fontsize=9.3, color="#17365D", fontweight="bold")
        ax.text(3.35, y, shorten(description, 38), va="center", fontsize=8.5, color="#37474F")
        ax.add_patch(
            patches.FancyBboxPatch(
                (8.0, y - 0.20),
                1.35,
                0.40,
                boxstyle="round,pad=0.03,rounding_size=0.08",
                facecolor=status_color[status],
                edgecolor="none",
            )
        )
        ax.text(8.675, y, status_label[status], ha="center", va="center", fontsize=7.7, color="white", fontweight="bold")
    ax.text(0.2, 12.05, "내부 구현 준비 1  |  진행 중 2  |  계획 9", fontsize=10.5, color="#17365D", fontweight="bold")
    save_figure(fig, path)


def chart_artifact_map(snap: Snapshot, path: Path) -> None:
    planned_paths = {
        row["location"]["planned_canonical_path"]
        for row in snap.register["artifacts"]
        if row["location"].get("planned_canonical_path")
    }
    current_paths = {
        row["location"]["canonical_path"]
        for row in snap.register["artifacts"]
        if row["location"].get("canonical_path")
    }
    fig, ax = plt.subplots(figsize=(12.2, 5.7))
    add_figure_title(
        fig,
        "257개 산출물 유형이 실제 파일로 모이는 방식",
        "산출물 유형 수와 파일 수를 같은 숫자로 읽으면 안 됩니다.",
    )
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5)
    ax.axis("off")
    nodes = [
        (0.4, 1.5, 2.25, 2.0, "257개", "챙겨야 할\n산출물 유형", "#17365D"),
        (3.25, 1.5, 2.25, 2.0, "40개", "관련 항목을 모은\n문서 묶음", "#1F6FA8"),
        (6.1, 1.5, 2.25, 2.0, f"{len(planned_paths)}개", "예정된 정본\n경로", "#007A5E"),
        (8.95, 2.25, 2.55, 1.25, f"{len(current_paths)}개", "현재 정본으로 연결", "#B95300"),
        (8.95, 0.85, 2.55, 1.05, f"{len(planned_paths-current_paths)}개", "실행 결과 대기 경로", "#7B8794"),
    ]
    for x, y, w, h, number, label, color in nodes:
        ax.add_patch(
            patches.FancyBboxPatch(
                (x, y),
                w,
                h,
                boxstyle="round,pad=0.08,rounding_size=0.14",
                facecolor="white",
                edgecolor=color,
                linewidth=2.0,
            )
        )
        ax.text(x + w / 2, y + h * 0.64, number, ha="center", va="center", fontsize=21, color=color, fontweight="bold")
        ax.text(x + w / 2, y + h * 0.28, label, ha="center", va="center", fontsize=9, color="#37474F")
    for start, end, y in ((2.65, 3.20, 2.5), (5.50, 6.05, 2.5), (8.35, 8.9, 2.8)):
        ax.annotate("", xy=(end, y), xytext=(start, y), arrowprops=dict(arrowstyle="->", color="#52616D", lw=2.2))
    ax.annotate("", xy=(9.0, 1.38), xytext=(8.35, 2.25), arrowprops=dict(arrowstyle="->", color="#52616D", lw=1.7))
    ax.text(
        0.4,
        0.25,
        "‘실행 결과 대기 경로’에는 계획·빈 양식 파일이 존재할 수 있지만, 실제 시험·사건 결과가 생긴 정본으로 인정되지는 않습니다.",
        fontsize=8.7,
        color="#52616D",
    )
    save_figure(fig, path)


def chart_artifact_lifecycle(snap: Snapshot, path: Path) -> None:
    fig = plt.figure(figsize=(13.0, 9.2))
    grid = fig.add_gridspec(2, 1, height_ratios=[0.8, 3.0], hspace=0.35)
    add_figure_title(
        fig,
        "257개 산출물 유형의 관리 상태",
        "이 그래프는 문서 작성·승인·실행 상태이며 프로젝트 기능 완료율이 아닙니다.",
    )
    order = ["APPROVED_BASELINED", "ACTIVE", "DRAFT", "PLANNED"]
    ax = fig.add_subplot(grid[0])
    left = 0
    for state in order:
        value = snap.lifecycle_counts[state]
        ax.barh([0], [value], left=left, height=0.5, color=LIFECYCLE_COLORS[state], edgecolor="white", label=f"{LIFECYCLE_LABELS[state]} {value}")
        ax.text(left + value / 2, 0, f"{value}\n{pct(value,257)}", ha="center", va="center", fontsize=8.7, color="white", fontweight="bold")
        left += value
    ax.set_xlim(0, 257)
    ax.set_yticks([])
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.grid(axis="x", color="#D9E1E8", linewidth=0.7)
    ax.legend(ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.35), frameon=False, fontsize=8.2)

    ax2 = fig.add_subplot(grid[1])
    y_positions = list(range(len(CATEGORY_ORDER)))
    left_values = [0] * len(CATEGORY_ORDER)
    for state in order:
        values = [snap.category_lifecycle[category][state] for category in CATEGORY_ORDER]
        ax2.barh(
            y_positions,
            values,
            left=left_values,
            height=0.64,
            color=LIFECYCLE_COLORS[state],
            edgecolor="white",
            label=LIFECYCLE_LABELS[state],
        )
        left_values = [left + value for left, value in zip(left_values, values)]
    labels = [f"{category}  {CATEGORY_INFO[category][0]}" for category in CATEGORY_ORDER]
    ax2.set_yticks(y_positions, labels)
    ax2.invert_yaxis()
    ax2.set_xlabel("산출물 유형 수")
    ax2.spines[["top", "right"]].set_visible(False)
    ax2.grid(axis="x", color="#E0E6EB", linewidth=0.7)
    for y, total in zip(y_positions, left_values):
        ax2.text(total + 0.25, y, str(total), va="center", fontsize=8, color="#37474F")
    save_figure(fig, path)


def chart_open_lanes(snap: Snapshot, path: Path) -> None:
    route_counts = snap.closure["summaries"]["current_queue_route_counts"]
    labels = ["내부 문서·내용 보완", "사실·책임자 승인·독립 확인", "내부 실행 필요", "실제 사건·외부 활동 필요"]
    values = [
        route_counts["INTERNAL_READY"],
        route_counts["EVIDENCE_FACT_PENDING"] + route_counts["OWNER_APPROVAL_PENDING"] + route_counts["ATTESTATION_REVIEW_PENDING"],
        route_counts["INTERNAL_RUN_REQUIRED"],
        route_counts["REAL_EVENT_PENDING"],
    ]
    colors = ["#1F6FA8", "#6F4C9B", "#D89000", "#59636E"]
    fig, ax = plt.subplots(figsize=(11.8, 6.0))
    add_figure_title(
        fig,
        "아직 닫히지 않은 산출물 131개의 남은 일",
        "126/257 ‘닫힘에 해당’ 투영은 문서 관점일 뿐 전체 프로젝트 49% 완료를 뜻하지 않습니다.",
    )
    y = list(range(len(labels)))
    bars = ax.barh(y, values, color=colors, height=0.62)
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlim(0, 70)
    ax.set_xlabel("열린 산출물 유형 수")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#E0E6EB", linewidth=0.8)
    for bar, value in zip(bars, values):
        ax.text(value + 0.8, bar.get_y() + bar.get_height() / 2, f"{value}  ({pct(value,131)})", va="center", fontsize=9, fontweight="bold", color="#37474F")
    ax.text(
        0,
        4.1,
        "세부: 사실 확인 6 + 담당자 승인 14 + 독립 확인 4 = 24",
        fontsize=8.8,
        color="#52616D",
    )
    save_figure(fig, path)


def chart_dependency_heatmap(snap: Snapshot, path: Path) -> None:
    matrix = [
        [snap.dependency_matrix.get((source, target), 0) for target in CATEGORY_ORDER]
        for source in CATEGORY_ORDER
    ]
    fig, ax = plt.subplots(figsize=(11.2, 9.0))
    add_figure_title(
        fig,
        "산출물 범주 사이의 연결관계",
        "앞 문서를 바꾸면 뒤 문서도 함께 확인해야 하는 방향 연결 698개를 13×13으로 묶었습니다.",
    )
    cmap = LinearSegmentedColormap.from_list("walksafe", ["#F7FAFC", "#B7D8E8", "#1F6FA8", "#17365D"])
    image = ax.imshow(matrix, cmap=cmap, aspect="auto")
    ax.set_xticks(range(len(CATEGORY_ORDER)), CATEGORY_ORDER, rotation=45, ha="right")
    ax.set_yticks(range(len(CATEGORY_ORDER)), CATEGORY_ORDER)
    ax.set_xlabel("영향을 받는 뒤쪽 산출물")
    ax.set_ylabel("먼저 확인할 앞쪽 산출물")
    for row in range(len(CATEGORY_ORDER)):
        for column in range(len(CATEGORY_ORDER)):
            value = matrix[row][column]
            if value:
                ax.text(column, row, str(value), ha="center", va="center", fontsize=7.2, color="white" if value >= 10 else "#37474F")
    colorbar = fig.colorbar(image, ax=ax, fraction=0.035, pad=0.03)
    colorbar.set_label("방향 연결 수")
    fig.text(
        0.04,
        0.03,
        "같은 범주 안 394개 · 다른 범주 사이 304개 · 중복·자기연결·순환 없음. 반대 방향 조회를 다시 더하지 않았습니다.",
        fontsize=8.5,
        color="#52616D",
    )
    save_figure(fig, path)


def chart_formal_tests(snap: Snapshot, path: Path) -> None:
    method_counts = snap.test_cases["summary"]["method_counts"]
    label_map = {
        "FUNCTIONAL": "기능 시험",
        "DEVICE_OR_FIELD": "실제 기기·현장",
        "FAILURE_OR_RECOVERY": "장애·복구",
        "INTEGRATION_OR_CONTRACT": "시스템 연결·약속",
        "ACCESSIBILITY_OR_USABILITY": "접근성·사용성",
        "PERFORMANCE_OR_CAPACITY": "성능·용량",
        "MEASUREMENT": "실제 측정",
        "ENGINEERING_REVIEW": "기술설계 검토",
        "INDEPENDENT_REVIEW": "독립 검토",
        "RECOVERY_DRILL": "복구 훈련",
    }
    ordered = sorted(method_counts.items(), key=lambda item: item[1])
    labels = [label_map[key] for key, _ in ordered]
    values = [value for _, value in ordered]
    fig, ax = plt.subplots(figsize=(11.8, 7.2))
    add_figure_title(
        fig,
        "정식 시험 279개의 구성",
        "미승인 초안에 적힌 시험은 모두 미실행입니다. 내부 자동시험 통과 기록과 정식 시험은 서로 다른 증거입니다.",
    )
    bars = ax.barh(range(len(labels)), values, color="#1F6FA8", height=0.62)
    ax.set_yticks(range(len(labels)), labels)
    ax.set_xlabel("계획된 정식 시험 수")
    ax.set_xlim(0, 132)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#E0E6EB", linewidth=0.8)
    for bar, value in zip(bars, values):
        ax.text(value + 1.2, bar.get_y() + bar.get_height() / 2, str(value), va="center", fontsize=9, fontweight="bold")
    ax.text(77, 1.2, "실행 0 / 합격 0 / 실패 0", fontsize=11, color="#963B4E", fontweight="bold")
    ax.text(77, 0.55, "현재 상태: 미승인 초안 · 279개 모두 미실행", fontsize=9.5, color="#52616D")
    save_figure(fig, path)


def chart_evidence_ladder(path: Path) -> None:
    steps = [
        ("1", "코드 존재", "파일·함수·설정이 있음"),
        ("2", "앱에 연결", "사용 흐름에서 실제 호출"),
        ("3", "내부 시험", "개발 중 자동검사 통과"),
        ("4", "정식 시험", "승인 환경·버전·원자료"),
        ("5", "검토·승인", "사람과 독립 검토 포함"),
        ("6", "서명·배포·복구", "운영 가능한 출시본"),
    ]
    fig, ax = plt.subplots(figsize=(12.4, 5.8))
    add_figure_title(
        fig,
        "‘코드가 있음’에서 ‘출시 가능’까지 필요한 증거",
        "현재 기능들은 1~3단계에 섞여 있으며, 정식 시험 0/279·출시 관문 0/5입니다.",
    )
    ax.set_xlim(0, 12.5)
    ax.set_ylim(0, 5.2)
    ax.axis("off")
    for index, (number, title, detail) in enumerate(steps):
        x = 0.25 + index * 2.05
        height = 1.0 + index * 0.46
        color = "#1F6FA8" if index < 3 else "#7B8794"
        face = "#E8F2F8" if index < 3 else "#F2F4F6"
        ax.add_patch(
            patches.FancyBboxPatch(
                (x, 0.65),
                1.65,
                height,
                boxstyle="round,pad=0.06,rounding_size=0.1",
                facecolor=face,
                edgecolor=color,
                linewidth=1.8,
                hatch="" if index < 3 else "++",
            )
        )
        ax.add_patch(patches.Circle((x + 0.24, 0.95 + height - 0.35), 0.16, color=color))
        ax.text(x + 0.24, 0.95 + height - 0.35, number, ha="center", va="center", color="white", fontsize=8, fontweight="bold")
        ax.text(x + 0.82, 0.95 + height - 0.35, title, ha="center", va="center", fontsize=9, fontweight="bold", color="#17365D")
        ax.text(x + 0.82, 0.95 + height - 0.87, textwrap.fill(detail, 10), ha="center", va="center", fontsize=7.8, color="#37474F")
        if index < len(steps) - 1:
            ax.annotate("", xy=(x + 2.0, 1.1 + height), xytext=(x + 1.68, 1.1 + height), arrowprops=dict(arrowstyle="->", color="#52616D", lw=1.8))
    ax.text(0.25, 0.15, "파란 단계: 일부 기능에서 내부 근거 존재  |  회색 단계: 프로젝트 전체의 정식 종료 근거가 아직 없음", fontsize=8.8, color="#52616D")
    save_figure(fig, path)


def integrity_counts(snap: Snapshot) -> tuple[int, int, int]:
    matching = 0
    mismatching = 0
    self_reference = 0
    for row in snap.register["artifacts"]:
        path_value = row["location"].get("canonical_path")
        if not path_value:
            continue
        expected = row["integrity"].get("sha256")
        if not expected:
            self_reference += 1
            continue
        actual_path = ROOT / path_value
        if actual_path.is_file() and sha256(actual_path) == expected:
            matching += 1
        else:
            mismatching += 1
    return matching, mismatching, self_reference


def chart_document_sync(snap: Snapshot, path: Path) -> None:
    matching, mismatching, self_reference = integrity_counts(snap)
    total = matching + mismatching + self_reference
    fig = plt.figure(figsize=(12.2, 6.5))
    grid = fig.add_gridspec(1, 2, wspace=0.33)
    add_figure_title(
        fig,
        "현재 등록부와 실제 파일의 동기화 상태",
        "기존 산출물 형식은 이 소개서에서 수정하지 않습니다. 아래 수치는 별도 정비가 필요한 현재 상태입니다.",
    )
    fig.subplots_adjust(top=0.78, bottom=0.2, left=0.08, right=0.97)
    ax = fig.add_subplot(grid[0])
    left = 0
    pieces = [
        ("파일 지문 일치", matching, "#007A5E"),
        ("파일 지문 불일치", mismatching, "#B95300"),
        ("자기참조 예외", self_reference, "#7B8794"),
    ]
    for label, value, color in pieces:
        ax.barh([0], [value], left=left, color=color, height=0.48, edgecolor="white", label=f"{label} {value}")
        if value >= 5:
            ax.text(left + value / 2, 0, f"{value}\n{pct(value,total)}", ha="center", va="center", fontsize=9, color="white", fontweight="bold")
        left += value
    ax.set_xlim(0, total)
    ax.set_yticks([])
    ax.set_title("DOC-01의 현재 파일 연결 182행", fontsize=11, color="#17365D", pad=14)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.24), frameon=False, fontsize=8)
    ax.text(0, -0.55, "같은 파일이 여러 산출물 유형을 담으므로 182는 파일 수가 아닙니다.", fontsize=8.5, color="#52616D")

    ax2 = fig.add_subplot(grid[1])
    values = [1, 6]
    labels = ["현재 일치", "현재 불일치"]
    colors = ["#007A5E", "#B95300"]
    bars = ax2.bar(labels, values, color=colors, width=0.55)
    ax2.set_ylim(0, 7)
    ax2.set_ylabel("분야별 생성기 묶음")
    ax2.set_title("결정론적 생성기 7개 검사", fontsize=11, color="#17365D", pad=14)
    ax2.spines[["top", "right"]].set_visible(False)
    ax2.grid(axis="y", color="#E0E6EB", linewidth=0.8)
    for bar, value in zip(bars, values):
        ax2.text(bar.get_x() + bar.get_width() / 2, value + 0.15, str(value), ha="center", fontsize=12, fontweight="bold")
    ax2.text(
        0.5,
        -1.25,
        "이 결과는 문서 동기화 문제이며\n기능 구현 완료율과는 별개입니다.",
        ha="center",
        fontsize=8.5,
        color="#52616D",
        transform=ax2.transData,
    )
    save_figure(fig, path)


def chart_roadmap(path: Path) -> None:
    phases = [
        ("1", "안전 핵심 연결", ["지원 기기·거리 실제 확인", "도착 확인·재경로 선택", "호출어·안전정지·TalkBack"]),
        ("2", "신고·개인정보", ["암호화 대기함·재시도", "관리자 검수·기관 전달", "동의·삭제·보관 영수증"]),
        ("3", "AI·서버·운영", ["승인 모델·독립 시험자료", "대용량 저장소·암호화 열쇠", "용량·비용·관제·복원"]),
        ("4", "정식 검증·출시", ["279개 정식 시험", "실제 사용자·현장 시험", "5개 관문·서명·단계 배포"]),
    ]
    fig, ax = plt.subplots(figsize=(12.5, 6.2))
    add_figure_title(
        fig,
        "남은 작업을 이해하기 위한 권장 단계",
        "세부 실행 순서는 새 통제 전환과 검토 뒤 확정해야 하며, 이 그림 자체는 실행 승인서가 아닙니다.",
    )
    ax.set_xlim(0, 12.5)
    ax.set_ylim(0, 5.5)
    ax.axis("off")
    colors = ["#1F6FA8", "#007A5E", "#B95300", "#6F4C9B"]
    for index, ((number, title, tasks), color) in enumerate(zip(phases, colors)):
        x = 0.35 + index * 3.05
        ax.add_patch(
            patches.FancyBboxPatch(
                (x, 1.0),
                2.55,
                3.25,
                boxstyle="round,pad=0.08,rounding_size=0.15",
                facecolor="white",
                edgecolor=color,
                linewidth=2,
            )
        )
        ax.add_patch(patches.Rectangle((x, 3.55), 2.55, 0.7, color=color))
        ax.text(x + 0.25, 3.9, number, ha="center", va="center", color="white", fontsize=13, fontweight="bold")
        ax.text(x + 1.35, 3.9, title, ha="center", va="center", color="white", fontsize=10, fontweight="bold")
        for task_index, task in enumerate(tasks):
            ax.text(x + 0.2, 3.12 - task_index * 0.75, "• " + textwrap.fill(task, 14), fontsize=8.5, color="#37474F", va="top")
        if index < len(phases) - 1:
            ax.annotate("", xy=(x + 2.98, 2.6), xytext=(x + 2.6, 2.6), arrowprops=dict(arrowstyle="->", color="#52616D", lw=2))
    ax.text(0.35, 0.35, "안전 영향이 큰 기능 → 민감자료와 신고 → 운영 기반 → 정식 검증 순으로 묶어 이해한 권장안", fontsize=8.8, color="#52616D")
    save_figure(fig, path)


def create_figures(snap: Snapshot, directory: Path) -> tuple[Path, list[FigureSpec]]:
    configure_matplotlib()
    cover = directory / "00-cover.png"
    create_cover_art(cover)
    specs: list[FigureSpec] = []

    def add(
        name: str,
        title: str,
        alt_text: str,
        caption: str,
        builder,
        *args,
    ) -> None:
        figure_path = directory / name
        builder(*args, figure_path)
        specs.append(FigureSpec(figure_path, title, alt_text, caption))

    add(
        "01-user-journey.png",
        "사용자가 경험하려는 전체 흐름",
        "앱 준비부터 목적지 선택, 걷기, 위험 알림, 손상 신고, 종료까지 일곱 단계를 화살표로 연결한 흐름도",
        "앱 준비→목적지 선택→보행 안내→주변 위험 확인→음성·진동 알림→손상 신고→안전 종료. 모든 단계는 공식 시험 전입니다.",
        chart_user_journey,
    )
    add(
        "02-architecture.png",
        "WalkSafe 정보 처리 구조",
        "사용자 Android 앱이 요청 검사 중간 서버와 본 서버를 거쳐 TMAP과 데이터베이스를 사용하고, 관리자 앱은 별도 경로로 본 서버에 연결되는 구조도",
        "사용자 앱·관리자 앱·중간 서버·본 서버의 현재 제품 경계. Web/PWA와 Python 음성 서버는 현행 제품 경로가 아닙니다.",
        chart_architecture,
    )
    add(
        "03-feature-status.png",
        "필요한 기능·정책 63개와 출시 관문 5개",
        "기능 63개를 일부 구현 32, 기준과 다름 16, 핵심 구현 없음 11, 시험 증거 부족 4, 공식 완료 0으로 나눈 누적 막대와 미실행 출시 관문 다섯 개",
        "기능·정책 63개 중 공식 완료 0, 일부 구현 32, 기준과 다름 16, 핵심 구현 없음 11, 시험·증거 부족 4. 출시 관문은 0/5입니다.",
        chart_feature_status,
        snap,
    )
    add(
        "04-epics.png",
        "12개 큰 기능 묶음",
        "제품 경계, 안전 보행, 계정 보안 등 열두 개 큰 기능 묶음을 내부 구현 준비 1, 진행 중 2, 계획 9로 표시한 목록",
        "12개 큰 기능 묶음의 작업 위치. 내부 구현 준비 1, 진행 중 2, 계획 9이며 프로젝트 완료율이 아닙니다.",
        chart_epics,
    )
    add(
        "05-artifact-map.png",
        "산출물 유형과 파일의 관계",
        "257개 산출물 유형이 40개 문서 묶음과 41개 예정 경로를 거쳐 37개 현재 정본 경로와 4개 실행 결과 대기 경로로 모이는 흐름도",
        "257은 파일 수가 아니라 챙겨야 할 결과물 종류입니다. 여러 유형이 같은 파일의 서로 다른 장에 들어갑니다.",
        chart_artifact_map,
        snap,
    )
    add(
        "06-artifact-lifecycle.png",
        "산출물 관리 상태",
        "257개 산출물을 승인 기준본 102, 계속 갱신 27, 초안 53, 계획 미실행 75로 나눈 전체 막대와 13개 범주별 누적 막대",
        "DOC-01 기준 산출물 유형 257개의 작성·승인·실행 상태. 기능·출시 완료율이 아닙니다.",
        chart_artifact_lifecycle,
        snap,
    )
    add(
        "07-open-lanes.png",
        "열린 산출물 131개의 남은 일",
        "내부 문서 보완 62, 사실과 승인과 독립 확인 24, 내부 실행 24, 실제 사건과 외부 활동 21을 나타낸 가로 막대",
        "열린 산출물 131개를 닫기 위해 필요한 작업 종류. ‘닫힘에 해당 126/257’은 문서 관점 투영일 뿐 전체 프로젝트 완료율이 아닙니다.",
        chart_open_lanes,
        snap,
    )
    add(
        "08-dependency-heatmap.png",
        "산출물 연결관계",
        "문서관리부터 종료까지 13개 범주의 앞뒤 연결 수를 색 농도와 숫자로 나타낸 13대13 표",
        "방향 연결 698개 중 같은 범주 안 394개, 다른 범주 사이 304개. 반대 방향 조회를 중복 계산하지 않았습니다.",
        chart_dependency_heatmap,
        snap,
    )
    add(
        "09-formal-tests.png",
        "정식 시험 279개 구성",
        "기능 시험 123, 실제 기기와 현장 54, 장애와 복구 46, 시스템 연결 32, 접근성 15 등 열 종류를 나타낸 가로 막대",
        "미승인 초안에 계획된 정식 시험 279개는 모두 미실행입니다. 내부 자동시험 결과와 합산하지 않습니다.",
        chart_formal_tests,
        snap,
    )
    add(
        "10-evidence-ladder.png",
        "코드에서 출시까지의 증거 사다리",
        "코드 존재, 앱 연결, 내부 시험, 정식 시험, 검토 승인, 서명 배포 복구의 여섯 단계를 계단처럼 표시",
        "코드가 있다는 사실만으로 출시할 수 없습니다. 정식 시험·사람 검토·서명·배포·복구 근거까지 이어져야 합니다.",
        chart_evidence_ladder,
    )
    add(
        "11-document-sync.png",
        "등록부와 파일 동기화",
        "DOC-01 파일 연결 182행 중 파일 지문 일치 58, 불일치 123, 자기참조 1과 생성기 검사 일치 1, 불일치 6을 비교",
        "현재 작업트리에서 DOC-01과 파일의 연결 및 생성기 동기화가 깨진 상태입니다. 이 소개서는 기존 산출물을 수정하지 않습니다.",
        chart_document_sync,
        snap,
    )
    add(
        "12-roadmap.png",
        "남은 작업 권장 단계",
        "안전 핵심 연결, 신고와 개인정보, AI와 서버와 운영, 정식 검증과 출시의 네 단계와 각 단계의 세 가지 대표 작업",
        "남은 작업을 이해하기 위한 권장 묶음입니다. 실제 실행 순서는 후속 통제 검토 뒤 확정해야 합니다.",
        chart_roadmap,
    )
    return cover, specs


def set_run_font(
    run,
    *,
    name: str = FONT_KO,
    size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm / 2.54 * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_picture_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_page_number(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_section(section, *, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(1.3)
        section.right_margin = Cm(1.3)
        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.3)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(1.75)
        section.right_margin = Cm(1.75)
        section.top_margin = Cm(1.55)
        section.bottom_margin = Cm(1.45)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)


def configure_document(doc: DocumentType) -> None:
    configure_section(doc.sections[0], landscape=False)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_KO
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    normal.font.size = Pt(9.7)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Title": (27, NAVY, 10, 8),
        "Subtitle": (12, MUTED, 4, 4),
        "Heading 1": (17, NAVY, 14, 7),
        "Heading 2": (13.5, TEAL, 11, 5),
        "Heading 3": (11.2, NAVY, 8, 4),
        "Heading 4": (10.2, MUTED, 7, 3),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = FONT_KO
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["Caption"].font.name = FONT_KO
    styles["Caption"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    styles["Caption"].font.size = Pt(8.3)
    styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)
    styles["Caption"].paragraph_format.keep_with_next = False
    styles["Caption"].paragraph_format.space_after = Pt(6)

    props = doc.core_properties
    props.title = "WalkSafe 시각장애인 보행 보조 시스템 산출물 소개서"
    props.subject = "제품 맥락, 기능 현황, 산출물 지도, 남은 작업"
    props.author = "WalkSafe 프로젝트"
    props.last_modified_by = "WalkSafe 프로젝트"
    props.created = BUILD_TIME
    props.modified = BUILD_TIME
    props.keywords = "WalkSafe, 산출물, 기능 현황, 비전공자 설명, DOCX"
    props.comments = "기존 산출물을 수정하지 않고 정본 자료를 사람이 읽기 쉽게 설명한 비정본 소개서"

    for section in doc.sections:
        update_header_footer(section)


def update_header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.clear()
    run = paragraph.add_run("WalkSafe | 산출물 소개서 · 비정본 설명 자료")
    set_run_font(run, size=7.8, color=MUTED)
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.clear()
    run = paragraph.add_run(f"기준일 {BUILD_DATE}  |  ")
    set_run_font(run, size=7.8, color=MUTED)
    add_page_number(paragraph)


def add_section(doc: DocumentType, *, landscape: bool) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=landscape)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True


def add_footnoted_run(paragraph, text: str, note: str, footnotes: list[Footnote]) -> None:
    run = paragraph.add_run(text)
    set_run_font(run)
    marker = f"[[[WS_FOOTNOTE:{len(footnotes) + 1:03d}]]]"
    marker_run = paragraph.add_run(marker)
    set_run_font(marker_run, size=7)
    footnotes.append(Footnote(marker, note))


def add_parts_paragraph(
    doc: DocumentType,
    parts: Sequence[str | tuple[str, str]],
    footnotes: list[Footnote],
    *,
    style: str | None = None,
    bold_lead: str | None = None,
    align: WD_ALIGN_PARAGRAPH | None = None,
    color: str | None = None,
) -> Any:
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    if bold_lead:
        run = paragraph.add_run(bold_lead)
        set_run_font(run, bold=True, color=color)
    for part in parts:
        if isinstance(part, tuple):
            add_footnoted_run(paragraph, part[0], part[1], footnotes)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, color=color)
    return paragraph


def add_bullet(doc: DocumentType, text: str, *, level: int = 0, color: str | None = None) -> Any:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(2.5)
    run = paragraph.add_run(text)
    set_run_font(run, color=color)
    return paragraph


def add_numbered(doc: DocumentType, text: str) -> Any:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(2.5)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_callout(
    doc: DocumentType,
    title: str,
    body: str,
    *,
    fill: str = LIGHT_BLUE,
    accent: str = BLUE,
) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=150, end=150)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), accent)
    borders.append(left)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    set_run_font(run, size=10, bold=True, color=accent)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(body)
    set_run_font(run, size=9.1, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(
    doc: DocumentType,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    *,
    widths_cm: Sequence[float] | None = None,
    font_size: float = 8.1,
    header_fill: str = NAVY,
    alternating: bool = True,
    first_col_bold: bool = False,
) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths_cm is None
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for index, (cell, header) in enumerate(zip(header_row.cells, headers)):
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, top=100, bottom=100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, size=font_size, bold=True, color="FFFFFF")
        if widths_cm:
            set_cell_width(cell, widths_cm[index])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        if alternating and row_index % 2:
            for cell in row.cells:
                set_cell_shading(cell, "F7F9FA")
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            set_cell_margins(cell, top=80, bottom=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, size=font_size, bold=first_col_bold and index == 0, color=INK)
            if widths_cm:
                set_cell_width(cell, widths_cm[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(
    doc: DocumentType,
    spec: FigureSpec,
    number: int,
    *,
    width_inches: float | None = None,
) -> None:
    section = doc.sections[-1]
    printable = float(section.page_width - section.left_margin - section.right_margin) / 914_400
    width = min(width_inches or printable - 0.12, printable - 0.12)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(spec.path), width=Inches(width))
    set_picture_alt_text(shape, spec.title, spec.alt_text)
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f"그림 {number}. {spec.title}. {spec.caption}")
    set_run_font(run, size=8.2, color=MUTED)


def add_heading(doc: DocumentType, text: str, level: int, *, page_break: bool = False) -> Any:
    heading = doc.add_heading(text, level=level)
    if page_break:
        heading.paragraph_format.page_break_before = True
    return heading


def format_status(status: str) -> str:
    return STATUS_LABELS[status]


def translate_blocker(code: str) -> str:
    if code in BLOCKER_LABELS:
        return BLOCKER_LABELS[code]
    if code.endswith("-INDEPENDENT-QA-UNASSIGNED"):
        return "독립 품질 검토 담당 미지정"
    if code.endswith("-REAL-TRIGGER-OR-EVIDENCE-PENDING"):
        return "실제 사건 또는 실행 증거 대기"
    if code.endswith("-NAMED-CANDIDATE-OR-SOURCE-PENDING"):
        return "대상 출시 후보 또는 원본 지정 대기"
    raise ValueError(f"untranslated artifact blocker: {code}")


def artifact_state_label(row: dict[str, Any]) -> str:
    blockers = row["state"].get("blockers") or []
    if "NOT_ACTIVE_UNDER_CURRENT_BASELINE" in blockers:
        return "현행 범위 제외\n(등록상 계획)"
    return LIFECYCLE_LABELS[row["state"]["lifecycle_status"]]


def artifact_next_step(row: dict[str, Any]) -> str:
    blockers = row["state"].get("blockers") or []
    if "NOT_ACTIVE_UNDER_CURRENT_BASELINE" in blockers:
        return "현행 범위에 다시 넣는 승인을 받은 경우에만 실행"
    state = row["state"]["lifecycle_status"]
    if state == "APPROVED_BASELINED":
        return "현재 정책·코드와 맞는지 변경 때 재확인"
    if state == "ACTIVE":
        return "새 결정·사건·변경이 생길 때 계속 갱신"
    if state == "DRAFT":
        if blockers:
            return shorten(" / ".join(translate_blocker(str(item)) for item in blockers), 75)
        return "내용·근거를 보완하고 지정 검토·승인"
    form = row.get("artifact_form", "")
    if form in {"GENERATED_EVIDENCE", "EXTERNAL_RECORD"}:
        return "실제 시험·사건·서명 뒤 결과와 원본 증거 연결"
    return "실제 실행 대상을 정하고 결과·증거 작성"


def bundle_state_summary(rows: Sequence[dict[str, Any]]) -> str:
    counts = Counter(row["state"]["lifecycle_status"] for row in rows)
    parts = [
        f"{LIFECYCLE_LABELS[state]} {counts[state]}"
        for state in ("APPROVED_BASELINED", "ACTIVE", "DRAFT", "PLANNED")
        if counts[state]
    ]
    return " · ".join(parts)


def bundle_paths(rows: Sequence[dict[str, Any]]) -> str:
    paths = sorted(
        {
            row["location"].get("canonical_path")
            or row["location"].get("planned_canonical_path")
            or "경로 미정"
            for row in rows
        }
    )
    return "\n".join(paths)


def register_file_integrity_breakdown(snap: Snapshot) -> tuple[int, int, int, int]:
    matching, mismatching, self_reference = integrity_counts(snap)
    return matching + mismatching + self_reference, matching, mismatching, self_reference


def build_document(
    snap: Snapshot,
    cover_art: Path,
    figures: Sequence[FigureSpec],
    output: Path,
) -> list[Footnote]:
    doc = Document()
    configure_document(doc)
    footnotes: list[Footnote] = []
    figure_by_name = {spec.path.name: spec for spec in figures}
    figure_number = {spec.path.name: index for index, spec in enumerate(figures, start=1)}

    # Cover
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_before = Pt(10)
    run = label.add_run("W A L K S A F E")
    set_run_font(run, size=11, bold=True, color=TEAL)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("시각장애인 보행 보조 시스템\n산출물 소개서")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("제품 맥락 · 기능 현황 · 산출물 지도 · 남은 작업")
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = picture_paragraph.add_run().add_picture(str(cover_art), width=Inches(6.75))
    set_picture_alt_text(
        shape,
        "WalkSafe 표지 그림",
        "휴대폰, 음성 파동, 보행 경로, 점자블록을 단순 도형으로 나타낸 WalkSafe 소개서 표지",
    )
    cover_rows = [
        ("문서 버전", "1.0.0"),
        ("기준일", BUILD_DATE),
        ("문서 상태", "비정본 설명 자료 · 사람 검토 필요"),
        ("대상", "hanium-dreamup-walksafe RC2 작업트리"),
        ("범위", "제품·기능·산출물·시험·출시 현황 설명"),
        ("제외", "기존 산출물 형식 통일·수정·승인 상태 변경"),
    ]
    add_table(doc, ["항목", "내용"], cover_rows, widths_cm=[4.2, 11.7], font_size=8.4, alternating=False)
    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = notice.add_run(
        "이 문서는 기존 257개 산출물을 바꾸지 않고 현재 정본 자료를 비전공자도 읽기 쉽게 설명한 자료입니다. "
        "안전성·출시 적격성을 승인하거나 보장하지 않습니다."
    )
    set_run_font(run, size=8.5, italic=True, color=BURGUNDY)
    doc.add_page_break()

    # Document control and reading guide
    add_heading(doc, "문서 통제 정보", 1)
    control_rows = [
        ("문서 목적", "WalkSafe가 무엇을 만들고 있으며, 기능과 산출물이 어디까지 준비됐는지 한 문서에서 찾게 한다."),
        ("주 독자", "팀원, 지도교수·멘토, 심사자, 비전공 협업자, 후속 담당자"),
        ("수치 기준", "최신 기능 차이 분석, 산출물 등록부, 미승인 정식 시험 계획, 산출물 폐쇄 원장, 현재 실행 포인터"),
        ("상태 경계", "코드 존재·내부 시험·정식 시험·승인·출시를 서로 다른 단계로 표시"),
        ("안전 경계", "흰지팡이·안내견·보호자를 대신하지 않으며 보행 안전을 보장하지 않음"),
        ("수정 경계", "기존 산출물, 산출물 등록부·변경 기록, 승인 기록, 생성 파일을 수정하지 않음"),
        ("승인 상태", "미승인 설명 자료. 커밋되지 않은 변경이 많은 진행 중 스냅샷이며 공식 제출·배포 전 사람 검토 필요"),
        ("갱신 시점", "기능 차이·개선 목록, 산출물 등록부, 정식 시험, 출시 관문 또는 제품 경계가 바뀔 때"),
    ]
    add_table(doc, ["통제 항목", "내용"], control_rows, widths_cm=[4.1, 12.0], font_size=8.5)
    add_callout(
        doc,
        "이 문서에서 가장 중요한 읽기 원칙",
        "‘일부 구현’, ‘산출물 기준 확정’, ‘내부 자동시험 통과’는 서로 다른 상태입니다. "
        "어느 한 수치도 전체 프로젝트 완료율로 단독 사용하지 않습니다.",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )
    add_parts_paragraph(
        doc,
        [
            "이 문서의 공식 수치는 ",
            ("Gap·Backlog r021", dict(GLOSSARY)["Gap·Backlog r021"]),
            ", ",
            ("DOC-01", dict(GLOSSARY)["DOC-01"]),
            ", ",
            ("폐쇄 원장 R011", dict(GLOSSARY)["폐쇄 원장 R011"]),
            ", ",
            ("체크포인트 seq39", dict(GLOSSARY)["체크포인트(checkpoint)"]),
            "를 기준으로 고정했습니다.",
        ],
        footnotes,
    )

    add_heading(doc, "수치의 기준 시점", 2)
    as_of_rows = [
        ("기능·정책 판정", "r021 · 2026-07-26"),
        ("정식 시험 계획(미승인 초안)", "2026-07-27 · 279개 · 전부 미실행"),
        ("산출물 DOC-01", "2026-07-28 · 257개 유형"),
        ("산출물 폐쇄 투영", "R011 · 열린 유형 131개"),
        (
            "프로젝트 실행 포인터",
            snap.checkpoint["goal_execution"]["validation_cutoff_at"] + " · checkpoint seq39",
        ),
        ("소개서 생성", snap.generated_at),
    ]
    add_table(doc, ["자료", "기준"], as_of_rows, widths_cm=[5.0, 11.1], font_size=8.4)

    # TOC marker is converted to a live TOC by LibreOffice UNO.
    toc_marker = doc.add_paragraph()
    toc_marker.paragraph_format.page_break_before = True
    toc_marker.add_run("[[[WS_TOC]]]")
    add_heading(doc, "그림 목록", 1)
    figure_rows = [(f"그림 {index}", spec.title) for index, spec in enumerate(figures, start=1)]
    add_table(doc, ["번호", "시각자료"], figure_rows, widths_cm=[2.6, 13.5], font_size=8.4)
    doc.add_page_break()

    # 1. Executive summary
    add_heading(doc, "1. 한눈에 보는 현재 상태", 1)
    dashboard_rows = [
        ("필요한 기능·공통 규칙", "63개", "공식 완료 0 · 일부 구현 32 · 기준과 다름 16 · 핵심 없음 11 · 증거 부족 4"),
        ("출시 전 필수 확인", "5개", "통과 0 · 모두 미실행"),
        ("산출물 유형", "257개", "승인 기준본 102 · 계속 갱신 27 · 초안 53 · 계획 75"),
        ("정식 시험", "279개", "미승인 초안 · 실행 0 · 합격 0 · 모두 미실행"),
        ("산출물 연결", "698개", "같은 범주 394 · 다른 범주 304"),
        ("출시 판정", "출시 불가", "서명·정식 시험·실기기·현장·관문 근거 없음"),
    ]
    add_table(doc, ["평가 축", "규모", "현재 상태"], dashboard_rows, widths_cm=[4.0, 2.2, 9.9], font_size=8.6)
    add_figure(
        doc,
        figure_by_name["03-feature-status.png"],
        figure_number["03-feature-status.png"],
        width_inches=6.65,
    )
    add_heading(doc, "단일 ‘전체 완료율’을 쓰지 않는 이유", 2)
    add_bullet(doc, "63개 기능은 크기와 남은 작업량이 서로 달라 단순 합산할 수 없습니다.")
    add_bullet(doc, "산출물이 승인됐어도 연결된 기능·시험·배포가 끝났다는 뜻은 아닙니다.")
    add_bullet(doc, "내부 자동시험은 개발 중 오류를 찾는 데 유용하지만 실제 사용자·현장 시험을 대신하지 않습니다.")
    add_bullet(doc, "출시 관문 5개는 코드 작성이 아니라 실제 측정·독립 검토·복구훈련이 필요한 별도 조건입니다.")
    axis_rows = [
        ("기능·정책", "공식 완료 0/63", "일부 구현 수를 완료율로 바꾸지 않음"),
        ("산출물 생명주기", "기준본 102/257", "문서 승인 상태이며 기능 완료율이 아님"),
        ("산출물 폐쇄 투영", "126/257 (49.0%)", "내용·적용 범위 관점. 전체 프로젝트 완료율 아님"),
        ("정식 시험", "합격 0/279", "미승인 초안이며 모두 미실행"),
        ("출시 관문", "통과 0/5", "모두 미실행·면제 없음"),
        ("생산 배포", "0건", "서명·단계 배포·복구 증거 없음"),
    ]
    add_table(doc, ["평가 축", "현재 수치", "정확한 해석"], axis_rows, widths_cm=[4.0, 3.3, 8.8], font_size=8.4)

    # 2. Product context
    add_heading(doc, "2. WalkSafe가 하려는 일", 1, page_break=True)
    add_parts_paragraph(
        doc,
        [
            "WalkSafe는 시각장애인이 일반 도심 보도를 걸을 때 가까운 위험을 알아차리고, 목적지까지 큰 이동 방향을 듣고, 손상된 점자블록 신고를 준비하도록 돕는 ",
            ("Android 앱", dict(GLOSSARY)["Android"]),
            " 중심 시스템입니다.",
        ],
        footnotes,
    )
    add_callout(
        doc,
        "제품의 약속",
        "① 가까운 위험을 알려준다. ② 목적지까지 큰 방향을 안내한다. ③ 손상 점자블록 신고를 돕는다. "
        "단, 안전을 보장하거나 흰지팡이·안내견·보호자를 대신하지 않는다.",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )
    add_figure(
        doc,
        figure_by_name["01-user-journey.png"],
        figure_number["01-user-journey.png"],
        width_inches=6.7,
    )
    add_heading(doc, "현재 제품에 포함되는 것과 포함되지 않는 것", 2)
    boundary_rows = [
        ("포함", "사용자 Android 앱", "카메라·길안내·음성·진동·신고·동의"),
        ("포함", "별도 관리자 Android 앱", "관리자 로그인·세션·복구와 향후 신고 운영"),
        ("포함", "요청 검사 중간 서버", "사용자 세션·동의·보행상태·전송 제한"),
        ("포함", "정보 처리 본 서버", "길찾기 중계·신고·관리자 보안·DB"),
        ("제외", "과거 Web/PWA", "현행 제품이 아닌 역사적 참고자료"),
        ("제외", "Python 음성 서버", "Android 제품 음성 경로가 아닌 실험용"),
        ("별도", "서버의 객체탐지 API", "사용자 Android 온디바이스 탐지의 필수 경로가 아님"),
    ]
    add_table(doc, ["구분", "대상", "쉬운 설명"], boundary_rows, widths_cm=[2.2, 4.5, 9.4], font_size=8.4)

    # 3. Current code capabilities
    add_heading(doc, "3. 현재 코드에서 확인된 핵심 동작", 1, page_break=True)
    add_callout(
        doc,
        "표를 읽는 방법",
        "아래의 ‘현재 코드에서 확인’은 소스와 내부 자동검사에서 동작 경로가 보인다는 뜻입니다. "
        "정식 시험·실기기·실사용자 검증을 통과했다는 뜻은 아닙니다.",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )
    capability_rows = [
        (
            "주변 물체 찾기",
            "휴대폰 카메라 영상에서 사람, 차량, 자전거, 턱, 점자블록 등 13종 후보를 찾습니다.",
            "AI 모델은 후보 상태이며 정식 출시 승인을 받지 않았습니다.",
        ),
        (
            "가까워지는 위험 판단",
            "지원되는 휴대폰에서는 물체까지의 대략적인 거리와 가까워지는 정도를 살펴 급한 위험을 먼저 고릅니다.",
            "생산 승인 휴대폰 목록과 실제 거리 시험이 없습니다.",
        ),
        (
            "음성·진동 알림",
            "위험이나 길안내를 음성으로 들려주고, 중요한 상황에는 진동도 사용합니다.",
            "‘길라잡이’ 호출어와 완전한 무버튼 사용은 미완성입니다.",
        ),
        (
            "목적지와 길안내",
            "목적지를 검색하고 큰 보행 경로를 안내하며, 길에서 벗어나면 다시 찾을 후보를 만듭니다.",
            "도착 확인과 새 길 선택을 사용자가 결정하는 흐름이 더 필요합니다.",
        ),
        (
            "손상 점자블록 신고 준비",
            "손상된 점자블록 후보를 만들고 사용자가 직접 신고 후보를 남길 수 있습니다.",
            "암호화된 영구 대기함·재부팅 후 재시도·기관 수신확인은 없습니다.",
        ),
        (
            "권한·동의·삭제 요청",
            "카메라 등 권한, 자료 수집 동의, 동의 철회와 계정 삭제 요청 흐름이 있습니다.",
            "모든 저장소를 실제로 삭제하고 결과 영수증을 남기는 중앙 처리가 미완성입니다.",
        ),
        (
            "보행 상태와 안전정지",
            "준비·보행·일시중지·안전정지·종료 상태를 나누고 비정상 종료 뒤 자동 재개를 막습니다.",
            "실제 휴대폰·화면잠금·장시간 보행 검증이 남았습니다.",
        ),
        (
            "관리자 계정 보호",
            "별도 관리자 앱에서 비밀번호와 추가 확인, 세션 조회·차단, 복구 흐름을 제공합니다.",
            "신고 검수·기관 전달·관제 화면과 실제 비공개 배포가 없습니다.",
        ),
    ]
    add_table(doc, ["사용자가 이해할 기능", "현재 코드에서 확인", "아직 남은 핵심"], capability_rows, widths_cm=[3.6, 6.4, 6.1], font_size=8.15)

    add_heading(doc, "꼭 유지해야 하는 기술 이름", 2)
    add_parts_paragraph(
        doc,
        [
            "주변 물체 후보는 휴대폰용 AI 형식인 ",
            ("TFLite", dict(GLOSSARY)["TFLite"]),
            "로 찾고, 지원되는 휴대폰에서는 ",
            ("ARCore Depth", dict(GLOSSARY)["ARCore Depth"]),
            "로 대략적인 거리를 살핍니다. 가까워지는 급함을 정할 때는 ",
            ("TTC", dict(GLOSSARY)["TTC"]),
            "도 참고합니다.",
        ],
        footnotes,
    )
    add_parts_paragraph(
        doc,
        [
            "큰 이동 경로는 ",
            ("TMAP", dict(GLOSSARY)["TMAP"]),
            "에서 받고, 화면을 보지 않아도 조작할 수 있도록 ",
            ("TalkBack", dict(GLOSSARY)["TalkBack"]),
            "을 사용합니다.",
        ],
        footnotes,
    )
    add_parts_paragraph(
        doc,
        [
            "거리 기술을 지원하지 않는 휴대폰용 후보 카메라 경로에는 ",
            ("CameraX", dict(GLOSSARY)["CameraX"]),
            "가 있으나 기본값으로 꺼져 있고 실기기 검증이 없습니다. 사용자의 음성을 명령으로 바꾸고 안내 문장을 읽을 때는 ",
            ("STT·TTS", dict(GLOSSARY)["STT·TTS"]),
            "를 사용합니다.",
        ],
        footnotes,
    )

    add_heading(doc, "공식 상태를 쉬운 말로 읽기", 2)
    state_rows = [
        (STATUS_LABELS["IMPLEMENTED"], "기능·정식 시험·증거가 모두 준비됨", "현재 0"),
        (STATUS_LABELS["PARTIAL"], "코드 일부는 있으나 연결·시험·승인이 더 필요", "현재 32"),
        (STATUS_LABELS["CONFLICTING"], "코드는 있지만 합의한 작동 방식과 다름", "현재 16"),
        (STATUS_LABELS["MISSING"], "필요한 핵심 구현을 확인하지 못함", "현재 11"),
        (STATUS_LABELS["EVIDENCE_MISSING"], "기능 후보는 있으나 공식 확인 자료가 없음", "현재 4"),
        (STATUS_LABELS["BLOCKED"], "실측·외부 검토·복구훈련을 기다림", "출시 관문 5"),
    ]
    add_table(doc, ["문서 표시", "쉬운 뜻", "현재"], state_rows, widths_cm=[3.9, 9.3, 2.9], font_size=8.5)

    add_heading(doc, "정본 판정과 현재 코드 관찰이 다른 예", 2)
    add_parts_paragraph(
        doc,
        [
            "앱과 서버가 어떤 주소와 자료 형식으로 통신하는지는 ",
            ("OpenAPI", dict(GLOSSARY)["OpenAPI"]),
            "라는 규격으로도 확인합니다.",
        ],
        footnotes,
    )
    drift_rows = [
        ("관리자 앱", "r021은 핵심 구현 없음", "현재 작업트리에는 별도 앱 골격과 인증 코드가 있음", "운영 화면·배포·시험을 포함해 다시 판정 필요"),
        ("중간 서버 API", "일부 개발 안내문은 4개", "현재 코드·OpenAPI는 /api/field-walk 포함 5개", "개발 안내문보다 코드·OpenAPI와 최신 경계 문서를 우선"),
        ("AI 모델 파일", "일부 개발 안내문은 통합 TFLite 부재", "현재 768 입력 TFLite 파일이 존재", "모델 존재와 출시 승인 여부를 분리"),
        ("다음 기능 포인터", "r021은 FP-048을 가리킴", "최신 인계서가 이 포인터를 오래된 값으로 판정", "새 통제 전환 뒤 우선순위를 확정"),
    ]
    add_table(doc, ["대상", "문서의 오래된 표현", "현재 관찰", "소개서의 처리"], drift_rows, widths_cm=[2.6, 4.3, 4.8, 4.4], font_size=7.9)

    # 4. Architecture
    add_heading(doc, "4. 시스템 구성과 정보 흐름", 1, page_break=True)
    add_parts_paragraph(
        doc,
        [
            "사용자 앱은 본 서버에 바로 모든 권한을 주지 않습니다. 요청은 먼저 보호된 중간 출입구인 ",
            ("Gateway와 API", dict(GLOSSARY)["Gateway·API"]),
            "를 거쳐 사용자·동의·보행 상태를 확인한 뒤 본 서버로 전달됩니다.",
        ],
        footnotes,
    )
    add_figure(
        doc,
        figure_by_name["02-architecture.png"],
        figure_number["02-architecture.png"],
        width_inches=6.65,
    )
    add_heading(doc, "사용자 휴대폰 안에서 일어나는 일", 2)
    runtime_rows = [
        ("1. 보기", "카메라 영상과 휴대폰 움직임·위치정보를 읽음"),
        ("2. 후보 찾기", "사람·차량·턱·점자블록 등 물체 후보를 찾음"),
        ("3. 거리 확인", "지원 기기에서 물체 영역의 대략적 거리와 신뢰도를 확인"),
        ("4. 위험 고르기", "여러 후보 중 더 급하고 믿을 만한 하나를 선택"),
        ("5. 알리기", "짧은 음성 문장과 상황에 맞는 진동을 제공"),
        ("6. 신고 준비", "손상 점자블록 후보와 필요한 최소 정보를 보관"),
    ]
    add_table(doc, ["단계", "쉬운 설명"], runtime_rows, widths_cm=[3.0, 13.1], font_size=8.6)
    add_parts_paragraph(
        doc,
        [
            "계정·신고 정보는 ",
            ("PostgreSQL·PostGIS", dict(GLOSSARY)["PostgreSQL·PostGIS"]),
            "에 저장할 수 있는 구조가 있으나, 큰 영상·음성 파일은 현재 암호화된 대용량 저장소가 아니라 로컬 파일 경로를 사용합니다.",
        ],
        footnotes,
    )

    # 5. Epics and implementation gaps
    add_heading(doc, "5. 필요한 기능 묶음과 남은 기능", 1, page_break=True)
    add_parts_paragraph(
        doc,
        [
            "프로젝트는 서로 관련된 기능을 크게 묶은 ",
            ("EPIC", dict(GLOSSARY)["EPIC"]),
            " 12개로 나눠 관리합니다. 각 묶음 아래에는 정한 요구와 현재 상태의 차이인 ",
            ("Gap", dict(GLOSSARY)["Gap"]),
            "이 연결됩니다.",
        ],
        footnotes,
    )
    add_figure(doc, figure_by_name["04-epics.png"], figure_number["04-epics.png"], width_inches=6.7)
    epic_rows = []
    status_korean = {
        "IMPLEMENTATION_READY": "내부 구현 준비",
        "IN_PROGRESS": "진행 중",
        "PLANNED": "계획",
    }
    for epic_id, title, description, state in EPIC_INFO:
        counts = EPIC_GAP_COUNTS[epic_id]
        distribution = " · ".join(
            f"{STATUS_LABELS[status]} {value}"
            for status, value in counts.items()
        )
        epic_rows.append((epic_id, title, description, status_korean[state], distribution))
    add_table(
        doc,
        ["묶음", "이름", "무엇을 하는가", "작업 위치", "연결 항목 상태"],
        epic_rows,
        widths_cm=[2.0, 3.2, 5.6, 2.3, 3.0],
        font_size=7.8,
    )
    add_heading(doc, "우선 보완할 기능", 2)
    priorities = [
        ("안전·실기기", "지원 휴대폰 목록, 실제 거리·저조도·흔들림·발열·배터리 확인"),
        ("음성·접근성", "‘길라잡이’ 호출어, 완전한 무버튼 보행, 안전정지 화면, TalkBack 전체 여정"),
        ("길안내", "사용자 확인 도착, 경로 이탈 확인, 새 길 선택, 지도 장애 처리"),
        ("자동신고", "암호화 대기함, 재부팅 뒤 재시도, 중복 방지, 서버 수신확인"),
        ("관리자·기관", "신고 검수·기각·기관 전달·수신증·사용자 처리상태"),
        ("개인정보·저장", "중앙 동의·삭제 처리, 저장소별 삭제 영수증, 암호화 대용량 저장소와 열쇠 관리"),
        ("AI", "재현 가능한 학습자료, 독립 시험자료, 약한 클래스 개선, 모델 승인·교체·복구"),
        ("운영·출시", "용량·비용·관제·복원, 앱 서명, 단계 배포, 정식 시험과 5개 출시 관문"),
    ]
    add_table(doc, ["우선 영역", "해야 할 일"], priorities, widths_cm=[3.5, 12.6], font_size=8.4)

    # 6. Artifact map
    add_heading(doc, "6. 산출물 체계와 경로 지도", 1, page_break=True)
    add_parts_paragraph(
        doc,
        [
            "산출물의 공식 위치와 상태는 ",
            ("정본", dict(GLOSSARY)["정본(canonical)"]),
            " 등록부인 DOC-01에서 관리합니다. 승인한 파일 묶음을 버전과 지문으로 고정한 시점을 ",
            ("기준선", dict(GLOSSARY)["기준선(baseline)"]),
            "이라고 부릅니다.",
        ],
        footnotes,
    )
    add_figure(doc, figure_by_name["05-artifact-map.png"], figure_number["05-artifact-map.png"], width_inches=6.7)
    add_callout(
        doc,
        "257개는 파일 수가 아닙니다",
        "257개는 챙겨야 할 결과물 ‘유형’입니다. 여러 유형이 한 문서의 서로 다른 장·표에 들어가므로 "
        "40개 묶음과 41개 예정 경로로 모입니다. 현재 정본 경로가 연결된 파일은 37개입니다.",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )
    add_heading(doc, "13개 산출물 범주", 2)
    category_rows = []
    for category in CATEGORY_ORDER:
        label, purpose, path_value = CATEGORY_INFO[category]
        counts = snap.category_lifecycle[category]
        total = sum(counts.values())
        status = " / ".join(
            str(counts[state]) for state in ("APPROVED_BASELINED", "ACTIVE", "DRAFT", "PLANNED")
        )
        category_rows.append((category, label, purpose, path_value, str(total), status))
    add_table(
        doc,
        ["코드", "쉬운 이름", "무엇을 보관하는가", "경로", "유형", "기준/갱신/초안/계획"],
        category_rows,
        widths_cm=[1.3, 2.5, 4.2, 4.8, 1.2, 2.3],
        font_size=7.5,
    )
    add_heading(doc, "권장 읽기 순서", 2)
    reading_rows = [
        ("DOC", "문서를 믿고 관리하는 방법"),
        ("MGT + DSC", "왜 만들고, 누구를 위해, 어디까지 만드는가"),
        ("REQ", "제품이 무엇을 해야 하는가"),
        ("DES", "그 기능을 어떤 구조로 만들 것인가"),
        ("DEV", "실제로 무엇을 만들었는가"),
        ("TST + SEC + AIML + WS", "제대로·안전하게 작동하는가"),
        ("REL + OPS", "배포하고 계속 운영할 수 있는가"),
        ("CLS", "무엇을 넘기고 어떻게 끝낼 것인가"),
    ]
    add_table(doc, ["순서", "읽어서 답할 질문"], reading_rows, widths_cm=[5.0, 11.1], font_size=8.5)
    add_heading(doc, "주요 통제 경로", 2)
    path_rows = [
        ("현재 실행 포인터", "docs/control/walksafe-project-continuation-checkpoint.json"),
        ("기능 정책 기준선", "docs/control/baselines/"),
        ("기능 Gap·개선 목록", "docs/control/audits/"),
        ("산출물 상태 정본", "docs/deliverables/00-control/artifact-register.json"),
        ("산출물 변경 이력", "docs/deliverables/00-control/artifact-change-log.json"),
        ("정식 시험", "docs/deliverables/06-testing/"),
        ("AI·데이터", "docs/deliverables/08-ai-ml-data/"),
        ("Goal 실행 결과", "docs/control/execution/goal-results/<goal-id>/"),
        ("정식 시험 실행 기록", "docs/deliverables/06-testing/evidence/executions/<run-id>.json"),
        ("민감 원본", "Git 밖의 암호화 저장소 · 저장소 ID와 SHA-256만 Git에 기록"),
    ]
    add_table(doc, ["용도", "경로"], path_rows, widths_cm=[5.0, 11.1], font_size=8.2)

    add_heading(doc, "40개 문서 묶음과 경로", 2, page_break=True)
    bundle_rows = []
    for bundle_id in sorted(snap.bundles, key=lambda value: (snap.bundles[value][0]["category"], value)):
        rows = snap.bundles[bundle_id]
        bundle_rows.append(
            (
                bundle_id,
                BUNDLE_PURPOSES.get(bundle_id, simplify(rows[0]["authoring_contract"]["purpose"])),
                bundle_paths(rows),
                str(len(rows)),
                bundle_state_summary(rows),
            )
        )
    add_table(
        doc,
        ["문서 묶음", "무엇을 하는가", "정본·예정 경로", "유형 수", "현재 상태"],
        bundle_rows,
        widths_cm=[3.1, 4.8, 4.8, 1.3, 2.1],
        font_size=7.2,
    )
    add_callout(
        doc,
        "이번 작업에서 하지 않은 것",
        "기존 산출물의 표지·헤더·본문 형식을 통일하거나 DOC-01의 해시·상태를 고치지 않았습니다. "
        "설계서·시험 결과·외부 서명 원본은 목적이 다르므로 본문 모양까지 같을 필요도 없습니다.",
        fill=LIGHT_GRAY,
        accent=GRAY,
    )

    # 7. Artifact progress
    add_heading(doc, "7. 산출물 현황과 남은 작업", 1, page_break=True)
    add_figure(
        doc,
        figure_by_name["06-artifact-lifecycle.png"],
        figure_number["06-artifact-lifecycle.png"],
        width_inches=6.7,
    )
    lifecycle_rows = [
        ("승인 기준본", "102/257", pct(102, 257), "승인된 문서 내용·파일 지문을 고정. 연결 기능 완료를 뜻하지 않음"),
        ("계속 갱신", "27/257", pct(27, 257), "결정·위험·진행·운영 사건이 생길 때 계속 기록"),
        ("초안", "53/257", pct(53, 257), "내용·근거·검토·승인이 더 필요"),
        ("계획·미실행", "75/257", pct(75, 257), "실제 시험·배포·서명·사건 결과가 아직 없음"),
    ]
    add_table(doc, ["상태", "분자/분모", "비율", "정확한 의미"], lifecycle_rows, widths_cm=[3.0, 2.4, 1.6, 9.1], font_size=8.4)
    add_figure(doc, figure_by_name["07-open-lanes.png"], figure_number["07-open-lanes.png"], width_inches=6.55)
    add_heading(doc, "산출물 ‘닫힘 49.0%’의 제한", 2)
    add_callout(
        doc,
        "49.0%는 프로젝트 완료율이 아닙니다",
        "폐쇄 원장 R011은 기준상 문제없음 124개와 현재 범위에서 제외 승인된 2개를 합쳐 126/257을 ‘닫힘에 해당’으로 봅니다. "
        "하지만 전체 산출물 완료 주장은 0/257이며, 기능·시험·출시 상태와 합산할 수 없습니다.",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )
    add_heading(doc, "현재 등록부·파일 동기화 문제", 2)
    add_figure(doc, figure_by_name["11-document-sync.png"], figure_number["11-document-sync.png"], width_inches=6.4)
    total_rows, matching, mismatching, self_ref = register_file_integrity_breakdown(snap)
    sync_rows = [
        ("DOC-01 현재 파일 연결 행", str(total_rows), "같은 파일이 여러 산출물 유형을 담으므로 파일 수가 아님"),
        ("실제 파일 지문과 일치", str(matching), "현재 파일과 등록 지문이 같음"),
        ("실제 파일 지문과 불일치", str(mismatching), "문서 변경과 DOC-01·생성 결과가 동기화되지 않음"),
        ("자기참조 예외", str(self_ref), "DOC-01이 자기 자신의 지문을 직접 갖지 않는 예외"),
        ("분야별 생성기 검사", "1/7 일치", "6개 묶음은 등록값과 현재 생성 결과가 어긋남"),
    ]
    add_table(doc, ["검사", "현재", "해석"], sync_rows, widths_cm=[4.6, 3.0, 8.5], font_size=8.3)

    # 8. Dependency and impact
    add_heading(doc, "8. 산출물 연결관계와 변경 영향", 1, page_break=True)
    add_parts_paragraph(
        doc,
        [
            "요구가 설계·코드·시험까지 이어지는 연결을 ",
            ("추적성", "요구가 설계·코드·시험·증거까지 빠짐없이 이어지는 연결을 뜻합니다."),
            "이라고 합니다. DOC-01에는 방향이 있는 유형 연결 698개가 있습니다.",
        ],
        footnotes,
    )
    add_figure(
        doc,
        figure_by_name["08-dependency-heatmap.png"],
        figure_number["08-dependency-heatmap.png"],
        width_inches=6.55,
    )
    top_flows = [
        ("REQ → DES", 29, "무엇을 만들지 바꾸면 어떻게 만들지 설계를 확인"),
        ("DES → DEV", 15, "설계를 바꾸면 코드·빌드·구현 증거를 확인"),
        ("DSC → REQ", 14, "사용자·문제 이해가 바뀌면 요구사항을 확인"),
        ("REQ → TST", 14, "요구를 바꾸면 합격조건과 시험을 확인"),
        ("TST → WS", 13, "일반 시험 결과가 WalkSafe 안전·인수 검증에 영향"),
        ("DES → SEC", 12, "구조·데이터 흐름이 바뀌면 보안·개인정보를 확인"),
        ("REL → CLS", 11, "출시 결과가 최종 인수·인계·종료에 영향"),
        ("DEV → REL", 10, "코드·빌드가 바뀌면 출시 후보를 다시 고정"),
        ("WS → REQ", 10, "현장·안전 검증 결과가 요구 보완으로 돌아감"),
    ]
    add_table(doc, ["주요 흐름", "연결 수", "쉬운 의미"], [(a, str(b), c) for a, b, c in top_flows], widths_cm=[3.0, 2.0, 11.1], font_size=8.3)
    add_callout(
        doc,
        "변경할 때의 기본 질문",
        "앞 문서가 바뀌었으면 연결된 뒤 문서의 내용·시험·증거도 같은 기준으로 다시 확인했는가?",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    # 9. Tests and release
    add_heading(doc, "9. 시험·실기기·출시 준비", 1, page_break=True)
    add_parts_paragraph(
        doc,
        [
            "승인된 환경과 정확한 앱·서버·AI 묶음으로 실행하고 원자료와 판정을 남기는 시험을 ",
            ("정식 시험", dict(GLOSSARY)["정식 시험(Formal Test)"]),
            "이라고 합니다. 현재 미승인 초안에는 279개가 계획돼 있고 실행 결과는 모두 미실행입니다.",
        ],
        footnotes,
    )
    add_figure(doc, figure_by_name["09-formal-tests.png"], figure_number["09-formal-tests.png"], width_inches=6.5)
    add_heading(doc, "내부 자동시험과 정식 시험의 차이", 2)
    comparison_rows = [
        ("목적", "개발 중 오류를 빨리 찾음", "출시 요구·안전·접근성·복구를 공식 판정"),
        ("대상", "일부 코드·모듈·가짜 의존성 포함", "정확한 앱·서버·AI·설정·실제 환경"),
        ("환경", "개발 PC·에뮬레이터·내부 DB", "승인된 실제 휴대폰·서버·현장·참여자"),
        ("증거", "테스트 로그·내부 영수증", "실행 목록·원자료·파일 지문·실제 결과·결함·승인"),
        ("현재", "Android·관리자·Gateway 등에 통과 기록 존재", "미승인 초안 · 0/279 실행"),
    ]
    add_table(doc, ["구분", "내부 자동시험", "정식 시험"], comparison_rows, widths_cm=[2.7, 6.5, 6.9], font_size=8.2)
    add_heading(doc, "출시 전 반드시 확인할 5개 관문", 2)
    add_parts_paragraph(
        doc,
        [
            "코드 작성만으로 닫을 수 없고 출시 전에 반드시 통과해야 하는 조건을 ",
            ("출시 관문", dict(GLOSSARY)["출시 관문(Release Gate)"]),
            "이라고 합니다.",
        ],
        footnotes,
    )
    gate_rows = [
        ("휴대폰 대기자료 실제 용량", "지원할 휴대폰마다 저장공간·대기자료 크기를 재서 한도를 정함", "미실행"),
        ("서버 용량상태 전달 규칙", "상태 전달 지연·조회 간격·정보 사용기한을 측정", "미실행"),
        ("무가림 원본 독립 검토", "민감한 원본 수집을 개인정보·법률 관점에서 독립 검토", "미실행"),
        ("실제 클라우드 비용", "실제 저장량·요청·복원 비용을 측정", "미실행"),
        ("관리자 폰 분실 복구", "분실을 가정해 기존 접속 차단과 안전 복귀를 실제 훈련", "미실행"),
    ]
    add_table(doc, ["관문", "무엇을 확인하는가", "현재"], gate_rows, widths_cm=[4.2, 9.6, 2.3], font_size=8.3)
    add_parts_paragraph(
        doc,
        [
            "특히 ",
            ("무가림 원본", dict(GLOSSARY)["무가림 원본"]),
            "은 주변 사람의 얼굴·번호판·목소리까지 포함할 수 있어 코드 구현과 별도로 독립 검토가 필요합니다.",
        ],
        footnotes,
    )
    add_figure(doc, figure_by_name["10-evidence-ladder.png"], figure_number["10-evidence-ladder.png"], width_inches=6.65)

    # 10. Roadmap
    add_heading(doc, "10. 앞으로 해야 할 일", 1, page_break=True)
    add_figure(doc, figure_by_name["12-roadmap.png"], figure_number["12-roadmap.png"], width_inches=6.65)
    roadmap_rows = [
        ("1. 안전 핵심 연결", "지원 기기·실제 거리·음성 호출어·안전정지·TalkBack·도착 확인", "실폰과 대상 사용자로 핵심 보행 흐름 확인"),
        ("2. 신고·개인정보", "암호화 대기함·재시도·기관 전달·동의·삭제·보관 영수증", "신고가 중복·유실되지 않고 사용자가 처리결과 확인"),
        ("3. AI·서버·운영", "독립 시험자료·승인 모델·대용량 저장·열쇠 관리·용량·비용·복원", "재현 가능한 모델과 운영 기반 확보"),
        ("4. 정식 검증·출시", "279개 시험·실기기·현장·접근성·보안·5개 관문·서명·단계 배포", "출시 또는 보류를 근거로 결정"),
    ]
    add_table(doc, ["단계", "대표 작업", "완료 판단"], roadmap_rows, widths_cm=[3.2, 7.5, 5.4], font_size=8.2)
    add_callout(
        doc,
        "우선순위에 관한 통제 주의",
        "r021 Backlog의 다음 기능 포인터는 최신 인계서에서 오래된 값으로 판정됐습니다. "
        "이 소개서의 단계는 이해를 돕는 권장 묶음이며 실제 착수 순서는 유효한 후속 통제 전환 뒤 확정해야 합니다.",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )

    # 11. How to use and limits
    add_heading(doc, "11. 이 소개서를 사용하는 방법", 1, page_break=True)
    add_numbered(doc, "먼저 1장의 다축 현황을 읽고, 단일 완료율로 합치지 않습니다.")
    add_numbered(doc, "사용자 기능을 확인할 때는 3장과 부록 A의 ‘해야 하는 일·현재·남은 일’을 함께 봅니다.")
    add_numbered(doc, "산출물 위치를 찾을 때는 6장의 범주·묶음 표와 부록 B·C를 사용합니다.")
    add_numbered(doc, "문서를 바꿀 때는 8장의 앞뒤 연결을 확인해 관련 요구·설계·시험·증거를 함께 검토합니다.")
    add_numbered(doc, "출시 판단에는 9장의 정식 시험과 5개 관문 결과만 사용하며 내부 자동시험으로 대체하지 않습니다.")
    add_heading(doc, "소개서 갱신 조건", 2)
    for item in [
        "기능 차이 분석·개선 목록의 개정판이 바뀔 때",
        "DOC-01의 유형·상태·경로·묶음이 바뀔 때",
        "정식 시험이 실행되거나 결과가 바뀔 때",
        "실제 기기·현장·접근성·보안 검토 증거가 생길 때",
        "출시 관문이 실행·통과·반려될 때",
        "제품 경계나 Web/PWA·음성 서버의 역할이 바뀔 때",
        "중요 개발 안내문과 현재 코드의 불일치가 해소되거나 새로 발견될 때",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "주장할 수 없는 것", 2)
    forbidden_claims = [
        ("“프로젝트 49% 완료”", "49.0%는 산출물 폐쇄 투영일 뿐입니다."),
        ("“기능 50.8% 완료”", "32/63은 일부 구현 상태의 수이며 작업량 기준 완료율이 아닙니다."),
        ("“AI 안전 검증 완료”", "현재 모델은 후보이며 독립 시험자료·실기기 평가·승인이 없습니다."),
        ("“실기기 시험 완료”", "과거 제한적 모델 동작 확인 기록은 전체 보행·접근성·현장 인수가 아닙니다."),
        ("“출시 준비 완료”", "정식 시험 0/279, 출시 관문 0/5, 서명·생산 배포 0입니다."),
    ]
    add_table(doc, ["금지할 표현", "이유"], forbidden_claims, widths_cm=[5.3, 10.8], font_size=8.4)
    add_heading(doc, "문서 자체의 한계", 2)
    add_bullet(doc, "현재 작업트리 코드가 r021 이후 바뀐 부분은 공식 상태에 자동 반영되지 않았습니다.")
    add_bullet(doc, "DOC-01과 실제 파일 지문·생성본의 동기화 문제가 있어 개별 파일의 최신성은 별도 확인해야 합니다.")
    add_bullet(doc, "실제 사용자·기기·현장·기관·클라우드 사건은 저장소 내부 문서만으로 만들거나 대신할 수 없습니다.")
    add_bullet(doc, "이 소개서는 기존 산출물 형식을 통일하지 않으며 승인·정본 상태를 바꾸지 않습니다.")

    # 12. Glossary
    add_heading(doc, "12. 쉬운 용어 설명", 1, page_break=True)
    glossary_rows = [(term, explanation) for term, explanation in GLOSSARY]
    add_table(doc, ["용어", "쉬운 설명"], glossary_rows, widths_cm=[5.0, 11.1], font_size=8.4)

    # Appendices in landscape
    add_section(doc, landscape=True)
    add_heading(doc, "부록 A. 기능·공통 규칙 63개와 출시 관문 5개", 1)
    add_callout(
        doc,
        "부록 표 읽기",
        "‘해야 하는 일’은 승인된 정책의 쉬운 설명, ‘r021 판정 당시 관찰’은 2026-07-26에 확인한 코드·자료, "
        "‘남은 일’은 종료 조건입니다. 현재 작업트리의 추가 코드는 재평가 전까지 공식 상태를 바꾸지 않습니다.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    domain_order: list[str] = []
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for assessment in snap.gaps["assessments"]:
        domain = assessment["domain"]
        if domain not in grouped:
            domain_order.append(domain)
        grouped[domain].append(assessment)
    for domain in domain_order:
        add_heading(doc, domain, 2)
        feature_rows = []
        for row in grouped[domain]:
            current = CURRENT_OBSERVATION_KO.get(
                row["gap_id"],
                simplify(row["current_implementation_in_plain_language"]),
            )
            remediation = simplify(row["remediation"])
            feature_rows.append(
                (
                    f"{row['gap_id']}\n{row['source_policy_id']}",
                    row["title"],
                    shorten(simplify(row["policy_in_plain_language"]), 190),
                    format_status(row["status"]),
                    shorten(current, 200),
                    shorten(remediation, 180),
                )
            )
        add_table(
            doc,
            ["Gap·정책 ID", "쉬운 기능명", "해야 하는 일", "공식 상태", "r021 판정 당시 관찰", "남은 일"],
            feature_rows,
            widths_cm=[3.0, 3.6, 5.6, 3.2, 6.0, 5.7],
            font_size=7.4,
        )

    add_heading(doc, "부록 B. 40개 문서 묶음 상세 목록", 1, page_break=True)
    bundle_appendix_rows = []
    for bundle_id in sorted(snap.bundles, key=lambda value: (snap.bundles[value][0]["category"], value)):
        rows = snap.bundles[bundle_id]
        type_ids = ", ".join(row["display_code"] for row in rows)
        next_steps = sorted({artifact_next_step(row) for row in rows})
        bundle_appendix_rows.append(
            (
                bundle_id,
                rows[0]["category"],
                BUNDLE_PURPOSES.get(bundle_id, simplify(rows[0]["authoring_contract"]["purpose"])),
                bundle_paths(rows),
                type_ids,
                bundle_state_summary(rows),
                shorten(" / ".join(next_steps), 170),
            )
        )
    add_table(
        doc,
        ["묶음 ID", "범주", "역할", "경로", "포함 유형", "상태", "주요 남은 일"],
        bundle_appendix_rows,
        widths_cm=[3.1, 1.5, 4.4, 5.2, 4.3, 3.5, 5.1],
        font_size=7.4,
    )

    add_heading(doc, "부록 C. 257개 산출물 유형 상세 목록", 1, page_break=True)
    add_callout(
        doc,
        "이름과 설명을 함께 읽는 방법",
        "유형 이름·ID·경로는 다른 공식 문서와 대조할 수 있도록 원래 표기를 유지했습니다. "
        "뜻은 바로 옆 ‘무엇을 하는가’ 열과 12장 용어 설명에서 쉬운 말로 확인할 수 있습니다.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    for category in CATEGORY_ORDER:
        add_heading(doc, f"{category} · {CATEGORY_INFO[category][0]}", 2)
        artifact_rows = []
        for row in snap.register["artifacts"]:
            if row["category"] != category:
                continue
            location = row["location"].get("canonical_path") or row["location"].get("planned_canonical_path") or "경로 미정"
            anchor = row["location"].get("coverage_anchor")
            if anchor:
                location += f"#{anchor}"
            artifact_rows.append(
                (
                    f"{row['display_code']}\n{row['artifact_instance_id']}",
                    row["title"],
                    shorten(simplify(row["authoring_contract"]["purpose"]), 185),
                    artifact_state_label(row),
                    location,
                    row["responsibility"]["content_owner_role"],
                    shorten(artifact_next_step(row), 145),
                )
            )
        add_table(
            doc,
            ["유형·인스턴스", "이름", "무엇을 하는가", "상태", "정본·예정 위치", "내용 책임", "다음 일"],
            artifact_rows,
            widths_cm=[3.1, 3.6, 5.5, 2.4, 5.2, 3.1, 4.3],
            font_size=7.4,
        )

    add_heading(doc, "부록 D. 근거 자료와 산식", 1, page_break=True)
    source_path_map = {
        "활성 체크포인트": CHECKPOINT_PATH,
        "기능 Gap r021": GAP_PATH,
        "개선 Backlog r021": BACKLOG_PATH,
        "산출물 등록부 DOC-01": DOC01_PATH,
        "정식 시험 목록": TEST_CASES_PATH,
        "산출물 폐쇄 원장 R011": CLOSURE_LEDGER_PATH,
    }
    source_rows = [
        (label, str(path.relative_to(ROOT)), snap.source_hashes[label])
        for label, path in source_path_map.items()
    ]
    add_table(doc, ["자료", "저장소 상대 경로", "SHA-256 파일 지문"], source_rows, widths_cm=[4.2, 11.1, 10.9], font_size=6.9)
    formula_rows = [
        ("기능 상태", "r021 기능 판정 63개에서 상태별 행 수 / 63", "출시 관문 5개 제외"),
        ("출시 관문", "‘남은 출시 관문’으로 분류된 5개 행: 상태 막힘·실행 미실행", "기능 구현률과 분리"),
        ("산출물 생명주기", "DOC-01 257행의 생명주기 상태별 수 / 257", "기능·프로젝트 완료율 아님"),
        ("산출물 폐쇄 투영", "(기준상 문제없음 124 + 현행 범위 제외 승인 2) / 257", "R011 투영, 전체 완료 주장 0"),
        ("열린 산출물", "내부 문서 보완 62 + 사실·승인·독립확인 24 + 내부 실행 24 + 실제 사건 21", "합계 131"),
        ("정식 시험", "정식 시험 목록 279행의 실행 상태", "합격 0, 미실행 279"),
        ("연결관계", "각 유형의 앞단 연결을 방향쌍으로 중복 제거", "698개, 뒤쪽 역조회 재합산 금지"),
    ]
    add_table(doc, ["지표", "산식", "해석 제한"], formula_rows, widths_cm=[4.2, 12.2, 9.8], font_size=7.1)
    add_callout(
        doc,
        "최종 확인",
        "이 문서는 기존 산출물 형식을 통일하거나 승인 상태를 바꾸지 않았습니다. "
        "내용·수치·출시 주장에 대한 사람 검토 뒤 사용하십시오.",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return footnotes


def validate_docx(
    path: Path,
    *,
    expected_footnotes: int,
    expected_figures: int,
) -> dict[str, Any]:
    if not path.is_file() or path.stat().st_size < 100_000:
        raise RuntimeError(f"DOCX is missing or unexpectedly small: {path}")
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise RuntimeError(f"corrupt DOCX ZIP member: {bad}")
        names = archive.namelist()
        if "word/footnotes.xml" not in names:
            raise RuntimeError("real Word footnotes are missing")
        xml_text_parts: list[str] = []
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                payload = archive.read(name)
                ElementTree.fromstring(payload)
                xml_text_parts.append(payload.decode("utf-8", errors="replace"))
        package_text = "\n".join(xml_text_parts)
        if "[[[WS_" in package_text:
            raise RuntimeError("unresolved TOC or footnote marker remains in DOCX")
        document_root = ElementTree.fromstring(archive.read("word/document.xml"))
        footnote_reference_nodes = document_root.findall(f".//{qn('w:footnoteReference')}")
        footnote_references = len(footnote_reference_nodes)
        referenced_ids = {
            int(node.attrib[qn("w:id")])
            for node in footnote_reference_nodes
        }
        footnote_root = ElementTree.fromstring(archive.read("word/footnotes.xml"))
        definition_ids = {
            int(node.attrib[qn("w:id")])
            for node in footnote_root
            if qn("w:id") in node.attrib
        }
        if footnote_references != expected_footnotes:
            raise RuntimeError(
                f"footnote reference count differs: {footnote_references} != {expected_footnotes}"
            )
        if len(referenced_ids) != expected_footnotes:
            raise RuntimeError(
                f"footnote reference IDs differ: {len(referenced_ids)} != {expected_footnotes}"
            )
        if not referenced_ids.issubset(definition_ids):
            raise RuntimeError("one or more footnote references lack definitions")
        media = [name for name in names if name.startswith("word/media/") and not name.endswith("/")]
        if len(media) != expected_figures + 1:
            raise RuntimeError(
                f"embedded image count differs: {len(media)} != {expected_figures + 1}"
            )
        if "ContentIndex" not in package_text and "TOC" not in package_text and "목차" not in package_text:
            raise RuntimeError("table of contents was not materialized")

    reopened = Document(path)
    text_parts = [paragraph.text for paragraph in reopened.paragraphs]
    for table in reopened.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    text = "\n".join(text_parts)
    required = [
        "시각장애인 보행 보조 시스템",
        "공식 완료 0",
        "정식 시험",
        "279개",
        "출시 관문",
        "257개 산출물 유형",
        "부록 A.",
        "부록 C.",
        "출시 불가",
    ]
    missing = [value for value in required if value not in text]
    if missing:
        raise RuntimeError(f"required content missing from DOCX: {missing}")
    forbidden = [
        "TODO",
        "TBD",
        "[[[WS_",
        "EVIDENCE_OR_EXTERNAL_VALUE_PENDING",
        "PENDING_ACTIVATION_EVALUATION",
        "READY25-",
        "authority boundaries",
        "NOT_ELIGIBLE",
        "(PARTIAL)",
        "(MISSING)",
        "(CONFLICTING)",
        "(EVIDENCE_MISSING)",
        "(BLOCKED)",
        "source_kind",
        "REMAINING_GATE",
        "lifecycle_status",
        "OK_BASELINE",
        "INTERNAL_READY",
        "execution_status",
        "NOT_RUN",
    ]
    found = [value for value in forbidden if value in text]
    if found:
        raise RuntimeError(f"forbidden or unresolved content found: {found}")
    if len(text) < 90_000:
        raise RuntimeError(f"document text is unexpectedly short: {len(text)} characters")
    return {
        "byte_length": path.stat().st_size,
        "sha256": sha256(path),
        "paragraph_count": len(reopened.paragraphs),
        "table_count": len(reopened.tables),
        "inline_shape_count": len(reopened.inline_shapes),
        "section_count": len(reopened.sections),
        "text_length": len(text),
        "footnote_count": expected_footnotes,
    }


def render_and_validate_pdf(docx_path: Path, workspace: Path) -> dict[str, Any]:
    profile = workspace / "lo-profile"
    outdir = workspace / "pdf"
    profile.mkdir()
    outdir.mkdir()
    command = [
        "libreoffice",
        "--headless",
        "--nologo",
        "--nodefault",
        "--nofirststartwizard",
        "--norestore",
        f"-env:UserInstallation=file://{profile}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(outdir),
        str(docx_path),
    ]
    result = subprocess.run(
        command,
        cwd=ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        timeout=240,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice PDF render failed: {result.stdout}")
    pdf_path = outdir / (docx_path.stem + ".pdf")
    if not pdf_path.is_file():
        raise RuntimeError(f"rendered PDF is missing: {result.stdout}")
    info = subprocess.run(
        ["pdfinfo", str(pdf_path)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=True,
    ).stdout
    pages_match = re.search(r"^Pages:\s+(\d+)", info, flags=re.MULTILINE)
    if not pages_match:
        raise RuntimeError("PDF page count is missing")
    pages = int(pages_match.group(1))
    if pages < 45:
        raise RuntimeError(f"rendered PDF is unexpectedly short: {pages} pages")
    text_path = workspace / "rendered.txt"
    subprocess.run(
        ["pdftotext", str(pdf_path), str(text_path)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.PIPE,
        text=True,
    )
    extracted = text_path.read_text(encoding="utf-8", errors="replace")
    for required in ("목차", "WalkSafe", "부록 A.", "부록 C.", "출시 관문"):
        if required not in extracted:
            raise RuntimeError(f"rendered PDF text is missing: {required}")
    if "[[[WS_" in extracted:
        raise RuntimeError("unresolved marker remains in rendered PDF")
    extracted_pages = extracted.split("\f")[:pages]
    near_empty_pages = [
        page_number
        for page_number, page_text in enumerate(extracted_pages, start=1)
        if len(re.sub(r"\s+", "", page_text)) < 35
    ]
    if near_empty_pages:
        raise RuntimeError(f"rendered PDF has near-empty pages: {near_empty_pages}")
    layout_path = workspace / "rendered-layout.txt"
    subprocess.run(
        ["pdftotext", "-layout", str(pdf_path), str(layout_path)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.PIPE,
        text=True,
    )
    layout_text = layout_path.read_text(encoding="utf-8", errors="replace")
    for title in (
        "부록 A. 기능·공통 규칙 63개와 출시 관문 5개",
        "부록 B. 40개 문서 묶음 상세 목록",
        "부록 C. 257개 산출물 유형 상세 목록",
        "부록 D. 근거 자료와 산식",
    ):
        toc_match = re.search(
            rf"{re.escape(title)}\.{{3,}}\s*(\d+)",
            layout_text,
        )
        if not toc_match:
            raise RuntimeError(f"TOC page entry is missing: {title}")
        actual_occurrences = [
            page_number
            for page_number, page_text in enumerate(extracted_pages, start=1)
            if title in page_text
        ]
        if not actual_occurrences:
            raise RuntimeError(f"rendered appendix heading is missing: {title}")
        toc_page = int(toc_match.group(1))
        actual_page = max(actual_occurrences)
        if toc_page != actual_page:
            raise RuntimeError(
                f"TOC page differs for {title}: {toc_page} != {actual_page}"
            )
    return {
        "pages": pages,
        "pdf_byte_length": pdf_path.stat().st_size,
        "pdf_text_length": len(extracted),
        "pdf_path": str(pdf_path),
    }


def verify_sources_unchanged(snap: Snapshot) -> None:
    source_path_map = {
        "활성 체크포인트": CHECKPOINT_PATH,
        "기능 Gap r021": GAP_PATH,
        "개선 Backlog r021": BACKLOG_PATH,
        "산출물 등록부 DOC-01": DOC01_PATH,
        "정식 시험 목록": TEST_CASES_PATH,
        "산출물 폐쇄 원장 R011": CLOSURE_LEDGER_PATH,
    }
    current = {label: sha256(path) for label, path in source_path_map.items()}
    if current != snap.source_hashes:
        changed = [
            label
            for label in current
            if current[label] != snap.source_hashes[label]
        ]
        raise RuntimeError(f"source changed during document build: {changed}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--skip-pdf-validation",
        action="store_true",
        help="skip temporary LibreOffice PDF rendering",
    )
    args = parser.parse_args()
    output = args.output.resolve()
    if not FINALIZER.is_file():
        raise RuntimeError(f"UNO finalizer is missing: {FINALIZER}")

    snap = snapshot()
    with tempfile.TemporaryDirectory(prefix="walksafe-deliverables-guide-") as raw:
        workspace = Path(raw)
        figure_dir = workspace / "figures"
        cover, figures = create_figures(snap, figure_dir)
        draft = workspace / "draft.docx"
        finalized = workspace / "final.docx"
        footnotes_path = workspace / "footnotes.json"
        footnotes = build_document(snap, cover, figures, draft)
        footnotes_path.write_text(
            json.dumps(
                [{"marker": item.marker, "note": item.note} for item in footnotes],
                ensure_ascii=False,
                indent=2,
            )
            + "\n",
            encoding="utf-8",
        )
        subprocess.run(
            [
                "/usr/bin/python3.14",
                str(FINALIZER),
                str(draft),
                str(finalized),
                str(footnotes_path),
            ],
            cwd=ROOT,
            check=True,
            timeout=180,
        )
        verify_sources_unchanged(snap)
        validation = validate_docx(
            finalized,
            expected_footnotes=len(footnotes),
            expected_figures=len(figures),
        )
        pdf_validation: dict[str, Any] | None = None
        if not args.skip_pdf_validation:
            pdf_validation = render_and_validate_pdf(finalized, workspace)
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(finalized, output)

    print(
        json.dumps(
            {
                "output": str(output.relative_to(ROOT)),
                "sha256": validation["sha256"],
                "byte_length": validation["byte_length"],
                "sections": validation["section_count"],
                "tables": validation["table_count"],
                "figures": len(figures),
                "footnotes": len(footnotes),
                "text_length": validation["text_length"],
                "pdf_validation": pdf_validation,
                "source_hashes": snap.source_hashes,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

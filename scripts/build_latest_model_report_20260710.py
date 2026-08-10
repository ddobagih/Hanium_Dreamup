#!/usr/bin/env python3
"""Build the evidence-backed report for the latest WalkSafe 13-class model.

The report intentionally reads immutable training/evaluation artifacts instead of
copying metrics from narrative documents. It produces plots, normalized evidence
CSVs, a Markdown report, and a DOCX report from the same in-memory values.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import math
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence
from xml.etree import ElementTree

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RUN_DIR = (
    ROOT
    / "runs/detect/"
    "walksafe_unified_aihub183_png_yolo26n_img768_e300_open150_strict768_20260701_b8w2"
)
RESULTS_CSV = RUN_DIR / "results.csv"
ARGS_YAML = RUN_DIR / "args.yaml"
CHECKPOINT = RUN_DIR / "weights/best.pt"
CANDIDATE_CHECKPOINTS = [
    ROOT
    / "model/artifacts/candidates/walksafe_13cls_yolo26n_img768_20260708/"
    "walksafe_13cls_yolo26n_img768_best_epoch270.pt",
    ROOT
    / "reports/walksafe_best_eval_20260708/final_model_candidate/"
    "walksafe_13cls_yolo26n_img768_best_epoch270.pt",
]
DATASET_DIR = (
    ROOT
    / "datasets/walksafe_unified_coco_aihub_13cls_reviewed_aihub183_png_20260627"
)
DATA_YAML = DATASET_DIR / "data.yaml"
MATERIALIZED_MANIFEST = DATASET_DIR / "materialized_manifest.csv"
PER_CLASS_CSV = ROOT / "reports/walksafe_best_eval_20260708/per_class_metrics_best_val.csv"
VAL_LOG = (
    ROOT
    / "reports/walksafe_best_eval_20260708/logs/"
    "yolo_val_best_img768_20260708_194125.log"
)
DEPTH500_CSV = (
    ROOT
    / "reports/walksafe_best_eval_20260708/inference_test/"
    "depth500_prediction_summary.csv"
)
OUTPUT_DIR = ROOT / "docs/model-data/latest_model_report_20260710"
ASSET_DIR = OUTPUT_DIR / "assets"
MARKDOWN_PATH = OUTPUT_DIR / "WalkSafe_최신_모델_종합보고서_20260710.md"
DOCX_PATH = OUTPUT_DIR / "WalkSafe_최신_모델_종합보고서_20260710.docx"

CHECKPOINT_SHA256 = "a38857e999e1dc1981fffe4c08e5a4bcb1f05be0c7241e9297d6150e913a5669"
AIHUB_183_CURRENT_URL = (
    "https://www.aihub.or.kr/aihubdata/data/view.do?"
    "aihubDataSe=data&currMenu=115&dataSetSn=183&topMenu="
)
CLASS_NAMES = [
    "person",
    "bicycle",
    "car",
    "motorcycle",
    "bus",
    "truck",
    "traffic light",
    "normal_tactile_block",
    "damaged_tactile_block",
    "crosswalk",
    "curb_step",
    "uneven_sidewalk",
    "e_scooter_obstruction",
]
CLASS_KO = [
    "사람",
    "자전거",
    "자동차",
    "오토바이",
    "버스",
    "트럭",
    "신호등",
    "정상 점자블록",
    "손상 점자블록",
    "횡단보도",
    "보도 단차",
    "불균일·파손 보도",
    "전동킥보드 장애물",
]

SOURCE_LABELS = {
    "coco": "COCO 2017",
    "aihub186_barrier_free_outdoor": "AIHub 186 베리어프리존 주행영상",
    "aihub513_road_facility": "AIHub 513 보행 안전 도로시설물",
    "aihub189_sidewalk_surface": "AIHub 189 인도보행 Surface",
    # This is a legacy local alias. The official AIHub catalog identifier is 572.
    "aihub183_abandoned_escooter": "AIHub 572 이륜자동차 안전 위험 시설물",
    "manual_escooter_obstruction": "AIHub 189 수동 승인 전동킥보드",
}

OFFICIAL_SOURCES = [
    {
        "catalog": "COCO 2017",
        "official_name": "COCO 2017 Train/Val",
        "local_key": "coco",
        "url": "https://cocodataset.org/#download",
        "note": "일반 객체 7종이 있는 이미지와 해당 bbox만 통합",
    },
    {
        "catalog": "AIHub 186",
        "official_name": "베리어프리존(장애물 없는 생활공간) 주행영상",
        "local_key": "aihub186_barrier_free_outdoor",
        "url": "https://aihub.or.kr/aihubdata/data/view.do?aihubDataSe=data&currMenu=115&dataSetSn=186&topMenu=",
        "note": "실외 동·서부산 객체 묶음",
    },
    {
        "catalog": "AIHub 513",
        "official_name": "보행 안전을 위한 도로 시설물 데이터",
        "local_key": "aihub513_road_facility",
        "url": "https://www.aihub.or.kr/aihubdata/data/view.do?aihubDataSe=data&currMenu=11&dataSetSn=513&topMenu=",
        "note": "TS8, TS9, VS1, VS2",
    },
    {
        "catalog": "AIHub 189",
        "official_name": "인도보행 영상",
        "local_key": "aihub189_sidewalk_surface / manual_escooter_obstruction",
        "url": "https://www.aihub.or.kr/aihubdata/data/view.do?aihubDataSe=realm&currMenu=115&dataSetSn=189&topMenu=100",
        "note": "Surface_1 및 수동 승인 Bbox/Polygon 7장",
    },
    {
        "catalog": "AIHub 572",
        "official_name": "이륜자동차 안전 위험 시설물 데이터",
        "local_key": "aihub183_abandoned_escooter (레거시 별칭)",
        "url": "https://www.aihub.or.kr/aihubdata/data/view.do?aihubDataSe=data&currMenu=115&dataSetSn=572&topMenu=",
        "note": "23.방치물(전동킥보드) 검수 승인분",
    },
]

SUBSETS = [
    (
        "AIHub 186",
        "train",
        "1.동부산실외_베리어프리객체_1, 2.서부산실외_베리어프리객체_1",
        57_073,
        136_986,
    ),
    (
        "AIHub 186",
        "val",
        "1.동부산실외_베리어프리객체_1, 2.서부산실외_베리어프리객체_1",
        7_147,
        17_250,
    ),
    ("AIHub 513", "train", "TS8/TS9 + TL8/TL9", 23_475, 23_475),
    ("AIHub 513", "val", "VS1/VS2 + VL1/VL2", 16_649, 40_567),
    ("AIHub 189", "train+val", "Surface_1.zip", 2_284, 4_097),
    ("AIHub 189", "train+val", "Bbox_1_new/Polygon_1_new 수동 승인", 7, 7),
    (
        "AIHub 572",
        "train",
        "TS_Bounding Box_23.방치물(전동킥보드) 승인분",
        14_291,
        17_655,
    ),
    (
        "AIHub 572",
        "val",
        "VS_Bounding Box_23.방치물(전동킥보드) 승인분",
        1_714,
        2_096,
    ),
]

CLASS_SOURCES = [
    "COCO 2017",
    "COCO 2017",
    "COCO 2017",
    "COCO 2017",
    "COCO 2017",
    "COCO 2017",
    "COCO 2017",
    "AIHub 186/513/189",
    "AIHub 186/513/189",
    "AIHub 186/189",
    "AIHub 186/513/189",
    "AIHub 186/513/189",
    "AIHub 572 + AIHub 189 수동 7장",
]


@dataclass(frozen=True)
class EpochRow:
    epoch: int
    time_s: float
    train_box: float
    train_cls: float
    train_dfl: float
    precision: float
    recall: float
    map50: float
    map5095: float
    val_box: float
    val_cls: float
    val_dfl: float
    lr: float


@dataclass(frozen=True)
class ClassMetric:
    class_id: int
    class_en: str
    class_ko: str
    definition_ko: str
    images: int
    instances: int
    precision: float
    recall: float
    map50: float
    map5095: float
    assessment: str


@dataclass(frozen=True)
class DatasetSummary:
    source_images: Counter[str]
    source_boxes: Counter[str]
    source_split_images: Counter[tuple[str, str]]
    split_boxes: Counter[str]
    class_images: Counter[int]
    class_boxes: Counter[int]

    @property
    def total_images(self) -> int:
        return sum(self.source_images.values())

    @property
    def total_boxes(self) -> int:
        return sum(self.source_boxes.values())

    @property
    def train_images(self) -> int:
        return sum(v for (source, split), v in self.source_split_images.items() if split == "train")

    @property
    def val_images(self) -> int:
        return sum(v for (source, split), v in self.source_split_images.items() if split == "val")


@dataclass(frozen=True)
class DepthSummary:
    total_images: int
    images_with_detection: int
    total_boxes: int
    class_boxes: dict[str, int]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--validate-only",
        action="store_true",
        help="Validate existing outputs without rebuilding them.",
    )
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_epochs() -> list[EpochRow]:
    rows: list[EpochRow] = []
    with RESULTS_CSV.open(encoding="utf-8", newline="") as stream:
        for raw in csv.DictReader(stream):
            rows.append(
                EpochRow(
                    epoch=int(raw["epoch"]),
                    time_s=float(raw["time"]),
                    train_box=float(raw["train/box_loss"]),
                    train_cls=float(raw["train/cls_loss"]),
                    train_dfl=float(raw["train/dfl_loss"]),
                    precision=float(raw["metrics/precision(B)"]),
                    recall=float(raw["metrics/recall(B)"]),
                    map50=float(raw["metrics/mAP50(B)"]),
                    map5095=float(raw["metrics/mAP50-95(B)"]),
                    val_box=float(raw["val/box_loss"]),
                    val_cls=float(raw["val/cls_loss"]),
                    val_dfl=float(raw["val/dfl_loss"]),
                    lr=float(raw["lr/pg0"]),
                )
            )
    return rows


def read_class_metrics() -> tuple[dict[str, float | int], list[ClassMetric]]:
    overall: dict[str, float | int] = {}
    metrics: list[ClassMetric] = []
    with PER_CLASS_CSV.open(encoding="utf-8", newline="") as stream:
        for raw in csv.DictReader(stream):
            if raw["class_en"] == "all":
                overall = {
                    "images": int(raw["images"]),
                    "instances": int(raw["instances"]),
                    "precision": float(raw["precision"]),
                    "recall": float(raw["recall"]),
                    "map50": float(raw["mAP50"]),
                    "map5095": float(raw["mAP50-95"]),
                }
                continue
            metrics.append(
                ClassMetric(
                    class_id=int(raw["class_id"]),
                    class_en=raw["class_en"],
                    class_ko=raw["class_ko"],
                    definition_ko=raw["definition_ko"],
                    images=int(raw["images"]),
                    instances=int(raw["instances"]),
                    precision=float(raw["precision"]),
                    recall=float(raw["recall"]),
                    map50=float(raw["mAP50"]),
                    map5095=float(raw["mAP50-95"]),
                    assessment=raw["assessment"],
                )
            )
    return overall, metrics


def read_dataset_summary() -> DatasetSummary:
    source_images: Counter[str] = Counter()
    source_boxes: Counter[str] = Counter()
    source_split_images: Counter[tuple[str, str]] = Counter()
    split_boxes: Counter[str] = Counter()
    class_images: Counter[int] = Counter()
    class_boxes: Counter[int] = Counter()
    with MATERIALIZED_MANIFEST.open(encoding="utf-8", newline="") as stream:
        for row in csv.DictReader(stream):
            source = row["source"]
            split = row["split"]
            boxes = int(row["boxes"])
            source_images[source] += 1
            source_boxes[source] += boxes
            source_split_images[source, split] += 1
            split_boxes[split] += boxes
            for class_id, class_name in enumerate(CLASS_NAMES):
                value = int(row[f"class_{class_id}_{class_name}_boxes"])
                class_boxes[class_id] += value
                if value:
                    class_images[class_id] += 1
    return DatasetSummary(
        source_images=source_images,
        source_boxes=source_boxes,
        source_split_images=source_split_images,
        split_boxes=split_boxes,
        class_images=class_images,
        class_boxes=class_boxes,
    )


def read_depth_summary() -> DepthSummary:
    class_boxes: dict[str, int] = {}
    total_images = images_with_detection = total_boxes = 0
    with DEPTH500_CSV.open(encoding="utf-8", newline="") as stream:
        for row in csv.DictReader(stream):
            if row["scope"] != "all_depth500":
                continue
            total_images = int(row["total_images"])
            images_with_detection = int(row["images_with_detection"])
            total_boxes = int(row["total_boxes"])
            class_boxes[row["class_en"]] = int(row["box_count"])
    return DepthSummary(total_images, images_with_detection, total_boxes, class_boxes)


def configure_plot_font() -> None:
    candidates = [
        Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Medium.ttc"),
    ]
    font_path = next((path for path in candidates if path.exists()), None)
    if font_path is None:
        raise RuntimeError("A Korean-capable font is required to build report plots")
    font_manager.fontManager.addfont(str(font_path))
    plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font_path)).get_name()
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 140
    plt.rcParams["savefig.dpi"] = 180


def apply_plot_style(ax: plt.Axes) -> None:
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#D7DBE0", linewidth=0.7, alpha=0.8)
    ax.set_axisbelow(True)


def build_plots(
    epochs: Sequence[EpochRow], metrics: Sequence[ClassMetric], dataset: DatasetSummary
) -> list[Path]:
    configure_plot_font()
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    xs = [row.epoch for row in epochs]
    best = max(epochs, key=lambda row: row.map5095)

    path = ASSET_DIR / "01_epoch_metrics.png"
    fig, ax = plt.subplots(figsize=(10, 5.5))
    ax.plot(xs, [row.precision for row in epochs], label="Precision", color="#0B6E4F", linewidth=1.7)
    ax.plot(xs, [row.recall for row in epochs], label="Recall", color="#E07A1F", linewidth=1.7)
    ax.plot(xs, [row.map50 for row in epochs], label="mAP50", color="#2864B4", linewidth=2.0)
    ax.plot(xs, [row.map5095 for row in epochs], label="mAP50-95", color="#8B3A3A", linewidth=2.0)
    ax.axvline(best.epoch, color="#20252B", linestyle="--", linewidth=1.0)
    ax.scatter([best.epoch], [best.map5095], color="#20252B", zorder=4)
    ax.annotate(
        f"best epoch {best.epoch}\nmAP50-95 {best.map5095:.4f}",
        xy=(best.epoch, best.map5095),
        xytext=(-105, -38),
        textcoords="offset points",
        arrowprops={"arrowstyle": "->", "color": "#20252B"},
        fontsize=9,
    )
    ax.set(title="Epoch별 validation 탐지 지표", xlabel="Epoch", ylabel="Score", ylim=(0, 1.0))
    apply_plot_style(ax)
    ax.legend(ncol=4, loc="lower right", frameon=False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    outputs.append(path)

    path = ASSET_DIR / "02_train_val_losses.png"
    fig, axes = plt.subplots(1, 3, figsize=(12, 4.2))
    loss_sets = [
        ("Box loss", [r.train_box for r in epochs], [r.val_box for r in epochs]),
        ("Class loss", [r.train_cls for r in epochs], [r.val_cls for r in epochs]),
        ("DFL loss", [r.train_dfl for r in epochs], [r.val_dfl for r in epochs]),
    ]
    for ax, (title, train_values, val_values) in zip(axes, loss_sets, strict=True):
        ax.plot(xs, train_values, label="train", color="#0B6E4F", linewidth=1.4)
        ax.plot(xs, val_values, label="val", color="#C4473A", linewidth=1.4)
        ax.set(title=title, xlabel="Epoch")
        apply_plot_style(ax)
        ax.legend(frameon=False)
    fig.suptitle("학습·검증 loss 비교", fontsize=15, fontweight="bold")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    outputs.append(path)

    path = ASSET_DIR / "03_class_performance.png"
    fig, ax = plt.subplots(figsize=(11, 7.6))
    ys = list(range(len(metrics)))
    width = 0.19
    series = [
        ("Precision", [m.precision for m in metrics], "#0B6E4F"),
        ("Recall", [m.recall for m in metrics], "#E07A1F"),
        ("mAP50", [m.map50 for m in metrics], "#2864B4"),
        ("mAP50-95", [m.map5095 for m in metrics], "#8B3A3A"),
    ]
    for offset, (label, values, color) in enumerate(series):
        ax.barh([y + (offset - 1.5) * width for y in ys], values, height=width, label=label, color=color)
    ax.set_yticks(ys, [f"{m.class_id}. {m.class_ko}" for m in metrics])
    ax.set_xlim(0, 1.0)
    ax.set(title="클래스별 validation 성능", xlabel="Score")
    ax.invert_yaxis()
    ax.grid(axis="x", color="#D7DBE0", linewidth=0.7)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(ncol=4, loc="lower right", frameon=False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    outputs.append(path)

    path = ASSET_DIR / "04_dataset_sources.png"
    ordered_sources = sorted(dataset.source_images, key=lambda source: dataset.source_images[source], reverse=True)
    labels = [SOURCE_LABELS[source] for source in ordered_sources]
    fig, axes = plt.subplots(1, 2, figsize=(12, 5.5))
    axes[0].barh(labels, [dataset.source_images[s] for s in ordered_sources], color="#2864B4")
    axes[0].set(title="출처별 이미지", xlabel="Images")
    axes[1].barh(labels, [dataset.source_boxes[s] for s in ordered_sources], color="#0B6E4F")
    axes[1].set(title="출처별 bbox", xlabel="Bounding boxes")
    for ax in axes:
        ax.invert_yaxis()
        ax.grid(axis="x", color="#D7DBE0", linewidth=0.7)
        ax.spines[["top", "right"]].set_visible(False)
        ax.ticklabel_format(style="plain", axis="x")
    fig.suptitle("최종 materialized dataset 출처 구성", fontsize=15, fontweight="bold")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    outputs.append(path)

    path = ASSET_DIR / "05_class_distribution.png"
    fig, ax = plt.subplots(figsize=(10, 6.8))
    values = [dataset.class_boxes[i] for i in range(len(CLASS_NAMES))]
    colors = ["#C4473A" if i in {10, 11} else "#486F8C" for i in range(len(values))]
    ax.barh([f"{i}. {CLASS_KO[i]}" for i in range(len(values))], values, color=colors)
    ax.set(title="학습 통합본 클래스별 bbox 분포", xlabel="Bounding boxes")
    ax.invert_yaxis()
    ax.grid(axis="x", color="#D7DBE0", linewidth=0.7)
    ax.spines[["top", "right"]].set_visible(False)
    for y, value in enumerate(values):
        ax.text(value, y, f" {value:,}", va="center", fontsize=8)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    outputs.append(path)

    path = ASSET_DIR / "06_precision_recall_risk.png"
    fig, ax = plt.subplots(figsize=(9.5, 6.2))
    sizes = [max(45, math.sqrt(m.instances) * 5) for m in metrics]
    colors = [m.map5095 for m in metrics]
    scatter = ax.scatter(
        [m.recall for m in metrics],
        [m.precision for m in metrics],
        s=sizes,
        c=colors,
        cmap="viridis",
        vmin=0,
        vmax=1,
        alpha=0.82,
        edgecolors="white",
        linewidths=0.7,
    )
    for metric in metrics:
        ax.annotate(
            str(metric.class_id),
            (metric.recall, metric.precision),
            xytext=(4, 4),
            textcoords="offset points",
            fontsize=8,
        )
    ax.axvline(0.5, color="#AAB0B6", linestyle="--", linewidth=0.8)
    ax.axhline(0.5, color="#AAB0B6", linestyle="--", linewidth=0.8)
    ax.set(
        title="클래스별 Precision-Recall 위치 (점 크기: 객체 수)",
        xlabel="Recall",
        ylabel="Precision",
        xlim=(0, 1),
        ylim=(0, 1),
    )
    apply_plot_style(ax)
    colorbar = fig.colorbar(scatter, ax=ax, pad=0.02)
    colorbar.set_label("mAP50-95")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    outputs.append(path)
    return outputs


def write_evidence_csvs(
    epochs: Sequence[EpochRow], metrics: Sequence[ClassMetric], dataset: DatasetSummary
) -> list[Path]:
    outputs: list[Path] = []
    path = OUTPUT_DIR / "epoch_milestones.csv"
    milestones = {1, 50, 100, 150, 200, 250, 270, 300}
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["epoch", "time_hours", "precision", "recall", "mAP50", "mAP50-95"])
        for row in epochs:
            if row.epoch in milestones:
                writer.writerow(
                    [
                        row.epoch,
                        f"{row.time_s / 3600:.4f}",
                        f"{row.precision:.5f}",
                        f"{row.recall:.5f}",
                        f"{row.map50:.5f}",
                        f"{row.map5095:.5f}",
                    ]
                )
    outputs.append(path)

    path = OUTPUT_DIR / "dataset_source_summary.csv"
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["source_key", "display_name", "train_images", "val_images", "total_images", "boxes"])
        for source in sorted(dataset.source_images, key=lambda value: SOURCE_LABELS[value]):
            writer.writerow(
                [
                    source,
                    SOURCE_LABELS[source],
                    dataset.source_split_images[source, "train"],
                    dataset.source_split_images[source, "val"],
                    dataset.source_images[source],
                    dataset.source_boxes[source],
                ]
            )
    outputs.append(path)

    path = OUTPUT_DIR / "class_distribution_and_metrics.csv"
    by_id = {metric.class_id: metric for metric in metrics}
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "class_id",
                "class_en",
                "class_ko",
                "dataset_images",
                "dataset_boxes",
                "evaluated_images",
                "evaluated_instances",
                "precision",
                "recall",
                "mAP50",
                "mAP50-95",
                "source_assessment",
                "report_assessment",
            ]
        )
        for class_id, class_name in enumerate(CLASS_NAMES):
            metric = by_id[class_id]
            writer.writerow(
                [
                    class_id,
                    class_name,
                    CLASS_KO[class_id],
                    dataset.class_images[class_id],
                    dataset.class_boxes[class_id],
                    metric.images,
                    metric.instances,
                    metric.precision,
                    metric.recall,
                    metric.map50,
                    metric.map5095,
                    metric.assessment,
                    report_assessment(metric),
                ]
            )
    outputs.append(path)
    return outputs


def markdown_table(headers: Sequence[str], rows: Iterable[Sequence[object]]) -> str:
    rendered = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        rendered.append("| " + " | ".join(str(value) for value in row) + " |")
    return "\n".join(rendered)


def pct(value: float, digits: int = 1) -> str:
    return f"{value * 100:.{digits}f}%"


def report_assessment(metric: ClassMetric) -> str:
    """Return a safety-oriented label while preserving the source CSV label."""
    if metric.class_id in {8, 12}:
        return "강함"
    if metric.class_id in {10, 11}:
        return "핵심 개선 필요"
    if metric.precision < 0.4:
        return "오탐 개선 필요"
    if metric.class_id == 6:
        return "취약"
    if metric.class_id in {1, 2, 3, 5}:
        return "보완 필요"
    return "상대적 양호"


def build_markdown(
    epochs: Sequence[EpochRow],
    overall: dict[str, float | int],
    metrics: Sequence[ClassMetric],
    dataset: DatasetSummary,
    depth: DepthSummary,
) -> str:
    best = max(epochs, key=lambda row: row.map5095)
    last = epochs[-1]
    source_rows = []
    for source in sorted(dataset.source_images, key=lambda value: dataset.source_images[value], reverse=True):
        source_rows.append(
            (
                f"`{source}`",
                SOURCE_LABELS[source],
                f"{dataset.source_split_images[source, 'train']:,}",
                f"{dataset.source_split_images[source, 'val']:,}",
                f"{dataset.source_images[source]:,}",
                f"{dataset.source_boxes[source]:,}",
            )
        )
    class_rows = []
    for metric in metrics:
        class_rows.append(
            (
                metric.class_id,
                f"{metric.class_ko}<br>`{metric.class_en}`",
                f"{dataset.class_images[metric.class_id]:,}",
                f"{dataset.class_boxes[metric.class_id]:,}",
                f"{metric.images:,}",
                f"{metric.instances:,}",
                pct(metric.precision),
                pct(metric.recall),
                pct(metric.map50),
                pct(metric.map5095),
                report_assessment(metric),
            )
        )
    official_rows = [
        (source["catalog"], f"[{source['official_name']}]({source['url']})", f"`{source['local_key']}`", source["note"])
        for source in OFFICIAL_SOURCES
    ]
    subset_rows = [
        (catalog, split, component, f"{images:,}", f"{boxes:,}")
        for catalog, split, component, images, boxes in SUBSETS
    ]
    mapping_rows = [
        (
            index,
            f"{CLASS_KO[index]} / `{CLASS_NAMES[index]}`",
            CLASS_SOURCES[index],
            metrics[index].definition_ko,
        )
        for index in range(len(CLASS_NAMES))
    ]
    milestone_rows = []
    for row in epochs:
        if row.epoch in {1, 50, 100, 150, 200, 250, 270, 300}:
            milestone_rows.append(
                (
                    row.epoch,
                    f"{row.time_s / 3600:.2f}",
                    f"{row.precision:.5f}",
                    f"{row.recall:.5f}",
                    f"{row.map50:.5f}",
                    f"{row.map5095:.5f}",
                )
            )
    return f"""# WalkSafe 최신 13클래스 객체탐지 모델 종합보고서

- 기준일: 2026-07-10
- 모델: YOLO26n, 입력 크기 768, 총 300 epoch
- 최종 후보: epoch {best.epoch}의 `best.pt`
- 판정: **후보 모델(PARTIAL), 현장 배포 확정 아님**

## 1. 결론 요약

최신 후보는 내부 validation에서 precision **{best.precision:.5f}**, recall **{best.recall:.5f}**, mAP50 **{best.map50:.5f}**, mAP50-95 **{best.map5095:.5f}**를 기록했다. 강한 클래스는 전동킥보드 장애물과 손상 점자블록이고, 가장 취약한 클래스는 보도 단차와 불균일·파손 보도다. 안전 안내의 핵심 클래스가 취약하므로 현재 weight를 실사용 최종 모델로 선언할 수 없다.

이번 보고서가 새로 바로잡는 출처 식별 사항이 있다. 로컬 폴더·manifest의 `aihub183_abandoned_escooter`와 학습 run 이름의 `aihub183`은 **레거시 내부 별칭**이다. [현재 공식 `dataSetSn=183`]({AIHUB_183_CURRENT_URL})의 제목은 `드론 이동체 인지 영상(전방 고정)`이며 이 학습 원천이 아니다. 원천 파일명 `23.방치물(전동킥보드)`과 공식 페이지를 대조하면 사용 원천의 카탈로그 번호는 **AIHub 572**이고 공식 자료명은 **이륜자동차 안전 위험 시설물 데이터**다. 재현 경로를 깨지 않기 위해 기존 파일명은 유지하되, 외부 문서에는 AIHub 572로 표기해야 한다.

## 2. 모델 식별과 재현 기준

| 항목 | 값 |
| --- | --- |
| 원 학습 run | `{RUN_DIR.relative_to(ROOT)}` |
| 학습 데이터 | `{DATA_YAML.relative_to(ROOT)}` |
| 후보 checkpoint | `{CHECKPOINT.relative_to(ROOT)}` |
| checkpoint SHA-256 | `{CHECKPOINT_SHA256}` |
| 구조 | YOLO26n, detect, 13 classes |
| 입력/학습 | imgsz 768, batch 8, SGD, AMP |
| 반복 | 300 epoch, cosine LR, close mosaic 20 |
| seed/재현 옵션 | seed 0, deterministic true |
| Ultralytics | checkpoint metadata 8.4.48 |
| 기록 누적시간 | {last.time_s / 3600:.2f}시간 (`results.csv`의 누적 `time`) |

`best.pt`와 별도 후보 복사본은 동일한 SHA-256을 가진다. checkpoint의 저장 후 metadata `epoch=-1`은 Ultralytics가 optimizer를 제거해 strip한 결과이므로, epoch 식별은 weight 내부 값이 아니라 원 run의 `results.csv` 최고점과 보존 경로로 확정한다.

## 3. 데이터 출처와 실제 사용 범위

### 3.1 공식 자료명·번호

{markdown_table(["공식 식별", "공식 자료명", "로컬 source key", "실제 사용 범위"], official_rows)}

AIHub 페이지에 적힌 전체 구축량과 이 프로젝트의 사용량은 다르다. 아래 수치는 공식 페이지의 전체 규모가 아니라 `materialized_manifest.csv`에 실제 들어간 이미지와 bbox만 집계한 값이다.

### 3.2 최종 통합본 구성

- train: **{dataset.train_images:,}장 / {dataset.split_boxes['train']:,} bbox**
- val: **{dataset.val_images:,}장 / {dataset.split_boxes['val']:,} bbox**
- 전체: **{dataset.total_images:,}장 / {dataset.total_boxes:,} bbox**
- 독립 test: **없음**

{markdown_table(["source key", "표시명", "train", "val", "전체 이미지", "bbox"], source_rows)}

![출처별 데이터 구성](assets/04_dataset_sources.png)

### 3.3 AIHub 내부에서 실제 선택한 묶음

{markdown_table(["출처", "split", "사용 컴포넌트", "이미지/행", "bbox"], subset_rows)}

- AIHub 186은 점자블록 정상/파손, 평면횡단보도, 연석·문턱·계단, 평탄성 D/E 및 파손 포장 계열을 변환했다.
- AIHub 513은 점자블럭, 턱낮추기·연석·보행자 계단, 보도·보도블록 불량 계열을 사용했다. 현재 통합본에는 이 source의 횡단보도 bbox가 없다.
- AIHub 189의 `Surface_1`은 surface/guide-block/crosswalk 후보를 bbox로 변환했다. `Depth_001~005`는 13클래스 YOLO 학습에 쓰지 않았다.
- AIHub 572는 `23.방치물(전동킥보드)` 원천 후보 중 사람 검수에서 `approve` 또는 bbox 수정 후 승인된 19,751개 bbox만 사용했다. 16,005개 고유 PNG로 구성된다.
- COCO 2017은 전체 원본을 무조건 넣은 것이 아니라, 사람·자전거·자동차·오토바이·버스·트럭·신호등 7종에 해당하는 이미지와 bbox를 매핑했다.

AIHub 572 공식 페이지의 구축 규모 표는 `23. 방치물(전동킥보드)`를 15,000장으로 안내하지만, local 승인·materialize 기록은 16,005개 고유 원천 이미지와 19,751개 bbox를 가진다. 현재 보존 근거만으로 이 1,005장 차이의 원인을 확정할 수 없다. 학습 재현에는 local manifest 수치를 사용하되, 이를 AIHub 공식 구축량이라고 표현하지 않는다.

### 3.4 탐지 클래스와 매핑

{markdown_table(["ID", "통합 클래스", "주요 원천", "운영 정의"], mapping_rows)}

![클래스별 bbox 분포](assets/05_class_distribution.png)

## 4. 학습 경과

### 4.1 Epoch별 성능

{markdown_table(["epoch", "누적시간(h)", "Precision", "Recall", "mAP50", "mAP50-95"], milestone_rows)}

![Epoch별 성능 지표](assets/01_epoch_metrics.png)

epoch {best.epoch}에서 mAP50-95가 최고 **{best.map5095:.5f}**에 도달했고, epoch 300은 **{last.map5095:.5f}**로 낮아졌다. 마지막 epoch의 precision은 {last.precision:.5f}로 소폭 높지만, 종합 localization 지표가 낮으므로 최종 후보는 epoch {best.epoch}이 타당하다.

### 4.2 Loss와 일반화 징후

![Train/validation loss](assets/02_train_val_losses.png)

후반부에도 train loss는 낮아지는 반면 validation box/class/DFL loss는 높은 수준에서 벌어진다. 이것만으로 원인을 하나로 단정할 수는 없지만, 과적합·source 분포 차이·클래스 정의 충돌·라벨 품질 문제를 우선 점검해야 하는 신호다. 최고점 이후 30 epoch를 더 학습해도 mAP50-95가 회복되지 않았다.

## 5. 최종 validation 결과

### 5.1 평가 경계

- intended val manifest: {dataset.val_images:,}장 / {dataset.split_boxes['val']:,} bbox
- 실제 `yolo val` 평가: **{int(overall['images']):,}장 / {int(overall['instances']):,} 객체**
- Ultralytics scan에서 제외: **{dataset.val_images - int(overall['images']):,}장**
- 제외 원인: 주로 AIHub 513 계열 truncated JPEG가 `corrupt image/label`로 판정됨
- 평가 imgsz: 768
- 독립 test가 아니므로 아래 수치는 내부 validation 성능이다.

### 5.2 클래스별 성능

{markdown_table(["ID", "클래스", "전체 이미지", "전체 bbox", "평가 이미지", "평가 객체", "P", "R", "mAP50", "mAP50-95", "판정"], class_rows)}

판정은 원 평가 CSV의 문구를 그대로 복사하지 않고 안전 관점으로 정규화했다. mAP가 높더라도 precision이 40% 미만이면 `오탐 개선 필요`, 보도 단차·불균일 보도는 낮은 recall 때문에 `핵심 개선 필요`로 표시한다. 따라서 원 CSV에서 높은 recall을 근거로 `양호`였던 정상 점자블록은 이 보고서에서 `오탐 개선 필요`다.

![클래스별 성능](assets/03_class_performance.png)

![Precision-Recall 위험 위치](assets/06_precision_recall_risk.png)

### 5.3 해석

- **강점:** `e_scooter_obstruction`은 P {pct(metrics[12].precision)}, R {pct(metrics[12].recall)}, mAP50-95 {pct(metrics[12].map5095)}다. `damaged_tactile_block`도 R {pct(metrics[8].recall)}, mAP50-95 {pct(metrics[8].map5095)}로 내부 validation에서 강하다.
- **치명적 미탐 위험:** `curb_step` recall {pct(metrics[10].recall)}, `uneven_sidewalk` recall {pct(metrics[11].recall)}로 실제 위험을 대부분 놓칠 수 있다.
- **오탐 위험:** `crosswalk` precision {pct(metrics[9].precision)}, `normal_tactile_block` precision {pct(metrics[7].precision)}다. 반복 음성 안내나 손상/정상 혼동 가능성을 현장 threshold 평가로 확인해야 한다.
- **일반 객체:** 사람·버스는 상대적으로 양호하지만 자전거·트럭·신호등은 recall 또는 엄격 mAP가 낮다.

## 6. 외부·실환경 시험의 현재 경계

기존 평가 폴더에는 클래스별 대표 샘플 69장과 AIHub 189 Depth RGB {depth.total_images:,}장 추론 기록이 있다. {depth.total_images:,}장 중 {depth.images_with_detection:,}장에서 {depth.total_boxes:,}개 예측이 발생했고, 주요 예측은 정상 점자블록 {depth.class_boxes['normal_tactile_block']:,}, 횡단보도 {depth.class_boxes['crosswalk']:,}, 보도 단차 {depth.class_boxes['curb_step']:,}, 불균일·파손 보도 {depth.class_boxes['uneven_sidewalk']:,}개다. 그러나 이 데이터에는 정답 bbox가 없으므로 이 숫자는 검출 발생량이지 정확도가 아니다.

대표 샘플도 클래스별로 고른 소규모 집합이라 독립 test를 대신할 수 없다. 실환경 성능을 주장하려면 프로젝트와 장소·시간·촬영자·연속 프레임이 겹치지 않는 고정 test set에 사람이 확정한 정답 bbox가 필요하다.

## 7. 서비스 연결 시 해석 주의

- 이 보고서의 정량 지표는 **imgsz 768** validation 결과다.
- backend unified runtime은 `backend/app/services/detect_v2.py`에서 **imgsz 768**로 추론한다.
- validation과 backend 입력 크기는 같지만 confidence, NMS, 카메라 흔들림·압축·network 지연이 다르므로 내부 validation을 현장 성능으로 그대로 인용하면 안 된다.
- epoch270 기반 unified float32 img768 13-class TFLite를 Android expected primary로 반영했다. asset/hash/tensor/class 계약과 SM-G981N 실제 load/invoke instrumentation 1/1은 확인했지만 PyTorch/TFLite box·class 동등성, 대화형 camera pipeline·FPS·실외 Device Field는 미검증이다.
- 실제 모델 기본 활성화, 모바일 브라우저/실폰 카메라 E2E, 지연시간·발열·배터리·야간 성능은 별도 검증 대상이다.

## 8. 개선 로드맵

### P0. 데이터·평가 무결성

1. truncated JPEG 2,331장을 원본에서 복구하거나 명시적으로 제거하고 manifest와 라벨을 다시 고정한다.
2. 촬영 연속 프레임·원천 zip·장소 단위 group split을 적용해 train/val 누수를 재감사한다.
3. 현재 val과 분리된 독립 test를 만든다. 단일 이미지뿐 아니라 낮/밤, 역광, 우천, 흔들림, 원거리·부분가림을 포함한다.
4. `aihub183` 내부 별칭과 공식 AIHub 572의 대응표를 provenance에 고정한다.

### P0. 핵심 약점 클래스 재정의

1. `curb_step`은 연석·문턱·계단·경사형 경계를 한 클래스에 묶은 범위가 넓다. 서비스 위험 정의에 맞춰 포함/제외 예시를 먼저 고정한다.
2. `uneven_sidewalk`은 평탄성 D/E, 포장 파손, 보도블록 파손이 섞여 크기·형태 분산이 크다. 손상 영역 bbox 정책을 통일한다.
3. 두 클래스가 같은 객체를 중복 또는 상충 라벨링한 표본을 source별로 층화 추출해 재검수한다.
4. 새 데이터 양을 늘리기 전에 승인 라벨의 일관성과 작은/과대 bbox 비율을 측정한다.

### P1. 오탐과 runtime 보정

1. `normal_tactile_block`과 `crosswalk` hard negative를 수집하고 손상/정상 점자블록 confusion을 별도 표로 관리한다.
2. 동일 checkpoint의 backend img768 전처리·threshold·NMS 회귀를 validation 조건과 비교한다.
3. 클래스별 confidence와 위험 후처리 threshold를 validation이 아니라 고정 test/현장 데이터에서 조정한다.
4. 처리 지연, 프레임 누락, TTS 빈도까지 포함한 end-to-end 안전 지표를 정의한다.

### P1/P2. 배포와 지속 개선

1. 적용된 13-class TFLite의 PyTorch 대비 box/class 수치 동등성 회귀를 수행한다.
2. 실폰에서 FPS, P95 latency, 메모리, 발열, 배터리, 768 입력 동작을 측정한다.
3. 개인정보를 제거한 실패 프레임과 사용자 피드백을 review queue로 축적하되 자동 재학습은 하지 않는다.
4. model registry에 checkpoint hash, dataset manifest hash, split policy, 평가 결과, 배포 승인자를 함께 기록한다.

## 9. 다음 모델의 최소 승인 조건

- 읽을 수 없는 평가 이미지 0장 또는 사전 확정된 제외 manifest
- 독립 test의 출처·그룹 분리와 라벨 검수 완료
- imgsz 768 backend 전처리·threshold·NMS 조건의 회귀 보고서
- 보도 단차·불균일 보도에 대해 팀이 사전에 합의한 recall 하한 충족
- 점자블록/횡단보도의 오탐이 음성 안내 정책에서 허용 가능한지 현장 확인
- 모바일 실기기 E2E와 정지·보행·야간 시나리오 통과
- PyTorch 후보와 배포 artifact의 hash·출력 동등성 기록

수치 하한은 이 문서에서 임의로 정하지 않는다. 안전 요구사항과 현장 실패 비용을 팀이 먼저 합의한 뒤 test set을 보지 않고 고정해야 한다.

## 10. 근거 파일과 재현

### 원 근거

- `{RESULTS_CSV.relative_to(ROOT)}`: 300 epoch 원 수치
- `{ARGS_YAML.relative_to(ROOT)}`: 학습 인자
- `{MATERIALIZED_MANIFEST.relative_to(ROOT)}`: 실제 데이터 provenance와 class count
- `{PER_CLASS_CSV.relative_to(ROOT)}`: 최종 validation 클래스별 수치
- `{VAL_LOG.relative_to(ROOT)}`: 2,331 corrupt 제외 및 실제 평가 집계
- `{DEPTH500_CSV.relative_to(ROOT)}`: 정답 없는 Depth 500장 추론 발생량
- `data_sources/manifests/walksafe_unified_13class_aihub183_png_append_20260627.json`: 전동킥보드 승인분 materialize 근거
- `data_sources/manifests/walksafe_aihub_source_usage_20260628.md`: zip/component provenance

### 생성 명령

```bash
.venv/bin/python scripts/build_latest_model_report_20260710.py
.venv/bin/python scripts/build_latest_model_report_20260710.py --validate-only
```

### 생성 산출물

- `{MARKDOWN_PATH.relative_to(ROOT)}`
- `{DOCX_PATH.relative_to(ROOT)}`
- `epoch_milestones.csv`
- `dataset_source_summary.csv`
- `class_distribution_and_metrics.csv`
- `assets/01_epoch_metrics.png` ~ `assets/06_precision_recall_risk.png`

## 11. 보고서 한계

이 보고서는 보존된 local artifact를 기준으로 재계산한 결과다. 학습 자체를 다시 실행하지 않았고, 독립 test가 없으며, corrupt 제외 전 원본 val 전체의 성능을 복원하지 않았다. 공식 AIHub 페이지는 자료명·카탈로그 번호·원천 라벨 의미 확인에만 사용했고, 프로젝트 반영량은 local manifest에서만 계산했다.
"""


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Noto Sans CJK KR"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_font(run)


def add_paragraph(document: Document, text: str = "", bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            set_run_font(run)
        if not paragraph.runs:
            set_run_font(paragraph.add_run(item))


def add_numbered(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(2)
        if paragraph.runs:
            paragraph.runs[0].text = item
            set_run_font(paragraph.runs[0])
        else:
            set_run_font(paragraph.add_run(item))


def add_table(
    document: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    font_size: float = 8.5,
) -> None:
    materialized = list(rows)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_properties.append(OxmlElement("w:tblHeader"))
    header_properties.append(OxmlElement("w:cantSplit"))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, "244A64")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, size=font_size, bold=True, color="FFFFFF")
    for row_index, values in enumerate(materialized):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "F2F5F7")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=font_size)
    document.add_paragraph()


def add_picture(document: Document, path: Path, width: float = 6.5, caption: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption_paragraph.add_run(caption)
        set_run_font(run, size=8, color="5B6570")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    for style_name, size in (("Normal", 9.5), ("Title", 26), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[style_name]
        style.font.name = "Noto Sans CJK KR"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK KR")
        style.font.size = Pt(size)
    document.styles["Heading 1"].font.color.rgb = RGBColor(0x18, 0x3C, 0x52)
    document.styles["Heading 2"].font.color.rgb = RGBColor(0x28, 0x64, 0x7D)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("WalkSafe | 최신 모델 종합보고서 | 2026-07-10")
    set_run_font(run, size=8, color="66717B")
    add_page_number(section.footer.paragraphs[0])


def build_docx(
    epochs: Sequence[EpochRow],
    overall: dict[str, float | int],
    metrics: Sequence[ClassMetric],
    dataset: DatasetSummary,
    depth: DepthSummary,
) -> None:
    best = max(epochs, key=lambda row: row.map5095)
    last = epochs[-1]
    document = Document()
    configure_document(document)

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(100)
    set_run_font(cover.add_run("WalkSafe"), size=18, bold=True, color="28647D")
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("최신 13클래스 객체탐지 모델\n종합보고서"), size=27, bold=True, color="183C52")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(28)
    set_run_font(subtitle.add_run("YOLO26n · imgsz 768 · 300 epoch"), size=13, color="486776")
    badge = document.add_paragraph()
    badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    badge.paragraph_format.space_before = Pt(20)
    set_run_font(badge.add_run("판정: 후보 모델(PARTIAL) / 현장 배포 확정 아님"), size=12, bold=True, color="A33D2D")
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(110)
    set_run_font(date.add_run("기준일 2026-07-10"), size=11, color="5B6570")
    document.add_page_break()

    add_heading(document, "1. 요약", 1)
    add_paragraph(
        document,
        f"epoch {best.epoch} 후보는 내부 validation에서 precision {best.precision:.5f}, "
        f"recall {best.recall:.5f}, mAP50 {best.map50:.5f}, mAP50-95 {best.map5095:.5f}를 기록했다. "
        "전동킥보드 장애물과 손상 점자블록은 강하지만, 보도 단차와 불균일·파손 보도 recall이 "
        "낮아 실사용 최종 모델로 확정할 수 없다.",
    )
    add_table(
        document,
        ["후보", "P", "R", "mAP50", "mAP50-95", "판정"],
        [[f"epoch {best.epoch}", f"{best.precision:.5f}", f"{best.recall:.5f}", f"{best.map50:.5f}", f"{best.map5095:.5f}", "PARTIAL"]],
        font_size=9,
    )
    add_heading(document, "출처 식별 정정", 2)
    add_paragraph(
        document,
        "로컬 `aihub183_abandoned_escooter`는 레거시 내부 별칭이다. 원천 `23.방치물(전동킥보드)`의 "
        "AIHub 공식 카탈로그 번호는 572이며 공식 자료명은 ‘이륜자동차 안전 위험 시설물 데이터’다. "
        "현재 공식 dataSetSn=183의 제목은 ‘드론 이동체 인지 영상(전방 고정)’이므로 이 학습 원천이 아니다. "
        "기존 학습 경로는 재현성을 위해 유지하고 외부 문서에는 AIHub 572로 표기한다.",
    )
    add_heading(document, "핵심 제한", 2)
    add_bullets(
        document,
        [
            "독립 test set이 없고 현재 수치는 내부 validation 결과다.",
            "intended val 28,747장 중 2,331장이 corrupt로 제외됐다.",
            "정량 평가와 backend unified runtime은 모두 imgsz 768이지만 전처리·threshold·NMS와 현장 입력 조건은 다르다.",
            "Android unified img768 13-class TFLite asset/runtime 계약은 적용됐지만 PyTorch 동등성·Device Field는 미검증이다.",
        ],
    )

    add_heading(document, "2. 모델 식별과 학습 설정", 1)
    add_table(
        document,
        ["항목", "값"],
        [
            ["원 학습 run", str(RUN_DIR.relative_to(ROOT))],
            ["학습 데이터", str(DATA_YAML.relative_to(ROOT))],
            ["checkpoint", str(CHECKPOINT.relative_to(ROOT))],
            ["SHA-256", CHECKPOINT_SHA256],
            ["모델/클래스", "YOLO26n detect / 13 classes"],
            ["입력/배치", "imgsz 768 / batch 8"],
            ["optimizer", "SGD, AMP, cosine LR"],
            ["반복/seed", "300 epoch / seed 0 / deterministic"],
            ["누적 기록시간", f"{last.time_s / 3600:.2f}시간"],
            ["Ultralytics", "8.4.48 (checkpoint metadata)"],
        ],
        font_size=8,
    )
    add_paragraph(
        document,
        "Ultralytics가 저장 후 optimizer를 제거한 checkpoint의 metadata epoch는 -1이다. 따라서 best epoch는 "
        "weight 내부 숫자가 아니라 원 run results.csv의 mAP50-95 최고점과 보존 경로로 확정했다.",
    )

    add_heading(document, "3. 데이터 출처", 1)
    add_heading(document, "3.1 공식 자료명과 내부 key", 2)
    add_table(
        document,
        ["공식 식별", "공식 자료명", "로컬 key", "사용 범위"],
        [[s["catalog"], s["official_name"], s["local_key"], s["note"]] for s in OFFICIAL_SOURCES],
        font_size=7.4,
    )
    add_paragraph(
        document,
        "AIHub 공식 페이지의 전체 구축량이 아니라 최종 materialized_manifest.csv에 실제 반영된 이미지와 "
        "bbox를 집계했다.",
    )
    add_table(
        document,
        ["출처", "train", "val", "전체 이미지", "bbox"],
        [
            [
                SOURCE_LABELS[source],
                f"{dataset.source_split_images[source, 'train']:,}",
                f"{dataset.source_split_images[source, 'val']:,}",
                f"{dataset.source_images[source]:,}",
                f"{dataset.source_boxes[source]:,}",
            ]
            for source in sorted(dataset.source_images, key=lambda value: dataset.source_images[value], reverse=True)
        ],
        font_size=7.8,
    )
    add_picture(document, ASSET_DIR / "04_dataset_sources.png", caption="그림 1. 최종 통합본 출처별 이미지와 bbox")

    add_heading(document, "3.2 실제 사용 component", 2)
    add_table(
        document,
        ["출처", "split", "사용 component", "이미지", "bbox"],
        [[a, b, c, f"{d:,}", f"{e:,}"] for a, b, c, d, e in SUBSETS],
        font_size=7.3,
    )
    add_bullets(
        document,
        [
            "AIHub 189 Depth_001~005는 YOLO 객체탐지 학습에 사용하지 않았다.",
            "AIHub 572 전동킥보드는 사람 승인 또는 bbox 수정 승인분 19,751개만 반영했다.",
            "AIHub 572 공식 규모 15,000장과 local 고유 원천 16,005장의 차이는 원인이 확정되지 않아 별도 한계로 기록한다.",
            "COCO는 일반 객체 7종의 이미지와 bbox만 매핑했다.",
            "train 167,759장, val 28,747장, 전체 196,506장/607,814 bbox이며 독립 test는 없다.",
        ],
    )

    add_heading(document, "4. 13개 탐지 클래스", 1)
    add_table(
        document,
        ["ID", "클래스", "주요 원천", "전체 이미지", "전체 bbox"],
        [
            [
                i,
                f"{CLASS_KO[i]} ({CLASS_NAMES[i]})",
                CLASS_SOURCES[i],
                f"{dataset.class_images[i]:,}",
                f"{dataset.class_boxes[i]:,}",
            ]
            for i in range(13)
        ],
        font_size=7.5,
    )
    add_picture(document, ASSET_DIR / "05_class_distribution.png", caption="그림 2. 클래스별 bbox 분포")

    add_heading(document, "5. 학습 경과", 1)
    add_picture(document, ASSET_DIR / "01_epoch_metrics.png", caption="그림 3. Epoch별 precision, recall, mAP")
    document.add_page_break()
    add_table(
        document,
        ["epoch", "누적시간(h)", "P", "R", "mAP50", "mAP50-95"],
        [
            [row.epoch, f"{row.time_s / 3600:.2f}", f"{row.precision:.5f}", f"{row.recall:.5f}", f"{row.map50:.5f}", f"{row.map5095:.5f}"]
            for row in epochs
            if row.epoch in {1, 50, 100, 150, 200, 250, 270, 300}
        ],
        font_size=8,
    )
    add_paragraph(
        document,
        f"mAP50와 mAP50-95 모두 epoch {best.epoch}에서 최고였다. epoch 300 mAP50-95는 "
        f"{last.map5095:.5f}로 최고점 {best.map5095:.5f}보다 낮아 epoch {best.epoch}의 best.pt를 후보로 선택했다.",
    )
    add_picture(document, ASSET_DIR / "02_train_val_losses.png", caption="그림 4. Train/validation loss 비교")
    add_paragraph(
        document,
        "후반 train loss 하락과 validation loss 격차는 과적합, source 분포 차이, 클래스 정의 충돌 또는 "
        "라벨 품질을 점검해야 하는 신호다. 그래프만으로 원인을 하나로 단정하지 않는다.",
    )

    add_heading(document, "6. Validation 성능", 1)
    add_table(
        document,
        ["평가 경계", "수치"],
        [
            ["intended val", f"{dataset.val_images:,}장 / {dataset.split_boxes['val']:,} bbox"],
            ["실제 평가", f"{int(overall['images']):,}장 / {int(overall['instances']):,} 객체"],
            ["corrupt 제외", f"{dataset.val_images - int(overall['images']):,}장"],
            ["전체 성능", f"P {best.precision:.5f} / R {best.recall:.5f} / mAP50 {best.map50:.5f} / mAP50-95 {best.map5095:.5f}"],
        ],
        font_size=8.5,
    )
    add_picture(
        document,
        ASSET_DIR / "03_class_performance.png",
        width=6.0,
        caption="그림 5. 클래스별 정량 성능",
    )
    add_paragraph(
        document,
        "판정은 안전 관점으로 정규화했다. mAP가 높더라도 precision 40% 미만은 ‘오탐 개선 필요’, "
        "보도 단차·불균일 보도는 낮은 recall 때문에 ‘핵심 개선 필요’로 표시한다. 원 CSV의 정상 "
        "점자블록 ‘양호’ 문구는 높은 recall만 강조하므로 종합 판정에 사용하지 않았다.",
    )
    add_table(
        document,
        ["ID", "클래스", "평가 객체", "P", "R", "mAP50", "mAP50-95", "판정"],
        [
            [
                m.class_id,
                m.class_ko,
                f"{m.instances:,}",
                pct(m.precision),
                pct(m.recall),
                pct(m.map50),
                pct(m.map5095),
                report_assessment(m),
            ]
            for m in metrics
        ],
        font_size=6.8,
    )
    add_picture(document, ASSET_DIR / "06_precision_recall_risk.png", caption="그림 6. 클래스별 Precision-Recall 위치")
    add_bullets(
        document,
        [
            f"강점: 전동킥보드 장애물 R {pct(metrics[12].recall)}, mAP50-95 {pct(metrics[12].map5095)}; 손상 점자블록 R {pct(metrics[8].recall)}.",
            f"미탐 위험: 보도 단차 R {pct(metrics[10].recall)}, 불균일·파손 보도 R {pct(metrics[11].recall)}.",
            f"오탐 위험: 정상 점자블록 P {pct(metrics[7].precision)}, 횡단보도 P {pct(metrics[9].precision)}.",
        ],
    )

    add_heading(document, "7. 외부 시험과 runtime 격차", 1)
    add_paragraph(
        document,
        f"AIHub 189 Depth RGB {depth.total_images:,}장에서는 {depth.images_with_detection:,}장에 "
        f"{depth.total_boxes:,}개 예측이 발생했다. 주요 예측은 정상 점자블록 "
        f"{depth.class_boxes['normal_tactile_block']:,}, 횡단보도 {depth.class_boxes['crosswalk']:,}, "
        f"보도 단차 {depth.class_boxes['curb_step']:,}, 불균일·파손 보도 "
        f"{depth.class_boxes['uneven_sidewalk']:,}개다. 이 데이터에는 정답 bbox가 없으므로 "
        "정확도가 아니라 검출 발생량으로만 해석한다.",
    )
    add_bullets(
        document,
        [
            "정량 validation과 backend unified runtime은 모두 imgsz 768이다.",
            "입력 크기가 같아도 validation 성능을 현장 backend 성능으로 직접 인용하면 안 된다.",
            "모바일 흔들림·압축·지연·야간·역광·발열·배터리는 평가하지 않았다.",
            "Android unified img768 13-class TFLite는 적용됐지만 PyTorch box/class 동등성·실기기 검증이 필요하다.",
        ],
    )

    document.add_page_break()
    add_heading(document, "8. 개선 로드맵", 1)
    add_heading(document, "P0. 데이터와 평가 무결성", 2)
    add_numbered(
        document,
        [
            "truncated JPEG 2,331장을 복구하거나 사전 제외 manifest로 고정한다.",
            "원천 zip·연속 프레임·장소 단위 group split 누수를 재감사한다.",
            "낮/밤·역광·우천·흔들림·가림을 포함한 독립 test set을 만든다.",
            "레거시 aihub183 별칭과 공식 AIHub 572 대응표를 provenance에 고정한다.",
        ],
    )
    add_heading(document, "P0. 핵심 약점 클래스", 2)
    add_numbered(
        document,
        [
            "curb_step의 연석·문턱·계단·경사형 경계 포함 기준을 서비스 위험 정의에 맞춰 좁힌다.",
            "uneven_sidewalk의 평탄성·포장·보도블록 손상 영역 bbox 정책을 통일한다.",
            "두 클래스의 중복·상충 라벨을 source별 층화 표본으로 재검수한다.",
            "새 데이터 추가 전에 승인 라벨 일관성과 작은/과대 bbox 비율을 측정한다.",
        ],
    )
    add_heading(document, "P1. 오탐·runtime·배포", 2)
    add_numbered(
        document,
        [
            "정상 점자블록·횡단보도 hard negative와 손상/정상 confusion을 별도 관리한다.",
            "동일 checkpoint의 backend img768 전처리·threshold·NMS 회귀를 기록한다.",
            "적용된 13-class TFLite의 PyTorch 대비 box/class 동등성 회귀를 수행한다.",
            "실폰 FPS, P95 latency, 메모리, 발열, 배터리와 현장 음성 안내 빈도를 측정한다.",
        ],
    )

    add_heading(document, "9. 다음 모델 승인 조건", 1)
    add_bullets(
        document,
        [
            "읽을 수 없는 평가 이미지 0장 또는 사전 확정 제외 manifest",
            "독립 test의 source/group 분리와 라벨 검수",
            "imgsz 768 backend 전처리·threshold·NMS 조건 회귀표",
            "팀이 사전 합의한 보도 단차·불균일 보도 recall 하한 충족",
            "모바일 정지·보행·야간 E2E 및 PyTorch/배포 artifact 동등성 검증",
        ],
    )
    add_paragraph(
        document,
        "승인 수치 하한은 이 보고서가 임의로 만들지 않는다. 안전 요구사항과 실패 비용을 먼저 합의하고 "
        "test set을 보지 않은 상태에서 고정해야 한다.",
    )

    add_heading(document, "10. 근거와 재현", 1)
    add_table(
        document,
        ["근거", "역할"],
        [
            [str(RESULTS_CSV.relative_to(ROOT)), "300 epoch 원 수치"],
            [str(ARGS_YAML.relative_to(ROOT)), "학습 인자"],
            [str(MATERIALIZED_MANIFEST.relative_to(ROOT)), "실제 데이터 provenance/count"],
            [str(PER_CLASS_CSV.relative_to(ROOT)), "클래스별 validation 수치"],
            [str(VAL_LOG.relative_to(ROOT)), "corrupt 제외와 평가 집계"],
            [str(DEPTH500_CSV.relative_to(ROOT)), "정답 없는 Depth 500장 추론 발생량"],
        ],
        font_size=7.5,
    )
    add_paragraph(document, ".venv/bin/python scripts/build_latest_model_report_20260710.py")
    add_paragraph(document, ".venv/bin/python scripts/build_latest_model_report_20260710.py --validate-only")
    add_heading(document, "공식 출처", 2)
    for source in OFFICIAL_SOURCES:
        add_paragraph(document, f"{source['catalog']} {source['official_name']}: {source['url']}")
    add_paragraph(document, f"현재 AIHub dataSetSn=183 비교 페이지(사용 원천 아님): {AIHUB_183_CURRENT_URL}")
    add_heading(document, "보고서 한계", 2)
    add_paragraph(
        document,
        "보존 artifact를 재집계했으며 학습을 다시 실행하지 않았다. 독립 test가 없고 corrupt 제외 전 val 전체 "
        "성능을 복원하지 않았다. 공식 페이지는 자료명·번호 확인에만 쓰고 프로젝트 반영량은 local manifest에서 계산했다.",
    )

    properties = document.core_properties
    properties.title = "WalkSafe 최신 13클래스 객체탐지 모델 종합보고서"
    properties.subject = "YOLO26n img768 epoch300 후보의 데이터, 학습, 평가, 개선 로드맵"
    properties.author = "WalkSafe project"
    properties.last_modified_by = "WalkSafe project"
    properties.keywords = "WalkSafe, YOLO26n, object detection, model report"
    document.save(DOCX_PATH)


def write_asset_readme(paths: Sequence[Path]) -> None:
    lines = [
        "# 최신 모델 보고서 시각자료",
        "",
        "모든 그래프는 `scripts/build_latest_model_report_20260710.py`가 local 원 근거에서 생성합니다.",
        "",
        "| 파일 | SHA-256 | 입력 근거 |",
        "| --- | --- | --- |",
    ]
    for path in paths:
        if "epoch" in path.name or "loss" in path.name:
            evidence = str(RESULTS_CSV.relative_to(ROOT))
        elif "class_performance" in path.name or "precision_recall" in path.name:
            evidence = str(PER_CLASS_CSV.relative_to(ROOT))
        else:
            evidence = str(MATERIALIZED_MANIFEST.relative_to(ROOT))
        lines.append(f"| `{path.name}` | `{sha256(path)}` | `{evidence}` |")
    lines.extend(
        [
            "",
            "그래프에는 원본 사진을 넣지 않아 얼굴·차량번호판 등 시각 개인정보가 포함되지 않습니다.",
            "",
        ]
    )
    (ASSET_DIR / "README.md").write_text("\n".join(lines), encoding="utf-8")


def validate_inputs(
    epochs: Sequence[EpochRow],
    overall: dict[str, float | int],
    metrics: Sequence[ClassMetric],
    dataset: DatasetSummary,
    depth: DepthSummary,
) -> None:
    required = [
        RESULTS_CSV,
        ARGS_YAML,
        CHECKPOINT,
        *CANDIDATE_CHECKPOINTS,
        DATA_YAML,
        MATERIALIZED_MANIFEST,
        PER_CLASS_CSV,
        VAL_LOG,
        DEPTH500_CSV,
    ]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise AssertionError(f"missing input files: {missing}")
    assert sha256(CHECKPOINT) == CHECKPOINT_SHA256, "checkpoint hash changed"
    assert all(sha256(path) == CHECKPOINT_SHA256 for path in CANDIDATE_CHECKPOINTS)
    assert len(epochs) == 300 and epochs[0].epoch == 1 and epochs[-1].epoch == 300
    best = max(epochs, key=lambda row: row.map5095)
    assert best.epoch == 270
    assert abs(best.precision - 0.62087) < 1e-7
    assert abs(best.recall - 0.56716) < 1e-7
    assert abs(best.map50 - 0.55403) < 1e-7
    assert abs(best.map5095 - 0.41651) < 1e-7
    assert len(metrics) == 13 and [metric.class_id for metric in metrics] == list(range(13))
    assert dataset.total_images == 196_506
    assert dataset.total_boxes == 607_814
    assert dataset.train_images == 167_759
    assert dataset.val_images == 28_747
    assert int(overall["images"]) == 26_416
    assert int(overall["instances"]) == 69_467
    log_text = VAL_LOG.read_text(encoding="utf-8", errors="replace")
    assert log_text.count("ignoring corrupt image/label") == 2_331
    assert "26416 images, 0 backgrounds, 2331 corrupt" in log_text
    assert "DEFAULT_UNIFIED_WALKSAFE_INFERENCE_IMGSZ = 768" in (
        ROOT / "backend/app/services/detect_v2.py"
    ).read_text(encoding="utf-8")
    assert depth.total_images == 500
    assert depth.images_with_detection == 242
    assert depth.total_boxes == 390
    assert depth.class_boxes["normal_tactile_block"] == 176
    assert depth.class_boxes["crosswalk"] == 110
    assert depth.class_boxes["curb_step"] == 74
    assert depth.class_boxes["uneven_sidewalk"] == 15


def validate_outputs() -> None:
    required = [
        MARKDOWN_PATH,
        DOCX_PATH,
        OUTPUT_DIR / "epoch_milestones.csv",
        OUTPUT_DIR / "dataset_source_summary.csv",
        OUTPUT_DIR / "class_distribution_and_metrics.csv",
        ASSET_DIR / "README.md",
    ] + [ASSET_DIR / f"{index:02d}_{name}.png" for index, name in enumerate(
        [
            "epoch_metrics",
            "train_val_losses",
            "class_performance",
            "dataset_sources",
            "class_distribution",
            "precision_recall_risk",
        ],
        start=1,
    )]
    missing = [str(path) for path in required if not path.is_file() or path.stat().st_size == 0]
    if missing:
        raise AssertionError(f"missing/empty outputs: {missing}")
    markdown = MARKDOWN_PATH.read_text(encoding="utf-8")
    for needle in [
        "AIHub 572",
        "레거시 내부 별칭",
        "15,000장",
        "16,005",
        "196,506",
        "607,814",
        "2,331",
        "imgsz 768",
        CHECKPOINT_SHA256,
        "독립 test",
        "오탐 개선 필요",
    ]:
        assert needle in markdown, f"missing report statement: {needle}"
    with zipfile.ZipFile(DOCX_PATH) as archive:
        bad_member = archive.testzip()
        assert bad_member is None, f"broken DOCX member: {bad_member}"
        names = archive.namelist()
        assert "word/document.xml" in names
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                ElementTree.fromstring(archive.read(name))
        document_xml = archive.read("word/document.xml").decode("utf-8")
        for needle in ["AIHub 572", "2,331", "mAP50-95", "PARTIAL"]:
            assert needle in document_xml
        all_xml = "\n".join(
            archive.read(name).decode("utf-8", errors="replace")
            for name in names
            if name.endswith(".xml") or name.endswith(".rels")
        )
        for forbidden in ["@gmail.com", "@naver.com", "/home/ddobagi", "2026. 00. 00"]:
            assert forbidden not in all_xml, f"forbidden DOCX content: {forbidden}"


def main() -> int:
    args = parse_args()
    epochs = read_epochs()
    overall, metrics = read_class_metrics()
    dataset = read_dataset_summary()
    depth = read_depth_summary()
    validate_inputs(epochs, overall, metrics, dataset, depth)
    if not args.validate_only:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        plots = build_plots(epochs, metrics, dataset)
        write_evidence_csvs(epochs, metrics, dataset)
        MARKDOWN_PATH.write_text(build_markdown(epochs, overall, metrics, dataset, depth), encoding="utf-8")
        build_docx(epochs, overall, metrics, dataset, depth)
        write_asset_readme(plots)
    validate_outputs()
    print(
        "latest model report PASS: "
        f"{dataset.total_images:,} images, {dataset.total_boxes:,} boxes, "
        f"best epoch {max(epochs, key=lambda row: row.map5095).epoch}, 6 plots, DOCX valid"
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (AssertionError, RuntimeError, OSError, ValueError) as exc:
        print(f"latest model report FAIL: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc

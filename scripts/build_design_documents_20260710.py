#!/usr/bin/env python3
"""Build the official readable Word pack from the reviewed Markdown sources."""

from __future__ import annotations

import sys

if __name__ == "__main__":
    _verified_submission_commit = getattr(
        sys.modules.get("__main__"), "_walksafe_submission_source_commit", None
    )
    _startup_flags = (
        sys.flags.isolated,
        sys.flags.no_site,
        sys.flags.ignore_environment,
        int(getattr(sys.flags, "safe_path", False)),
        sys.flags.dont_write_bytecode,
    )
    if (
        any(value != 1 for value in _startup_flags)
        or "site" in sys.modules
        or not isinstance(_verified_submission_commit, str)
        or len(_verified_submission_commit) != 40
        or any(character not in "0123456789abcdef" for character in _verified_submission_commit)
    ):
        raise SystemExit(
            "WalkSafe submission CLI requires Python -I -S -B through the submission runner"
        )

import argparse
import hashlib
import json
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from submission_build_io import (
    atomic_output_path,
    canonical_android_device_evidence,
    submission_build_lock,
    verification_snapshot_policy_boundary,
)
from submission_manifest_policy import (
    ALL_SUBMISSION_GENERATED_PATHS,
    DESIGN_BUILD_INPUT_PATHS,
    DESIGN_OUTPUT_NAMES,
    build_tool_provenance,
    require_exact_file_set,
    require_no_unexpected_files,
    require_clean_source_revision,
    submission_toolchain_attestation,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs/submission/design_documents"
ASSET_DIR = ROOT / "docs/submission/form_materials/assets"
FACTS_PATH = ROOT / "docs/submission/form_materials/09_제출_사실_기준.json"
FONT_KO = "Malgun Gothic"
FONT_CODE = "D2Coding"

with FACTS_PATH.open(encoding="utf-8") as handle:
    FACTS = json.load(handle)
if FACTS.get("schema_version") != "walksafe.submission_facts.v2":
    raise SystemExit(f"unexpected submission facts schema: {FACTS.get('schema_version')!r}")
BUILD_DATE = FACTS.get("as_of_date")
if not isinstance(BUILD_DATE, str):
    raise SystemExit("canonical facts must declare as_of_date")
try:
    FIXED_TIME = datetime.strptime(BUILD_DATE, "%Y-%m-%d").replace(tzinfo=timezone.utc)
except ValueError as exc:
    raise SystemExit("canonical facts as_of_date must use YYYY-MM-DD") from exc


def android_device_evidence_labels(snapshot: object) -> tuple[str, str]:
    try:
        _, _, current_label, historical_label = canonical_android_device_evidence(snapshot)
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc
    return current_label, historical_label


def verification_snapshot_contract(snapshot: object) -> str:
    if not isinstance(snapshot, dict):
        raise SystemExit("canonical facts must contain a verification snapshot")
    try:
        return verification_snapshot_policy_boundary(
            snapshot.get("status"),
            snapshot.get("executed_at"),
            snapshot.get("policy"),
        )
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc


VERIFICATION_SCOPE_BOUNDARY = verification_snapshot_contract(
    FACTS.get("verification_snapshot")
)
ANDROID_DEVICE_CURRENT_LABEL, ANDROID_DEVICE_HISTORICAL_LABEL = (
    android_device_evidence_labels(FACTS.get("verification_snapshot"))
)
ALLOWED_FIXED_TEST_COUNTS = {
    int(value)
    for value in re.findall(
        r"(?<!\d)(\d{1,4})\s*(?:passed|tests?)(?![A-Za-z])",
        "\n".join(
            str(run.get("details", ""))
            for run in FACTS.get("verification_snapshot", {}).get("runs", [])
            if isinstance(run, dict)
        ),
        flags=re.IGNORECASE,
    )
}


@dataclass(frozen=True)
class DocumentSpec:
    output_name: str
    title: str
    source: Path
    landscape: bool = False
    figures: tuple[tuple[str, str, float], ...] = ()
    page_break_before_headings: tuple[str, ...] = ()


SPECS = (
    DocumentSpec(
        "01_요구사항_정의서.docx",
        "요구사항 정의서",
        ROOT / "docs/submission/deliverables/요구사항_정의서.md",
        landscape=True,
        figures=(("problem_solution_map.png", "사용자 문제·설계 대응·검증할 효과의 연결", 9.5),),
    ),
    DocumentSpec(
        "02_유스케이스_정의서.docx",
        "유스케이스 정의서",
        ROOT / "docs/submission/deliverables/유스케이스_정의서.md",
        figures=(
            ("use_case_swimlane.png", "보행 사용자·시스템·관리자의 핵심 유스케이스와 책임 분리", 8.0),
            ("risk_processing_flow.png", "보행 위험 처리 흐름 — 객체별 3 frame·700ms 안정화, 보조 투영 ROI, WebXR 표시 전용 경계", 8.0),
            ("navigation_state_flow.png", "TMAP 전역 경로·Web/미지원 Android 비계량 advisory·Android metric steering·복귀 흐름", 8.0),
        ),
        page_break_before_headings=("5. 상태 요약",),
    ),
    DocumentSpec(
        "03_요구사항_기능_추적표.docx",
        "요구사항·기능 추적표",
        ROOT / "docs/submission/drafts/요구사항_추적표.md",
        landscape=True,
        figures=(("evidence_results.png", "요구사항에 연결되는 계층별 자동검증 결과와 현장 검증 경계", 9.5),),
    ),
    DocumentSpec(
        "04_서비스_구성도_및_흐름도.docx",
        "서비스 구성도 및 흐름도",
        ROOT / "docs/submission/deliverables/서비스_구성도_및_흐름도.md",
        figures=(
            ("system_architecture.png", "전체 시스템 구성 — Web/PWA 주 사용자 앱, Android 연구·ARCore 미지원 제한 모드, PostGIS 5개 application table과 사람의 외부 제출", 8.0),
            ("risk_processing_flow.png", "카메라 입력부터 객체별 위험 후보 선별까지의 처리 흐름", 8.0),
            ("report_csv_flow.png", "신고 검수 및 기관 수동 후속 흐름 — primary ID·duplicate_count 유지와 다른 후보 ID의 admin 경계", 8.0),
            ("navigation_state_flow.png", "TMAP 전역 경로와 platform별 비계량 advisory·metric steering의 조건부 결합", 8.0),
        ),
    ),
    DocumentSpec(
        "05_화면설계서_UIUX_정의서.docx",
        "화면설계서 및 UI/UX 정의서",
        ROOT / "docs/submission/deliverables/화면설계서_UIUX_정의서.md",
        figures=(
            ("ui_state_map.png", "보행 보조의 정상·위험·신고·길안내·오류 상태와 접근성 피드백", 8.0),
            ("ui_screen_storyboard.png", "SC-00~SC-11 화면별 정보 구조·상태·복구 행동 스토리보드", 8.0),
            ("web_main_mobile.png", "2026-07-11 Web/PWA local documentation fixture — Web frame/report 개별 동의·세션 UI 확인용이며 실모델·실폰 근거 아님", 3.0),
            ("web_main_desktop.png", "2026-07-11 Web/PWA local documentation fixture — Web frame/report 개별 동의·보행 UI 확인용이며 실모델·실폰 근거 아님", 8.0),
            ("admin_desktop.png", "2026-07-11 관리자 local documentation fixture — 정확 격자·기관·공개용 export 분리 UI", 8.0),
        ),
    ),
    DocumentSpec(
        "06_엔티티관계도_테이블정의서.docx",
        "엔티티 관계도 및 테이블 정의서",
        ROOT / "docs/submission/deliverables/엔티티관계도_테이블정의서.md",
        landscape=True,
        figures=(("reports_erd.png", "reports·report_status_audits·report_export_audits·report_read_audits·actor_rate_limit_events(60초 non-audit)와 DB 밖 durable journal/image 경계", 9.5),),
    ),
    DocumentSpec(
        "07_기능처리도_알고리즘명세서.docx",
        "기능 처리도 및 알고리즘 명세서",
        ROOT / "docs/submission/deliverables/기능처리도_알고리즘명세서.md",
        landscape=True,
        figures=(
            ("risk_processing_flow.png", "위험 후보 판정 및 안내 처리 — 객체별 안정화·경로 조건·WebXR 위험 미연결", 9.5),
            ("risk_timeline.png", "한 객체가 안내되기까지의 frame·시간·경로 조건", 9.5),
            ("navigation_state_flow.png", "TMAP 전역 경로·Web/미지원 Android 비계량 advisory·Android metric steering·복귀 알고리즘", 9.5),
            ("report_csv_flow.png", "손상 점자블록 신고와 운영 처리 — primary ID 계약·admin lookup·named-review agency 분리", 9.5),
        ),
    ),
    DocumentSpec(
        "08_프로그램목록_핵심소스코드_개발환경.docx",
        "프로그램 목록·핵심 소스코드·개발환경",
        ROOT / "docs/submission/deliverables/프로그램목록_핵심소스코드_개발환경.md",
        landscape=True,
        figures=(
            ("system_architecture.png", "프로그램 구성요소 배치 — Web/PWA BFF·named actor·PostGIS 5개 application table·Release 대기 gate", 9.5),
            ("evidence_results.png", "서비스별 자동검증·빌드·실기기 TFLite invoke 결과와 범위", 9.5),
            ("scope_change_map.png", "원안 목표·확인한 제약·현재 구현 결정·남은 증거의 연결", 9.5),
        ),
        page_break_before_headings=("6. 핵심 Gap",),
    ),
)

IMPLEMENTATION_BINDINGS = (
    ROOT / "apps/web/app/_walksafe/hooks/useDetectionV2.ts",
    ROOT / "apps/web/app/_walksafe/hooks/useRiskFeedback.ts",
    ROOT / "apps/web/app/_walksafe/hooks/useNavigationGuidance.ts",
    ROOT / "apps/web/app/_walksafe/hooks/useVoiceCommands.ts",
    ROOT / "apps/web/app/_walksafe/voice-intent-executor.ts",
    ROOT / "apps/web/app/_walksafe/tactile-route-policy.ts",
    ROOT / "apps/web/app/_walksafe/absolute-heading.ts",
    ROOT / "apps/web/app/_walksafe/feedback.ts",
    ROOT / "apps/web/app/_walksafe/camera-policy.ts",
    ROOT / "apps/web/app/_walksafe/detection-availability.ts",
    ROOT / "apps/web/app/_walksafe/motion-projection.ts",
    ROOT / "apps/web/app/_walksafe/nonmetric-hazard-advisory.ts",
    ROOT / "apps/web/app/_walksafe/navigation-destination.ts",
    ROOT / "apps/web/app/api/_gateway-auth.ts",
    ROOT / "apps/web/app/api/_backend.ts",
    ROOT / "apps/web/lib/auto-report-v2.ts",
    ROOT / "apps/web/lib/detect-api-v2.ts",
    ROOT / "apps/web/lib/navigation-api.ts",
    ROOT / "backend/app/api/detect.py",
    ROOT / "backend/app/api/health.py",
    ROOT / "backend/app/api/navigation.py",
    ROOT / "backend/app/api/reports.py",
    ROOT / "backend/app/main.py",
    ROOT / "backend/app/config.py",
    ROOT / "backend/app/field_test_security.py",
    ROOT / "backend/app/uploads.py",
    ROOT / "backend/app/services/inference_process.py",
    ROOT / "backend/app/services/report_policy.py",
    ROOT / "backend/app/services/report_read_audit.py",
    ROOT / "backend/app/services/report_storage.py",
    ROOT / "backend/app/services/actor_rate_limit.py",
    ROOT / "backend/app/services/tmap_pedestrian.py",
    ROOT / "backend/app/services/yolo_inference_adapter.py",
    ROOT / "backend/app/openapi_contract.py",
    ROOT / "backend/alembic/versions/202607130003_report_read_audits.py",
    ROOT / "backend/alembic/versions/202607130004_shared_actor_rate_limits.py",
    ROOT / "backend/alembic/versions/202607160001_report_duplicate_check_audit.py",
    ROOT / "contracts/walksafe.openapi.json",
    ROOT / "backend/tests/conftest.py",
    ROOT / "apps/android/app/build.gradle.kts",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/MainActivity.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/device/AndroidLocalTactileCapability.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/feedback/AndroidNonMetricObstacleAdvisoryPolicy.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/navigation/AndroidVoiceCommand.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/navigation/BackendWalkingRouteClient.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/navigation/RouteNavigator.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/navigation/TactileRoutePolicy.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/report/ReportPrivacyConsentSession.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/inference/TfliteAndroidFrameDetector.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/feedback/AndroidFeedbackActuator.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/feedback/NavigationSpeechDelivery.kt",
    ROOT / "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/feedback/UtteranceCallbackRegistry.kt",
    ROOT / "apps/android/app/src/main/assets/model-config/two_model_runtime.json",
    ROOT / "voice/intents.py",
    ROOT / "voice/server.py",
    ROOT / "voice/tts.py",
    ROOT / "apps/web/app/_walksafe/server-v2-privacy.ts",
    ROOT / "scripts/run_walksafe_web_single_instance_20260713.py",
    ROOT / "scripts/run_walksafe_test_layers_20260711.sh",
    ROOT / "scripts/check_walksafe_test_database_20260713.py",
    ROOT / "scripts/check_android_apk_model_asset_20260713.py",
    ROOT / "scripts/check_web_runtime_trace_scope_20260713.py",
    ROOT / "scripts/check_walksafe_backup_source_20260713.py",
    ROOT / "scripts/walksafe_environment_identity.py",
    ROOT / "scripts/backup_walksafe_data_20260711.sh",
    ROOT / "scripts/restore_walksafe_backup_drill_20260711.sh",
    ROOT / "scripts/check_walksafe_release_evidence_20260711.py",
    ROOT / "scripts/walksafe_external_check_receipt.py",
    ROOT / ".github/workflows/quality.yml",
)

BUILD_INPUTS = tuple(ROOT / relative for relative in sorted(DESIGN_BUILD_INPUT_PATHS))


FORBIDDEN_PATTERNS = {
    "local home path": re.compile(r"/home/", re.IGNORECASE),
    "local account": re.compile(r"ddobagi", re.IGNORECASE),
    "email": re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}"),
    "placeholder date": re.compile(r"2026\.\s*00\.\s*00"),
    "sample instruction": re.compile(r"예시\s*내용을|샘플\s*텍스트|홍길동"),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def source_revision() -> dict[str, object]:
    return require_clean_source_revision(ROOT, ALL_SUBMISSION_GENERATED_PATHS)


def build_tools() -> dict[str, object]:
    return build_tool_provenance(("Pillow", "python-docx"))


def implementation_records() -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for implementation in IMPLEMENTATION_BINDINGS:
        if implementation.is_symlink() or not implementation.is_file():
            raise FileNotFoundError(f"missing implementation binding: {implementation}")
        records.append(
            {
                "file": implementation.relative_to(ROOT).as_posix(),
                "sha256": sha256(implementation),
                "bytes": implementation.stat().st_size,
            }
        )
    return records


def build_input_records() -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for path in BUILD_INPUTS:
        if path.is_symlink() or not path.is_file():
            raise FileNotFoundError(f"missing build input: {path}")
        records.append(
            {
                "file": path.relative_to(ROOT).as_posix(),
                "sha256": sha256(path),
                "bytes": path.stat().st_size,
            }
        )
    return records


def figure_asset_names() -> tuple[str, ...]:
    return tuple(sorted({file_name for spec in SPECS for file_name, _, _ in spec.figures}))


def asset_records() -> list[dict[str, object]]:
    records = []
    for file_name in figure_asset_names():
        path = ASSET_DIR / file_name
        if not path.is_file():
            raise FileNotFoundError(f"missing figure: {path}")
        records.append(
            {
                "file": str(path.relative_to(ROOT)),
                "sha256": sha256(path),
                "bytes": path.stat().st_size,
            }
        )
    return records


def set_run_font(run, name: str = FONT_KO, size: float | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def configure_document(doc: DocumentType, spec: DocumentSpec) -> None:
    section = doc.sections[0]
    if spec.landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(1.35)
        section.right_margin = Cm(1.35)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(1.75)
        section.right_margin = Cm(1.75)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_KO
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(4)

    heading_specs = {
        "Title": (24, "17324D", 12, 8),
        "Subtitle": (11, "4B5D6B", 4, 4),
        "Heading 1": (16, "17324D", 14, 6),
        "Heading 2": (13, "205B57", 11, 4),
        "Heading 3": (11, "17324D", 8, 3),
        "Heading 4": (10, "37474F", 6, 2),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = FONT_KO
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KO)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section_item in doc.sections:
        header = section_item.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("WalkSafe Assist | 상세 설계 산출물")
        set_run_font(run, size=8)
        run.font.color.rgb = RGBColor.from_string("667780")

        footer = section_item.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"{spec.title}  |  {BUILD_DATE}  |  ")
        set_run_font(run, size=8)
        run.font.color.rgb = RGBColor.from_string("667780")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)

    props = doc.core_properties
    props.title = spec.title
    props.subject = "2026 한이음 드림업 WalkSafe Assist 상세 설계 산출물"
    props.author = "WalkSafe Assist"
    props.last_modified_by = "WalkSafe Assist"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.keywords = "WalkSafe Assist, Web/PWA, 설계 산출물, PARTIAL"
    props.comments = "Markdown 근거 문서에서 재현 가능하게 생성된 공식 열람본"


def add_inline_markdown(paragraph, text: str, *, table: bool = False) -> None:
    token_pattern = re.compile(
        r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*]+\*(?!\*))"
    )
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=8.3 if table else None)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, FONT_CODE, 8.0 if table else 8.5)
            run.font.color.rgb = RGBColor.from_string("245A73")
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=8.3 if table else None)
            run.bold = True
        elif token.startswith("["):
            label, target = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label} ({target})")
            set_run_font(run, size=8.3 if table else None)
            run.font.color.rgb = RGBColor.from_string("205B57")
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=8.3 if table else None)
            run.italic = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=8.3 if table else None)


def split_table_row(line: str) -> list[str]:
    content = line.strip()
    if content.startswith("|"):
        content = content[1:]
    if content.endswith("|"):
        content = content[:-1]
    values: list[str] = []
    current: list[str] = []
    escaped = False
    for char in content:
        if escaped:
            current.append(char)
            escaped = False
        elif char == "\\":
            escaped = True
        elif char == "|":
            values.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    values.append("".join(current).strip())
    return values


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_table(doc: DocumentType, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])

    for row_index, (word_row, values) in enumerate(zip(table.rows, normalized)):
        if row_index > 0 and sum(len(value) for value in values) < 450:
            prevent_row_split(word_row)
        for cell, value in zip(word_row.cells, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if row_index == 0:
                set_cell_shading(cell, "17324D")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EEF4F3")
            add_inline_markdown(paragraph, value, table=True)
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def clean_mermaid_text(value: str) -> str:
    text = re.sub(r"<br\s*/?>|\\n", " / ", value, flags=re.IGNORECASE)
    text = text.replace("`", "").replace('"', "").strip()
    text = re.sub(r"\s+", " ", text)
    return text.strip(" [](){}")


def mermaid_aliases(lines: list[str]) -> dict[str, str]:
    aliases: dict[str, str] = {}
    for raw in lines:
        line = raw.strip()
        actor = re.match(r"(?:actor|participant)\s+(\w+)\s+as\s+(.+)$", line)
        if actor:
            aliases[actor.group(1)] = clean_mermaid_text(actor.group(2))
        for pattern in (
            r"\b([A-Za-z][A-Za-z0-9_]*)\s*\[\((.*?)\)\]",
            r"\b([A-Za-z][A-Za-z0-9_]*)\s*\(\((.*?)\)\)",
            r"\b([A-Za-z][A-Za-z0-9_]*)\s*\{(.*?)\}",
            r"\b([A-Za-z][A-Za-z0-9_]*)\s*\[(.*?)\]",
        ):
            for match in re.finditer(pattern, line):
                aliases.setdefault(match.group(1), clean_mermaid_text(match.group(2)))
    return aliases


def strip_mermaid_node_declarations(line: str) -> str:
    for pattern in (
        r"\b([A-Za-z][A-Za-z0-9_]*)\s*\[\((.*?)\)\]",
        r"\b([A-Za-z][A-Za-z0-9_]*)\s*\(\((.*?)\)\)",
        r"\b([A-Za-z][A-Za-z0-9_]*)\s*\{(.*?)\}",
        r"\b([A-Za-z][A-Za-z0-9_]*)\s*\[(.*?)\]",
    ):
        line = re.sub(pattern, lambda match: match.group(1), line)
    return line


def mermaid_summary_rows(lines: list[str]) -> list[list[str]]:
    aliases = mermaid_aliases(lines)
    meaningful = [line.strip() for line in lines if line.strip()]
    is_sequence = bool(meaningful and meaningful[0].startswith("sequenceDiagram"))
    rows: list[list[str]] = [["순서", "출발", "조건·처리", "도착"]]

    if is_sequence:
        branch = ""
        for line in meaningful[1:]:
            if line.startswith("alt "):
                branch = clean_mermaid_text(line[4:])
                continue
            if line.startswith("else"):
                branch = clean_mermaid_text(line[4:])
                continue
            if line == "end":
                branch = ""
                continue
            note = re.match(r"Note\s+over\s+([^:]+):\s*(.+)$", line, flags=re.IGNORECASE)
            if note:
                message = clean_mermaid_text(note.group(2))
                if branch:
                    message = f"[{branch}] {message}"
                note_actors = [
                    aliases.get(actor.strip(), actor.strip())
                    for actor in note.group(1).split(",")
                ]
                rows.append([str(len(rows)), "·".join(note_actors), message, "참고"])
                continue
            edge = re.match(r"(\w+)\s*-{1,2}>>\s*(\w+)\s*:\s*(.+)$", line)
            if not edge:
                continue
            message = clean_mermaid_text(edge.group(3))
            if branch:
                message = f"[{branch}] {message}"
            rows.append(
                [
                    str(len(rows)),
                    aliases.get(edge.group(1), edge.group(1)),
                    message,
                    aliases.get(edge.group(2), edge.group(2)),
                ]
            )
        return rows

    arrow_pattern = re.compile(r"(-->|\-\.\s*(.*?)\s*\.->)(?:\|([^|]+)\|)?")
    for raw in meaningful[1:]:
        if raw.startswith(("classDef ", "class ", "subgraph ", "end")):
            continue
        line = strip_mermaid_node_declarations(raw)
        matches = list(arrow_pattern.finditer(line))
        for index, match in enumerate(matches):
            left_start = matches[index - 1].end() if index else 0
            right_end = matches[index + 1].start() if index + 1 < len(matches) else len(line)
            left_ids = re.findall(r"\b[A-Za-z][A-Za-z0-9_]*\b", line[left_start : match.start()])
            right_ids = re.findall(r"\b[A-Za-z][A-Za-z0-9_]*\b", line[match.end() : right_end])
            if not left_ids or not right_ids:
                continue
            left_id, right_id = left_ids[-1], right_ids[0]
            label = clean_mermaid_text(match.group(3) or match.group(2) or "다음 단계")
            rows.append(
                [
                    str(len(rows)),
                    aliases.get(left_id, left_id),
                    label,
                    aliases.get(right_id, right_id),
                ]
            )
    return rows


def add_mermaid_summary(doc: DocumentType, lines: list[str]) -> None:
    rows = mermaid_summary_rows(lines)
    add_callout(
        doc,
        "편집 가능한 Mermaid 정의는 근거 Markdown에 보존한다. "
        f"공식 열람본은 아래 {max(len(rows) - 1, 0)}개 연결을 단계표로 풀어 검토·인쇄 가독성을 확보했다.",
    )
    if len(rows) > 1:
        add_table(doc, rows)


def add_code_paragraph(doc: DocumentType, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.25)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F5F6")
    p_pr.append(shading)
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, FONT_CODE, 8.0)
    run.font.color.rgb = RGBColor.from_string("263238")


def add_code_block(doc: DocumentType, lines: Iterable[str], language: str) -> None:
    materialized = list(lines)
    if language.lower() == "mermaid":
        add_mermaid_summary(doc, materialized)
        return
    if len(materialized) <= 28:
        add_code_paragraph(doc, materialized)
        return

    chunks = [materialized[index : index + 24] for index in range(0, len(materialized), 24)]
    for index, chunk in enumerate(chunks, start=1):
        marker = doc.add_paragraph()
        if index > 1:
            marker.paragraph_format.page_break_before = True
        marker.paragraph_format.keep_with_next = True
        marker.paragraph_format.space_after = Pt(4)
        run = marker.add_run(f"핵심 코드 발췌 {index}/{len(chunks)}" + (" · 계속" if index > 1 else ""))
        set_run_font(run, size=9.5)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("205B57")
        add_code_paragraph(doc, chunk)


def add_callout(doc: DocumentType, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "E7F0EE")
    p_pr.append(shading)
    add_inline_markdown(paragraph, text)


def add_cover(doc: DocumentType, spec: DocumentSpec) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(44)

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = label.add_run("2026 한이음 드림업 | W A L K S A F E  A S S I S T")
    set_run_font(run, size=10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("205B57")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(spec.title)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("상세 설계 산출물 · 공식 열람본")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = (
        ("작성일 / 사실 기준", f"{BUILD_DATE} / {FACTS['as_of_date']} canonical facts"),
        ("제품 상태", "PARTIAL · 코드/정적/단위 검증과 현장/운영 검증을 구분"),
        ("주 사용자 앱", "Web/PWA"),
        (
            "연구 보조",
            "Android ARCore·unified img768 TFLite 연구 경로 · "
            f"{ANDROID_DEVICE_CURRENT_LABEL} · {ANDROID_DEVICE_HISTORICAL_LABEL} · "
            "대화형 Device Field 대기",
        ),
        ("기관 신고", "named admin 검수 → agency CSV 최소자료 준비 → 사람이 기관 외부 채널에 수동 제출"),
        ("근거 원문", str(spec.source.relative_to(ROOT))),
    )
    for row, value in zip(table.rows, values):
        prevent_row_split(row)
        for index, (cell, text_value) in enumerate(zip(row.cells, value)):
            set_cell_margins(cell, top=125, bottom=125)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_markdown(paragraph, text_value, table=True)
            if index == 0:
                set_cell_shading(cell, "17324D")
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            else:
                set_cell_shading(cell, "F2F5F6")

    notice = doc.add_paragraph()
    notice.paragraph_format.space_before = Pt(14)
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = notice.add_run(
        "이 문서는 구현 근거와 미검증 항목을 함께 기록한다. "
        "CODE·AUTO·CONTROLLED_E2E 근거는 FIELD·RELEASE를 대신하지 않는다. "
        "독립 test·sequence-safe split·실제 외출용 실폰·기관 접수·backup→restore→retention 시행 전에는 운영 완료로 해석하지 않는다."
    )
    set_run_font(run, size=8.5)
    run.italic = True
    run.font.color.rgb = RGBColor.from_string("586B73")
    doc.add_page_break()


def add_figures(doc: DocumentType, spec: DocumentSpec) -> None:
    if not spec.figures:
        return
    doc.add_heading("시각자료", level=1)
    section = doc.sections[-1]
    printable_width_inches = float(
        section.page_width - section.left_margin - section.right_margin
    ) / 914_400
    for index, (file_name, caption, width_inches) in enumerate(spec.figures, start=1):
        asset = ASSET_DIR / file_name
        if not asset.is_file():
            raise FileNotFoundError(f"missing figure: {asset}")
        if index > 1:
            doc.add_page_break()
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        render_width_inches = min(width_inches, printable_width_inches - 0.08)
        run.add_picture(str(asset), width=Inches(render_width_inches))
        caption_paragraph = doc.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_paragraph.add_run(f"그림 {index}. {caption}")
        set_run_font(caption_run, size=8.5)
        caption_run.bold = True
        caption_run.font.color.rgb = RGBColor.from_string("37474F")
    doc.add_page_break()


def add_markdown(
    doc: DocumentType,
    source_text: str,
    title: str,
    *,
    page_break_before_headings: tuple[str, ...] = (),
) -> None:
    lines = source_text.splitlines()
    index = 0
    skipped_source_title = False
    in_code = False
    code_language = ""
    code_lines: list[str] = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if in_code:
            if stripped.startswith("```"):
                add_code_block(doc, code_lines, code_language)
                in_code = False
                code_language = ""
                code_lines = []
            else:
                code_lines.append(line.rstrip())
            index += 1
            continue

        if stripped.startswith("```"):
            in_code = True
            code_language = stripped[3:].strip()
            index += 1
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            source_level = len(heading_match.group(1))
            heading_text = heading_match.group(2).strip()
            if source_level == 1 and not skipped_source_title:
                skipped_source_title = True
            else:
                target_level = min(max(source_level - 1, 1), 4)
                heading = doc.add_heading(heading_text, level=target_level)
                if heading_text in page_break_before_headings:
                    heading.paragraph_format.page_break_before = True
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows = [split_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(doc, rows)
            continue

        if not stripped:
            index += 1
            continue

        if re.fullmatch(r"-{3,}", stripped):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            p_pr = paragraph._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "B7C7CC")
            border.append(bottom)
            p_pr.append(border)
            index += 1
            continue

        if stripped.startswith(">"):
            callout_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                callout_lines.append(lines[index].strip()[1:].strip())
                index += 1
            text = " ".join(callout_lines)
            if "Status: DRAFT SNAPSHOT" in text:
                text = (
                    "이 추적표의 Markdown 원문은 drafts 경로에 보존되어 있으며, "
                    f"본 DOCX가 {BUILD_DATE} 작성 공식 열람본이다. 구현 사실 기준은 "
                    f"{FACTS['as_of_date']} canonical facts이며 최신 구현 판정은 "
                    "docs/status/current_status.md를 함께 확인한다."
                )
            add_callout(doc, text)
            continue

        bullet_match = re.match(r"^\s*[-*]\s+(?:\[[ xX]\]\s+)?(.+)$", line)
        number_match = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if bullet_match or number_match:
            text = (bullet_match or number_match).group(1)
            style = "List Bullet" if bullet_match else "List Number"
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline_markdown(paragraph, text)
            index += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        next_index = index + 1
        while next_index < len(lines) and not lines[next_index].strip():
            next_index += 1
        if next_index < len(lines) and lines[next_index].strip().startswith("```"):
            paragraph.paragraph_format.keep_with_next = True
        add_inline_markdown(paragraph, stripped.rstrip("  "))
        index += 1

    if in_code:
        raise ValueError(f"unterminated code fence in {title}")


def normalize_docx(path: Path) -> None:
    """Rewrite the package with fixed member timestamps for stable build artifacts."""
    with tempfile.TemporaryDirectory(prefix="walksafe-docx-normalize-") as temp_dir:
        temp_path = Path(temp_dir) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temp_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, date_time=(2026, 7, 13, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, source.read(name))
        shutil.move(temp_path, path)


def document_text(doc: DocumentType) -> str:
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def validate_docx(path: Path, spec: DocumentSpec) -> dict[str, object]:
    if not path.is_file() or path.stat().st_size < 20_000:
        raise ValueError(f"missing or unexpectedly small DOCX: {path}")

    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f"corrupt ZIP member in {path.name}: {bad_member}")
        package_text_parts = ["\n".join(archive.namelist())]
        for name in archive.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                data = archive.read(name)
                ElementTree.fromstring(data)
                package_text_parts.append(data.decode("utf-8", errors="replace"))
        embedded_hashes = sorted(
            sha256_bytes(archive.read(name))
            for name in archive.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        )

    expected_asset_hashes = [sha256(ASSET_DIR / file_name) for file_name, _, _ in spec.figures]
    if embedded_hashes != sorted(expected_asset_hashes):
        raise ValueError(f"embedded figures do not match current assets in {path.name}")

    reopened = Document(path)
    text = document_text(reopened)
    if spec.title not in text:
        raise ValueError(f"title missing from {path.name}")
    if "Web/PWA" not in text or "PARTIAL" not in text:
        raise ValueError(f"scope/status marker missing from {path.name}")
    if BUILD_DATE not in text or "FIELD·RELEASE" not in text:
        raise ValueError(f"current-date/evidence-boundary marker missing from {path.name}")
    if (
        ANDROID_DEVICE_CURRENT_LABEL not in text
        or ANDROID_DEVICE_HISTORICAL_LABEL not in text
    ):
        raise ValueError(f"Android current/historical device evidence boundary missing from {path.name}")
    if "CSV" not in text or "수동 제출" not in text:
        raise ValueError(f"manual CSV flow marker missing from {path.name}")
    if len(text.strip()) < 1_000:
        raise ValueError(f"document text is unexpectedly short: {path.name}")
    if len(reopened.inline_shapes) != len(spec.figures):
        raise ValueError(
            f"figure count mismatch in {path.name}: "
            f"{len(reopened.inline_shapes)} != {len(spec.figures)}"
        )
    printable_width = min(
        section.page_width - section.left_margin - section.right_margin
        for section in reopened.sections
    )
    if any(shape.width > printable_width for shape in reopened.inline_shapes):
        raise ValueError(f"figure exceeds printable page width in {path.name}")
    ordered_shape_hashes = []
    for shape in reopened.inline_shapes:
        relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        ordered_shape_hashes.append(sha256_bytes(reopened.part.related_parts[relationship_id].blob))
    if ordered_shape_hashes != expected_asset_hashes:
        raise ValueError(f"figure order/relationship mismatch in {path.name}")

    body_children = list(reopened.element.body.iterchildren())
    adjacent_captions = []
    for index, child in enumerate(body_children[:-1]):
        if child.tag == qn("w:p") and child.xpath(".//w:drawing"):
            next_child = body_children[index + 1]
            if next_child.tag != qn("w:p"):
                raise ValueError(f"figure caption is not adjacent in {path.name}")
            adjacent_captions.append(Paragraph(next_child, reopened).text)
    if len(adjacent_captions) != len(spec.figures):
        raise ValueError(f"figure/caption adjacency count mismatch in {path.name}")
    for index, (_, caption, _) in enumerate(spec.figures, start=1):
        expected_caption = f"그림 {index}. {caption}"
        if adjacent_captions[index - 1] != expected_caption:
            raise ValueError(f"figure caption missing from {path.name}: {expected_caption}")

    expected_orientation = WD_ORIENT.LANDSCAPE if spec.landscape else WD_ORIENT.PORTRAIT
    if not reopened.sections or any(section.orientation != expected_orientation for section in reopened.sections):
        raise ValueError(f"actual section orientation mismatch in {path.name}")
    scan_text = text + "\n" + "\n".join(package_text_parts)
    for label, pattern in FORBIDDEN_PATTERNS.items():
        if pattern.search(scan_text):
            raise ValueError(f"{label} found in {path.name}")
    for match in re.finditer(
        r"(?i)(?<![\w.])(\d+)\s*(?:개\s*)?(?:tests?|테스트|passed)(?![\w.])",
        text,
    ):
        if int(match.group(1)) not in ALLOWED_FIXED_TEST_COUNTS:
            raise ValueError(f"unrefreshed fixed test-count claim found in {path.name}: {match.group(0)}")
    if re.search(r"로컬\s+PostGIS\s+\d+", text):
        raise ValueError(f"brittle fixed PostGIS test-count claim found in {path.name}")
    for stale_phrase in ("2026-07-10 기준 공식 열람본",):
        if stale_phrase in text:
            raise ValueError(f"stale claim found in {path.name}: {stale_phrase}")
    for phrase in FACTS["forbidden_claims"]:
        if phrase.casefold() in text.casefold():
            raise ValueError(f"canonical forbidden claim found in {path.name}: {phrase}")

    expected_uc_ids = {f"UC-{number:02d}" for number in range(1, 10)}
    expected_req_ids = {f"REQ-{number:03d}" for number in range(1, 15)}
    actual_uc_ids = set(re.findall(r"(?<![A-Za-z0-9-])UC-\d{2}(?!\d)", text))
    actual_req_ids = set(re.findall(r"(?<![A-Za-z0-9-])REQ-\d{3}(?!\d)", text))
    if spec.output_name.startswith(("02_", "03_")) and actual_uc_ids != expected_uc_ids:
        raise ValueError(
            f"use-case ID set mismatch in {path.name}: "
            f"missing={sorted(expected_uc_ids - actual_uc_ids)}, extra={sorted(actual_uc_ids - expected_uc_ids)}"
        )
    if spec.output_name.startswith(("01_", "03_")) and actual_req_ids != expected_req_ids:
        raise ValueError(
            f"requirement ID set mismatch in {path.name}: "
            f"missing={sorted(expected_req_ids - actual_req_ids)}, extra={sorted(actual_req_ids - expected_req_ids)}"
        )
    if spec.output_name.startswith(("01_", "03_")):
        missing_titles = [item["title"] for item in FACTS["requirements"] if item["title"] not in text]
        if missing_titles:
            raise ValueError(f"canonical requirement titles missing from {path.name}: {missing_titles}")
    if spec.output_name.startswith(("02_", "03_")):
        missing_titles = [item["title"] for item in FACTS["use_cases"] if item["title"] not in text]
        if missing_titles:
            raise ValueError(f"canonical use-case titles missing from {path.name}: {missing_titles}")
    if spec.output_name.startswith("06_"):
        expected_tables = [table["name"] for table in FACTS["database"]["tables"]]
        if any(table not in text for table in expected_tables):
            raise ValueError(f"canonical database table missing from {path.name}: {expected_tables}")
    if spec.output_name.startswith(("04_", "07_")):
        if "primary" not in text or "duplicate_count" not in text:
            raise ValueError(f"primary report response contract missing from {path.name}")
    for heading_text in spec.page_break_before_headings:
        matching = [paragraph for paragraph in reopened.paragraphs if paragraph.text == heading_text]
        if len(matching) != 1 or not matching[0].paragraph_format.page_break_before:
            raise ValueError(f"required page-break heading missing from {path.name}: {heading_text}")

    return {
        "file": path.name,
        "sha256": sha256(path),
        "bytes": path.stat().st_size,
        "paragraphs": len(reopened.paragraphs),
        "tables": len(reopened.tables),
        "inline_shapes": len(reopened.inline_shapes),
    }


def build_document(spec: DocumentSpec) -> Path:
    source_text = spec.source.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc, spec)
    add_cover(doc, spec)
    add_figures(doc, spec)
    doc.add_heading("상세 내용", level=1)
    add_markdown(
        doc,
        source_text,
        spec.title,
        page_break_before_headings=spec.page_break_before_headings,
    )
    while doc.paragraphs:
        paragraph = doc.paragraphs[-1]
        if paragraph.text.strip() or paragraph._p.xpath(".//w:br | .//w:drawing"):
            break
        paragraph._element.getparent().remove(paragraph._element)

    output_path = OUTPUT_DIR / spec.output_name
    with atomic_output_path(output_path) as temporary:
        doc.save(temporary)
        normalize_docx(temporary)
    return output_path


def build_all() -> list[dict[str, object]]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    require_no_unexpected_files(OUTPUT_DIR, DESIGN_OUTPUT_NAMES)
    results = []
    for spec in SPECS:
        if not spec.source.is_file():
            raise FileNotFoundError(f"missing source: {spec.source}")
        path = build_document(spec)
        result = validate_docx(path, spec)
        result["source"] = str(spec.source.relative_to(ROOT))
        result["source_sha256"] = sha256(spec.source)
        result["orientation"] = "landscape" if spec.landscape else "portrait"
        results.append(result)

    manifest = {
        "schema_version": "walksafe.design-documents.v2",
        "build_date": BUILD_DATE,
        "facts": {
            "file": str(FACTS_PATH.relative_to(ROOT)),
            "sha256": sha256(FACTS_PATH),
            "schema_version": FACTS["schema_version"],
        },
        "product_status": "PARTIAL",
        "primary_app": "Web/PWA",
        "android_role": "ARCore/depth/TFLite research support",
        "agency_submission": "admin-reviewed export with allow-field check followed by manual external submission",
        "assets": asset_records(),
        **source_revision(),
        "build_tools": build_tools(),
        "submission_toolchain": submission_toolchain_attestation(ROOT),
        "build_inputs": build_input_records(),
        "implementation_bindings": implementation_records(),
        "supporting_files": [
            {
                "file": "docs/submission/design_documents/README.md",
                "sha256": sha256(OUTPUT_DIR / "README.md"),
                "bytes": (OUTPUT_DIR / "README.md").stat().st_size,
            }
        ],
        "documents": results,
    }
    manifest_path = OUTPUT_DIR / "BUILD_MANIFEST.json"
    with atomic_output_path(manifest_path) as temporary:
        temporary.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    require_exact_file_set(OUTPUT_DIR, DESIGN_OUTPUT_NAMES)
    return results


def validate_existing() -> list[dict[str, object]]:
    require_exact_file_set(OUTPUT_DIR, DESIGN_OUTPUT_NAMES)
    manifest_path = OUTPUT_DIR / "BUILD_MANIFEST.json"
    if not manifest_path.is_file():
        raise FileNotFoundError(f"missing build manifest: {manifest_path}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    expected_metadata = {
        "schema_version": "walksafe.design-documents.v2",
        "build_date": BUILD_DATE,
        "product_status": "PARTIAL",
        "primary_app": "Web/PWA",
        "android_role": "ARCore/depth/TFLite research support",
        "agency_submission": "admin-reviewed export with allow-field check followed by manual external submission",
    }
    for key, expected in expected_metadata.items():
        if manifest.get(key) != expected:
            raise ValueError(f"manifest {key} mismatch: {manifest.get(key)!r} != {expected!r}")
    expected_facts = {
        "file": str(FACTS_PATH.relative_to(ROOT)),
        "sha256": sha256(FACTS_PATH),
        "schema_version": FACTS["schema_version"],
    }
    if manifest.get("facts") != expected_facts:
        raise ValueError("manifest facts binding does not match canonical submission facts")

    for key, expected in source_revision().items():
        if manifest.get(key) != expected:
            raise ValueError(f"manifest {key} mismatch: {manifest.get(key)!r} != {expected!r}")
    if manifest.get("source_dirty") is not False:
        raise ValueError("design manifest records dirty submission source")
    if manifest.get("build_tools") != build_tools():
        raise ValueError("manifest build tools do not match the validation environment")
    if manifest.get("submission_toolchain") != submission_toolchain_attestation(ROOT):
        raise ValueError("manifest submission toolchain does not match the pinned environment")
    if manifest.get("implementation_bindings") != implementation_records():
        raise ValueError("manifest implementation bindings do not match current code/config")
    if manifest.get("build_inputs") != build_input_records():
        raise ValueError("manifest build inputs do not match current builders")
    expected_supporting_files = [
        {
            "file": "docs/submission/design_documents/README.md",
            "sha256": sha256(OUTPUT_DIR / "README.md"),
            "bytes": (OUTPUT_DIR / "README.md").stat().st_size,
        }
    ]
    if manifest.get("supporting_files") != expected_supporting_files:
        raise ValueError("manifest supporting files do not match design README")

    document_entries = manifest.get("documents")
    if not isinstance(document_entries, list):
        raise ValueError("manifest documents must be a list")
    document_by_name = {entry.get("file"): entry for entry in document_entries if isinstance(entry, dict)}
    expected_names = {spec.output_name for spec in SPECS}
    if len(document_by_name) != len(document_entries) or set(document_by_name) != expected_names:
        raise ValueError("manifest document set is missing, duplicated, or unexpected")

    results = []
    for spec in SPECS:
        result = validate_docx(OUTPUT_DIR / spec.output_name, spec)
        current = {
            **result,
            "source": str(spec.source.relative_to(ROOT)),
            "source_sha256": sha256(spec.source),
            "orientation": "landscape" if spec.landscape else "portrait",
        }
        recorded = document_by_name[spec.output_name]
        for key, expected in current.items():
            if recorded.get(key) != expected:
                raise ValueError(
                    f"manifest document mismatch for {spec.output_name}.{key}: "
                    f"{recorded.get(key)!r} != {expected!r}"
                )
        results.append(result)

    asset_entries = manifest.get("assets")
    if not isinstance(asset_entries, list):
        raise ValueError("manifest assets must be a list")
    asset_by_name = {entry.get("file"): entry for entry in asset_entries if isinstance(entry, dict)}
    current_assets = asset_records()
    expected_asset_names = {entry["file"] for entry in current_assets}
    if len(asset_by_name) != len(asset_entries) or set(asset_by_name) != expected_asset_names:
        raise ValueError("manifest asset set is missing, duplicated, or unexpected")
    for current in current_assets:
        recorded = asset_by_name[current["file"]]
        for key, expected in current.items():
            if recorded.get(key) != expected:
                raise ValueError(
                    f"manifest asset mismatch for {current['file']}.{key}: "
                    f"{recorded.get(key)!r} != {expected!r}"
                )
    return results


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--validate-only",
        action="store_true",
        help="validate existing DOCX files without rebuilding them",
    )
    args = parser.parse_args()
    with submission_build_lock(ROOT):
        require_clean_source_revision(ROOT, ALL_SUBMISSION_GENERATED_PATHS)
        submission_toolchain_attestation(ROOT)
        results = validate_existing() if args.validate_only else build_all()
    total_bytes = sum(int(result["bytes"]) for result in results)
    print(
        f"design documents PASS: {len(results)} DOCX, "
        f"{total_bytes:,} bytes, sources preserved, product status PARTIAL"
    )
    for result in results:
        print(
            f"- {result['file']}: {result['tables']} tables, "
            f"{result['inline_shapes']} figures, sha256={str(result['sha256'])[:12]}..."
        )


if __name__ == "__main__":
    main()

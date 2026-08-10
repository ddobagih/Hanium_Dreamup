#!/usr/bin/env python3
"""Build the 24-slide WalkSafe Assist midterm design presentation.

The user-provided 2026 application-software template is the immutable base.
Only deliverables marked required in the template's ``일반`` column are expanded.
"""

from __future__ import annotations

import hashlib
import io
import os
import shutil
import tempfile
import textwrap
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    "/home/ddobagi/Downloads/[서식2] 2026년 한이음 드림업 제작설계서(택1, 개인정보 기재x)/"
    "[서식2] 2026년_제작설계서_응용소프트웨어(개인정보 기재x).pptx"
)
OUTPUT = Path("/home/ddobagi/Downloads/제작설계서_완성본_개선본.pptx")
EXPECTED_TEMPLATE_SHA256 = "75d2579807ea758185556b6d4349b82e71eb84451940c592434a55e1dadb3d1b"

ASSETS = ROOT / "docs/submission/form_materials/assets"
WEB_MOBILE = ASSETS / "web_main_mobile.png"
WEB_DESKTOP = ASSETS / "web_main_desktop.png"
ADMIN_DESKTOP = ASSETS / "admin_desktop.png"
DAMAGED_TACTILE = (
    ROOT / "ai_tasks/walksafe_v3_relabel_20260521/images/wsv3_img_0021_51453cd280.jpg"
)

FONT = "Noto Sans CJK KR"
CODE_FONT = "DejaVu Sans Mono"

NAVY = RGBColor(22, 40, 63)
BLUE = RGBColor(59, 90, 168)
TEAL = RGBColor(15, 139, 141)
GREEN = RGBColor(42, 157, 143)
GOLD = RGBColor(242, 177, 52)
CORAL = RGBColor(217, 93, 57)
RED = RGBColor(180, 53, 47)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(92, 105, 120)
LINE = RGBColor(214, 220, 229)
PALE = RGBColor(246, 248, 251)
WHITE = RGBColor(255, 255, 255)
SOFT_BLUE = RGBColor(235, 240, 252)
SOFT_TEAL = RGBColor(232, 247, 245)
SOFT_GOLD = RGBColor(254, 247, 225)
SOFT_CORAL = RGBColor(253, 238, 234)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_package(path: Path) -> None:
    """Remove stale template counts and normalize OOXML member metadata."""
    app_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    with tempfile.TemporaryDirectory(prefix="walksafe-ppt-normalize-") as directory:
        normalized = Path(directory) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/app.xml":
                    root = ET.fromstring(data)
                    slides = root.find(f"{{{app_ns}}}Slides")
                    notes = root.find(f"{{{app_ns}}}Notes")
                    if slides is not None:
                        slides.text = "24"
                    if notes is not None:
                        notes.text = "0"
                    for stale_tag in ("HeadingPairs", "TitlesOfParts"):
                        stale = root.find(f"{{{app_ns}}}{stale_tag}")
                        if stale is not None:
                            root.remove(stale)
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                info = zipfile.ZipInfo(name, date_time=(2026, 7, 12, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, data)
        shutil.move(normalized, path)


def delete_slide(prs: Presentation, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)
    for rel_id, rel in list(slide.part.rels.items()):
        if not rel.reltype.endswith("/slideLayout"):
            slide.part.drop_rel(rel_id)


def rect(slide, x, y, w, h, *, fill=WHITE, line=LINE, radius=False):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def line(slide, x1, y1, x2, y2, *, color=LINE, width=1.5):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def textbox(
    slide,
    value,
    x,
    y,
    w,
    h,
    *,
    size=15,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font=FONT,
    margin=0.04,
    fit=False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    if fit:
        frame.fit_text(font_family=font, max_size=Pt(size))
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.05
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullets(slide, values, x, y, w, h, *, size=14, color=INK, bullet_color=None, gap=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.02)
    for index, value in enumerate(values):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(gap)
        paragraph.line_spacing = 1.1
        paragraph.level = 0
        run = paragraph.add_run()
        run.text = f"• {value}"
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = bullet_color or color
    return box


def label(slide, value, x, y, w, *, fill=TEAL, color=WHITE, size=10.5):
    rect(slide, x, y, w, 0.34, fill=fill, line=fill, radius=True)
    return textbox(
        slide,
        value,
        x + 0.08,
        y + 0.055,
        w - 0.16,
        0.22,
        size=size,
        color=color,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )


def header(slide, section: str, conclusion: str, number: int, logo_blob: bytes) -> None:
    rect(slide, 0, 0, 10, 7.5, fill=WHITE, line=WHITE)
    rect(slide, 0.12, 0, 3.4, 0.72, fill=BLUE, line=BLUE)
    rect(slide, 0, 0, 0.14, 0.72, fill=TEAL, line=TEAL)
    textbox(slide, section, 0.36, 0.19, 3.0, 0.3, size=15, color=WHITE, bold=True)
    line(slide, 3.57, 0.57, 8.05, 0.57, color=BLUE, width=1.1)
    slide.shapes.add_picture(io.BytesIO(logo_blob), Inches(8.23), Inches(0.11), width=Inches(1.64), height=Inches(0.38))
    textbox(slide, f"{number:02d}", 9.15, 0.54, 0.55, 0.22, size=9.5, color=CORAL, bold=True, align=PP_ALIGN.RIGHT)
    textbox(slide, conclusion, 0.52, 0.92, 8.95, 0.56, size=21, color=NAVY, bold=True, valign=MSO_ANCHOR.MIDDLE)
    line(slide, 0.52, 1.5, 9.48, 1.5, color=LINE, width=1)
    textbox(slide, "WalkSafe Assist · 중간 제작설계서 · 근거 기준 2026-07-11", 0.52, 7.13, 7.4, 0.18, size=7.8, color=MUTED)
    textbox(slide, str(number), 8.8, 7.1, 0.65, 0.2, size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, *, accent=TEAL, fill=PALE, title_size=14, body_size=11.5):
    rect(slide, x, y, w, h, fill=fill, line=LINE, radius=True)
    rect(slide, x, y, 0.08, h, fill=accent, line=accent)
    textbox(slide, title, x + 0.2, y + 0.16, w - 0.35, 0.35, size=title_size, color=accent, bold=True)
    textbox(slide, body, x + 0.2, y + 0.61, w - 0.35, h - 0.75, size=body_size, color=INK)


def step_card(slide, x, y, w, h, number, title, body, *, accent=TEAL):
    rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True)
    rect(slide, x, y, 0.1, h, fill=accent, line=accent)
    rect(slide, x + 0.17, y + 0.17, 0.36, 0.36, fill=accent, line=accent, radius=True)
    textbox(slide, str(number), x + 0.19, y + 0.23, 0.32, 0.17, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, title, x + 0.64, y + 0.17, w - 0.82, 0.32, size=14, color=NAVY, bold=True)
    textbox(slide, body, x + 0.2, y + 0.67, w - 0.4, h - 0.83, size=11.2, color=MUTED)


def add_table(slide, rows, x, y, w, h, *, widths=None, font_size=11.2, header_fill=NAVY, highlight_col=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if widths:
        total = sum(widths)
        for i, weight in enumerate(widths):
            table.columns[i].width = Inches(w * weight / total)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.025)
            cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = header_fill
            elif highlight_col is not None and col_index == highlight_col:
                cell.fill.fore_color.rgb = SOFT_GOLD
            else:
                cell.fill.fore_color.rgb = WHITE if row_index % 2 else PALE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER if row_index == 0 or col_index == 0 else PP_ALIGN.LEFT
                paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size - 0.5 if row_index == 0 else font_size)
                    run.font.bold = row_index == 0 or (highlight_col is not None and col_index == highlight_col)
                    run.font.color.rgb = WHITE if row_index == 0 else INK
    return shape


def add_picture_fill(slide, path: Path, x, y, w, h, *, line_color=LINE):
    with Image.open(path) as image:
        image_w, image_h = image.size
    frame_ratio = w / h
    image_ratio = image_w / image_h
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if image_ratio > frame_ratio:
        visible = frame_ratio / image_ratio
        picture.crop_left = (1 - visible) / 2
        picture.crop_right = (1 - visible) / 2
    else:
        visible = image_ratio / frame_ratio
        picture.crop_top = (1 - visible) / 2
        picture.crop_bottom = (1 - visible) / 2
    picture.line.color.rgb = line_color
    picture.line.width = Pt(1)
    return picture


def add_picture_fit(slide, path: Path, x, y, w, h, *, line_color=LINE):
    with Image.open(path) as image:
        image_w, image_h = image.size
    scale = min(w / image_w, h / image_h)
    draw_w = image_w * scale
    draw_h = image_h * scale
    rect(slide, x, y, w, h, fill=WHITE, line=line_color, radius=True)
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - draw_w) / 2),
        Inches(y + (h - draw_h) / 2),
        width=Inches(draw_w),
        height=Inches(draw_h),
    )


def flow(slide, items, x, y, total_w, h, *, colors=None, body_size=10.8):
    gap = 0.28
    count = len(items)
    card_w = (total_w - gap * (count - 1)) / count
    for i, (title, body) in enumerate(items):
        accent = colors[i] if colors else [BLUE, TEAL, GOLD, CORAL][i % 4]
        xx = x + i * (card_w + gap)
        step_card(slide, xx, y, card_w, h, i + 1, title, body, accent=accent)
        if i < count - 1:
            textbox(slide, "→", xx + card_w, y + h / 2 - 0.2, gap, 0.35, size=19, color=MUTED, bold=True, align=PP_ALIGN.CENTER)


def phone_frame(slide, x, y, w, h, *, title="WalkSafe Assist", mode="보행 보조"):
    rect(slide, x, y, w, h, fill=RGBColor(10, 18, 28), line=NAVY, radius=True)
    rect(slide, x + 0.12, y + 0.12, w - 0.24, 0.62, fill=RGBColor(25, 35, 45), line=RGBColor(25, 35, 45), radius=True)
    textbox(slide, title, x + 0.3, y + 0.24, w - 0.65, 0.22, size=12, color=WHITE, bold=True)
    textbox(slide, mode, x + 0.3, y + 0.47, w - 0.65, 0.16, size=8.5, color=RGBColor(190, 199, 210))


def code_excerpt(relative: str, start_marker: str, line_count: int) -> str:
    lines = (ROOT / relative).read_text(encoding="utf-8").splitlines()
    start = next(i for i, value in enumerate(lines) if start_marker in value)
    return textwrap.dedent("\n".join(lines[start : start + line_count]))


def build() -> None:
    required = [TEMPLATE, WEB_MOBILE, WEB_DESKTOP, ADMIN_DESKTOP, DAMAGED_TACTILE]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise SystemExit("missing input: " + ", ".join(missing))
    actual_hash = sha256(TEMPLATE)
    if actual_hash != EXPECTED_TEMPLATE_SHA256:
        raise SystemExit(f"unexpected template hash: {actual_hash}")

    prs = Presentation(TEMPLATE)
    if len(prs.slides) != 35 or (prs.slide_width, prs.slide_height) != (9144000, 6858000):
        raise SystemExit("unexpected template shape")

    cover_logo = next(
        shape.image.blob
        for shape in prs.slides[0].shapes
        if shape.shape_type == 13 and shape.left / 914400 > 8
    )
    while len(prs.slides) > 24:
        delete_slide(prs, len(prs.slides) - 1)
    for slide in prs.slides:
        clear_slide(slide)

    core = prs.core_properties
    core.title = "WalkSafe Assist 중간 제작설계서"
    core.subject = "2026 한이음 드림업 제작설계서"
    core.author = "WalkSafe Assist 프로젝트"
    core.last_modified_by = "WalkSafe Assist 프로젝트"
    core.comments = "개인정보 없이 코드 및 검증 근거로 작성"
    core.created = datetime(2026, 7, 12, 0, 0, 0)
    core.modified = datetime(2026, 7, 12, 0, 0, 0)

    slides = list(prs.slides)

    # 1. Cover
    slide = slides[0]
    rect(slide, 0, 0, 10, 7.5, fill=NAVY, line=NAVY)
    rect(slide, 0, 0, 0.18, 7.5, fill=TEAL, line=TEAL)
    slide.shapes.add_picture(io.BytesIO(cover_logo), Inches(8.05), Inches(0.3), width=Inches(1.64), height=Inches(0.38))
    textbox(slide, "2026 한이음 드림업 · 중간 제작설계서", 0.68, 0.55, 5.9, 0.34, size=13, color=GOLD, bold=True)
    textbox(slide, "WalkSafe Assist", 0.68, 1.08, 5.7, 0.5, size=25, color=RGBColor(101, 225, 211), bold=True)
    textbox(slide, "현재 위험은 선별해 알리고,", 0.68, 1.78, 6.2, 0.65, size=30, color=WHITE, bold=True)
    textbox(slide, "손상 시설은 개선 자료로 남깁니다.", 0.68, 2.48, 7.5, 0.65, size=26, color=WHITE, bold=True)
    textbox(slide, "AI를 이용한 카메라 기반 시각장애인 보행 지원 시스템", 0.7, 3.32, 7.8, 0.36, size=13.5, color=RGBColor(205, 214, 224))
    cover_cards = [
        (0.7, "현재 위험", "경로와 관련된 후보만\n짧은 음성·진동으로 안내", TEAL),
        (3.72, "목적지까지", "TMAP 전역 경로와\n음성 제어를 한 흐름으로", GOLD),
        (6.74, "시설 개선", "손상 점자블록을\n위치·이미지로 검수", CORAL),
    ]
    for x, title, body, accent in cover_cards:
        card(slide, x, 4.25, 2.55, 1.72, title, body, accent=accent, fill=RGBColor(29, 49, 72), title_size=16, body_size=12.5)
        for paragraph in slide.shapes[-1].text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = WHITE
    label(slide, "주제품 Web/PWA · Android는 ARCore·TFLite 연구 보조", 0.7, 6.35, 5.6, fill=TEAL)
    textbox(slide, "기준일 2026. 07. 11.", 0.7, 6.86, 2.2, 0.25, size=10.5, color=RGBColor(190, 200, 212))

    # 2. Official deliverables matrix, general column highlighted.
    slide = slides[1]
    header(slide, "수행 단계별 주요 산출물", "일반 분야 필수 산출물 9종만 설계 범위로 확정", 2, cover_logo)
    matrix = [
        ["단계", "일반 분야 필수 산출물", "표시", "수록 장"],
        ["요구사항", "요구사항 정의서", "○", "3–4"],
        ["아키텍처", "서비스 구성도(시스템 구성도)", "○", "5–6"],
        ["기능 설계", "메뉴 구성도", "○", "7–8"],
        ["", "화면 설계서", "○", "9–14"],
        ["", "엔티티 관계도", "○", "15"],
        ["", "기능 처리도(기능 흐름도)", "○", "16–18"],
        ["개발/구현", "프로그램 목록", "○", "19"],
        ["", "테이블 정의서", "○", "20–21"],
        ["", "핵심 소스코드", "○", "22–23"],
    ]
    add_table(slide, matrix, 0.55, 1.82, 6.65, 4.78, widths=[1.4, 3.45, 0.7, 0.8], font_size=10.7, highlight_col=2)
    card(slide, 7.42, 1.72, 2.05, 2.2, "작성 대상", "요구사항\n서비스 구성\n메뉴·화면\nERD·기능 처리\n프로그램·테이블·코드", accent=GOLD, fill=SOFT_GOLD, title_size=14, body_size=11.5)
    card(slide, 7.42, 4.15, 2.05, 1.55, "구성 원칙", "한 장에 한 질문\n결론형 제목\n편집 가능한 도형", accent=TEAL, fill=SOFT_TEAL, title_size=14, body_size=11.5)
    textbox(slide, "공식 산출물 표의 일반 열에서\n○ 표시 항목만 발췌", 7.47, 6.05, 1.95, 0.52, size=9.3, color=MUTED, align=PP_ALIGN.CENTER)

    # 3. Requirements 1
    slide = slides[2]
    header(slide, "요구사항 정의서 1/2", "사용자는 보고, 선별된 위험을 듣고, 말로 제어", 3, cover_logo)
    groups = [
        (0.62, "보고 구분", "REQ-001 카메라 위험 감지\nREQ-002 실탐지·데모 구분", BLUE),
        (2.92, "선별 안내", "REQ-004 객체 안정 추적\nREQ-006 TTS·진동 피드백", TEAL),
        (5.22, "말로 제어", "REQ-007 목적지·길안내·신고", GOLD),
        (7.52, "기록·연구", "REQ-005 손상 시설 신고\nREQ-003 Android 연구 보조", CORAL),
    ]
    for x, title, body, accent in groups:
        card(slide, x, 1.88, 2.08, 2.15, title, body, accent=accent, fill=WHITE, title_size=15, body_size=11.4)
    flow(slide, [
        ("카메라", "전방 객체 후보"),
        ("안정화", "3 frame·700ms"),
        ("위험 판단", "경로·접근 근거"),
        ("피드백", "짧은 음성·진동"),
    ], 0.7, 4.47, 8.6, 1.42, colors=[BLUE, TEAL, GOLD, CORAL], body_size=10.5)
    label(slide, "구현·자동검증 확보 / 실폰 카메라·GPS·마이크와 실외 사용자 과업 검증 전", 1.0, 6.32, 7.95, fill=NAVY)

    # 4. Requirements 2
    slide = slides[3]
    header(slide, "요구사항 정의서 2/2", "신고는 저장에서 끝나지 않고 검수·반출·추적까지 이어짐", 4, cover_logo)
    rows = [
        ["묶음", "요구사항", "현재 근거", "다음 확인"],
        ["신뢰 저장", "REQ-008 durable journal·시작 조정\nREQ-009 정확히 5개 application table", "코드·격리 DB", "운영 복구·보존"],
        ["검수·반출", "REQ-010 list/detail/image read audit\nREQ-014 네 가지 독립 기본 OFF 동의", "코드·자동검증", "조직 IAM·실제 운영"],
        ["설치·길안내", "REQ-011 PWA·오프라인 경계\nREQ-012 목적지·보행 경로", "빌드·통제 시험", "실폰 HTTPS·실외 과업"],
        ["모델 추적", "REQ-013 모델·데이터·클래스 순서", "registry·hash", "독립 test·재분할"],
    ]
    add_table(slide, rows, 0.62, 1.82, 8.76, 3.65, widths=[1.25, 3.45, 1.55, 1.8], font_size=11.2)
    cards = [
        (0.7, "실패를 숨기지 않음", "권한·네트워크·서버 오류는\n‘위험 없음’이 아니라 판단 불가", CORAL),
        (3.55, "사람 검수 유지", "AI 후보가 기관으로\n자동 제출되지 않음", GOLD),
        (6.4, "근거를 함께 기록", "모델·출처·촬영 시각·위치로\n검수 가능성을 확보", TEAL),
    ]
    for x, title, body, accent in cards:
        card(slide, x, 5.56, 2.6, 1.17, title, body, accent=accent, fill=PALE, title_size=12.5, body_size=9.4)

    # 5. Architecture
    slide = slides[4]
    header(slide, "서비스 구성도 1/2", "Web/PWA가 사용자 흐름을, Backend가 신뢰 경계를 담당", 5, cover_logo)
    nodes = [
        (0.55, 2.0, 1.45, 2.05, "사용자 단말", "카메라\nGPS·방향\n마이크\nTTS·진동", BLUE),
        (2.35, 1.78, 1.85, 2.5, "Web/PWA", "frame/report 별도 동의\n철회 시 in-flight 취소\n단일 replica\nprocess lock", TEAL),
        (4.65, 1.78, 1.85, 2.5, "FastAPI", "인증 OpenAPI\nTMAP-only navigation\nshared DB limiter\nnonblocking ready", CORAL),
        (6.95, 1.78, 2.25, 2.5, "저장·외부", "PostGIS 5 application tables\nDB sslmode=verify-full\ngssencmode=disable\nDB 밖 journal/image\nTMAP-only", GOLD),
    ]
    for x, y, w, h, title, body, accent in nodes:
        card(slide, x, y, w, h, title, body, accent=accent, fill=WHITE, title_size=14, body_size=11.2)
    for x in [2.04, 4.34, 6.64]:
        textbox(slide, "→", x, 2.73, 0.28, 0.4, size=20, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    card(slide, 0.7, 4.78, 2.35, 1.15, "Admin/Ops", "list/detail/image actor·resource·time·purpose 감사", accent=GREEN, fill=SOFT_TEAL, title_size=12.5, body_size=8.8)
    card(slide, 3.42, 4.78, 2.35, 1.15, "Android 연구", "ARCore depth·unified TFLite", accent=BLUE, fill=SOFT_BLUE, title_size=13.5, body_size=10.5)
    card(slide, 6.14, 4.78, 2.95, 1.15, "기관 외부 채널", "관리자가 검수한 자료를 사람이 제출", accent=CORAL, fill=SOFT_CORAL, title_size=13.5, body_size=10.5)
    line(slide, 1.88, 4.78, 4.65, 4.25, color=GREEN, width=1.3)
    line(slide, 4.6, 4.78, 3.2, 4.25, color=BLUE, width=1.3)
    line(slide, 7.6, 4.78, 5.6, 4.25, color=CORAL, width=1.3)
    label(slide, "Voice: 전용 token · pre-copy auth/Content-Length/global·actor·IP rate · copy byte · post-copy duration/inference queue · disconnect · worker/replica=1+lock", 0.9, 6.3, 8.2, fill=NAVY)

    # 6. Service flow
    slide = slides[5]
    header(slide, "서비스 구성도 2/2", "한 번의 관찰은 위험 안내와 시설 신고의 두 경로로 분기", 6, cover_logo)
    flow(slide, [
        ("입력", "Web frame/report·Android report\ntelemetry 각각 기본 OFF"),
        ("분석", "객체 후보\n인스턴스 안정화"),
        ("판단", "경로·접근\n신뢰조건 확인"),
    ], 0.62, 1.95, 5.45, 1.6, colors=[BLUE, TEAL, GOLD])
    textbox(slide, "↘", 5.95, 2.35, 0.6, 0.45, size=24, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    card(slide, 6.55, 1.82, 2.82, 1.65, "현재 위험", "짧은 TTS·진동\n우선순위·반복 제한", accent=CORAL, fill=SOFT_CORAL, title_size=16, body_size=12)
    textbox(slide, "↘", 5.95, 4.38, 0.6, 0.45, size=24, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    card(slide, 6.55, 3.88, 2.82, 1.65, "손상 점자블록", "위치·이미지·모델 근거\n관리자 검수 자료", accent=TEAL, fill=SOFT_TEAL, title_size=16, body_size=12)
    flow(slide, [
        ("관리자", "근거 확인\n상태 변경"),
        ("자료", "기관용 최소정보\n반출 감사"),
        ("사람", "외부 채널\n수동 제출"),
    ], 0.62, 4.2, 5.45, 1.6, colors=[GREEN, GOLD, CORAL])
    label(slide, "탐지 존재만으로 경고하지 않고, AI 후보가 사람 검수 없이 기관으로 나가지 않음", 1.05, 6.33, 7.9, fill=NAVY)

    # 7. Web menu
    slide = slides[6]
    header(slide, "메뉴 구성도 1/2", "Web/PWA는 한 화면에서 다섯 보행 과업만 제어", 7, cover_logo)
    menus = [
        ("시작·상태", "Web frame/report 별도 동의\n철회 시 in-flight 취소", BLUE),
        ("보행 보조", "탐지 · 위험 선별\nTTS · 진동", TEAL),
        ("음성", "목적지 · 후보 선택\n취소 · 신고", GOLD),
        ("길안내", "전역 경로 · 이탈\n재탐색 · 도착", CORAL),
        ("설치·연결", "선택형 PWA\n오프라인 화면", GREEN),
    ]
    center_x, center_y = 4.2, 3.45
    rect(slide, center_x, center_y, 1.6, 0.75, fill=NAVY, line=NAVY, radius=True)
    textbox(slide, "WalkSafe\nAssist", center_x + 0.18, center_y + 0.13, 1.25, 0.45, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    positions = [(0.65, 1.85), (3.02, 1.78), (6.92, 1.85), (1.6, 4.75), (6.0, 4.75)]
    for (title, body, accent), (x, y) in zip(menus, positions):
        card(slide, x, y, 2.35, 1.38, title, body, accent=accent, fill=WHITE, title_size=14, body_size=10.8)
        line(slide, center_x + 0.8, center_y + 0.38, x + 1.17, y + 0.7, color=accent, width=1.2)
    label(slide, "실패 상태마다 원인과 다음 행동을 함께 표시", 3.05, 6.44, 3.9, fill=TEAL)

    # 8. Admin menu
    slide = slides[7]
    header(slide, "메뉴 구성도 2/2", "Admin은 근거 확인 뒤 상태와 반출을 결정", 8, cover_logo)
    flow(slide, [
        ("필터·목록", "상태·유형·출처\n날짜·반경"),
        ("상세", "이미지·위치·bbox\n모델·신고 방식"),
        ("상태", "신규 → 검토\n→ 처리 완료"),
        ("반출", "internal·minimum\nagency + 감사"),
    ], 0.62, 1.92, 8.76, 1.76, colors=[BLUE, TEAL, GOLD, CORAL])
    rows = [
        ["메뉴", "판단 기준", "허용되는 행동"],
        ["목록", "검수 대상을 조건으로 좁힘", "purpose 포함 read audit 뒤 조회"],
        ["상세", "동일 촬영 근거와 위치 품질 확인", "detail/image 각각 read audit"],
        ["상태", "담당자·이전/다음 상태·시각", "new → reviewed → resolved"],
        ["반출", "검토·최소정보·행 수 제한", "CSV·JSON·GeoJSON 준비"],
    ]
    add_table(slide, rows, 0.82, 4.12, 8.36, 2.05, widths=[1.2, 3.4, 2.8], font_size=11)
    label(slide, "기관 전송 버튼 없음 · 검수된 자료를 사람이 외부 채널에 제출", 2.0, 6.45, 6.0, fill=NAVY)

    # 9. Permission screen
    slide = slides[8]
    header(slide, "화면 설계서 1/6", "보행 보조는 권한과 연결 상태를 확인한 뒤 시작", 9, cover_logo)
    add_picture_fill(slide, WEB_MOBILE, 0.68, 1.8, 3.05, 4.92)
    label(slide, "개발용 fixture 화면", 0.86, 6.25, 1.85, fill=BLUE)
    flow(slide, [
        ("고지", "수집 범위·보존\n카메라·음성 원본 제외"),
        ("권한", "카메라·GPS·마이크\n명시적 요청"),
        ("복구", "거부·오프라인 원인\n재요청 버튼"),
    ], 4.08, 2.1, 5.15, 1.65, colors=[BLUE, GOLD, TEAL])
    card(slide, 4.08, 4.15, 5.15, 1.42, "세션 종료 규칙", "백그라운드 전환 또는 세션 종료 시 카메라·길안내·음성을 중지하고 이전 탐지 상태를 폐기", accent=CORAL, fill=SOFT_CORAL, title_size=14, body_size=11.5)
    label(slide, "실폰 권한 흐름·TalkBack·브라우저별 동작은 현장 검증 전", 4.45, 6.2, 4.4, fill=NAVY)

    # 10. Main assist screen
    slide = slides[9]
    header(slide, "화면 설계서 2/6", "현재 위험·위치·방향·길안내를 한 화면에서 확인", 10, cover_logo)
    add_picture_fit(slide, WEB_DESKTOP, 0.55, 1.8, 6.55, 4.82)
    label(slide, "개발용 fake-v2 합성 입력", 0.77, 6.23, 2.2, fill=BLUE)
    cards = [
        (7.35, 1.82, "카메라", "미리보기·bbox\n탐지 출처", BLUE),
        (7.35, 2.98, "상태 카드", "현재 위험·GPS\n방향·신고", TEAL),
        (7.35, 4.14, "조작", "음성 안내·명령\n보행 시작·중지", GOLD),
        (7.35, 5.30, "복구", "권한 재요청\n연결 실패 표시", CORAL),
    ]
    for x, y, title, body, accent in cards:
        card(slide, x, y, 1.95, 1.03, title, body, accent=accent, fill=WHITE, title_size=12.5, body_size=8.9)

    # 11. Risk screen
    slide = slides[10]
    header(slide, "화면 설계서 3/6", "경고는 근거가 충분한 가장 중요한 위험만 전달", 11, cover_logo)
    phone_frame(slide, 0.7, 1.78, 3.45, 4.92, mode="전방 위험 안내")
    rect(slide, 0.95, 2.72, 2.95, 0.82, fill=RGBColor(75, 25, 25), line=CORAL, radius=True)
    textbox(slide, "전방 장애물. 멈추세요.", 1.15, 2.97, 2.55, 0.3, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    card(slide, 0.96, 3.88, 1.34, 1.0, "현재 위험", "보행자\n신뢰도 80%", accent=CORAL, fill=RGBColor(28, 35, 40), title_size=10.5, body_size=9.2)
    card(slide, 2.52, 3.88, 1.38, 1.0, "행동", "멈춤\n전방 확인", accent=GOLD, fill=RGBColor(28, 35, 40), title_size=10.5, body_size=9.2)
    label(slide, "화면 동작 예시", 1.55, 5.4, 1.74, fill=BLUE)
    flow(slide, [
        ("안정화", "3 frame·700ms"),
        ("위험 판단", "경로·접근·거리"),
        ("우선 안내", "위험 > 명령 > 길안내"),
    ], 4.52, 2.05, 4.75, 1.62, colors=[BLUE, GOLD, CORAL])
    card(slide, 4.52, 4.06, 4.75, 1.52, "오래되거나 실패한 탐지", "stale·frozen·server unavailable 상태는 ‘위험 없음’으로 바꾸지 않고 안내와 신고 판단을 중지", accent=TEAL, fill=SOFT_TEAL, title_size=14, body_size=11.2)
    label(slide, "고위험만 진동 · 반복 경고는 cooldown", 5.28, 6.18, 3.2, fill=NAVY)

    # 12. Navigation screen
    slide = slides[11]
    header(slide, "화면 설계서 4/6", "목적지는 말로 찾고, 후보를 고른 뒤 안내를 시작", 12, cover_logo)
    phone_frame(slide, 0.68, 1.78, 3.2, 4.93, mode="음성 목적지·길안내")
    label(slide, "음성 명령", 0.97, 2.65, 1.15, fill=BLUE)
    textbox(slide, "“서울역으로 안내해줘”", 0.98, 3.16, 2.55, 0.45, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 0.98, 3.82, 2.58, 1.4, fill=RGBColor(28, 38, 48), line=TEAL, radius=True)
    textbox(slide, "1. 서울역 1번 출구\n2. 서울역 버스환승센터", 1.18, 4.15, 2.18, 0.72, size=11.5, color=WHITE)
    label(slide, "번호를 말해 선택", 1.25, 5.58, 1.95, fill=TEAL)
    flow(slide, [
        ("검색", "TMAP 장소 검색"),
        ("선택", "후보 번호 확인"),
        ("안내", "계단 회피 보행 경로"),
        ("복귀", "이탈 시 재탐색"),
    ], 4.25, 1.98, 5.05, 1.64, colors=[BLUE, TEAL, GOLD, CORAL], body_size=9.8)
    card(slide, 4.25, 4.04, 5.05, 1.52, "정상 점자블록의 역할", "신뢰도·최신성·GPS·방향·경로 조건을 모두 통과한 동안만 짧은 local steering을 제공하고, 조건 미달이면 TMAP으로 복귀", accent=TEAL, fill=SOFT_TEAL, title_size=14, body_size=11)
    label(slide, "임의 자동 선택 없음 · 시작·다음 안내·재탐색·중지를 음성으로 제어", 4.47, 6.18, 4.62, fill=NAVY)

    # 13. Report screen
    slide = slides[12]
    header(slide, "화면 설계서 5/6", "손상 점자블록은 촬영 근거가 맞을 때만 신고", 13, cover_logo)
    add_picture_fill(slide, DAMAGED_TACTILE, 0.7, 1.8, 3.25, 4.82)
    label(slide, "개인정보 정제 학습·검수 데이터 예시", 0.83, 6.22, 2.98, fill=BLUE, size=9.2)
    flow(slide, [
        ("후보", "손상 점자블록\n신뢰도 0.70 이상"),
        ("촬영", "동일 이미지·위치\n촬영 시각"),
        ("전송", "진행·완료·실패\n상태 표시"),
    ], 4.3, 2.0, 4.98, 1.68, colors=[GOLD, TEAL, BLUE], body_size=10)
    card(slide, 4.3, 4.08, 4.98, 1.48, "중복 처리", "동일 class·10m·1분 후보는 건수만 알리고 자동 병합하지 않으며, 관리자 상세에서 근거를 비교", accent=CORAL, fill=SOFT_CORAL, title_size=14, body_size=11.2)
    label(slide, "자동 신고는 기본 무음 · 사용자가 음성 요청한 경우만 짧게 결과 안내", 4.55, 6.18, 4.5, fill=NAVY)

    # 14. Admin screen
    slide = slides[13]
    header(slide, "화면 설계서 6/6", "운영자는 신고 근거를 확인한 뒤 상태와 반출을 결정", 14, cover_logo)
    add_picture_fit(slide, ADMIN_DESKTOP, 0.55, 1.8, 6.55, 4.82)
    label(slide, "실제 Admin UI · 로컬 fixture 자료", 0.75, 6.22, 2.62, fill=BLUE)
    cards = [
        (7.32, 1.82, "1 필터", "상태·유형·출처\n날짜·반경", BLUE),
        (7.32, 2.98, "2 검수", "이미지·좌표·bbox\n모델·신고 방식", TEAL),
        (7.32, 4.14, "3 상태", "담당자·사유·이력", GOLD),
        (7.32, 5.30, "4 반출", "검토 완료 건만\n기관용 자료 준비", CORAL),
    ]
    for x, y, title, body, accent in cards:
        card(slide, x, y, 1.98, 1.03, title, body, accent=accent, fill=WHITE, title_size=12, body_size=8.8)

    # 15. ERD
    slide = slides[14]
    header(slide, "엔티티 관계도", "신고·3종 감사·공유 제한을 정확히 5개 application table로 분리", 15, cover_logo)
    card(slide, 0.7, 1.88, 4.25, 3.72, "reports", "id UUID · status\nclass_name · confidence · bbox\ncaptured_at · source\nlatitude · longitude · accuracy · heading\nlocation Point(4326)\nimage_path · image_mime_type\nmetadata JSONB", accent=TEAL, fill=SOFT_TEAL, title_size=17, body_size=12.1)
    card(slide, 5.55, 1.72, 3.8, 0.92, "report_status_audits", "상태 전이·actor·time · APPEND-ONLY", accent=GOLD, fill=SOFT_GOLD, title_size=12, body_size=7.5)
    card(slide, 5.55, 2.82, 3.8, 0.92, "report_export_audits", "audit/requested ID·profile·hash · APPEND-ONLY", accent=CORAL, fill=SOFT_CORAL, title_size=12, body_size=7.2)
    card(slide, 5.55, 3.92, 3.8, 0.92, "report_read_audits", "list/detail/image·actor/resource/time/purpose · APPEND-ONLY", accent=BLUE, fill=SOFT_BLUE, title_size=12, body_size=6.8)
    card(slide, 5.55, 5.02, 3.8, 0.92, "actor_rate_limit_events", "shared atomic 60초 window · non-audit", accent=TEAL, fill=SOFT_TEAL, title_size=12, body_size=7.3)
    line(slide, 4.95, 2.18, 5.55, 2.18, color=GOLD, width=2)
    line(slide, 4.95, 4.38, 5.55, 4.38, color=BLUE, width=2)
    rect(slide, 1.2, 5.86, 3.2, 0.78, fill=SOFT_BLUE, line=LINE, radius=True)
    rect(slide, 1.2, 5.86, 0.08, 0.78, fill=BLUE, line=BLUE)
    textbox(slide, "DB 밖 이미지 파일", 1.43, 6.02, 1.55, 0.25, size=11.2, color=BLUE, bold=True)
    textbox(slide, "durable journal · startup reconciliation", 2.77, 6.04, 1.45, 0.22, size=7.5, color=MUTED, align=PP_ALIGN.RIGHT)
    label(slide, "status/export/read 3종 감사만 변경 금지", 5.62, 6.15, 3.65, fill=NAVY)

    # 16. Risk flow
    slide = slides[15]
    header(slide, "기능 처리도 1/3", "위험은 감지 즉시가 아니라 안정화·경로 판단 뒤 안내", 16, cover_logo)
    flow(slide, [
        ("새 프레임", "기본 450ms 간격\n완료 frame만 사용"),
        ("인스턴스", "bbox·class별\n동일 객체 추적"),
        ("안정화", "서로 다른 3 frame\n700ms 지속"),
        ("위험 근거", "경로·접근·TTC\n신뢰 거리"),
        ("피드백", "우선순위·cooldown\nTTS·고위험 진동"),
    ], 0.55, 2.0, 8.9, 1.83, colors=[BLUE, BLUE, TEAL, GOLD, CORAL], body_size=9.8)
    rows = [
        ["근거 부족", "처리"],
        ["stale·frozen·server unavailable", "판단 중지 · 기존 위험 해제"],
        ["bbox 변화만 있는 stop 후보", "독립된 신뢰 거리 없으면 warning으로 제한"],
        ["경로 밖 또는 접근 근거 없음", "화면 표시만 · TTS·진동 없음"],
    ]
    add_table(slide, rows, 1.05, 4.25, 7.9, 1.62, widths=[3.3, 4.3], font_size=10.8)
    label(slide, "탐지됨 ≠ 경고함", 3.73, 6.32, 2.55, fill=NAVY)

    # 17. Report flow
    slide = slides[16]
    header(slide, "기능 처리도 2/3", "신고는 후보·서버·사람의 세 신뢰 단계를 통과", 17, cover_logo)
    flow(slide, [
        ("후보 gate", "손상 class·conf≥0.70\nGPS≤15m·3 frame·700ms"),
        ("동일 촬영", "image·위치·시각\nmodel metadata"),
        ("서버 재검증", "class·model·GPS\n시각·MIME"),
        ("원자 저장", "journal→rename→DB\nstartup reconciliation"),
        ("사람 검수", "근거·중복 확인\n기관용 자료 준비"),
    ], 0.55, 1.95, 8.9, 1.88, colors=[GOLD, TEAL, BLUE, GREEN, CORAL], body_size=9.5)
    card(slide, 0.9, 4.28, 2.65, 1.38, "자동 신고", "조건 충족 시 무음 저장\n완료 TTS 없음", accent=TEAL, fill=SOFT_TEAL, title_size=14, body_size=10.8)
    card(slide, 3.68, 4.28, 2.65, 1.38, "음성 요청 신고", "사용자가 명시하면\n성공·실패를 짧게 안내", accent=GOLD, fill=SOFT_GOLD, title_size=14, body_size=10.8)
    card(slide, 6.46, 4.28, 2.65, 1.38, "외부 제출", "자동 API 없음\n관리자가 수동 제출", accent=CORAL, fill=SOFT_CORAL, title_size=14, body_size=10.8)
    label(slide, "중복 후보는 자동 병합하지 않고 관리자 검수 정보로만 사용", 2.45, 6.2, 5.1, fill=NAVY)

    # 18. Navigation flow
    slide = slides[17]
    header(slide, "기능 처리도 3/3", "TMAP 전역 경로는 유지하고 점자블록은 짧게 보조", 18, cover_logo)
    flow(slide, [
        ("목적지", "음성 검색\n후보 번호 선택"),
        ("전역 경로", "TMAP 보행 경로\n계단 회피"),
        ("진행 판단", "GPS·heading\nguide point"),
        ("이탈", "이전 안내 중지\n재탐색"),
    ], 0.62, 1.95, 8.76, 1.65, colors=[BLUE, TEAL, GOLD, CORAL])
    card(slide, 0.85, 4.05, 3.75, 1.55, "정상 점자블록 local steering", "3 frame·700ms · conf≥0.55 · age≤1,200ms · GPS≤25m · heading≤35° · corridor 조건을 모두 확인", accent=TEAL, fill=SOFT_TEAL, title_size=14, body_size=10.9)
    card(slide, 5.15, 4.05, 3.75, 1.55, "즉시 TMAP 복귀", "손상 · 미검출 · 오래된 frame · 위치/방향 부정확 · 경로 밖이면 점자블록 보조를 해제", accent=CORAL, fill=SOFT_CORAL, title_size=14, body_size=10.9)
    label(slide, "점자블록을 전체 경로망이나 안전 보장으로 표현하지 않음", 2.15, 6.18, 5.7, fill=NAVY)

    # 19. Program list
    slide = slides[18]
    header(slide, "프로그램 목록", "모듈은 카메라 입력부터 검수까지 한 책임씩 연결", 19, cover_logo)
    rows = [
        ["영역", "대표 프로그램", "책임"],
        ["Capture", "page.tsx · useCamera.ts · useDetectionV2.ts", "카메라·센서·탐지 연결"],
        ["Risk", "risk-instance-tracker.ts · risk-evaluator.ts", "안정화·경로 위험 판단"],
        ["Route", "useNavigationGuidance.ts · tactile-route-policy.ts", "TMAP 진행·점자블록 보조"],
        ["Voice", "useVoiceCommands.ts · voice-intent-executor.ts", "목적지·신고·길안내 명령"],
        ["Report", "reports.py · report_storage.py", "journal 저장·시작 조정"],
        ["Admin", "admin/page · report_read_audit.py", "필터·상태·반출·조회 감사"],
        ["Model", "detect_v2.py · yolo_inference_adapter.py", "13-class 추론"],
        ["Android", "MainActivity.kt · depth/ · inference/", "ARCore·TFLite 연구 보조"],
    ]
    add_table(slide, rows, 0.52, 1.82, 8.96, 4.95, widths=[1.1, 4.5, 2.7], font_size=10.7)

    # 20. Reports table
    slide = slides[19]
    header(slide, "테이블 정의서 1/2", "reports는 신고 한 건의 탐지·위치·근거를 함께 저장", 20, cover_logo)
    groups = [
        (0.65, "식별·상태", "id UUID\nstatus\ncaptured_at\ncreated_at\nupdated_at", BLUE),
        (2.85, "탐지", "class_name\nconfidence\nbbox x/y/w/h\nsource", TEAL),
        (5.05, "위치", "latitude·longitude\nPoint(4326)\naccuracy·heading\nGiST index", GOLD),
        (7.25, "근거", "image_path·MIME\nmetadata JSONB\nmodel_key·trigger\nauto_reported", CORAL),
    ]
    for x, title, body, accent in groups:
        card(slide, x, 1.9, 2.05, 2.7, title, body, accent=accent, fill=WHITE, title_size=15, body_size=11.3)
    rows = [
        ["정합성", "DB 제약"],
        ["confidence·bbox·GPS 값 범위", "CHECK constraint"],
        ["위·경도와 Point(4326) 일치", "coordinate consistency"],
        ["상태·source·MIME 허용값", "enum/format validation"],
    ]
    add_table(slide, rows, 1.25, 5.05, 7.5, 1.42, widths=[3.8, 3.2], font_size=10.6)

    # 21. Audit tables
    slide = slides[20]
    header(slide, "테이블 정의서 2/2", "3종 append-only 감사와 non-audit rate event를 구분", 21, cover_logo)
    card(slide, 0.6, 1.72, 4.25, 1.85, "report_status_audits", "report_id · previous/next · actor · note/reason · created_at\n역할: 상태 전이와 담당자 추적 · APPEND-ONLY", accent=GOLD, fill=SOFT_GOLD, title_size=15, body_size=10)
    card(slide, 5.15, 1.72, 4.25, 1.85, "report_export_audits", "audit/requested ID · actor · profile · rows hash · filters · time\n역할: 반출 목적·범위 추적 · APPEND-ONLY", accent=CORAL, fill=SOFT_CORAL, title_size=15, body_size=9.7)
    card(slide, 0.6, 4.02, 4.25, 1.85, "report_read_audits", "actor · purpose · resource type/id · details · created_at\n역할: list/detail/image 민감 조회 추적 · APPEND-ONLY", accent=BLUE, fill=SOFT_BLUE, title_size=15, body_size=9.7)
    card(slide, 5.15, 4.02, 4.25, 1.85, "actor_rate_limit_events", "actor digest · rate group · observed_at\n역할: PostgreSQL 공유 원자 60초 window · NON-AUDIT", accent=TEAL, fill=SOFT_TEAL, title_size=15, body_size=10)
    textbox(slide, "application table은 reports 포함 정확히 5개 · audit trigger는 status/export/read에만 적용", 1.2, 6.35, 7.6, 0.25, size=9.1, color=MUTED, align=PP_ALIGN.CENTER)

    # 22. Web code
    slide = slides[21]
    header(slide, "핵심 소스코드 1/2", "위험 코드는 손상 시설을 경고가 아닌 신고로 분기", 22, cover_logo)
    source = code_excerpt("apps/web/app/_walksafe/risk-evaluator.ts", "export function evaluateTwoModelDetectionRisk", 14)
    rect(slide, 0.62, 1.8, 5.65, 4.95, fill=RGBColor(248, 250, 252), line=LINE, radius=True)
    textbox(slide, source, 0.82, 2.0, 5.25, 4.5, size=10, color=INK, font=CODE_FONT)
    card(slide, 6.58, 1.9, 2.72, 1.42, "입력", "TwoModelDetection\nRiskEvaluationContext", accent=BLUE, fill=SOFT_BLUE, title_size=14, body_size=10.5)
    card(slide, 6.58, 3.55, 2.72, 1.42, "판정", "손상 점자블록이면\nalertable=false\nreportable=true", accent=TEAL, fill=SOFT_TEAL, title_size=14, body_size=10.5)
    card(slide, 6.58, 5.2, 2.72, 1.42, "출력", "경고 대신 신고 후보\n기본 사용자 발화 없음", accent=CORAL, fill=SOFT_CORAL, title_size=14, body_size=10.5)
    textbox(slide, "apps/web/app/_walksafe/risk-evaluator.ts", 0.78, 6.78, 5.4, 0.18, size=8.2, color=MUTED)

    # 23. Backend code
    slide = slides[22]
    header(slide, "핵심 소스코드 2/2", "신고 코드는 대표 ID를 지키고 내부 중복 후보를 숨김", 23, cover_logo)
    source = code_excerpt("backend/app/api/reports.py", "def _field_creation_response", 15)
    rect(slide, 0.62, 1.8, 5.65, 4.95, fill=RGBColor(248, 250, 252), line=LINE, radius=True)
    textbox(slide, source, 0.82, 2.0, 5.25, 4.5, size=9.6, color=INK, font=CODE_FONT)
    card(slide, 6.58, 1.9, 2.72, 1.42, "대표 신고", "primary report id와\nduplicate_count 유지", accent=BLUE, fill=SOFT_BLUE, title_size=14, body_size=10.5)
    card(slide, 6.58, 3.55, 2.72, 1.42, "응답 최소화", "다른 중복 후보 ID는\n사용자 응답에서 제거", accent=TEAL, fill=SOFT_TEAL, title_size=14, body_size=10.5)
    card(slide, 6.58, 5.2, 2.72, 1.42, "저장 안전", "durable journal→rename→DB\n시작 시 미완료 write 조정", accent=CORAL, fill=SOFT_CORAL, title_size=14, body_size=9.8)
    textbox(slide, "backend/app/api/reports.py · backend/app/uploads.py", 0.78, 6.78, 5.4, 0.18, size=8.2, color=MUTED)

    # 24. Status
    slide = slides[23]
    header(slide, "개발 현황 및 다음 단계", "연결 구현은 확보했고, 다음 증거는 실폰·실외·운영", 24, cover_logo)
    status = [
        (0.65, "구현", "Web/PWA · Backend · Admin\n음성 길안내 · 신고 흐름\nPostGIS 5 application tables", TEAL, SOFT_TEAL),
        (3.55, "검증", "Web build·정책 자동시험\n격리 DB 기능 회귀\nAndroid TFLite invoke 1/1", GOLD, SOFT_GOLD),
        (6.45, "다음", "실폰 camera·GPS·mic\nTTS·진동·실외 길안내\nIAM·복구·보존·기관 접수", CORAL, SOFT_CORAL),
    ]
    for x, title, body, accent, fill in status:
        card(slide, x, 2.0, 2.55, 3.35, title, body, accent=accent, fill=fill, title_size=19, body_size=13)
    label(slide, "현재 전체 상태: 중간 구현 · 실사용 안전과 운영 배포를 보장하지 않음", 1.45, 5.88, 7.1, fill=NAVY)
    textbox(slide, "주제품 Web/PWA · Android 연구 보조 · 기관 제출은 관리자 검수 후 사람의 외부 행위", 1.05, 6.5, 7.9, 0.3, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)

    # Remove comments/notes/people-like presentation relationships.
    for rel_id, rel in list(prs.part.rels.items()):
        reltype = rel.reltype.lower()
        if any(token in reltype for token in ("authors", "comment", "notesmaster", "person")):
            prs.part.drop_rel(rel_id)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix="walksafe-midterm-", suffix=".pptx", dir=OUTPUT.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        prs.save(temporary)
        normalize_package(temporary)
        os.replace(temporary, OUTPUT)
    finally:
        temporary.unlink(missing_ok=True)

    if sha256(TEMPLATE) != EXPECTED_TEMPLATE_SHA256:
        raise SystemExit("source template changed while building")
    print(f"generated {OUTPUT} ({OUTPUT.stat().st_size} bytes, sha256={sha256(OUTPUT)})")


if __name__ == "__main__":
    build()

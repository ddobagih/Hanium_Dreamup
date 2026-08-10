#!/usr/bin/env python3
"""Build the template-faithful, practical-deliverable WalkSafe design deck."""

from __future__ import annotations

from copy import deepcopy
import os
from pathlib import Path
import tempfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

import build_midterm_design_ppt_template_faithful_20260712 as base


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path("/home/ddobagi/Downloads/제작설계서_완성본_실무산출물_완성본.pptx")
DATE_TEXT = "2026.07.12"

REQ_NAMES = {
    "REQ-001": "Web/PWA 카메라 기반 위험 감지",
    "REQ-002": "실탐지와 fake/demo 결과의 명시적 분리",
    "REQ-003": "Android ARCore·TFLite 연구 보조",
    "REQ-004": "객체별 안정 추적과 경로 위험 선별",
    "REQ-005": "손상 점자블록 전용 자동·음성 요청 신고",
    "REQ-006": "짧은 TTS·진동 위험 피드백",
    "REQ-007": "음성 신고·목적지·길안내 제어",
    "REQ-008": "신고 이미지·metadata의 신뢰 경계와 원자 저장",
    "REQ-009": "PostGIS 위치 저장과 검수 추적",
    "REQ-010": "관리자 검수·상태관리·기관용 자료 준비",
    "REQ-011": "설치 가능한 PWA와 안전한 offline 경계",
    "REQ-012": "목적지 검색과 보행 경로 안내",
    "REQ-013": "모델·데이터·class order 추적",
    "REQ-014": "개인정보·인증·반출·보존 경계",
}


SCREEN_SPECS = [
    {
        "id": "SC-00",
        "name": "field/admin 역할 인증",
        "menu": "앱 실행 > 현장 인증 / 관리자 인증",
        "overview": "현장 사용자와 관리자를 역할별 세션으로 분리한다.",
        "entry": "로그인되지 않은 / 또는 /admin 접속",
        "output": "12시간 HttpOnly field/admin 세션",
        "kind": "auth",
        "rows": [
            ("UI-001", "역할별 로그인", "현장 사용자와 관리자를 분리 인증", "route에 따라 field/admin session API 호출 후 세션 발급", "역할 세션 상호 대체 금지", ["REQ-014"]),
            ("UI-002", "인증 상태 안내", "제출·실패·잠금·만료 상태 표시", "invalid·locked·config-missing·expired를 구분해 안내", "운영 설정 누락은 503", ["REQ-014"]),
            ("UI-003", "로그아웃", "현재 actor 세션 종료", "cookie 삭제 후 field 보행 resource까지 정리", "token 원문 미표시", ["REQ-014"]),
        ],
    },
    {
        "id": "SC-01",
        "name": "Web/PWA 보행 보조 메인",
        "menu": "현장 인증 > 보행 보조 홈 > 실시간 보행",
        "overview": "카메라 영상과 탐지·길안내·신고 상태를 한 화면에 제공한다.",
        "entry": "field 인증, Web frame 기본 OFF 동의 및 카메라 권한 승인",
        "output": "실시간 overlay와 보행 보조 상태",
        "kind": "main",
        "asset": base.WEB_DESKTOP,
        "rows": [
            ("UI-004", "카메라 프리뷰", "후면 카메라 영상을 주 화면에 표시", "Web frame 동의와 명시적 권한 뒤 stream 연결, 새 frame만 capture", "미동의·철회 시 in-flight 취소", ["REQ-001", "REQ-014"]),
            ("UI-005", "탐지 overlay", "bbox·class·confidence 표시", "strict parser 통과 결과만 표시하고 stale은 제외", "WebXR depth는 위험 입력 아님", ["REQ-001"]),
            ("UI-006", "탐지 출처", "real과 fake/demo를 명시적으로 구분", "mode·source·model 상태를 화면에 표시", "real 실패 시 fake 전환 금지", ["REQ-002"]),
            ("UI-007", "보행 보조 중지", "카메라·센서·음성·길안내 종료", "stop·hidden·pagehide·인증 상실 시 resource 해제", "재시작 시 재초기화", ["REQ-001", "REQ-014"]),
        ],
    },
    {
        "id": "SC-02",
        "name": "권한·연결 상태",
        "menu": "보행 보조 홈 > 설정·세션 > 권한·연결",
        "overview": "권한·센서·네트워크 실패를 안전 상태와 구분해 표시한다.",
        "entry": "보행 보조 홈 진입 또는 장치 상태 변경",
        "output": "camera·GPS·heading·mic·API 상태",
        "kind": "status",
        "rows": [
            ("UI-008", "카메라 권한", "권한 상태와 오류 원인 안내", "prompt·denied·unavailable 구분, 버튼으로 재요청", "비HTTPS·장치 점유 안내", ["REQ-001"]),
            ("UI-009", "위치·방향", "GPS 정확도와 absolute heading 표시", "sensor listener로 갱신하고 낮은 정확도를 분리", "없어도 탐지는 시작 가능", ["REQ-001", "REQ-012"]),
            ("UI-010", "마이크 상태", "음성 명령의 권한·지원 여부 표시", "녹음 시작 후 거부·미지원·중단을 오류로 전환", "사전 PermissionState 아님", ["REQ-007"]),
            ("UI-011", "온라인·API", "offline/backend/model 불가를 구분", "offline banner·API unavailable 표시, 안전 요청 차단", "shell 외 기능 network-only", ["REQ-002", "REQ-011"]),
        ],
    },
    {
        "id": "SC-03",
        "name": "위험 안내 상태",
        "menu": "실시간 보행 > 객체·경로 위험 > TTS·진동",
        "overview": "객체 추적 근거가 충분한 경우에만 짧은 위험 피드백을 제공한다.",
        "entry": "검증된 v2 탐지 결과 수신",
        "output": "display-only 또는 TTS·진동 위험 안내",
        "kind": "risk",
        "rows": [
            ("UI-012", "객체 안정 추적", "동일 객체의 frame 이력 관리", "source·model·class+IoU로 1:1 연결, 3 frame·700ms 안정화", "2회 빈 frame에서 해제", ["REQ-004"]),
            ("UI-013", "경로 위험 판정", "존재가 아닌 차단·접근 위험 선별", "bbox·ROI·motion·route context로 active 위험 선택", "bbox TTC 단독 STOP 금지", ["REQ-004"]),
            ("UI-014", "TTS 위험 안내", "방향·유형·행동을 짧게 발화", "risk>interaction>navigation 우선순위와 객체별 cooldown", "medium은 주의/피하기", ["REQ-006"]),
            ("UI-015", "high 진동", "high에서만 강한 피드백 제공", "정지 안내와 haptic, 별도 진동 cooldown 적용", "단말 미지원 시 TTS·UI", ["REQ-006"]),
        ],
    },
    {
        "id": "SC-06",
        "name": "목적지 검색·보행 길안내",
        "menu": "보행 보조 홈 > 목적지 길안내 > TMAP 도보 안내",
        "overview": "목적지 후보를 선택하고 TMAP 계단 회피 경로를 음성으로 안내한다.",
        "entry": "field 세션, 유효 GPS, 목적지 검색어",
        "output": "경로·다음 안내·이탈·도착 상태",
        "kind": "navigation",
        "rows": [
            ("UI-024", "목적지 검색", "음성 검색어와 현재 위치로 POI 조회", "TMAP 후보의 provider·좌표를 strict 검증", "텍스트 검색창 없음", ["REQ-007", "REQ-012"]),
            ("UI-025", "후보 선택", "후보·더 보기·재검색·취소 제공", "여러 결과나 거리 미상은 명시 선택", "잘못된 번호는 no-op", ["REQ-012"]),
            ("UI-026", "도보 길안내", "TMAP 경로와 다음 guide 표시", "STAIR_AVOID route·guide point·현재 단계를 갱신", "provider fallback 없음", ["REQ-012"]),
            ("UI-027", "이탈·도착", "경로 진행과 예외 상태 관리", "accuracy gate·제한 재탐색·2-sample 도착 판정", "점자블록은 국소 보조", ["REQ-012"]),
        ],
    },
    {
        "id": "SC-04",
        "name": "손상 점자블록 신고",
        "menu": "보행 보조 홈 > 손상 점자블록 신고",
        "overview": "손상 점자블록만 동일 촬영 근거로 자동 또는 음성 요청 신고한다.",
        "entry": "Web report 동의 뒤 damaged class 탐지 또는 명시 음성 요청",
        "output": "primary report ID와 duplicate count",
        "kind": "report",
        "rows": [
            ("UI-016", "신고 후보 gate", "손상 점자블록만 후보로 선별", "conf·GPS·3 frame·700ms·gap·jump·cooldown 검사", "일반 객체 신고 금지", ["REQ-005"]),
            ("UI-017", "자동·음성 요청", "자동 또는 명시 음성으로 신고", "Web report/Android report 각각 기본 OFF 동의 뒤 동일 snapshot 사용", "철회 시 in-flight 취소", ["REQ-005", "REQ-007", "REQ-014"]),
            ("UI-018", "서버 검증·저장", "입력과 파일/DB를 재검증·저장", "durable journal→원자 rename→DB commit, startup reconciliation", "PostGIS·audit 연계", ["REQ-008", "REQ-009"]),
            ("UI-019", "전송 결과", "대기·전송·성공·실패 표시", "aria-live·haptic; 음성 요청은 짧은 TTS", "offline queue 없음", ["REQ-005"]),
        ],
    },
    {
        "id": "SC-05",
        "name": "음성 제어 상태",
        "menu": "보행 보조 홈 > 음성 제어",
        "overview": "음성을 strict intent로 해석해 신고·목적지·길안내를 fail-closed 실행한다.",
        "entry": "사용자가 음성 입력 버튼 실행",
        "output": "허용 action 실행과 TTS·aria-live 결과",
        "kind": "voice",
        "rows": [
            ("UI-020", "음성 녹음", "명령 녹음 상태 제공", "MediaRecorder 시작·중지와 업로드·분석 상태 표시", "mic 거부·미지원 처리", ["REQ-007"]),
            ("UI-021", "STT·intent 검증", "응답을 실행 전 엄격 검증", "execute·should_execute=true·confidence/score≥0.70·acoustic gate 확인", "malformed·낮은 신뢰도는 실행 안 함", ["REQ-007"]),
            ("UI-022", "명령 실행", "수락된 intent를 hook 또는 navigation executor로 분기", "신고·음성·반복·현재 위치는 hook; 목적지·길안내는 callback/local fallback 처리", "unknown·부정문·범위 밖은 재안내", ["REQ-007"]),
            ("UI-023", "음성 우선순위", "위험·상호작용·길안내 중재", "risk>interaction>navigation; 위험은 녹음 취소 가능", "일반 TTS는 녹음 중 보류", ["REQ-006", "REQ-007"]),
        ],
    },
    {
        "id": "SC-11",
        "name": "사용자 설정·현장 기록·PWA",
        "menu": "보행 보조 홈 > 설정·세션 > PWA·현장 기록",
        "overview": "Web frame·Web report·Android report·telemetry 동의와 PWA 상태를 각각 관리한다.",
        "entry": "설정 패널 확장",
        "output": "로컬 설정·동의 상태·PWA 설치/업데이트 상태",
        "kind": "settings",
        "rows": [
            ("UI-040", "음성 안내 설정", "위험 TTS on/off", "설정 저장 후 즉시 feedback 정책에 반영", "OFF여도 화면 상태 유지", ["REQ-006"]),
            ("UI-041", "긴급 연락처", "최대 3개 연락처 로컬 저장", "추가·삭제·저장·초기화", "자동 외부 전송 없음", ["REQ-014"]),
            ("UI-042", "보폭 측정", "움직임 기반 보폭 측정·재보정", "motion permission·sample·confidence·저장값 표시", "길안내 ‘보’ 표현에 사용", ["REQ-012"]),
            ("UI-043", "용도별 동의", "Web frame/report·Android report·telemetry 각각 기본 OFF", "각 scope를 별도 동의·철회하고 철회 시 in-flight 취소", "telemetry만 최대 7일", ["REQ-014"]),
            ("UI-044", "PWA 설치·갱신", "설치·update waiting 상태 관리", "명시 install, 승인 뒤 waiting SW 적용", "안전 API offline 성공 금지", ["REQ-011"]),
        ],
    },
    {
        "id": "SC-07",
        "name": "Android 연구·검증 화면",
        "menu": "메뉴 외 · Android 연구 보조 앱",
        "overview": "주제품 메뉴와 분리된 앱에서 ARCore depth와 TFLite 정합을 검증한다.",
        "entry": "ARCore 지원 Android, camera 권한, 검증된 model asset",
        "output": "bbox·depth overlay, load/invoke 상태, app-private 현장 로그",
        "kind": "android",
        "rows": [
            ("UI-045", "ARCore 카메라·깊이", "카메라와 Raw/Full Depth 입력을 표시", "ARCore session·camera preview·depth snapshot 연결", "Web/PWA 주제품과 분리", ["REQ-003"]),
            ("UI-046", "unified TFLite 탐지", "13-class 기기 내 탐지와 bbox 표시", "artifact hash·class order·tensor 계약 확인 후 invoke", "unified 계약 실패 때만 legacy 연구 fallback", ["REQ-002", "REQ-003", "REQ-013"]),
            ("UI-047", "bbox-depth overlay", "탐지 bbox와 depth sample 정합 표시", "유효 sample·거리·confidence·source를 개발 overlay로 표시", "실기기 정합 검증 필요", ["REQ-003", "REQ-004"]),
            ("UI-048", "Device gate·현장 로그", "연구 기능 증거를 app-private 로그로 기록", "명시 시작·종료, metadata-only JSONL과 TTS/haptic/report 후보 상태 저장", "raw image·depth·GPS 미기록", ["REQ-003", "REQ-007"]),
        ],
    },
    {
        "id": "SC-08",
        "name": "Admin 신고 목록·필터·위치 분포",
        "menu": "관리자 인증 > 신고 관리 > 목록·상세 검수",
        "overview": "동일 필터로 신고 목록·요약·위치 분포를 조회한다.",
        "entry": "admin 세션 인증",
        "output": "페이지 목록·집계·정적 Heatmap",
        "kind": "admin-list",
        "asset": base.ADMIN_DESKTOP,
        "rows": [
            ("UI-028", "신고 목록", "최신순 50건 단위 조회", "actor·report_list·time·purpose read audit 뒤 pagination", "purpose 누락·감사 실패 시 거부", ["REQ-009", "REQ-010"]),
            ("UI-029", "구조화 필터", "status·class·source·model·trigger·auto_reported·날짜·반경·Fake/Demo", "목록·summary·반출에 동일 query 적용", "페이지와 전체 집계 구분", ["REQ-010"]),
            ("UI-030", "운영 요약", "상태·source·위치 유무 집계", "전체 필터 기준 status·bounds·Top 5 표시", "위치 없음 별도 집계", ["REQ-009", "REQ-010"]),
            ("UI-031", "위치 분포", "0.001° 격자 Heatmap", "cell count·중심·bounds·150m 반경 필터", "외부 basemap 미사용", ["REQ-009", "REQ-010"]),
        ],
    },
    {
        "id": "SC-09",
        "name": "Admin 신고 상세·상태 변경",
        "menu": "신고 관리 > 목록·상세 검수 > 상태 변경",
        "overview": "신고 근거를 확인하고 낙관적 잠금으로 상태·메모를 갱신한다.",
        "entry": "목록에서 신고 선택",
        "output": "상태 변경과 append-only 감사 이력",
        "kind": "admin-detail",
        "rows": [
            ("UI-032", "신고 상세", "이미지·위치·confidence·metadata 확인", "detail/image 각각 actor·resource·time·purpose read audit", "감사 실패 시 민감 조회 거부", ["REQ-009", "REQ-013"]),
            ("UI-033", "검토 표시", "중복·위치 품질·Fake/Demo 확인", "duplicate IDs·review flag·성능 제외 표시", "기관 적격성과 성능 분리", ["REQ-010"]),
            ("UI-034", "상태 변경", "new→reviewed→resolved 처리", "note·reason·expected_updated_at PATCH와 audit 저장", "resolved→new 금지", ["REQ-009", "REQ-010"]),
            ("UI-035", "충돌·세션 오류", "동시 수정과 인증 만료 처리", "409 재조회·재확인, 만료 시 admin gate 복귀", "자동 덮어쓰기 금지", ["REQ-014"]),
        ],
    },
    {
        "id": "SC-10",
        "name": "Admin 자료 내보내기",
        "menu": "신고 관리 > 상태·위치 분포·내보내기",
        "overview": "목적별 반출 프로필을 적용하고 생성 이력을 append-only로 남긴다.",
        "entry": "관리자 필터와 반출 profile 선택",
        "output": "CSV·JSON·GeoJSON·manifest와 반출 감사",
        "kind": "admin-export",
        "rows": [
            ("UI-036", "내부 자료", "검수용 전체 필드 다운로드", "internal CSV·JSON·GeoJSON 생성", "최대 10,000행", ["REQ-010"]),
            ("UI-037", "최소정보", "공유용 위치·민감정보 축약", "minimum profile 적용 후 다운로드", "redaction 혼합 금지", ["REQ-014"]),
            ("UI-038", "기관용 묶음", "적격 신고의 agency 자료 준비", "named reviewer·reviewed·damage·위치 정확도·non-fake gate", "기관 자동 제출 없음", ["REQ-010", "REQ-014"]),
            ("UI-039", "반출 감사", "누가 어떤 조건으로 생성했는지 기록", "append-only 반출 감사와 동일 audit_id manifest", "외부 receipt는 수동", ["REQ-009", "REQ-014"]),
        ],
    },
]


PROGRAMS = [
    ("보행 안전", "F-001", "위험 감지"),
    ("보행 안전", "F-002", "거리·접근/경로 보조"),
    ("보행 안전", "F-003", "위험 안내"),
    ("시설 신고", "F-004", "자동/음성 요청 신고"),
    ("음성 제어", "F-005", "음성 제어·신고"),
    ("운영 관리", "F-006", "운영자 검수"),
    ("운영 관리", "F-007", "Export·기관용 자료 생성"),
    ("길안내", "F-008", "목적지 안내"),
    ("모델 연구", "F-009", "모델 연구·후보 관리"),
    ("운영 보호", "F-010", "개인정보/현장 기록"),
    ("PWA", "F-011", "PWA 설치·offline 경계"),
]


FLOW_SPECS = [
    ("F-001", "WSA-PRG-F001", "Web/PWA 위험 감지", "field 실탐지와 기본 fake contract를 분리하고 검증된 real 결과만 위험 판단에 전달한다.", ["field 세션·새 frame", "/detect/v2\nstrict parse", "freshness·중복 확인", "real 객체 선별", "후속 판단 전달"], "실탐지·응답 유효?", "fake/demo → 데모 표시만; malformed·stale·503 → unavailable; real→fake 자동 전환 없음", ["REQ-001", "REQ-002"]),
    ("F-002", "WSA-PRG-F002", "객체 추적·위험 분석", "동일 객체의 안정화·경로 근거를 모두 확인해 표시 후보와 경고 후보를 구분한다.", ["유효 detection", "IoU·3f/700ms·ROI·접근", "RiskDecision 생성", "display/alert 분리", "후속 피드백 전달"], "경고 근거 충분?", "track 불연속·근거 부족·stale 거리 → display-only", ["REQ-004"]),
    ("F-003", "WSA-PRG-F003", "TTS·진동 위험 안내", "우선순위와 cooldown에 따라 짧은 TTS와 high 진동을 출력한다.", ["alertable 후보", "level·우선순위", "짧은 방향 TTS", "high만 진동", "UI 상태 갱신"], "피드백 출력 가능?", "alertable 없음·cooldown·음성 OFF → 시각 표시, 중복 발화 억제", ["REQ-006"]),
    ("F-004", "WSA-PRG-F004", "손상 점자블록 신고", "검증한 동일 촬영 근거를 durable journal과 PostGIS 경계로 저장한다.", ["damaged 후보", "용도별 동의 gate", "동일 snapshot", "journal→rename→DB", "startup reconciliation"], "신고 gate 통과?", "미동의·GPS·MIME 실패 → 미전송/4xx; corrupt journal → startup fail-closed", ["REQ-005", "REQ-008", "REQ-009", "REQ-014"]),
    ("F-005", "WSA-PRG-F005", "음성 명령 처리", "상위 실행 gate를 통과한 intent만 hook 직접 처리 또는 navigation executor로 분기한다.", ["음성 녹음", "STT·schema·실행 gate", "hook/nav intent\n분기", "callback/local\nfallback", "TTS·aria-live"], "실행 gate 통과?", "mic·timeout·malformed·낮은 신뢰도 → 실행 안 함; 부정문·범위 밖 번호 → 재안내", ["REQ-007"]),
    ("F-006", "WSA-PRG-F006", "관리자 신고 검수", "신고 근거와 next state·expected_updated_at을 검증해 상태·감사 이력을 원자 갱신한다.", ["admin 세션 확인", "상세·next state\nversion 입력", "상태·audit\ntransaction", "commit", "재조회·화면 갱신"], "인증·전이·버전 유효?", "401/403·409 충돌 → 재조회; 422 전이 거부; DB 실패 → rollback", ["REQ-009", "REQ-010"]),
    ("F-007", "WSA-PRG-F007", "기관용 자료 생성·반출", "검수된 신고를 기관용 파일로 생성해 관리자가 내려받으며 외부 제출은 시스템 범위 밖이다.", ["필터·profile", "role·row cap", "필드·좌표 변환", "audit·file\nmanifest", "관리자 다운로드"], "profile gate 통과?", "10,000행 초과·agency 부적격·audit 실패 → 반출 보류", ["REQ-009", "REQ-010", "REQ-014"]),
    ("F-008", "WSA-PRG-F008", "TMAP 보행 길안내", "TMAP 전역 경로와 조건부 점자블록 보조를 결합해 음성으로 안내한다.", ["GPS·목적지", "TMAP POI 후보", "명시 후보 선택", "TMAP route·진행", "tactile 보조·복귀·도착"], "후보·GPS 유효?", "GPS/TMAP 실패 → 안내 pause; tactile gate 실패 → TMAP 유지", ["REQ-012"]),
    ("F-009", "WSA-PRG-F009", "모델·Android 연구 관리", "release gate와 backend field·Android 연구 runtime의 서로 다른 실패 경계를 검증한다.", ["dataset\nmodel 후보", "hash·split\nclass·tensor", "registry 상태 기록", "backend / Android\n경로 분리", "연구 결과 보관"], "release gate 통과?", "FAIL: deployment_eligible=false\nbackend field: fail-closed\nAndroid: unified 계약 실패 시 legacy", ["REQ-002", "REQ-003", "REQ-013"]),
    ("F-010", "WSA-PRG-F010", "인증·개인정보·현장 기록", "역할 세션과 네 가지 독립 동의·조회 감사·배포 보호 경계를 적용한다.", ["field/admin 로그인", "Web frame/report 동의", "Android report/telemetry 동의", "철회→in-flight 취소", "read audit·배포 gate"], "역할 session·동의 유효?", "역할 교차·미동의·purpose 누락·감사 실패 → 401/403/거부", ["REQ-014"]),
    ("F-011", "WSA-PRG-F011", "PWA 설치·업데이트 관리", "명시적으로 활성화된 PWA의 설치·업데이트·offline shell을 관리한다.", ["PWA flag 확인", "SW 등록", "shell cache", "waiting worker", "사용자 승인 적용"], "PWA·SW 사용 가능?", "기본 OFF·등록 실패 → 일반 Web; 안전 API는 항상 network-only", ["REQ-011"]),
]


CODE_SPECS = [
    ("핵심 소스코드(1)", "F-002·F-003", "apps/web/app/_walksafe/risk-evaluator.ts", "export function isStableTracking", "탐지 결과의 안정성과 거리 근거를 먼저 검증해 raw detection이 곧바로 위험 경고로 사용되지 않도록 막는 핵심 판단부다.", ["stable_frames≥3 AND stable_ms≥700", "model_estimate confidence·finite·range 검증", "검증 실패 시 거리 근거 제거"], ["REQ-004", "REQ-006"]),
    ("핵심 소스코드(2)", "F-004", "backend/app/api/reports.py", "def _commit_new_report_with_image", "durable journal·원자 rename·PostGIS commit과 startup reconciliation으로 process death 경계까지 신고 증거를 조정한다.", ["journal을 먼저 durable write", "rename 뒤 DB commit", "startup에서 미완료 write 조정"], ["REQ-005", "REQ-008", "REQ-009"]),
    ("핵심 소스코드(3)", "F-005·F-008", "apps/web/app/_walksafe/voice-intent-executor.ts", "export async function executeNavigationVoiceIntent", "상위 실행 gate를 통과한 navigation intent를 분기하고 선택적 callback 또는 local fallback 결과를 UI·TTS·진동에 반영한다.", ["목적지 설정·후보 선택·취소 분기", "길안내 시작·재탐색·중지·다음 안내", "목적지 없음·잘못된 번호·일부 callback 부재는 실패 안내"], ["REQ-007", "REQ-012"]),
    ("핵심 소스코드(4)", "F-008", "apps/web/app/_walksafe/tactile-route-policy.ts", "export function evaluateTactileRoutePolicy", "점자블록을 전역 경로로 과장하지 않고 엄격한 gate를 통과한 짧은 구간에서만 TMAP 보조로 사용한다.", ["TMAP 경로가 기준", "freshness·GPS·heading·corridor gate", "실패·손상 탐지 시 즉시 TMAP 복귀"], ["REQ-012"]),
    ("핵심 소스코드(5)", "F-006·F-007·F-010", "apps/web/app/api/_backend.ts", "export function createBackendActorAssertion", "field/admin 역할 세션을 짧은 HMAC actor assertion으로 백엔드에 연결해 정확 위치·반출 행위 주체를 추적한다.", ["same-origin BFF에서만 서명", "30초 만료와 role scope", "secret·token 원문을 UI/로그에 노출하지 않음"], ["REQ-009", "REQ-010", "REQ-014"]),
    ("핵심 소스코드(6)", "F-009", "apps/android/app/src/main/java/kr/co/hanium/dreamup/walksafe/inference/TfliteAndroidFrameDetector.kt", "fun validateTensorContract", "모델 이름이 아니라 입력·출력 tensor와 class 계약을 확인해 잘못된 Android 연구 artifact 실행을 막는다.", ["[1,768,768,3] 입력 계약", "[1,300,6] 출력 계약", "hash·class 불일치 시 unified 배포 금지"], ["REQ-002", "REQ-003", "REQ-013"]),
]


def clone_slide(prs: Presentation, source_slide):
    """Clone a slide inside one package while preserving its relationship IDs."""
    destination = prs.slides.add_slide(source_slide.slide_layout)
    for shape in list(destination.shapes):
        base.remove_shape(destination, shape)
    for shape in source_slide.shapes:
        destination.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")
    for rel_id, rel in source_slide.part.rels.items():
        reltype = rel.reltype.lower()
        if "slidelayout" in reltype or "notesslide" in reltype:
            continue
        new_rel_id = destination.part.rels._add_relationship(rel.reltype, rel._target, rel.is_external)
        if new_rel_id != rel_id:
            raise RuntimeError(f"relationship ID changed while cloning: {rel_id} -> {new_rel_id}")
    return destination


def reorder_slides(prs: Presentation, ordered_slides) -> None:
    slide_id_list = prs.slides._sldIdLst
    element_by_part = {
        prs.part.rels[slide_id.rId].target_part: slide_id
        for slide_id in list(slide_id_list)
    }
    if len(ordered_slides) != len(element_by_part) or len({slide.part for slide in ordered_slides}) != len(ordered_slides):
        raise RuntimeError("invalid final slide order")
    for slide_id in list(slide_id_list):
        slide_id_list.remove(slide_id)
    for slide in ordered_slides:
        slide_id_list.append(element_by_part[slide.part])


def add_arrowhead(shape) -> None:
    line = shape._element.spPr.get_or_add_ln()
    for element in list(line):
        if element.tag.endswith("tailEnd"):
            line.remove(element)
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "sm")
    line.append(tail)


def diagram_connector(slide, x1, y1, x2, y2, *, color=base.MUTED, width=1.3, dash=False, elbow=False):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    if dash:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_arrowhead(connector)
    return connector


def requirement_text(requirement_ids) -> str:
    return "\n".join(f"{requirement_id} {REQ_NAMES[requirement_id]}" for requirement_id in requirement_ids)


def fit_table_rows(table_shape, total_height: float, body_count: int, *, header_height=0.42, header_font=7.0) -> None:
    table_shape.table.rows[0].height = Inches(header_height)
    body_height = (total_height - header_height) / body_count
    for row_index in range(1, len(table_shape.table.rows)):
        table_shape.table.rows[row_index].height = Inches(body_height)
    for cell in table_shape.table.rows[0].cells:
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(header_font)


def draw_wireframe(slide, spec, x, y, w, h) -> None:
    asset = spec.get("asset")
    if asset:
        base.add_picture_fit(slide, asset, x, y, w, h)
        base.ppt_rect(slide, x + 0.08, y + 0.08, 0.62, 0.26, fill=base.BLUE, line_color=base.BLUE, radius=True)
        base.textbox(slide, spec["id"], x + 0.12, y + 0.12, 0.54, 0.15, size=7.2, color=base.WHITE, bold=True, align=PP_ALIGN.CENTER)
        return

    base.ppt_rect(slide, x, y, w, h, fill=base.PALE, line_color=base.LINE, radius=True)
    base.ppt_rect(slide, x, y, w, 0.34, fill=base.BLUE, line_color=base.BLUE)
    base.textbox(slide, spec["name"], x + 0.12, y + 0.07, w - 0.24, 0.20, size=8.4, color=base.WHITE, bold=True)
    kind = spec["kind"]

    if kind == "auth":
        base.ppt_rect(slide, x + 0.65, y + 0.52, w - 1.30, 1.28, fill=base.WHITE, line_color=base.LINE, radius=True)
        base.textbox(slide, "WalkSafe 로그인", x + 0.90, y + 0.66, w - 1.80, 0.22, size=9.0, color=base.INK, bold=True, align=PP_ALIGN.CENTER)
        base.ppt_rect(slide, x + 0.90, y + 1.02, w - 1.80, 0.25, fill=base.PALE, line_color=base.LINE)
        base.ppt_rect(slide, x + 0.90, y + 1.40, w - 1.80, 0.27, fill=base.BLUE, line_color=base.BLUE, radius=True)
        base.textbox(slide, "역할 인증", x + 0.96, y + 1.45, w - 1.92, 0.15, size=7.2, color=base.WHITE, bold=True, align=PP_ALIGN.CENTER)
    elif kind == "status":
        labels = [("카메라", base.BLUE), ("GPS·방향", base.TEAL), ("마이크", base.GOLD), ("API", base.CORAL)]
        for index, (label, color) in enumerate(labels):
            yy = y + 0.50 + index * 0.37
            base.ppt_rect(slide, x + 0.25, yy, w - 0.50, 0.27, fill=base.WHITE, line_color=color, radius=True)
            base.textbox(slide, label, x + 0.37, yy + 0.06, 0.80, 0.14, size=7.2, color=color, bold=True)
            base.textbox(slide, "정상 / 확인 필요", x + 1.25, yy + 0.06, w - 1.62, 0.14, size=6.8, color=base.MUTED, align=PP_ALIGN.RIGHT)
    elif kind == "risk":
        base.ppt_rect(slide, x + 0.20, y + 0.48, w - 0.40, 1.05, fill=RGBColor(49, 58, 68), line_color=base.LINE)
        base.ppt_rect(slide, x + 0.70, y + 0.68, 0.78, 0.60, fill=RGBColor(49, 58, 68), line_color=base.CORAL, width=1.6)
        base.textbox(slide, "사람 · 접근 중", x + 0.78, y + 0.74, 0.64, 0.16, size=6.6, color=base.WHITE, bold=True, align=PP_ALIGN.CENTER)
        base.ppt_rect(slide, x + 0.35, y + 1.67, w - 0.70, 0.26, fill=base.LIGHT_CORAL, line_color=base.CORAL, radius=True)
        base.textbox(slide, "전방 위험 · 멈추세요 · 진동", x + 0.45, y + 1.72, w - 0.90, 0.14, size=7.1, color=base.CORAL, bold=True, align=PP_ALIGN.CENTER)
    elif kind == "navigation":
        base.ppt_rect(slide, x + 0.18, y + 0.46, 1.18, 1.35, fill=base.WHITE, line_color=base.LINE)
        for index, value in enumerate(("1 서울역", "2 시청역", "3 남대문")):
            base.ppt_rect(slide, x + 0.30, y + 0.62 + index * 0.33, 0.94, 0.23, fill=base.PALE, line_color=base.LINE)
            base.textbox(slide, value, x + 0.36, y + 0.67 + index * 0.33, 0.82, 0.13, size=6.6, color=base.INK)
        base.ppt_rect(slide, x + 1.52, y + 0.46, w - 1.70, 1.35, fill=base.LIGHT_BLUE, line_color=base.BLUE)
        diagram_connector(slide, x + 1.75, y + 1.55, x + w - 0.43, y + 0.72, color=base.BLUE, width=2.0)
        base.textbox(slide, "다음 안내\n좌회전 준비", x + 1.72, y + 0.62, w - 2.02, 0.45, size=7.2, color=base.BLUE, bold=True, align=PP_ALIGN.CENTER)
    elif kind == "report":
        stages = [("후보", base.GOLD), ("검증", base.BLUE), ("저장", base.TEAL), ("결과", base.CORAL)]
        for index, (label, color) in enumerate(stages):
            xx = x + 0.18 + index * ((w - 0.36) / 4)
            base.ppt_rect(slide, xx, y + 0.72, 0.62, 0.62, fill=base.WHITE, line_color=color, radius=True)
            base.textbox(slide, label, xx + 0.05, y + 0.92, 0.52, 0.16, size=7.2, color=color, bold=True, align=PP_ALIGN.CENTER)
            if index < 3:
                base.textbox(slide, "→", xx + 0.65, y + 0.90, 0.20, 0.20, size=10, color=base.MUTED, bold=True, align=PP_ALIGN.CENTER)
        base.textbox(slide, "primary ID · duplicate count", x + 0.50, y + 1.58, w - 1.0, 0.20, size=7.0, color=base.INK, bold=True, align=PP_ALIGN.CENTER)
    elif kind == "voice":
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.28), Inches(y + 0.62), Inches(0.82), Inches(0.82))
        circle.fill.solid(); circle.fill.fore_color.rgb = base.BLUE; circle.line.color.rgb = base.BLUE
        base.textbox(slide, "MIC", x + 0.40, y + 0.91, 0.58, 0.16, size=8.2, color=base.WHITE, bold=True, align=PP_ALIGN.CENTER)
        base.ppt_rect(slide, x + 1.35, y + 0.53, w - 1.60, 0.50, fill=base.WHITE, line_color=base.LINE, radius=True)
        base.textbox(slide, "“서울역으로 안내해줘”", x + 1.48, y + 0.68, w - 1.86, 0.18, size=7.0, color=base.INK)
        base.ppt_rect(slide, x + 1.35, y + 1.18, w - 1.60, 0.48, fill=base.LIGHT_BLUE, line_color=base.BLUE, radius=True)
        base.textbox(slide, "intent: set_destination", x + 1.48, y + 1.32, w - 1.86, 0.18, size=7.0, color=base.BLUE, bold=True)
    elif kind == "settings":
        settings = ["음성 안내", "긴급 연락처", "보폭 측정", "현장 기록", "PWA 업데이트"]
        for index, label in enumerate(settings):
            yy = y + 0.47 + index * 0.29
            base.textbox(slide, label, x + 0.30, yy + 0.04, 1.35, 0.14, size=6.8, color=base.INK)
            base.ppt_rect(slide, x + w - 0.78, yy, 0.48, 0.20, fill=base.LIGHT_BLUE if index % 2 == 0 else base.PALE, line_color=base.BLUE, radius=True)
    elif kind == "android":
        base.ppt_rect(slide, x + 0.18, y + 0.46, 1.60, 1.38, fill=RGBColor(45, 53, 62), line_color=base.LINE)
        base.ppt_rect(slide, x + 0.55, y + 0.70, 0.70, 0.72, fill=RGBColor(45, 53, 62), line_color=base.GOLD, width=1.5)
        base.textbox(slide, "person\n2.1m", x + 0.62, y + 0.87, 0.56, 0.28, size=6.6, color=base.WHITE, bold=True, align=PP_ALIGN.CENTER)
        labels = [("ARCore Depth", base.BLUE), ("TFLite 13-class", base.TEAL), ("Device Gate", base.GOLD), ("JSONL Log", base.CORAL)]
        for index, (label, color) in enumerate(labels):
            yy = y + 0.48 + index * 0.34
            base.ppt_rect(slide, x + 1.97, yy, w - 2.17, 0.26, fill=base.WHITE, line_color=color, radius=True)
            base.textbox(slide, label, x + 2.07, yy + 0.06, w - 2.37, 0.14, size=6.7, color=color, bold=True)
    elif kind == "admin-list":
        base.ppt_rect(slide, x + 0.20, y + 0.48, w - 0.40, 0.28, fill=base.PALE, line_color=base.LINE)
        base.textbox(slide, "상태 · 유형 · 날짜 · 반경", x + 0.30, y + 0.55, w - 0.60, 0.14, size=6.7, color=base.MUTED)
        for index in range(4):
            yy = y + 0.88 + index * 0.25
            base.ppt_rect(slide, x + 0.20, yy, 1.82, 0.19, fill=base.WHITE, line_color=base.LINE)
            base.ppt_rect(slide, x + 2.15, yy, w - 2.35, 0.19, fill=base.LIGHT_TEAL, line_color=base.TEAL)
    elif kind == "admin-detail":
        base.ppt_rect(slide, x + 0.20, y + 0.48, 1.28, 1.10, fill=RGBColor(62, 69, 78), line_color=base.LINE)
        base.textbox(slide, "신고 이미지", x + 0.42, y + 0.92, 0.84, 0.18, size=7.0, color=base.WHITE, bold=True, align=PP_ALIGN.CENTER)
        for index, value in enumerate(("위치·bbox", "중복·품질", "new → reviewed", "메모·사유")):
            yy = y + 0.49 + index * 0.31
            base.ppt_rect(slide, x + 1.68, yy, w - 1.90, 0.24, fill=base.WHITE, line_color=base.LINE)
            base.textbox(slide, value, x + 1.80, yy + 0.05, w - 2.14, 0.14, size=6.8, color=base.INK)
    elif kind == "admin-export":
        profiles = [("internal", base.BLUE), ("minimum", base.TEAL), ("agency", base.CORAL)]
        for index, (label, color) in enumerate(profiles):
            xx = x + 0.22 + index * 1.02
            base.ppt_rect(slide, xx, y + 0.58, 0.86, 0.78, fill=base.WHITE, line_color=color, radius=True)
            base.textbox(slide, label, xx + 0.06, y + 0.75, 0.74, 0.18, size=6.9, color=color, bold=True, align=PP_ALIGN.CENTER)
            base.textbox(slide, "CSV·JSON\nGeoJSON", xx + 0.08, y + 1.00, 0.70, 0.25, size=6.1, color=base.MUTED, align=PP_ALIGN.CENTER)
        base.textbox(slide, "audit + manifest · 기관 제출은 수동", x + 0.36, y + 1.60, w - 0.72, 0.18, size=6.8, color=base.INK, bold=True, align=PP_ALIGN.CENTER)


def build_screen_slide(slide, spec, page_index: int, total: int) -> None:
    base.clean_to_template_chrome(slide, f"화면 설계서({spec['id']})", f"{page_index}/{total}")
    draw_wireframe(slide, spec, 0.30, 1.42, 3.20, 2.12)
    base.panel(slide, 3.70, 1.42, 6.00, 2.12, "화면 기본정보", accent=base.BLUE, fill=base.WHITE)
    info = [
        ("화면 ID", spec["id"]),
        ("화면명", spec["name"]),
        ("메뉴 경로", spec["menu"]),
        ("화면 개요", spec["overview"]),
        ("진입 조건", spec["entry"]),
        ("주요 출력", spec["output"]),
    ]
    for index, (label, value) in enumerate(info):
        row = index // 2
        column = index % 2
        xx = 3.92 + column * 2.86
        yy = 1.98 + row * 0.48
        base.textbox(slide, label, xx, yy, 0.74, 0.18, size=7.2, color=base.BLUE, bold=True)
        base.textbox(slide, value, xx + 0.76, yy - 0.01, 1.98, 0.34, size=7.3, color=base.INK)

    rows = [["기능 번호", "기능명", "기능설명", "처리내용", "비고", "요구사항명"]]
    for function_id, name, description, process, note, requirement_ids in spec["rows"]:
        rows.append([function_id, name, description, process, note, requirement_text(requirement_ids)])
    table_height = 2.92
    table_shape = base.add_table(
        slide, rows, 0.30, 3.78, 9.40, table_height,
        widths=[0.72, 1.00, 1.48, 2.16, 1.35, 2.69], font_size=6.8, header_fill=base.BLUE,
    )
    fit_table_rows(table_shape, table_height, len(rows) - 1, header_font=7.2)


def metadata_field(slide, label, value, x, y, label_w, value_w) -> None:
    base.ppt_rect(slide, x, y, label_w, 0.36, fill=base.BLUE, line_color=base.BLUE)
    base.textbox(slide, label, x + 0.05, y + 0.08, label_w - 0.10, 0.18, size=7.4, color=base.WHITE, bold=True, align=PP_ALIGN.CENTER)
    base.ppt_rect(slide, x + label_w, y, value_w, 0.36, fill=base.WHITE, line_color=base.LINE)
    base.textbox(slide, value, x + label_w + 0.08, y + 0.07, value_w - 0.16, 0.20, size=7.5, color=base.INK, bold=True)


def flow_node(slide, x, y, w, h, value, *, kind="process", color=base.BLUE):
    shape_type = {
        "start": MSO_AUTO_SHAPE_TYPE.FLOWCHART_TERMINATOR,
        "decision": MSO_AUTO_SHAPE_TYPE.FLOWCHART_DECISION,
        "process": MSO_AUTO_SHAPE_TYPE.FLOWCHART_PROCESS,
    }[kind]
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = base.WHITE if kind != "start" else color
    shape.line.color.rgb = color
    shape.line.width = Pt(1.4)
    text_color = base.WHITE if kind == "start" else base.INK
    base.textbox(slide, value, x + 0.08, y + 0.16, w - 0.16, h - 0.26, size=8.2 if kind != "decision" else 7.6, color=text_color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shape


def build_flow_slide(slide, spec, page_index: int, total: int) -> None:
    function_id, program_id, program_name, overview, stages, decision, exception, requirement_ids = spec
    base.clean_to_template_chrome(slide, "기능처리도(기능흐름도)", f"{page_index}/{total}")
    metadata_field(slide, "프로그램 ID", program_id, 0.30, 1.38, 0.90, 1.35)
    metadata_field(slide, "프로그램명", program_name, 2.67, 1.38, 0.86, 2.20)
    metadata_field(slide, "작성일", DATE_TEXT, 5.87, 1.38, 0.66, 1.15)
    metadata_field(slide, "Page", f"{page_index}/{total}", 7.82, 1.38, 0.52, 1.36)
    metadata_field(slide, "개요", overview, 0.30, 1.82, 0.58, 8.82)

    y = 3.02
    flow_node(slide, 0.20, y, 0.62, 0.76, "시작", kind="start", color=base.BLUE)
    process_positions = [(1.00, 1.05), (2.27, 1.05), (4.84, 1.05), (6.11, 1.05), (7.38, 1.05)]
    process_colors = [base.BLUE, base.TEAL, base.CORAL, base.TEAL, base.BLUE]
    for index, (xx, ww) in enumerate(process_positions):
        flow_node(slide, xx, y, ww, 0.76, stages[index], kind="process", color=process_colors[index])
    flow_node(slide, 3.54, y - 0.11, 1.08, 0.98, decision, kind="decision", color=base.GOLD)
    flow_node(slide, 8.75, y, 0.82, 0.76, "종료", kind="start", color=base.TEAL)

    arrow_pairs = [(0.82, 1.00), (2.05, 2.27), (3.32, 3.54), (4.62, 4.84), (5.89, 6.11), (7.16, 7.38), (8.43, 8.75)]
    for x1, x2 in arrow_pairs:
        diagram_connector(slide, x1, y + 0.38, x2, y + 0.38, color=base.MUTED, width=1.15)
    base.textbox(slide, "예", 4.64, y + 0.05, 0.18, 0.16, size=7.0, color=base.TEAL, bold=True, align=PP_ALIGN.CENTER)

    diagram_connector(slide, 4.08, y + 0.87, 4.08, 4.72, color=base.CORAL, width=1.2, elbow=True)
    base.textbox(slide, "아니오", 4.16, 4.20, 0.42, 0.17, size=7.0, color=base.CORAL, bold=True)
    base.ppt_rect(slide, 2.45, 4.72, 2.50, 0.90, fill=base.LIGHT_CORAL, line_color=base.CORAL, radius=True, width=1.2)
    base.textbox(slide, "예외 분기", 2.64, 4.84, 2.12, 0.18, size=8.2, color=base.CORAL, bold=True, align=PP_ALIGN.CENTER)
    base.textbox(slide, exception, 2.62, 5.09, 2.16, 0.40, size=6.8, color=base.INK, align=PP_ALIGN.CENTER)
    diagram_connector(slide, 4.95, 5.17, 5.25, 5.17, color=base.CORAL, width=1.2)
    flow_node(slide, 5.25, 4.79, 2.10, 0.76, "예외 종료·상태 표시", kind="start", color=base.CORAL)
    base.ppt_rect(slide, 7.70, 4.72, 1.90, 0.90, fill=base.PALE, line_color=base.LINE, radius=True)
    base.textbox(slide, "연계 요구사항", 7.84, 4.85, 1.62, 0.18, size=7.5, color=base.BLUE, bold=True, align=PP_ALIGN.CENTER)
    base.textbox(slide, " · ".join(requirement_ids), 7.84, 5.16, 1.62, 0.18, size=7.0, color=base.INK, align=PP_ALIGN.CENTER)


def build_program_list(slide) -> None:
    base.clean_to_template_chrome(slide, "프로그램 - 목록", "전체 기능")
    rows = [["기능 분류", "기능번호", "기능명"], *PROGRAMS]
    table_shape = base.add_table(
        slide, rows, 0.52, 1.42, 8.96, 5.34,
        widths=[2.2, 1.3, 5.46], font_size=10.0, header_fill=base.BLUE,
    )
    fit_table_rows(table_shape, 5.34, len(PROGRAMS), header_height=0.46, header_font=9.2)
    for row_index in range(1, len(table_shape.table.rows)):
        table_shape.table.cell(row_index, 1).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


def compact_code(relative: str, marker: str, line_count=22, width=90) -> str:
    source = base.code_excerpt(relative, marker, line_count)
    lines = []
    for line in source.splitlines():
        lines.append(line if len(line) <= width else line[: width - 3] + "...")
    return "\n".join(lines)


def build_code_slide(slide, spec, page_index: int, total: int) -> None:
    title, function_ids, relative, marker, reason, rules, requirement_ids = spec
    base.clean_to_template_chrome(slide, title, f"{page_index}/{total}")
    base.ppt_rect(slide, 0.30, 1.42, 6.02, 5.32, fill=base.PALE, line_color=base.LINE)
    base.ppt_rect(slide, 0.30, 1.42, 6.02, 0.38, fill=base.BLUE, line_color=base.BLUE)
    base.textbox(slide, relative, 0.45, 1.52, 5.72, 0.18, size=7.4, color=base.WHITE, bold=True)
    base.textbox(slide, compact_code(relative, marker), 0.50, 1.98, 5.62, 4.52, size=7.15, color=base.INK, font=base.CODE_FONT, margin=0.02)

    base.panel(slide, 6.52, 1.42, 3.18, 5.32, "핵심 선정 사유", accent=base.TEAL, fill=base.WHITE)
    base.textbox(slide, "연계 기능", 6.78, 2.00, 0.78, 0.18, size=7.5, color=base.TEAL, bold=True)
    base.textbox(slide, function_ids, 7.64, 1.99, 1.75, 0.22, size=8.2, color=base.INK, bold=True)
    base.textbox(slide, "왜 핵심인가", 6.78, 2.43, 1.06, 0.18, size=7.5, color=base.TEAL, bold=True)
    base.textbox(slide, reason, 6.78, 2.72, 2.62, 1.03, size=8.3, color=base.INK)
    base.textbox(slide, "핵심 보호 규칙", 6.78, 3.98, 1.24, 0.18, size=7.5, color=base.TEAL, bold=True)
    for index, rule in enumerate(rules, 1):
        base.ppt_rect(slide, 6.78, 4.28 + (index - 1) * 0.50, 0.28, 0.28, fill=base.LIGHT_TEAL, line_color=base.TEAL, radius=True)
        base.textbox(slide, str(index), 6.84, 4.34 + (index - 1) * 0.50, 0.16, 0.14, size=7.0, color=base.TEAL, bold=True, align=PP_ALIGN.CENTER)
        base.textbox(slide, rule, 7.16, 4.27 + (index - 1) * 0.50, 2.22, 0.38, size=7.3, color=base.INK)
    base.ppt_rect(slide, 6.78, 5.96, 2.62, 0.48, fill=base.LIGHT_BLUE, line_color=base.BLUE, radius=True)
    base.textbox(slide, "요구사항  " + " · ".join(requirement_ids), 6.91, 6.10, 2.36, 0.18, size=7.2, color=base.BLUE, bold=True, align=PP_ALIGN.CENTER)


REPORT_FIELDS = [
    ("PK", "id", "UUID NN"),
    ("IDX", "status", "VARCHAR(16) NN DEFAULT 'new'"),
    ("", "class_id", "INTEGER NN"),
    ("IDX", "class_name", "VARCHAR(64) NN"),
    ("", "confidence", "DOUBLE PRECISION NN"),
    ("", "bbox_x", "DOUBLE PRECISION NN"),
    ("", "bbox_y", "DOUBLE PRECISION NN"),
    ("", "bbox_width", "DOUBLE PRECISION NN"),
    ("", "bbox_height", "DOUBLE PRECISION NN"),
    ("", "captured_at", "TIMESTAMPTZ NN"),
    ("", "source", "VARCHAR(16) NN"),
    ("", "latitude", "DOUBLE PRECISION NULL"),
    ("", "longitude", "DOUBLE PRECISION NULL"),
    ("", "accuracy_m", "DOUBLE PRECISION NULL"),
    ("", "heading", "DOUBLE PRECISION NULL"),
    ("GIST", "location", "geometry(Point,4326) NULL"),
    ("", "image_path", "VARCHAR(255) NN"),
    ("", "image_content_type", "VARCHAR(128) NN"),
    ("GIN", "metadata", "JSONB NN"),
    ("IDX", "created_at", "TIMESTAMPTZ NN DEFAULT now()"),
    ("", "updated_at", "TIMESTAMPTZ NN DEFAULT now()"),
]

STATUS_AUDIT_FIELDS = [
    ("PK", "id", "UUID NN"),
    ("LREF", "report_id", "UUID NN INDEX"),
    ("", "previous_status", "VARCHAR(16) NN"),
    ("", "next_status", "VARCHAR(16) NN"),
    ("IDX", "actor_id", "VARCHAR(64) NN"),
    ("", "note", "VARCHAR(500) NULL"),
    ("", "resolution_reason", "VARCHAR(500) NULL"),
    ("IDX", "created_at", "TIMESTAMPTZ NN DEFAULT now()"),
]

EXPORT_AUDIT_FIELDS = [
    ("PK", "id", "UUID NN"),
    ("UQ", "audit_id", "UUID NN UNIQUE"),
    ("IDX", "requested_audit_id", "UUID NULL"),
    ("IDX", "actor_id", "VARCHAR(64) NN"),
    ("", "export_format", "VARCHAR(16) NN"),
    ("", "profile", "VARCHAR(16) NN"),
    ("", "aggregate", "VARCHAR(16) NULL"),
    ("", "row_count", "INTEGER NN CHECK ≥0"),
    ("", "location_precision", "VARCHAR(32) NN"),
    ("", "rows_sha256", "CHAR(64) NULL"),
    ("", "filters", "JSONB NN"),
    ("IDX", "created_at", "TIMESTAMPTZ NN DEFAULT now()"),
]

READ_AUDIT_FIELDS = [
    ("PK", "id", "UUID NN"),
    ("IDX", "actor_id", "VARCHAR(64) NN"),
    ("IDX", "purpose", "VARCHAR(64) NN"),
    ("IDX", "resource_type", "list/detail/image"),
    ("", "resource_id", "VARCHAR(160) NN"),
    ("", "details", "JSONB NN"),
    ("IDX", "created_at", "TIMESTAMPTZ NN DEFAULT now()"),
]

RATE_EVENT_FIELDS = [
    ("PK", "id", "BIGINT IDENTITY"),
    ("IDX", "actor_digest", "CHAR(64) NN"),
    ("IDX", "rate_group", "VARCHAR(32) NN"),
    ("IDX", "observed_at", "TIMESTAMPTZ NN"),
]


def entity_card(slide, x, y, w, h, title, fields, *, accent=base.BLUE, append_only=False, font_size=6.4):
    base.ppt_rect(slide, x, y, w, h, fill=base.WHITE, line_color=accent, width=1.3)
    base.ppt_rect(slide, x, y, w, 0.38, fill=accent, line_color=accent)
    base.textbox(slide, title, x + 0.12, y + 0.08, w - 0.24, 0.20, size=9.0, color=base.WHITE, bold=True)
    if append_only:
        base.ppt_rect(slide, x + w - 1.20, y + 0.07, 1.05, 0.22, fill=base.WHITE, line_color=base.WHITE, radius=True)
        base.textbox(slide, "APPEND-ONLY", x + w - 1.14, y + 0.11, 0.93, 0.12, size=5.8, color=accent, bold=True, align=PP_ALIGN.CENTER)
    row_height = (h - 0.42) / len(fields)
    for index, (key, name, datatype) in enumerate(fields):
        yy = y + 0.40 + index * row_height
        fill = base.WHITE if index % 2 == 0 else base.PALE
        base.ppt_rect(slide, x + 0.02, yy, w - 0.04, row_height, fill=fill, line_color=base.LINE, width=0.35)
        if key:
            base.textbox(slide, key, x + 0.06, yy + 0.03, 0.34, row_height - 0.04, size=font_size - 0.6, color=accent, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        base.textbox(slide, name, x + 0.42, yy + 0.02, w * 0.39, row_height - 0.03, size=font_size, color=base.INK, bold=key == "PK", valign=MSO_ANCHOR.MIDDLE)
        base.textbox(slide, datatype, x + 0.46 + w * 0.39, yy + 0.02, w - (0.55 + w * 0.39), row_height - 0.03, size=font_size - 0.35, color=base.MUTED, valign=MSO_ANCHOR.MIDDLE)


def build_erd_overview(slide) -> None:
    base.clean_to_template_chrome(slide, "엔티티 관계도 - ERD", "정확히 5개 application table")
    entity_card(
        slide, 0.25, 1.45, 3.25, 3.90, "reports",
        [("PK", "id", "UUID"), ("", "status", "new/reviewed/resolved"), ("", "class·bbox", "탐지 근거"), ("", "location", "Point(4326)"), ("", "image_path", "DB 밖 파일 경로"), ("", "metadata", "JSONB"), ("", "created/updated", "TIMESTAMPTZ")],
        accent=base.TEAL, font_size=7.1,
    )
    entity_card(
        slide, 3.75, 1.45, 2.85, 2.25, "report_status_audits",
        [("PK", "id", "UUID"), ("LREF", "report_id", "FK 없음"), ("", "previous/next", "상태 전이"), ("", "actor_id", "문자열 주체"), ("", "note·reason", "검수 근거"), ("", "created_at", "TIMESTAMPTZ")],
        accent=base.GOLD, append_only=True, font_size=6.2,
    )
    entity_card(
        slide, 6.85, 1.45, 2.85, 2.25, "report_export_audits",
        [("PK", "id", "UUID"), ("UQ", "audit_id", "manifest 연계"), ("", "requested_id", "요청 연계"), ("", "actor·profile", "반출 정책"), ("", "rows_sha256", "결과 hash"), ("", "created_at", "TIMESTAMPTZ")],
        accent=base.CORAL, append_only=True, font_size=6.2,
    )
    entity_card(
        slide, 3.75, 4.00, 2.85, 2.25, "report_read_audits",
        [("PK", "id", "UUID"), ("", "actor·purpose", "조회 주체·목적"), ("", "resource_type", "list/detail/image"), ("", "resource_id", "대상"), ("", "details", "JSONB"), ("", "created_at", "TIMESTAMPTZ")],
        accent=base.BLUE, append_only=True, font_size=6.2,
    )
    entity_card(
        slide, 6.85, 4.00, 2.85, 2.25, "actor_rate_limit_events",
        [("PK", "id", "BIGINT"), ("", "actor_digest", "HMAC digest"), ("", "rate_group", "API group"), ("", "observed_at", "60초 window")],
        accent=base.TEAL, append_only=False, font_size=6.4,
    )
    base.ppt_rect(slide, 0.48, 5.63, 2.82, 0.64, fill=RGBColor(238, 240, 243), line_color=base.MUTED, radius=True)
    base.textbox(slide, "DB 밖 image · durable journal · startup reconciliation", 0.65, 5.83, 2.48, 0.20, size=6.2, color=base.INK, bold=True, align=PP_ALIGN.CENTER)
    diagram_connector(slide, 3.50, 2.52, 3.75, 2.52, color=base.GOLD, width=1.4, dash=True)
    diagram_connector(slide, 3.50, 4.55, 3.75, 4.55, color=base.BLUE, width=1.4, dash=True)
    base.textbox(slide, "status/export/read 3종만 append-only · rate events는 non-audit 공유 원자 60초 제한", 1.15, 6.56, 7.70, 0.18, size=6.8, color=base.INK, bold=True, align=PP_ALIGN.CENTER)


def build_table_definition(slide) -> None:
    base.clean_to_template_chrome(slide, "테이블 정의서 - ERD", "5 tables · 52 columns")
    entity_card(slide, 0.15, 1.36, 3.55, 5.35, "reports · 21 columns", REPORT_FIELDS, accent=base.TEAL, font_size=5.25)
    entity_card(slide, 3.95, 1.36, 2.75, 2.45, "report_status_audits · 8", STATUS_AUDIT_FIELDS, accent=base.GOLD, append_only=True, font_size=5.15)
    entity_card(slide, 6.95, 1.36, 2.75, 2.45, "report_export_audits · 12", EXPORT_AUDIT_FIELDS, accent=base.CORAL, append_only=True, font_size=4.75)
    entity_card(slide, 3.95, 4.08, 2.75, 2.63, "report_read_audits · 7", READ_AUDIT_FIELDS, accent=base.BLUE, append_only=True, font_size=5.15)
    entity_card(slide, 6.95, 4.08, 2.75, 2.63, "actor_rate_limit_events · 4", RATE_EVENT_FIELDS, accent=base.TEAL, append_only=False, font_size=5.15)
    diagram_connector(slide, 3.70, 2.58, 3.95, 2.58, color=base.GOLD, width=1.2, dash=True)
    diagram_connector(slide, 3.70, 5.15, 3.95, 5.15, color=base.BLUE, width=1.2, dash=True)
    base.textbox(slide, "PK 5 · status/export/read 3종만 append-only · rate events는 non-audit 60초 shared limiter", 3.88, 6.79, 5.90, 0.18, size=5.8, color=base.INK, bold=True, align=PP_ALIGN.CENTER)


def apply_review_corrections(originals) -> None:
    """Correct reviewed facts without changing the official template layout."""
    for shape in originals[0].shapes:
        text = getattr(shape, "text", "")
        if "2026. 07. 11." in text:
            frame = shape.text_frame
            frame.clear()
            frame.word_wrap = True
            for index, line in enumerate((
                "2026. 07. 12.",
                "WalkSafe Assist · 중간 제작설계서",
                "주제품 Web/PWA · Android 연구 보조",
            )):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.space_after = Pt(0)
                run = paragraph.add_run()
                run.text = line
                run.font.name = base.FONT
                run.font.size = Pt(16)
                run.font.color.rgb = base.MUTED

    for shape in originals[5].shapes:
        text = getattr(shape, "text", "")
        if text.strip() == "YOLO26":
            base.set_text(shape, "Field AI", size=8.4, color=base.CORAL, bold=True, align=PP_ALIGN.CENTER)
        elif text.strip() == "13종 탐지":
            base.set_text(shape, "YOLO26 · 13종", size=6.4, color=base.INK, align=PP_ALIGN.CENTER)
        elif "/detect/v2 — YOLO26 13종 객체 탐지" in text:
            base.set_text(
                shape,
                "• /detect/v2 — field: unified YOLO26 13종 / local 기본: fake\n"
                "• real 실패 — unavailable · fake 자동 전환 없음\n"
                "• /reports/v2 — 손상 신고·중복 판정\n"
                "• /navigation/* — TMAP POI·계단 회피 경로",
                size=8.0,
                color=base.INK,
            )

    for shape in originals[6].shapes:
        if getattr(shape, "text", "").strip() == "상태·지도·내보내기":
            base.set_text(shape, "상태·위치 분포·내보내기", size=6.6, color=base.INK, bold=True, align=PP_ALIGN.CENTER)


def synchronize_requirement_trace(slide) -> None:
    """Make REQ-009 trace the report persistence function as well as admin functions."""
    tables = [shape.table for shape in slide.shapes if getattr(shape, "has_table", False)]
    if len(tables) != 1:
        raise RuntimeError("REQ-008~014 table not found")
    table = tables[0]
    for row_index in range(1, len(table.rows)):
        if table.cell(row_index, 0).text.strip() != "REQ-009":
            continue
        cell = table.cell(row_index, 2)
        cell.text = "F-004·F-006·F-007"
        for paragraph in cell.text_frame.paragraphs:
            paragraph.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = base.FONT
                run.font.size = Pt(7.5)
                run.font.color.rgb = base.INK
        return
    raise RuntimeError("REQ-009 row not found")


def build() -> None:
    with tempfile.TemporaryDirectory(prefix="walksafe-practical-base-") as directory:
        base_path = Path(directory) / "base.pptx"
        original_output = base.OUTPUT
        try:
            base.OUTPUT = base_path
            base.build()
        finally:
            base.OUTPUT = original_output

        prs = Presentation(base_path)
        originals = list(prs.slides)
        if len(originals) != 19:
            raise RuntimeError("base deck must have 19 slides")

        synchronize_requirement_trace(originals[3])
        apply_review_corrections(originals)

        screen_slides = originals[7:11] + [clone_slide(prs, originals[7]) for _ in range(8)]
        flow_slides = originals[12:14] + [clone_slide(prs, originals[12]) for _ in range(9)]
        code_slides = originals[16:18] + [clone_slide(prs, originals[16]) for _ in range(4)]

        for page_index, (slide, spec) in enumerate(zip(screen_slides, SCREEN_SPECS), 1):
            build_screen_slide(slide, spec, page_index, len(SCREEN_SPECS))
        build_erd_overview(originals[11])
        for page_index, (slide, spec) in enumerate(zip(flow_slides, FLOW_SPECS), 1):
            build_flow_slide(slide, spec, page_index, len(FLOW_SPECS))
        build_program_list(originals[14])
        build_table_definition(originals[15])
        for page_index, (slide, spec) in enumerate(zip(code_slides, CODE_SPECS), 1):
            build_code_slide(slide, spec, page_index, len(CODE_SPECS))

        final_order = (
            originals[:7]
            + screen_slides
            + [originals[11]]
            + flow_slides
            + [originals[14], originals[15]]
            + code_slides
            + [originals[18]]
        )
        reorder_slides(prs, final_order)
        if len(prs.slides) != 40:
            raise RuntimeError(f"final slide count mismatch: {len(prs.slides)}")

        core = prs.core_properties
        core.title = "WalkSafe Assist 실무 제작설계서 — 공식 템플릿 준수본"
        core.subject = "화면·프로그램·ERD·기능흐름·핵심코드 완전 추적본"
        core.comments = "REQ-001~014, UI-001~048, F-001~011 실무 산출물"

        for rel_id, rel in list(prs.part.rels.items()):
            if any(token in rel.reltype.lower() for token in ("authors", "comment", "notesmaster", "person")):
                prs.part.drop_rel(rel_id)
        for slide in prs.slides:
            for rel_id, rel in list(slide.part.rels.items()):
                if any(token in rel.reltype.lower() for token in ("notesslide", "comment", "person")):
                    slide.part.drop_rel(rel_id)

        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(prefix="walksafe-practical-", suffix=".pptx", dir=OUTPUT.parent, delete=False) as handle:
            temporary = Path(handle.name)
        try:
            prs.save(temporary)
            base.normalize_package(temporary, len(prs.slides))
            os.replace(temporary, OUTPUT)
        finally:
            temporary.unlink(missing_ok=True)

    print(f"generated {OUTPUT} ({OUTPUT.stat().st_size} bytes, sha256={base.sha256(OUTPUT)})")


if __name__ == "__main__":
    build()

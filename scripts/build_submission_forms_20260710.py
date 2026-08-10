#!/usr/bin/env python3
"""Build the two 2026 Hanium submission forms from verified local materials.

The bracket-prefixed files under templates/ are immutable source templates.  This
script writes the unbracketed working copies and their build receipt; promotion into
docs/submission/final/ happens after independent render and privacy QA.
"""

from __future__ import annotations

import sys

if __name__ == "__main__":
    _verified_submission_commit = getattr(
        sys.modules.get("__main__"), "_walksafe_submission_source_commit", None
    )
    _startup_flags = (
        sys.flags.isolated,
        sys.flags.no_site,
        sys.flags.ignore_environment,
        int(getattr(sys.flags, "safe_path", False)),
        sys.flags.dont_write_bytecode,
    )
    if (
        any(value != 1 for value in _startup_flags)
        or "site" in sys.modules
        or not isinstance(_verified_submission_commit, str)
        or len(_verified_submission_commit) != 40
        or any(character not in "0123456789abcdef" for character in _verified_submission_commit)
    ):
        raise SystemExit(
            "WalkSafe submission CLI requires Python -I -S -B through the submission runner"
        )

import hashlib
import json
import shutil
import tempfile
import textwrap
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor as DocxRGBColor
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt
from PIL import Image

from submission_build_io import (
    atomic_output_path,
    canonical_android_device_evidence,
    submission_build_lock,
    verification_snapshot_policy_boundary,
)
from submission_manifest_policy import (
    ALL_SUBMISSION_GENERATED_PATHS,
    FORM_BUILD_INPUT_PATHS,
    FORM_TEMPLATE_FILES,
    build_tool_provenance,
    file_record,
    record_bundle_sha256,
    require_exact_file_set,
    require_no_unexpected_files,
    require_clean_source_revision,
    submission_toolchain_attestation,
)


REPO_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = REPO_ROOT / "templates"
MATERIAL_DIR = REPO_ROOT / "docs/submission/form_materials"
ASSET_DIR = MATERIAL_DIR / "assets"

DOCX_SOURCE = TEMPLATE_DIR / "[서식1] 2026 한이음 드림업 개발보고서 양식.docx"
PPTX_SOURCE = TEMPLATE_DIR / "[서식2] 2026년_제작설계서_일반(개인정보 기재x).pptx"
DOCX_OUTPUT = TEMPLATE_DIR / "개발보고서 양식.docx"
PPTX_OUTPUT = TEMPLATE_DIR / "제작설계서_일반.pptx"
FORM_BUILD_MANIFEST = TEMPLATE_DIR / "SUBMISSION_BUILD_MANIFEST.json"
ALLOWED_TEMPLATE_FILES = FORM_TEMPLATE_FILES

EXPECTED_HASHES = {
    DOCX_SOURCE: "b8d52b810760c54c61646566ea3ee78208f43d7b1c0e3b65d53d99dc8cde22e3",
    PPTX_SOURCE: "140628cdd6b160c8d86a2919ce01620c274c54362bd71c9e82a77670a7621336",
}

FACTS_PATH = MATERIAL_DIR / "09_제출_사실_기준.json"
REQUIRED_MATERIALS = (
    MATERIAL_DIR / "00_작성기준_및_가정.md",
    MATERIAL_DIR / "01_주장_근거_검증_매트릭스.md",
    MATERIAL_DIR / "02_기능_진척도_산정표.md",
    MATERIAL_DIR / "03_수행일정_계획대비실적.md",
    MATERIAL_DIR / "04_개발보고서_본문원고.md",
    MATERIAL_DIR / "05_제작설계서_34장_원고.md",
    MATERIAL_DIR / "06_외부근거_출처표.md",
    MATERIAL_DIR / "07_화면_시각자료_목록.md",
    MATERIAL_DIR / "08_도식_정의_및_검증.md",
    FACTS_PATH,
)
BUILD_SOURCES = (
    REPO_ROOT / "scripts/build_submission_forms_20260710.py",
    REPO_ROOT / "scripts/submission_build_io.py",
    REPO_ROOT / "scripts/submission_manifest_policy.py",
)

def load_submission_facts() -> dict[str, object]:
    with FACTS_PATH.open(encoding="utf-8") as handle:
        facts = json.load(handle)
    if facts.get("schema_version") != "walksafe.submission_facts.v2":
        raise SystemExit(f"unexpected submission facts schema: {facts.get('schema_version')!r}")
    return facts


FACTS = load_submission_facts()
PROJECT = FACTS["project"]
MODEL = FACTS["model_data"]
RUNTIME = FACTS["runtime"]
DATABASE = FACTS["database"]
REQUIREMENTS = FACTS["requirements"]
USE_CASES = FACTS["use_cases"]
VERIFICATION = FACTS["verification_snapshot"]
VERIFICATION_COUNTS = VERIFICATION["counts"]
NON_METRIC_CAPABILITY_SUMMARY = (
    "Capability: Web CAMERA_NON_METRIC_ADVISORY는 IMU 선택, low 좌/중앙/우, "
    "3 distinct frame+700ms이고 Android CAMERA_IMU_NON_METRIC은 fresh IMU 필수, "
    "reports=false다. 실패하면 TMAP_ONLY다. 거리·N보·STOP/high·local steering·안전 경로·"
    "경로 변경·자동 신고는 금지한다. 기존 Web alertable risk와 명시 동의 damaged-report는 "
    "유지한다. 동일 심각도는 3A bounded sequential handoff와 전달 직전 recheck를 적용한다. "
    "Web 실폰·ARCore 미지원 Android Field는 UNVERIFIED_OPEN이다."
)
EXPECTED_VERIFICATION_COUNT_KEYS = {
    "unit_python",
    "functional_python",
    "integration_python",
    "backend_full",
    "voice",
    "android_jvm",
    "web_trace_count",
    "web_trace_files",
}
if set(VERIFICATION_COUNTS) != EXPECTED_VERIFICATION_COUNT_KEYS or any(
    isinstance(value, bool) or not isinstance(value, int) or value < 1
    for value in VERIFICATION_COUNTS.values()
):
    raise SystemExit("verification_snapshot.counts must contain the exact positive integer count contract")
VERIFICATION_STATUS = VERIFICATION.get("status")
VERIFICATION_EXECUTED_AT = VERIFICATION.get("executed_at")
if not isinstance(VERIFICATION_STATUS, str) or not isinstance(VERIFICATION_EXECUTED_AT, str):
    raise SystemExit("verification_snapshot must declare status and executed_at")


def verification_snapshot_contract(snapshot: object) -> str:
    if not isinstance(snapshot, dict):
        raise SystemExit("canonical facts must contain a verification snapshot")
    try:
        return verification_snapshot_policy_boundary(
            snapshot.get("status"),
            snapshot.get("executed_at"),
            snapshot.get("policy"),
        )
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc


def android_device_evidence(
    snapshot: dict[str, object],
) -> tuple[dict[str, object], dict[str, object], str, str]:
    try:
        return canonical_android_device_evidence(snapshot)
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc


(
    ANDROID_DEVICE_CURRENT,
    ANDROID_DEVICE_HISTORICAL,
    ANDROID_DEVICE_CURRENT_LABEL,
    ANDROID_DEVICE_HISTORICAL_LABEL,
) = android_device_evidence(VERIFICATION)
ANDROID_DEVICE_CURRENT_METRIC = "NOT RUN"
VERIFICATION_SCOPE_LABEL = f"{VERIFICATION_EXECUTED_AT} · {VERIFICATION_STATUS}"
VERIFICATION_SCOPE_BOUNDARY = verification_snapshot_contract(VERIFICATION)
VERIFICATION_SCOPE_NOTE = f"{VERIFICATION_SCOPE_LABEL} · {VERIFICATION_SCOPE_BOUNDARY}"

PROJECT_NAME = PROJECT["project_name"]
SERVICE_NAME = PROJECT["service_name"]
VALUE_PROPOSITION = PROJECT["value_proposition"]
FACTS_AS_OF = datetime.strptime(FACTS["as_of_date"], "%Y-%m-%d")
AS_OF_DATE = f"{FACTS_AS_OF.year}. {FACTS_AS_OF.month}. {FACTS_AS_OF.day}."
FIXED_CORE_TIME = FACTS_AS_OF
NEUTRAL_AUTHOR = "WalkSafe Assist 프로젝트"
FONT_KO = "Malgun Gothic"

DOC_ASSETS = {
    "architecture": ASSET_DIR / "system_architecture.png",
    "web_mobile": ASSET_DIR / "web_main_mobile.png",
    "web_desktop": ASSET_DIR / "web_main_desktop.png",
    "admin": ASSET_DIR / "admin_desktop.png",
    "risk_flow": ASSET_DIR / "risk_processing_flow.png",
    "report_flow": ASSET_DIR / "report_csv_flow.png",
    "erd": ASSET_DIR / "reports_erd.png",
    "value_flow": ASSET_DIR / "value_flow.png",
    "model_quality": ASSET_DIR / "model_quality_gate.png",
    "problem_solution": ASSET_DIR / "problem_solution_map.png",
    "use_case_swimlane": ASSET_DIR / "use_case_swimlane.png",
    "ui_states": ASSET_DIR / "ui_state_map.png",
    "navigation_flow": ASSET_DIR / "navigation_state_flow.png",
    "risk_timeline": ASSET_DIR / "risk_timeline.png",
    "evidence_results": ASSET_DIR / "evidence_results.png",
    "scope_change": ASSET_DIR / "scope_change_map.png",
    "adoption_roadmap": ASSET_DIR / "adoption_roadmap.png",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_ooxml(path: Path) -> None:
    """Rewrite an Office package with fixed member metadata for deterministic hashes."""
    with tempfile.TemporaryDirectory(prefix="walksafe-ooxml-normalize-") as temp_dir:
        temp_path = Path(temp_dir) / path.name
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temp_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for name in sorted(source.namelist()):
                info = zipfile.ZipInfo(name, date_time=FIXED_CORE_TIME.timetuple()[:6])
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, source.read(name))
        shutil.move(temp_path, path)


def verify_inputs() -> None:
    missing = [
        REPO_ROOT / relative
        for relative in FORM_BUILD_INPUT_PATHS
        if not (REPO_ROOT / relative).is_file()
    ]
    if missing:
        raise SystemExit("missing submission material: " + ", ".join(str(path) for path in missing))

    for path, expected in EXPECTED_HASHES.items():
        actual = sha256(path)
        if actual != expected:
            raise SystemExit(f"source template changed: {path.name} ({actual})")

    expected_requirement_ids = [f"REQ-{index:03d}" for index in range(1, 15)]
    expected_use_case_ids = [f"UC-{index:02d}" for index in range(1, 10)]
    if [item["id"] for item in REQUIREMENTS] != expected_requirement_ids:
        raise SystemExit("submission facts requirements must be ordered REQ-001..REQ-014")
    if [item["id"] for item in USE_CASES] != expected_use_case_ids:
        raise SystemExit("submission facts use cases must be ordered UC-01..UC-09")
    if any(item["status"] != "PARTIAL" for item in [*REQUIREMENTS, *USE_CASES]):
        raise SystemExit("all current requirement/use-case statuses must remain PARTIAL")
    if [item["name"] for item in DATABASE["tables"]] != [
        "reports",
        "report_status_audits",
        "report_export_audits",
        "report_read_audits",
        "actor_rate_limit_events",
    ]:
        raise SystemExit("submission database truth must contain the exact five canonical tables")


def write_build_manifest() -> None:
    inputs = [
        file_record(REPO_ROOT, REPO_ROOT / relative)
        for relative in sorted(FORM_BUILD_INPUT_PATHS)
    ]
    artifacts = [file_record(REPO_ROOT, path) for path in (DOCX_OUTPUT, PPTX_OUTPUT)]
    manifest = {
        "schema_version": "walksafe.submission-forms-build.v1",
        **require_clean_source_revision(REPO_ROOT, ALL_SUBMISSION_GENERATED_PATHS),
        "build_tools": build_tool_provenance(("Pillow", "python-docx", "python-pptx")),
        "submission_toolchain": submission_toolchain_attestation(REPO_ROOT),
        "allowed_template_files": sorted(ALLOWED_TEMPLATE_FILES),
        "input_bundle_sha256": record_bundle_sha256(inputs),
        "inputs": inputs,
        "artifacts": artifacts,
    }
    with atomic_output_path(FORM_BUILD_MANIFEST) as temporary:
        temporary.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )


def set_docx_run_font(run, *, size: float = 10, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = FONT_KO
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_KO)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = DocxRGBColor.from_string(color)


def set_docx_paragraph(
    paragraph,
    text: str,
    *,
    size: float = 10,
    bold: bool = False,
    align: WD_ALIGN_PARAGRAPH | None = None,
    line_spacing: float = 1.3,
    space_after: float = 2,
) -> None:
    paragraph.text = ""
    run = paragraph.add_run(text)
    set_docx_run_font(run, size=size, bold=bold)
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        paragraph.alignment = align


def clear_docx_cell(cell) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)


def fill_docx_cell(
    cell,
    text: str,
    *,
    size: float = 9,
    bold: bool = False,
    align: WD_ALIGN_PARAGRAPH | None = None,
    line_spacing: float = 1.2,
) -> None:
    clear_docx_cell(cell)
    set_docx_paragraph(
        cell.paragraphs[0],
        text,
        size=size,
        bold=bold,
        align=align,
        line_spacing=line_spacing,
        space_after=0,
    )


def fill_docx_bullets(cell, lines: list[str], *, size: float = 9, compact: bool = False) -> None:
    clear_docx_cell(cell)
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        set_docx_paragraph(
            paragraph,
            f"• {line}",
            size=size,
            line_spacing=1.15 if compact else 1.25,
            space_after=1,
        )


def shade_docx_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def repeat_docx_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for existing in list(tr_pr.findall(qn("w:tblHeader"))):
        tr_pr.remove(existing)
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_docx_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for existing in list(tr_pr.findall(qn("w:cantSplit"))):
        tr_pr.remove(existing)
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def clear_docx_row_height(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for existing in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(existing)


def remove_docx_trailing_paragraphs_after(table) -> None:
    element = table._tbl.getnext()
    while element is not None and element.tag == qn("w:p"):
        next_element = element.getnext()
        element.getparent().remove(element)
        element = next_element


def clear_docx_template(document: Document) -> None:
    for paragraph in document.paragraphs:
        paragraph.text = ""
    for table in document.tables:
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                key = cell._tc
                if key in seen:
                    continue
                seen.add(key)
                clear_docx_cell(cell)


def add_docx_picture(cell, path: Path, *, width: float, caption: str) -> None:
    clear_docx_cell(cell)
    picture_paragraph = cell.paragraphs[0]
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = cell.add_paragraph()
    set_docx_paragraph(
        caption_paragraph,
        caption,
        size=8,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        line_spacing=1.1,
        space_after=0,
    )


def add_docx_picture_and_bullets(
    cell,
    path: Path,
    *,
    width: float,
    caption: str,
    bullets: list[str],
    bullet_size: float = 8.5,
) -> None:
    add_docx_picture(cell, path, width=width, caption=caption)
    for line in bullets:
        paragraph = cell.add_paragraph()
        set_docx_paragraph(
            paragraph,
            f"• {line}",
            size=bullet_size,
            line_spacing=1.15,
            space_after=1,
        )


def insert_docx_page_break_after(table) -> None:
    paragraph_element = OxmlElement("w:p")
    table._tbl.addnext(paragraph_element)
    paragraph = Paragraph(paragraph_element, table._parent)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def replace_docx_table(document: Document, old_table, *, rows: int, columns: int):
    new_table = document.add_table(rows=rows, cols=columns)
    # The official Korean template has no English-named "Table Grid" style.
    # Reuse the source table's internal style so generation remains template-safe.
    new_table.style = old_table.style
    new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    old_table._tbl.addprevious(new_table._tbl)
    old_table._tbl.getparent().remove(old_table._tbl)
    return new_table


def style_docx(document: Document) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.text:
                set_docx_run_font(run, size=run.font.size.pt if run.font.size else 10, bold=run.bold)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = 1
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.text:
                            set_docx_run_font(run, size=run.font.size.pt if run.font.size else 9, bold=run.bold)
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            set_docx_paragraph(
                paragraph,
                f"{SERVICE_NAME} | 근거 기준 {AS_OF_DATE}",
                size=8,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                line_spacing=1.0,
                space_after=0,
            )


def build_docx() -> None:
    document = Document(DOCX_SOURCE)
    if len(document.tables) != 30:
        raise SystemExit(f"unexpected DOCX table count: {len(document.tables)}")
    clear_docx_template(document)

    evaluation_bars = {
        6: "※ 평가항목 : 기획력 (필요성, 차별성)",
        10: "※ 평가항목 : 기술력 (기능구체성, 난이도, 완성도)",
        21: "※ 평가항목 : 수행능력 (문서완성도, 문제해결능력, 수행충실성)",
        27: "※ 평가항목 : 기획력 (활용가능성)",
    }
    for table_index, label in evaluation_bars.items():
        fill_docx_cell(
            document.tables[table_index].cell(0, 0),
            label,
            size=9,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    core = document.core_properties
    core.title = f"{PROJECT_NAME} 개발보고서"
    core.subject = "2026 한이음 드림업 개발보고서"
    core.author = NEUTRAL_AUTHOR
    core.last_modified_by = NEUTRAL_AUTHOR
    core.comments = "코드 및 검증 근거 기준 작성본"
    core.created = FIXED_CORE_TIME
    core.modified = FIXED_CORE_TIME

    headings = {
        22: "I. 프로젝트 개요",
        24: "1. 프로젝트 소개",
        29: "2. 개발 배경 및 필요성",
        34: "3. 프로젝트 특·장점",
        41: "II. 프로젝트 내용",
        44: "1. 프로젝트 구성도",
        51: "2. 프로젝트 기능",
        52: "1) 전체 기능 목록",
        59: "2) S/W 주요 기능",
        63: "3) H/W 주요 기능",
        68: "3. 주요 적용 기술",
        72: "4. 프로젝트 개발 환경",
        75: "5. 기타 사항",
        78: "III. 프로젝트 수행 내용",
        80: "프로젝트 수행일정",
        84: "2. 프로젝트 추진 과정에서의 문제점 및 해결방안",
        85: "1) 프로젝트 관리 측면",
        88: "2) 프로젝트 개발 측면",
        91: "3. 프로젝트를 통해 배우거나 느낀 점",
        98: "Ⅳ. 기대효과 및 활용분야",
        101: "1. 프로젝트의 기대효과",
        108: "2. 프로젝트의 활용분야",
    }
    for index, text in headings.items():
        level_one = text.startswith(("I.", "II.", "III.", "Ⅳ."))
        set_docx_paragraph(
            document.paragraphs[index],
            text,
            size=15 if level_one else 12,
            bold=True,
            line_spacing=1.2,
            space_after=4,
        )
        if index in {41, 51, 75, 98}:
            document.paragraphs[index].paragraph_format.page_break_before = True
    # Cover and one-page summary.
    fill_docx_cell(document.tables[0].cell(0, 0), "양식", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(document.tables[1].cell(0, 0), "2026년 한이음 드림업", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(document.tables[1].cell(1, 0), "개 발 보 고 서", size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(document.tables[1].cell(2, 0), f"{SERVICE_NAME}\n{AS_OF_DATE}", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(document.tables[2].cell(0, 0), "프로젝트명", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(document.tables[2].cell(0, 1), PROJECT_NAME, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(document.tables[3].cell(0, 0), "요 약 본", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    summary = document.tables[4]
    summary_rows = {
        0: ("프로젝트 정보", ""),
        1: ("프로젝트명", PROJECT_NAME),
        6: ("성과 목표", "■ 실용화   □ 앱등록   □ 프로그램등록   □ 특허   □ 기술이전"),
        7: ("프로젝트 소개", f"{VALUE_PROPOSITION} 보행 중에는 필요한 위험만 짧게 안내하고, 손상 시설은 위치·이미지·검수 이력을 갖춘 자료로 전환한다."),
        8: ("개발 배경 및 필요성", "고정 경로 정보만으로는 파손·단절·방치 장애물처럼 지금 달라지는 보행환경을 설명하기 어렵다. 화면을 계속 보기 어려운 사용자의 현재 위험 인지와 시설 신고 과정을 하나의 흐름으로 연결할 필요가 있다."),
        9: ("프로젝트 특·장점", "① 객체별 3프레임·700ms와 경로 근거로 알림을 선별하고 ② 전역 provider는 TMAP-only로 유지하며 Web·ARCore 미지원 Android 비계량 advisory와 Android metric strict fallback을 분리하고 ③ 손상 후보는 관리자 검수 후 기관용 자료로 분리한다."),
        10: ("주요 기능", f"Web/PWA 카메라 탐지·TTS/진동·음성 제어·TMAP 길안내와 FastAPI/PostGIS 신고·Admin 검수를 연결했다. {NON_METRIC_CAPABILITY_SUMMARY}"),
        11: ("기대효과 및 활용 분야", "사용자에게는 보조적 위험 안내, 복지기관에는 안전 프로토콜을 적용한 실증·교육, 시설관리자에게는 검수 가능한 손상 자료를 제공하는 것이 목표다. 오경고/분·지연·과업완료율·검수 승인율로 효과를 측정한다."),
    }
    for row_index, (label, value) in summary_rows.items():
        if row_index == 0:
            fill_docx_cell(summary.cell(row_index, 0), label, size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0)
            continue
        fill_docx_cell(summary.cell(row_index, 0), label, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0)
        fill_docx_cell(summary.cell(row_index, 1), value, size=10, line_spacing=1.6)
    checkbox_rows = {
        2: ("주제 영역", "□ 생활", "□ 업무", "■ 공공/교통", "□ 금융/핀테크"),
        3: ("", "□ 의료", "□ 교육", "□ 유통/쇼핑", "□ 엔터테인먼트"),
        4: ("기술 분야", "■ 소프트웨어", "■ 인공지능", "□ 스마트 디바이스", "□ 방송·콘텐츠"),
        5: ("", "□ 디지털융합", "□ 차세대통신", "□ 사이버보안", ""),
    }
    for row_index, values in checkbox_rows.items():
        for col_index, value in enumerate(values):
            fill_docx_cell(
                summary.cell(row_index, col_index),
                value,
                size=9,
                bold=col_index == 0,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                line_spacing=1.6,
            )
    insert_docx_page_break_after(summary)

    fill_docx_cell(document.tables[5].cell(0, 0), "본    문", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_bullets(document.tables[7].cell(0, 0), [
        "대표 장면: 사용자가 목적지를 말하면 TMAP 전역 경로가 시작되고, 카메라의 새 프레임에서 안정적으로 확인된 경로 위험만 짧은 음성·진동으로 안내한다.",
        "normal_tactile_block은 3프레임·700ms, confidence·freshness·GPS·heading·camera corridor 조건을 모두 만족한 짧은 구간에서만 진행 방향을 보조하며, 조건이 깨지면 즉시 TMAP 안내로 복귀한다.",
        "damaged_tactile_block은 일반 위험 발화와 분리한다. 동일 snapshot·GPS·model gate를 서버가 다시 확인하고 PostGIS에 저장한 뒤 관리자가 검수해 기관용 최소자료로 준비한다.",
        "핵심 기술은 단순 YOLO 화면 표시가 아니라 객체별 안정 추적, 보수적 위험 정책, 음성 우선순위, 원자 저장·감사 이력, 사용자와 관리자의 책임 분리다.",
        "현재 성숙도: CODE·AUTO·격리 E2E와 Android unified TFLite 실기기 load/invoke 근거를 보유했다. 사용자 효과·실외 성능·운영 접수는 측정 전이며 완료로 주장하지 않는다.",
    ], size=10)
    fill_docx_bullets(document.tables[8].cell(0, 0), [
        "배경 근거(사고율 아님): 2023년 한국장애인개발원 조사—편의시설 적정 79.2%, 안내시설 적정 57.2%.",
        "사용자 문제 가설: 고정 경로와 장면 전체 설명만으로는 지금 경로를 막는 위험의 우선순위를 알기 어렵고, 보행 중 화면 확인과 복잡한 신고 입력은 부담이 될 수 있다.",
        "설계 대응: 현재 경로 위험 후보만 선별해 짧게 안내하고, 손상 시설 후보는 보행 안내와 분리해 위치·이미지·검수 이력이 있는 자료로 남긴다.",
        "검증할 효과: 실폰 지연 p50/p95, class별 오경고/분·누락률, 목적지 과업완료율, TTS 이해도, 신고 검수 승인율을 별도로 측정한다.",
    ], size=10)
    fill_docx_bullets(document.tables[9].cell(0, 0), [
        "선별 안내: 객체 존재를 그대로 읽지 않고 서로 다른 3 frame·700ms, 현재/4초 보조 ROI, 경로·접근 신호를 조합해 불필요한 경고를 억제한다.",
        "경로 결합: TMAP이 목적지까지의 전역 경로를 유지하고, 정상 점자블록은 엄격한 조건을 만족할 때만 카메라에 보이는 단거리 방향을 우선한다.",
        "책임 있는 신고: 손상 점자블록은 report-only로 분기해 서버 재검증→원자 저장→관리자 검수→기관용 최소자료 흐름을 거친다. 관리자 검수 전에는 외부 제출하지 않는다.",
        "안전한 경계: trusted metric depth가 없으면 bbox TTC만으로 STOP을 만들지 않으며, WebXR 중앙 거리는 객체 위험 판단에 연결하지 않는다.",
        "비교 범위: Google Lookout 공식 도움말·TMAP API·행정안전부 안전신문고 안내의 기능 초점을 기준으로, ‘현재 위험 선별+경로+검수 자료’의 결합을 차별점으로 삼는다.",
    ], size=10)

    add_docx_picture(
        document.tables[11].cell(0, 0),
        DOC_ASSETS["architecture"],
        width=6.2,
        caption="시스템 구성도: Web/PWA 주 앱, FastAPI/PostGIS, Admin, TMAP, 모델·데이터, Android 연구 보조",
    )

    functions = [
        ("REQ-001·002·004 위험 감지·선별", "새 frame 직렬 요청, real/fake 표시, 객체별 안정화·ROI·접근·TTC 경계. 확보: CODE·AUTO / 남음: class별 현장 품질", "60%*"),
        ("REQ-006·007 피드백·음성", "위험>상호작용>길안내와 13개 intent schema(12개 실행 intent+unknown). 확보: CODE·AUTO / 남음: 실폰 청취·진동·mic", "60%*"),
        ("REQ-012 목적지·보행 경로", "TMAP-only 전역 경로. Web은 IMU 선택 비계량 advisory, ARCore 미지원 Android는 fresh IMU 필수 비계량 advisory, Android metric steering은 route-bound strict gate. 확보: CODE·AUTO·provider smoke / 남음: 두 비계량 Field·실외 과업", "60%*"),
        ("REQ-005·008 손상 시설 신고", "동일 snapshot·GPS·model 재검증, durable journal·원자 rename·PostGIS commit·시작 조정. 확보: CODE·AUTO·통제 E2E / 남음: 실폰 신고", "75%*"),
        ("REQ-009·010·014 검수·운영 보호", "5개 application table, 상태/반출/조회 감사, 네 가지 독립 동의와 기관용 자료. 확보: CODE·AUTO(REQ-009 격리 DB) / 남음: 조직 IAM·복구/보존 시행·실제 접수", "60%*"),
        ("REQ-011 PWA lifecycle", "opt-in 설치·사용자 승인 update·offline shell. 확보: CODE·AUTO·headless / 남음: 실폰 standalone·안정 HTTPS", "75%*"),
        ("REQ-013 13-class 모델·데이터", "epoch270 내부 val·artifact hash. 확보: DATA·MODEL_EVAL·HASH(AUTO 상당) / 남음: sequence-safe split·독립 test", "60%*"),
        ("REQ-003 Android 연구 보조", "unified img768 TFLite와 SM-G981N load/invoke. 확보: CODE·AUTO / 남음: camera pipeline·FPS·전체 Device FIELD", "60%*"),
    ]
    function_table = document.tables[12]
    while len(function_table.rows) < len(functions) + 3:
        function_table.add_row()
    headers = ("구분", "기능", "설명", "현재진척도(%)")
    for col, value in enumerate(headers):
        fill_docx_cell(function_table.cell(0, col), value, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_docx_cell(function_table.cell(0, col), "D9EAD3")
    repeat_docx_header(function_table.rows[0])
    function_row_indexes = [1, 2, 3, 4, 6, 7, 8, 9]
    for row_index, (name, description, progress) in zip(function_row_indexes, functions, strict=True):
        fill_docx_cell(function_table.cell(row_index, 0), "S/W", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_docx_cell(function_table.cell(row_index, 1), name, size=8.5, bold=True)
        fill_docx_cell(function_table.cell(row_index, 2), description, size=8)
        fill_docx_cell(function_table.cell(row_index, 3), progress, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(
        function_table.cell(5, 0),
        "* 관리용 산식: 요구·설계 10 + CODE 30 + AUTO 20(DATA·MODEL_EVAL·HASH 포함) + 통제 E2E/격리·headless 15 + FIELD 15 + RELEASE 10. 그룹은 공통 최저 gate, 각 REQ는 1회만 산정한다.",
        size=7.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    hw_row = 10
    fill_docx_cell(function_table.cell(hw_row, 0), "H/W", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(function_table.cell(hw_row, 1), "별도 제작 없음", size=8.5)
    fill_docx_cell(function_table.cell(hw_row, 2), "기성 스마트폰을 실행환경으로 사용하며 전용 H/W는 제작하지 않음", size=8)
    fill_docx_cell(function_table.cell(hw_row, 3), "해당 없음", size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in function_table.rows:
        keep_docx_row_together(row)
    clear_docx_row_height(function_table.rows[5])

    sw_table = replace_docx_table(document, document.tables[13], rows=4, columns=3)
    sw_table.autofit = False
    column_widths = (1.15, 2.0, 3.05)
    for row in sw_table.rows:
        for cell, width in zip(row.cells, column_widths, strict=True):
            cell.width = Inches(width)
    for col, value in enumerate(("기능", "설명", "구현 화면·도식")):
        fill_docx_cell(sw_table.cell(0, col), value, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_docx_cell(sw_table.cell(0, col), "D9E2F3")
    repeat_docx_header(sw_table.rows[0])
    sw_rows = [
        ("보행 UI", "Web frame/report 개별 동의·탐지·위험·음성·길안내 상태", "web_desktop", "구현 UI—합성 입력 · 레이아웃/정책 확인용"),
        ("검수 UI", "신고 필터·상세·상태·반출", "admin", "구현 UI—local fixture · 운영 DB 근거 아님"),
        ("책임 흐름", "신고 후보부터 기관용 자료까지", "report_flow", "서버 재검증·journal/시작 조정·관리자 검수·사람의 외부 제출"),
    ]
    for row_index, (name, description, asset_key, caption) in enumerate(sw_rows, start=1):
        fill_docx_cell(sw_table.cell(row_index, 0), name, size=8.5, bold=True)
        fill_docx_cell(sw_table.cell(row_index, 1), description, size=8)
        add_docx_picture(sw_table.cell(row_index, 2), DOC_ASSETS[asset_key], width=2.8, caption=caption)
    for row in sw_table.rows:
        keep_docx_row_together(row)
    fill_docx_cell(document.tables[14].cell(0, 0), "캡처는 개인정보 없는 합성 정지 입력으로 정상·오류 상태와 레이아웃을 확인한 구현 증거다. 별도 자동검증 결과와 함께 보며 실외 성능·사용자 효과로 확대하지 않는다.", size=9)

    hw_table = document.tables[15]
    while len(hw_table.rows) > 2:
        hw_table._tbl.remove(hw_table.rows[-1]._tr)
    for col, value in enumerate(("기능/부품", "설명", "프로젝트 실물사진")):
        fill_docx_cell(hw_table.cell(0, col), value, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index in range(1, len(hw_table.rows)):
        for col in range(3):
            fill_docx_cell(hw_table.cell(row_index, col), "", size=9)
    fill_docx_cell(hw_table.cell(1, 0), "해당 없음", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(hw_table.cell(1, 1), "기성 스마트폰 실행환경", size=8.5)
    fill_docx_cell(hw_table.cell(1, 2), "전용 H/W·대체 이미지 없음", size=8.5)
    repeat_docx_header(hw_table.rows[0])
    for row in hw_table.rows:
        keep_docx_row_together(row)
        clear_docx_row_height(row)
    fill_docx_cell(document.tables[16].cell(0, 0), "H/W가 없는 경우에도 공식 표 구조를 유지했다. 스마트폰 camera·GPS·orientation/motion·microphone·speaker·vibration은 구매·제작 H/W가 아니라 실행환경이다.", size=8.5)
    fill_docx_cell(
        document.tables[17].cell(0, 0),
        "별도 제작 H/W는 없다. Android ARCore/TFLite 코드는 연구 보조이며 "
        f"{ANDROID_DEVICE_CURRENT_LABEL}. {ANDROID_DEVICE_HISTORICAL_LABEL}. "
        "대화형 Device Field PASS는 없다.",
        size=9,
    )

    add_docx_picture_and_bullets(document.tables[18].cell(0, 0), DOC_ASSETS["risk_timeline"], width=5.9, caption="핵심 기술: 한 객체가 안내되기까지의 frame·시간·경로 조건", bullets=[
        "탐지→안내: 최대변 960 JPEG의 새 frame만 기본 450ms 간격으로 보내고, 서로 다른 3 frame·700ms와 freshness·ROI·경로·접근 조건을 통과해야 안내한다.",
        "보수적 판단: 한 번의 빈 detection은 직전 후보를 유지하고 2회 연속 비면 해제한다. trusted metric depth 없이 bbox TTC만으로 high/STOP을 만들지 않는다.",
        "음성·길안내: 위험>상호작용>길안내>비계량 advisory 순이다. 전역 provider는 TMAP-only다. Web은 IMU 선택, ARCore 미지원 Android는 fresh IMU 필수 비계량 tier이며 Android metric supplier는 strict gate 실패 시 TMAP으로 복귀한다.",
        NON_METRIC_CAPABILITY_SUMMARY,
        "신고: damaged_tactile_block만 conf≥0.70, GPS≤15m, 3 frame·700ms, frame gap≤1.8초, 위치 도약≤20m를 통과한다. 응답은 primary 신고 ID와 duplicate_count를 유지하고 다른 후보 ID를 제거한다. 기관 API는 없으며 named admin 검수 후 수동 제출한다. 서버는 durable journal·startup reconciliation으로 crash 경계를 복구한다.",
        "모델 근거: 196,506장/607,814 bbox, 내부 val mAP50 0.55403·mAP50-95 0.41651. damaged recall 0.889인 반면 curb 0.195·uneven 0.0765여서 독립 test·sequence-safe split 전에는 배포 적격이 아니다.",
        "Android 연구: unified img768 13-class를 적용했다. "
        f"{ANDROID_DEVICE_CURRENT_LABEL}. {ANDROID_DEVICE_HISTORICAL_LABEL}. "
        "camera pipeline·FPS·실외 FIELD는 미검증이다.",
    ], bullet_size=7.9)

    env_rows = [
        ("S/W 개발환경", "OS", "Linux 개발환경; Android 연구 단말"),
        ("S/W 개발환경", "개발환경(IDE)", "저장소 비종속 개발환경; 특정 IDE 사용 실적은 확정하지 않음"),
        ("S/W 개발환경", "개발도구", "Git, npm, Gradle, Docker Compose, Alembic, pytest"),
        ("S/W 개발환경", "개발언어", "TypeScript/React, Python, Kotlin, SQL"),
        ("S/W 개발환경", "기타사항", "Next.js 16.2.6, React 19.2.6, FastAPI 0.128.8, Ultralytics 8.4.48"),
        ("H/W 구성장비", "디바이스", "기성 스마트폰; 전용 제작 H/W 없음"),
        ("H/W 구성장비", "센서", "camera, GPS, orientation/motion, microphone, speaker, vibration"),
        ("H/W 구성장비", "통신", "HTTPS 대상 REST API; 안정 domain/Release 대기"),
        ("H/W 구성장비", "언어", "Android 연구: Kotlin 2.2.21"),
        (
            "H/W 구성장비",
            "기타사항",
            "ARCore/LiteRT 연구 보조; unified img768 asset; "
            f"{ANDROID_DEVICE_CURRENT_LABEL}; {ANDROID_DEVICE_HISTORICAL_LABEL}; "
            "대화형 Device Field 대기",
        ),
        ("프로젝트 관리환경", "형상관리", "Git/GitHub remote, runtime config와 checkpoint SHA-256"),
        ("프로젝트 관리환경", "의사소통관리", "daylog·문서·issue 성격 기록"),
        ("프로젝트 관리환경", "기타사항", "current/source/snapshot/archive 분류와 책임 README"),
    ]
    env_table = document.tables[19]
    for col, value in enumerate(("구분", "항목", "상세내용", "상세내용")):
        fill_docx_cell(env_table.cell(0, col), value, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_docx_cell(env_table.cell(0, col), "D9EAD3")
    for row_index, (category, item, detail) in enumerate(env_rows, start=1):
        fill_docx_cell(env_table.cell(row_index, 0), category, size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_docx_cell(env_table.cell(row_index, 1), item, size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_docx_cell(env_table.cell(row_index, 2), detail, size=7.5)
    fill_docx_cell(env_table.cell(14, 0), "버전은 lockfile·runtime·artifact 기준이며 변경 시 재검증", size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_docx_picture_and_bullets(document.tables[20].cell(0, 0), DOC_ASSETS["evidence_results"], width=5.9, caption=f"계층별 자동검증 결과와 증명 범위 — {VERIFICATION_SCOPE_LABEL}", bullets=[
        VERIFICATION_SCOPE_NOTE,
        f"정책·단위: DB 없는 Python unit {VERIFICATION_COUNTS['unit_python']}, Web policy·voice callback, Android JVM {VERIFICATION_COUNTS['android_jvm']}/{VERIFICATION_COUNTS['android_jvm']}을 각 suite로 확인했다.",
        f"통합·빌드: 격리 PostGIS functional {VERIFICATION_COUNTS['functional_python']}와 integration {VERIFICATION_COUNTS['integration_python']}, Web production build·NFT {VERIFICATION_COUNTS['web_trace_count']} traces/{VERIFICATION_COUNTS['web_trace_files']:,} unique files, Voice {VERIFICATION_COUNTS['voice']}를 확인했다.",
        f"실기기 현재 경계: {ANDROID_DEVICE_CURRENT_LABEL}. "
        f"보존한 실기기 근거: {ANDROID_DEVICE_HISTORICAL_LABEL}. "
        "이는 camera→inference→feedback 전체 Field 결과가 아니다.",
        "실행 경계: Web/PWA server-v2는 server-sampled 보조다. MediaRecorder→faster-whisper STT→browser speechSynthesis TTS를 사용하며, WebXR은 중앙 거리 표시 전용이고 객체 위험 판단에는 연결되지 않는다.",
        "신고·저장 경계: application table은 reports·report_status_audits·report_export_audits·report_read_audits·actor_rate_limit_events의 정확히 5개다. 상태/반출/조회 감사 3종만 append-only이고 rate event는 공유 원자 60초 제한용 일반 table이다.",
        "운영·개인정보 경계: 목록/상세/image 조회는 actor·resource·time·purpose 감사를 남긴다. Web frame·Web report·Android report·telemetry는 각각 기본 OFF이며 철회 시 in-flight를 취소한다.",
        "배포 경계: 원격 DB는 sslmode=verify-full·gssencmode=disable, Backend는 PostgreSQL 공유 limiter와 nonblocking daemon single-flight readiness/TMAP failure cooldown, Web은 WALKSAFE_WEB_REPLICAS=1/process lock이다. Voice는 전용 VOICE_SERVICE_TOKEN을 쓰고 body copy 전 auth·Content-Length·global/actor/IP rate, copy 중 byte 상한, copy 후 duration·inference queue, disconnect와 worker/replica=1/process lock을 강제하며 OpenAPI는 인증되고 navigation은 TMAP-only다.",
        "데이터·운영 blocker: truncated 2,331장 제외, content hash·sequence-safe split·독립 test 없음, 실폰·실외 접근성, 조직 IAM·복구/보존 시행, 기관 receipt가 남아 있다.",
    ], bullet_size=8.2)

    schedule = document.tables[22]
    fill_docx_cell(schedule.cell(0, 0), "프로젝트 기간(한이음 드림업 사이트 기준)", size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_docx_cell(schedule.cell(0, 2), "2026. 4. 1. ~ 2026. 10. 30.", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    schedule_rows = [
        ("계획", "요구사항·프로젝트 계획", [4], []),
        ("분석", "솔루션·데이터·환경 분석", [4, 5], []),
        ("설계", "아키텍처·DB", [5, 6], [6]),
        ("설계", "UI/UX 접근성", [5, 6], [6]),
        ("개발·개선", "13-class 모델·데이터", [6, 7], [6, 7]),
        ("개발·개선", "Web/PWA·FastAPI·PostGIS", [6, 7, 8], [6, 7]),
        ("개발·개선", "sequence-safe split·실폰 실증", [8, 9], []),
        ("테스트", "접근성·운영·기관 절차", [9, 10], []),
        ("종료 (성과등록)", "최종 산출물·제출", [10], []),
    ]
    for col, value in enumerate(("구분", "추진내용", "4월", "5월", "6월", "7월", "8월", "9월", "10월", "11월", "12월")):
        fill_docx_cell(schedule.cell(3, col), value, size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for offset, (stage, task, planned_months, actual_months) in enumerate(schedule_rows, start=4):
        fill_docx_cell(schedule.cell(offset, 0), stage, size=7.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        fill_docx_cell(schedule.cell(offset, 1), task, size=7.5)
        for month in range(4, 13):
            cell = schedule.cell(offset, month - 2)
            marker = ""
            if month in planned_months:
                marker = "●" if month in actual_months else "○"
            fill_docx_cell(cell, marker, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            if month in planned_months:
                shade_docx_cell(cell, "D9EAD3" if marker == "●" else "FFF2CC")

    fill_docx_bullets(document.tables[23].cell(0, 0), [
        "범례: ● 저장소·산출물 생성 시점으로 확인한 실적 / ○ 원안 계획. 4~5월 활동은 당시 완료 근거가 없어 실적으로 추정하지 않았고, 11~12월은 공식 수행기간 밖이라 비웠다.",
        "확인 실적: 요구·설계, 13-class 내부 평가, Web/Backend/Admin, PWA lifecycle, Android invoke. 8~10월 조건: 재분할·독립 test, 실폰/실외 과업, IAM·복구/보존, 기관 receipt.",
    ], size=7.0, compact=True)

    fill_docx_bullets(document.tables[24].cell(0, 0), [
        "원안 목표 대조: 정확도 90% 수준·경보 1초·안정 도메인·재학습 15% 개선·기관 연계는 기준일 측정/운영 증거가 없어 완료로 산정하지 않았다.",
        "정밀 측위: 구현 부재 확인 → 4초 보조 ROI로 축소 → 과장된 충돌 예측 차단.",
        "자동 MLOps: 승격 workflow 부재 → 수동 train·eval·registry·rollback으로 재정의.",
        "기관·offline: 자동 API 미확정 → 관리자 검수·agency 수동 제출; app shell만 cache하고 안전 API는 network-only로 명시.",
    ], size=6.8, compact=True)
    fill_docx_bullets(document.tables[25].cell(0, 0), [
        "학습 재현성 — 증상: 768 resume 중 640 혼합·OOM / 원인: 설정 회귀·worker 메모리 / 조치: 혼합 run 보존·중단, workers 4→2, strict 768 resume / 결과: 300 epoch·epoch270 artifact hash 확보.",
        "데이터 일반화 — 증상: e-scooter 내부 지표 과도 / 원인: AIHub 572 sequence 599개 train/val 겹침 / 조치: leakage 공개·deployment_eligible=false / 결과: 일반화 주장 차단, 재분할 기준 확정.",
        "검증환경 drift — 증상: PWA field-session 503 / 원인: actor·backend token·session secret 혼용 / 조치: 역할·비밀 분리 / 결과: installability·SW update·offline shell 재검증.",
        "안전 경계 — 증상: WebXR 중앙 거리와 bbox 위험 연결 가능성 / 원인: 좌표 정합 미검증 / 조치: display-only·위험 context null / 결과: 근거 없는 N보·STOP 안내 차단.",
    ], size=9)
    fill_docx_bullets(document.tables[26].cell(0, 0), [
        "기능 수보다 사용자가 언제 어떤 안내를 받고, 시스템이 언제 침묵해야 하는지 정의하는 것이 중요했다.",
        "평균 모델 지표보다 curb·uneven 같은 고위험 class recall과 데이터 분할 품질이 제품 적합성을 더 잘 드러냈다.",
        "오류를 ‘위험 없음’으로 바꾸지 않는 fail-closed UI, 위험>상호작용>길안내 우선순위, 사람의 검수 경계가 접근성·공공 서비스의 핵심 설계였다.",
        "설정·artifact hash·격리 검증·증거등급을 함께 남기면 실패를 숨기지 않고 다음 의사결정으로 전환할 수 있었다.",
    ], size=9)
    add_docx_picture_and_bullets(document.tables[28].cell(0, 0), DOC_ASSETS["adoption_roadmap"], width=3.9, caption="구현 산출물에서 사용자·운영 효과를 검증하는 단계", bullets=[
        "통제 실폰→당사자/전문가→시설관리 순으로 지표를 확인하며, 사고 감소·비용 절감은 측정하지 않았다.",
    ], bullet_size=6.8)
    final_usage_cell = document.tables[29].cell(0, 0)
    fill_docx_bullets(final_usage_cell, [
        "사용자 지표: class별 오경고/분·누락률, 지연 p50/p95, 목적지 과업완료율, 위험 안내 이해도.",
        "운영 지표: 신고 검수 승인율, 처리시간 중앙값, 실제 기관 접수 receipt 건수.",
        "실증 단계: 통제 실폰(camera·GPS·mic·TTS·진동) → 당사자·복지기관(이해도·과업·접근성) → 시설관리(reviewed 자료·IAM·복구·보존·접수). 검증 전에는 안전 효과를 주장하지 않는다.",
    ], size=6.6, compact=True)
    source_paragraph = final_usage_cell.add_paragraph()
    set_docx_paragraph(
        source_paragraph,
        "공식 출처(접근 2026-07-12) ① 한국장애인개발원, 「2023년 장애인편의시설 설치 현황조사」 — koddi.or.kr/data/news_view.jsp?brdNum=7421441  ② Google Android Accessibility Help, 「Lookout」 — support.google.com/accessibility/android/answer/9031274  ③ TMAP Mobility API 공식 안내 — tmapmobility.com/service/corporate/api  ④ 행정안전부, 「안전신문고 안내」 — mois.go.kr/frt/sub/a06/b10/safetyReport/screen.do",
        size=5.0,
        line_spacing=1.0,
        space_after=0,
    )

    remove_docx_trailing_paragraphs_after(document.tables[29])
    style_docx(document)
    with atomic_output_path(DOCX_OUTPUT) as temporary:
        document.save(temporary)
        normalize_ooxml(temporary)


# PPT palette: neutral ink, teal for implemented paths, coral for gaps, yellow for cautions.
PPT_INK = RGBColor(31, 41, 55)
PPT_MUTED = RGBColor(75, 85, 99)
PPT_TEAL = RGBColor(15, 118, 110)
PPT_CORAL = RGBColor(224, 90, 71)
PPT_YELLOW = RGBColor(245, 190, 61)
PPT_LIGHT = RGBColor(244, 247, 246)
PPT_WHITE = RGBColor(255, 255, 255)
PPT_BORDER = RGBColor(209, 213, 219)


def ppt_rect(slide, x, y, w, h, *, fill=PPT_WHITE, line=None):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        PptxInches(x),
        PptxInches(y),
        PptxInches(w),
        PptxInches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def ppt_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    bold: bool = False,
    color=PPT_INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font: str = FONT_KO,
    margin: float = 0.04,
):
    box = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(w), PptxInches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = PptxInches(margin)
    frame.margin_right = PptxInches(margin)
    frame.margin_top = PptxInches(margin)
    frame.margin_bottom = PptxInches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = PptxPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def ppt_bullets(slide, lines: list[str], x: float, y: float, w: float, h: float, *, size: float = 17, color=PPT_INK):
    box = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(w), PptxInches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = PptxInches(0.08)
    frame.margin_right = PptxInches(0.04)
    frame.margin_top = PptxInches(0.04)
    frame.margin_bottom = PptxInches(0.04)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {line}"
        paragraph.font.name = FONT_KO
        paragraph.font.size = PptxPt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = PptxPt(7)
        paragraph.level = 0
    return box


def ppt_title(slide, title: str, number: int, *, section: str = "2026 한이음 드림업 제작설계서") -> None:
    ppt_rect(slide, 0, 0, 10, 0.7, fill=PPT_INK)
    ppt_rect(slide, 0, 0, 0.13, 0.7, fill=PPT_TEAL)
    ppt_text(slide, title, 0.34, 0.12, 7.9, 0.42, size=22, bold=True, color=PPT_WHITE, valign=MSO_ANCHOR.MIDDLE)
    ppt_text(slide, f"{number:02d}", 8.65, 0.12, 0.85, 0.42, size=16, bold=True, color=PPT_YELLOW, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
    ppt_text(slide, section, 0.32, 7.15, 6.8, 0.2, size=8, color=PPT_MUTED)
    ppt_text(slide, f"사실 기준 {FACTS['as_of_date']}", 7.2, 7.15, 2.45, 0.2, size=8, color=PPT_MUTED, align=PP_ALIGN.RIGHT)


def ppt_label(slide, text: str, x: float, y: float, w: float, *, fill=PPT_TEAL, color=PPT_WHITE):
    ppt_rect(slide, x, y, w, 0.36, fill=fill)
    return ppt_text(slide, text, x + 0.06, y + 0.04, w - 0.12, 0.24, size=11, bold=True, color=color, valign=MSO_ANCHOR.MIDDLE)


def ppt_metric(slide, value: str, label: str, x: float, y: float, w: float, *, accent=PPT_TEAL):
    ppt_rect(slide, x, y, w, 1.08, fill=PPT_LIGHT, line=PPT_BORDER)
    ppt_rect(slide, x, y, 0.08, 1.08, fill=accent)
    ppt_text(slide, value, x + 0.2, y + 0.13, w - 0.3, 0.4, size=21, bold=True, color=accent)
    ppt_text(slide, label, x + 0.16, y + 0.56, w - 0.24, 0.38, size=9.2, color=PPT_MUTED, align=PP_ALIGN.CENTER)


def ppt_add_picture_fit(slide, path: Path, x: float, y: float, w: float, h: float, *, border: bool = True):
    with Image.open(path) as image:
        width, height = image.size
    scale = min(w / width, h / height)
    draw_w = width * scale
    draw_h = height * scale
    draw_x = x + (w - draw_w) / 2
    draw_y = y + (h - draw_h) / 2
    if border:
        ppt_rect(slide, x, y, w, h, fill=PPT_WHITE, line=PPT_BORDER)
    return slide.shapes.add_picture(
        str(path), PptxInches(draw_x), PptxInches(draw_y), width=PptxInches(draw_w), height=PptxInches(draw_h)
    )


def ppt_table(slide, rows: list[list[str]], x: float, y: float, w: float, h: float, *, font_size: float = 12, widths: list[float] | None = None):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), PptxInches(x), PptxInches(y), PptxInches(w), PptxInches(h))
    table = table_shape.table
    if widths:
        total = sum(widths)
        for index, weight in enumerate(widths):
            table.columns[index].width = PptxInches(w * weight / total)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.margin_left = PptxInches(0.05)
            cell.margin_right = PptxInches(0.05)
            cell.margin_top = PptxInches(0.03)
            cell.margin_bottom = PptxInches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = PPT_INK if row_index == 0 else (PPT_LIGHT if row_index % 2 else PPT_WHITE)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER if row_index == 0 or col_index == 0 else PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = FONT_KO
                    run.font.size = PptxPt(font_size if row_index else font_size - 0.5)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = PPT_WHITE if row_index == 0 else PPT_INK
    return table_shape


def source_excerpt(relative_path: str, marker: str, line_count: int) -> str:
    path = REPO_ROOT / relative_path
    lines = path.read_text(encoding="utf-8").splitlines()
    try:
        start = next(index for index, line in enumerate(lines) if marker in line)
    except StopIteration as exc:
        raise SystemExit(f"source marker not found: {relative_path}: {marker}") from exc
    return textwrap.dedent("\n".join(lines[start : start + line_count]))


def verification_summary_lines() -> list[str]:
    runs = VERIFICATION.get("runs")
    if not isinstance(runs, list) or not runs:
        return [
            "최종 빌드 전 자동검증 재실행 대기",
            "고정 test 수는 실행일·격리 여부와 함께만 표시",
            "화면은 local fake-v2/fixture이며 실모델·실폰 근거 아님",
        ]
    lines: list[str] = []
    compact_names = {
        "Backend·root 격리 회귀": "Backend/root 격리 회귀",
        "Web/PWA·Voice·Android JVM/build": "Web/PWA·Voice·Android build",
        "모델·데이터 integrity": "모델·데이터 무결성",
        "Android unified TFLite export": "Android unified export",
    }
    for run in runs:
        name = str(run.get("name", "검증"))
        result = str(run.get("result", "UNKNOWN"))
        lines.append(f"{compact_names.get(name, name)}: {result}")
    return lines


def prepare_presentation() -> Presentation:
    presentation = Presentation(PPTX_SOURCE)
    if len(presentation.slides) != 34:
        raise SystemExit(f"unexpected PPTX slide count: {len(presentation.slides)}")
    for slide_index, slide in enumerate(presentation.slides):
        if slide_index == 1:
            # Slide 2 is the official required/optional deliverables matrix.
            # Keep its shapes and image relationships intact.
            continue
        for shape in list(slide.shapes):
            slide.shapes._spTree.remove(shape._element)
        for relationship_id, relationship in list(slide.part.rels.items()):
            if not relationship.reltype.endswith("/slideLayout"):
                slide.part.drop_rel(relationship_id)
    for relationship_id, relationship in list(presentation.part.rels.items()):
        relationship_type = relationship.reltype.lower()
        if any(token in relationship_type for token in ("authors", "comment", "notesmaster")):
            presentation.part.drop_rel(relationship_id)
    core = presentation.core_properties
    core.title = f"{PROJECT_NAME} 제작설계서"
    core.subject = "2026 한이음 드림업 제작설계서"
    core.author = NEUTRAL_AUTHOR
    core.last_modified_by = NEUTRAL_AUTHOR
    core.comments = "코드 및 검증 근거 기준 작성본"
    core.created = FIXED_CORE_TIME
    core.modified = FIXED_CORE_TIME
    return presentation


def build_pptx() -> None:
    presentation = prepare_presentation()
    slides = list(presentation.slides)
    for slide_index, slide in enumerate(slides):
        if slide_index == 1:
            continue
        ppt_rect(slide, 0, 0, 10, 7.5, fill=PPT_WHITE)

    # 1. Cover
    slide = slides[0]
    ppt_rect(slide, 0, 0, 10, 7.5, fill=PPT_INK)
    ppt_rect(slide, 0, 0, 0.16, 7.5, fill=PPT_TEAL)
    ppt_text(slide, "2026 한이음 드림업 · 제작설계서", 0.7, 0.55, 5.4, 0.4, size=14, bold=True, color=PPT_YELLOW)
    ppt_text(slide, SERVICE_NAME, 0.7, 1.15, 5.4, 0.55, size=24, bold=True, color=RGBColor(94, 234, 212))
    ppt_text(slide, PROJECT_NAME, 0.7, 1.62, 5.55, 0.22, size=9.5, color=PPT_WHITE)
    ppt_text(slide, "현재 보행 위험은\n선별해 알려주고,", 0.7, 1.85, 5.45, 1.25, size=30, bold=True, color=PPT_WHITE)
    ppt_text(slide, "손상 보행시설은\n검수 가능한 자료로 남깁니다.", 0.7, 3.12, 5.6, 1.25, size=25, bold=True, color=PPT_WHITE)
    ppt_label(slide, "구현 연결 완료 · 사용자/현장 실증 전", 0.7, 4.75, 3.55, fill=PPT_TEAL)
    ppt_bullets(slide, ["Web/PWA 위험 선별·음성 길안내", "PostGIS 신고·관리자 검수", "Android unified TFLite 연구 검증"], 0.7, 5.3, 5.45, 1.25, size=13.5, color=PPT_WHITE)
    cover_cards = [
        (6.15, 1.35, "지금의 위험", "객체별 안정화와 경로 근거로 필요한 후보만 선별", PPT_TEAL),
        (6.15, 3.05, "목적지까지", "TMAP 전역 경로 + 조건부 점자블록 단거리 steering", PPT_YELLOW),
        (6.15, 4.75, "시설 개선으로", "손상 후보를 저장·검수해 기관용 자료로 준비", PPT_CORAL),
    ]
    for x, y, title, body, accent in cover_cards:
        ppt_rect(slide, x, y, 3.15, 1.25, fill=RGBColor(34, 50, 68), line=accent)
        ppt_rect(slide, x, y, 0.08, 1.25, fill=accent)
        ppt_text(slide, title, x + 0.22, y + 0.18, 2.7, 0.3, size=16, bold=True, color=accent)
        ppt_text(slide, body, x + 0.22, y + 0.53, 2.65, 0.5, size=11.5, color=PPT_WHITE)
    ppt_text(slide, AS_OF_DATE, 0.7, 6.82, 1.7, 0.3, size=11, color=PPT_WHITE)

    # 2. Official deliverables matrix is preserved from the immutable template.

    # 3. Need
    slide = slides[2]
    ppt_title(slide, "보행 중 정보 공백과 시설 신고 부담을 함께 해결", 3, section="환경 분석 · 필요성")
    ppt_add_picture_fit(slide, DOC_ASSETS["problem_solution"], 0.55, 0.92, 8.9, 5.75)
    ppt_text(slide, "배경 근거: 2023년 안내시설 적정 설치율 57.2% · 사고율이나 앱 효과로 확대하지 않음", 0.7, 6.7, 8.6, 0.25, size=10.5, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 4. Differentiation
    slide = slides[3]
    ppt_title(slide, "장면 설명·경로 안내·신고 사이의 빈틈을 연결", 4, section="환경 분석 · 차별성")
    ppt_table(slide, [
        ["접근 방식", "공식 설명 기준의 초점", "WalkSafe가 연결하는 지점"],
        ["Google Lookout", "객체·텍스트·주변 설명", "현재 경로 위험 후보를 별도로 선별"],
        ["TMAP API", "장소 검색·보행자 경로", "전역 경로와 카메라 앞 단거리 방향 결합"],
        ["안전신문고", "사용자가 사진·위치·내용 입력", "손상 후보를 검수 가능한 자료로 구조화"],
        ["WalkSafe Assist", "선별 안내 + 경로 + 검수 자료", "즉시 안내와 사후 시설 개선 흐름을 분리·연결"],
    ], 0.45, 1.02, 9.1, 4.8, font_size=12.5, widths=[1.35, 2.65, 3.15])
    ppt_label(slide, "차별성: 현재 위험 선별 · 조건부 경로 보조 · 사람 검수 후 자료화", 0.7, 6.08, 6.1, fill=PPT_TEAL)
    ppt_text(slide, "공식 기능 초점 비교이며 경쟁 서비스의 기능 부재를 단정하지 않습니다.", 0.75, 6.55, 8.4, 0.25, size=10.5, color=PPT_MUTED)

    # 5. Data collection and integration
    slide = slides[4]
    ppt_title(slide, "6개 출처를 13개 보행 관련 class로 통합", 5, section="데이터 수집·처리 정의서")
    ppt_table(slide, [
        ["로컬 반영 출처", "이미지", "bbox", "역할"],
        ["COCO 2017", "73,866", "365,681", "일반 객체"],
        ["AIHub 186 베리어프리존", "64,220", "154,236", "보행 환경"],
        ["AIHub 513 도로시설물", "40,124", "64,042", "보행 시설"],
        ["AIHub 189 Surface", "2,284", "4,097", "보도 표면"],
        ["AIHub 572 승인 subset", "16,005", "19,751", "전동킥보드"],
        ["AIHub 189 수동 승인", "7", "7", "전동킥보드 보완"],
    ], 0.45, 0.98, 9.1, 5.35, font_size=10.7, widths=[2.6, 1.0, 1.0, 1.8])
    ppt_text(slide, "train 167,759장 / val 28,747장 · 총 196,506장 / 607,814 bbox · 독립 test 없음", 0.62, 6.48, 8.8, 0.35, size=12, bold=True, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_text(slide, "Provenance: 저장소의 AIHub183 표기는 레거시 별칭이며 공식 출처 ID로 사용하지 않음", 0.65, 6.82, 8.7, 0.22, size=8.5, color=PPT_MUTED)

    # 6. Data/model quality gate
    slide = slides[5]
    ppt_title(slide, "내부 validation과 배포 승인을 분리", 6, section="데이터 품질·모델 gate")
    ppt_add_picture_fit(slide, DOC_ASSETS["model_quality"], 0.5, 0.98, 9.0, 5.45)
    ppt_text(slide, "후보 epoch270: P 0.62087 · R 0.56716 · mAP50 0.55403 · mAP50-95 0.41651 · sequence 599개 겹침", 0.55, 6.48, 8.9, 0.28, size=10.3, bold=True, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_label(slide, "UC-08 모델 개선과 Android 연구 지원 · deployment_eligible = false", 0.7, 6.82, 6.0, fill=PPT_CORAL)

    # 7. Requirements 1
    slide = slides[6]
    ppt_title(slide, "사용자 기능은 구현 단계와 현장 검증을 분리", 7, section="요구사항 정의서 1/2")
    ppt_table(slide, [
        ["ID", "요구사항", "확인 단계", "다음 검증"],
        ["REQ-001", "Web/PWA 카메라 기반 위험 감지", "통제 E2E", "실폰 지연·오경고"],
        ["REQ-002", "실탐지와 fake/demo 결과의 명시적 분리", "통제 E2E", "승인 모델·Release"],
        ["REQ-003", "Android ARCore·TFLite 연구 보조", "실기기 invoke", "전체 Device FIELD"],
        ["REQ-004", "객체별 안정 추적과 경로 위험 선별", "자동검증", "class별 현장 품질"],
        ["REQ-005", "손상 점자블록 전용 자동·음성 요청 신고", "통제 E2E", "실폰 camera/GPS/mic"],
        ["REQ-006", "짧은 TTS·진동 위험 피드백", "자동검증", "청취성·진동·TalkBack"],
        ["REQ-007", "음성 신고·목적지·길안내 제어", "자동검증", "모바일 mic·소음"],
    ], 0.35, 0.98, 9.3, 5.72, font_size=11.2, widths=[0.8, 3.75, 1.35, 2.0])
    ppt_label(slide, "구현·자동검증 확보 / 사용자·실외 FIELD 대기", 0.55, 6.82, 4.1, fill=PPT_TEAL)

    # 8. Requirements 2
    slide = slides[7]
    ppt_title(slide, "신고·운영·데이터의 책임 경계까지 추적", 8, section="요구사항 정의서 2/2")
    ppt_table(slide, [
        ["ID", "요구사항", "확인 단계", "다음 검증"],
        ["REQ-008", "신고 이미지·metadata의 신뢰 경계와 원자 저장", "통제 E2E", "운영 복구·보존"],
        ["REQ-009", "PostGIS 위치 저장과 검수 추적", "격리 DB", "운영 lifecycle"],
        ["REQ-010", "관리자 검수·상태관리·기관용 자료 준비", "자동검증", "조직 IAM·접수"],
        ["REQ-011", "설치 가능한 PWA와 안전한 offline 경계", "headless", "실폰·안정 HTTPS"],
        ["REQ-012", "목적지 검색과 보행 경로 안내", "provider smoke", "실외 과업·복귀"],
        ["REQ-013", "모델·데이터·class order 추적", "내부 평가·hash", "독립 test·재분할"],
        ["REQ-014", "개인정보·인증·반출·보존 경계", "정책 자동검증", "조직 운영·시행"],
    ], 0.35, 0.98, 9.3, 5.72, font_size=11.2, widths=[0.8, 3.75, 1.35, 2.0])
    ppt_label(slide, "종합 상태 PARTIAL · 외부 제출은 사람의 책임", 0.55, 6.82, 4.15, fill=PPT_YELLOW, color=PPT_INK)

    # 9. Use cases
    slide = slides[8]
    ppt_title(slide, "보행 사용자·시스템·관리자의 책임을 분리", 9, section="유스케이스 정의서")
    ppt_add_picture_fit(slide, DOC_ASSETS["use_case_swimlane"], 0.45, 0.9, 9.1, 5.8)
    ppt_text(slide, "lane 내부 UC-01~09와 화살표가 주체별 행동·검수 책임·외부 제출 경계를 표시", 0.65, 6.72, 8.7, 0.25, size=10.2, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 10. Service scenario
    slide = slides[9]
    ppt_title(slide, "한 번의 보행을 안내와 시설 자료화로 연결", 10, section="서비스 시나리오")
    ppt_add_picture_fit(slide, DOC_ASSETS["value_flow"], 0.55, 0.98, 8.9, 5.65)
    ppt_text(slide, "오탐이 곧 외부 신고가 되지 않도록 Admin human review를 책임 경계로 둡니다.", 0.75, 6.68, 8.5, 0.3, size=11, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 11. Architecture
    slide = slides[10]
    ppt_title(slide, "Web/PWA부터 PostGIS·Admin까지 연결", 11, section="서비스 구성도")
    ppt_add_picture_fit(slide, DOC_ASSETS["architecture"], 0.55, 0.95, 8.9, 5.7)
    ppt_text(slide, "실선은 구현 연결 · 점선은 Android 연구 보조와 사람의 기관 외부 제출", 0.75, 6.68, 8.5, 0.3, size=11, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 12. Main flow
    slide = slides[11]
    ppt_title(slide, "탐지됨이 곧 경고함을 뜻하지 않음", 12, section="서비스 흐름도")
    ppt_add_picture_fit(slide, DOC_ASSETS["risk_flow"], 0.6, 0.95, 8.8, 5.7)
    ppt_text(slide, "UC-01 보행 위험 감지 · UC-02 TTS·진동 위험 안내 — 탐지 존재만으로 경고하지 않는 fail-closed 처리", 0.75, 6.68, 8.5, 0.3, size=10.5, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 13. UI/UX
    slide = slides[12]
    ppt_title(slide, "위험 발화 우선순위와 실패 상태를 제품 규칙으로 고정", 13, section="UI/UX 정의서")
    priority_cards = [
        (0.7, "1 위험", "즉시 발화\n녹음·STT upload 취소\n낮은 우선순위 TTS 중지", PPT_CORAL),
        (3.65, "2 상호작용", "사용자 요청·확인\n위험 발화와 겹치지 않음\nunknown이면 재질문", PPT_YELLOW),
        (6.6, "3 길안내", "전역 TMAP-only\nWeb/미지원 Android 비계량 advisory\nAndroid metric strict fallback", PPT_TEAL),
    ]
    for x, title, body, color in priority_cards:
        ppt_rect(slide, x, 1.2, 2.7, 2.15, fill=PPT_LIGHT, line=color)
        ppt_text(slide, title, x + 0.2, 1.48, 2.3, 0.35, size=18, bold=True, color=color, align=PP_ALIGN.CENTER)
        ppt_text(slide, body, x + 0.2, 2.0, 2.3, 1.0, size=12.5, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_table(slide, [
        ["실패 상태", "사용자 표시", "안전한 대체 행동"],
        ["camera·권한", "원인+재시도", "탐지 중지·기존 위험 해제"],
        ["GPS·TMAP", "위치/경로 오류", "tactile 우선 해제·경로 재요청"],
        ["mic·STT", "시간초과/unknown", "수동 조작·재발화"],
        ["network·server", "unavailable", "오프라인 안전 기능처럼 위장하지 않음"],
    ], 0.7, 3.75, 8.6, 2.35, font_size=11.5, widths=[1.7, 2.8, 3.5])
    ppt_label(slide, "UC-04 음성 제어와 명시적 신고 · 위험 > 상호작용 > 길안내", 0.75, 6.35, 5.7, fill=PPT_INK)
    ppt_text(slide, "색상+문구 병기 · aria-live/TalkBack 순서 · 48px 주요 조작 · 반복 알림 cooldown", 0.75, 6.78, 8.5, 0.22, size=9.5, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 14. Web menu
    slide = slides[13]
    ppt_title(slide, "Web/PWA는 다섯 가지 핵심 보행 과업에 집중", 14, section="메뉴 구성도 1/2")
    ppt_table(slide, [
        ["사용 흐름", "기능", "실패·경계 표시"],
        ["시작", "Web frame·Web report·Android report·telemetry 각각 기본 OFF", "철회 시 관련 in-flight 취소·hidden"],
        ["보행 보조", "탐지→위험 선별→TTS·진동", "fake/real·stale·unavailable"],
        ["음성", "목적지·신고·길안내 intent", "mic/STT timeout·unknown"],
        ["길안내", "TMAP 경로·tactile steering·복귀", "GPS stale·off-route·provider"],
        ["설치", "opt-in PWA·update 승인", "app shell만 offline"],
    ], 0.6, 1.05, 8.8, 4.95, font_size=13, widths=[1.35, 3.0, 2.6])
    ppt_label(slide, "보행 보조 중지·인증 상실·pagehide 시 camera·sensor·voice·navigation 종료", 0.75, 6.25, 7.0, fill=PPT_YELLOW, color=PPT_INK)
    ppt_text(slide, "UC-06 Web/PWA 설치·실행", 0.75, 6.7, 8.5, 0.22, size=9.8, bold=True, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 15. Admin menu and status
    slide = slides[14]
    ppt_title(slide, "Admin은 확인·결정·반출·외부 제출을 순서대로 수행", 15, section="메뉴 구성도 2/2")
    admin_steps = [
        (0.55, "1", "필터·상세", "이미지·위치·모델·trigger 확인", PPT_TEAL),
        (2.8, "2", "검수 결정", "new → reviewed\nnote·reason 감사", PPT_TEAL),
        (5.05, "3", "agency 반출", "reviewed·damage\nGPS≤15m·non-fake\n최대 10,000행", PPT_YELLOW),
        (7.3, "4", "사람 제출", "별도 기관 채널 제출\n후속 조치 뒤 resolved", PPT_CORAL),
    ]
    for x, number, title, body, color in admin_steps:
        ppt_rect(slide, x, 1.45, 2.05, 2.25, fill=PPT_LIGHT, line=color)
        ppt_label(slide, number, x + 0.15, 1.62, 0.38, fill=color, color=PPT_WHITE if color != PPT_YELLOW else PPT_INK)
        ppt_text(slide, title, x + 0.18, 2.12, 1.7, 0.34, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
        ppt_text(slide, body, x + 0.18, 2.58, 1.7, 0.75, size=11.5, color=PPT_MUTED, align=PP_ALIGN.CENTER)
        if x < 7.3:
            ppt_text(slide, "→", x + 2.02, 2.25, 0.28, 0.4, size=22, bold=True, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_table(slide, [
        ["책임 경계", "자동 처리", "사람이 확인하는 것"],
        ["중복", "10m·±60초 후보 태그", "병합 여부·상태 결정"],
        ["기관용", "reviewed·damage·GPS≤15m·non-fake 검사", "외부 제출·접수번호 기록"],
        ["감사", "status/export/read append-only 기록", "목록·상세·image의 actor/resource/time/purpose"],
    ], 0.65, 4.25, 8.7, 2.0, font_size=11.5, widths=[1.4, 3.1, 3.0])
    ppt_text(slide, "기관 자동 API 접수는 구현하지 않았습니다.", 0.75, 6.55, 8.5, 0.25, size=10.5, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 16. Web screen design
    slide = slides[15]
    ppt_title(slide, "보행 UI는 상태 변화와 대체 행동을 먼저 설계", 16, section="화면 설계서 1/2")
    ppt_add_picture_fit(slide, DOC_ASSETS["ui_states"], 0.55, 0.95, 8.9, 5.6)
    ppt_text(slide, "입력·출력 목록보다 사용자가 ‘현재 무엇이 가능하고 다음에 무엇을 해야 하는지’를 알 수 있게 설계", 0.75, 6.62, 8.5, 0.3, size=10.8, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 17. Admin screen design
    slide = slides[16]
    ppt_title(slide, "운영 화면은 근거 확인 뒤에만 상태·반출을 허용", 17, section="화면 설계서 2/2")
    ppt_table(slide, [
        ["화면 상태", "운영자 과업", "성공 피드백", "실패·보호"],
        ["목록", "구조화 필터·50건 page", "조건·건수+read audit", "purpose 누락·권한 거부"],
        ["상세", "이미지·좌표·model·trigger 확인", "상태 이력+read audit", "없는 파일·잘못된 ID"],
        ["상태 변경", "new→reviewed→resolved", "actor·note·reason 감사", "역할·전이 검증"],
        ["집계", "0.001° cell Top 5·bounds", "집계와 exact source 분리", "minimum profile과 혼용 거부"],
        ["반출", "internal/minimum/agency 선택", "audit_id·manifest", "10,000행 초과·부적격 거부"],
    ], 0.45, 1.0, 9.1, 4.95, font_size=12, widths=[1.25, 2.75, 2.4, 2.1])
    ppt_label(slide, "UC-05 운영자 신고 검수 · 조직 IdP/RBAC·실제 기관 운영은 후속 RELEASE", 0.7, 6.25, 7.2, fill=PPT_YELLOW, color=PPT_INK)

    # 18. Real Web captures
    slide = slides[17]
    ppt_title(slide, "구현 화면 — Web/PWA 합성 입력", 18)
    ppt_add_picture_fit(slide, DOC_ASSETS["web_mobile"], 0.55, 1.0, 2.75, 5.7)
    ppt_add_picture_fit(slide, DOC_ASSETS["web_desktop"], 3.55, 1.0, 5.9, 4.25)
    ppt_text(slide, "모바일 viewport", 0.75, 6.72, 2.35, 0.23, size=10, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_bullets(slide, ["개인정보 없는 정지 입력", "fake-v2 source를 화면에 명시", "권한·상태·레이아웃 확인용", "실모델·실외 성능 근거와 분리"], 3.75, 5.48, 5.3, 1.22, size=12.5)

    # 19. Real Admin
    slide = slides[18]
    ppt_title(slide, "구현 화면 — Admin 로컬 시험자료", 19)
    ppt_add_picture_fit(slide, DOC_ASSETS["admin"], 0.6, 1.0, 8.8, 5.55)
    ppt_text(slide, "local fixture · 개인정보 없음 · 운영 PostGIS/보안 검증과 분리", 0.75, 6.67, 8.5, 0.3, size=11, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 20. ERD
    slide = slides[19]
    ppt_title(slide, "신고·3종 감사·공유 제한을 5개 application table로 분리", 20, section="엔티티 관계도")
    ppt_add_picture_fit(slide, DOC_ASSETS["erd"], 0.6, 0.95, 8.8, 5.65)
    ppt_text(slide, "reports + status/export/read append-only audits + actor_rate_limit_events(60초, non-audit) · DB 밖 journal/image", 0.75, 6.68, 8.5, 0.3, size=10.2, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 21. Risk processing
    slide = slides[20]
    ppt_title(slide, "새 프레임·시간·경로 gate를 모두 통과해야 안내", 21, section="기능 처리도 1/2")
    ppt_add_picture_fit(slide, DOC_ASSETS["risk_timeline"], 0.55, 0.9, 8.9, 4.95)
    ppt_text(slide, "UC-07 목적지 검색과 보행 경로 안내", 0.7, 5.9, 8.6, 0.22, size=10.2, bold=True, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_label(slide, "비계량: 3 distinct frame+700ms · Web IMU 선택 · Android fresh IMU 필수 — 실패하면 TMAP_ONLY", 0.7, 6.12, 8.6, fill=PPT_TEAL)
    ppt_text(slide, "구현 임계값은 사용자 실험 최적값이 아니며 실외 오경고/분·누락률·지연 p50/p95를 후속 측정", 0.7, 6.62, 8.6, 0.28, size=10.1, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 22. Report flow
    slide = slides[21]
    ppt_title(slide, "신고는 신뢰 gate·원자 저장·사람 검수를 통과", 22, section="기능 처리도 2/2")
    ppt_add_picture_fit(slide, DOC_ASSETS["report_flow"], 0.55, 0.95, 8.9, 5.45)
    ppt_text(slide, "UC-03 손상 점자블록 자동 신고 · primary ID+duplicate_count 반환 · 다른 후보 ID는 admin 전용 · 기관 접수는 사람의 외부 행위", 0.7, 6.5, 8.6, 0.42, size=10.1, color=PPT_MUTED, align=PP_ALIGN.CENTER)

    # 23. Risk algorithm
    slide = slides[22]
    ppt_title(slide, "class마다 안내·신고·표시 정책을 분리", 23, section="알고리즘 명세서")
    ppt_table(slide, [
        ["분류", "정책", "출력"],
        ["damaged tactile", "report-only", "자동/음성 요청 경로"],
        ["normal tactile", "기존 risk 기본 no-alert", "비계량 low 좌/중앙/우 가능"],
        ["crosswalk", "advisory 제외", "TMAP 정렬 보조/표시"],
        ["curb·uneven", "surface warning", "안정화+경로 조건 후 경고 후보"],
        ["e-scooter", "path obstacle", "안정화+경로 조건 후 경고 후보"],
        ["그 외 일반 객체", "stable + path/approach/TTC", "근거 부족 시 display-only"],
    ], 0.6, 1.05, 8.8, 4.65, font_size=12.5, widths=[1.5, 2.4, 2.0])
    ppt_bullets(slide, ["위험 > 상호작용 > 길안내 > advisory; advisory는 TMAP을 선점하지 않음", NON_METRIC_CAPABILITY_SUMMARY, "Web 13 rule intents callback dispatch와 Android 목적지/후보 번호·취소·다음 안내·중지 action"], 0.75, 5.85, 8.3, 1.18, size=9.8)

    # 24. Report algorithm
    slide = slides[23]
    ppt_title(slide, "캡처 신뢰·중복·반출 최소화를 세 경계로 관리", 24, section="데이터 처리 상세")
    data_panels = [
        (0.65, "1. 신고 신뢰", "손상 class만\n신뢰도 ≥0.70\nGPS 오차 ≤15m\n3개 frame·700ms\n위치 도약 ≤20m\nimage·model·시각 재검증", PPT_TEAL),
        (3.55, "2. 저장·중복", "durable journal\n원자 rename→DB commit\nstartup reconciliation\n같은 class·10m·1분 후보", PPT_YELLOW),
        (6.45, "3. 검수·반출", "목록/상세/image 조회 감사\nactor·resource·time·purpose\n내부·최소·기관용\n반출 감사 append-only\n사람이 기관에 제출", PPT_CORAL),
    ]
    for x, title, body, color in data_panels:
        ppt_rect(slide, x, 1.35, 2.55, 4.35, fill=PPT_LIGHT, line=color)
        ppt_rect(slide, x, 1.35, 0.1, 4.35, fill=color)
        ppt_text(slide, title, x + 0.22, 1.7, 2.1, 0.35, size=17, bold=True, color=color, align=PP_ALIGN.CENTER)
        ppt_text(slide, body, x + 0.16, 2.25, 2.23, 2.95, size=12.8, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_label(slide, "오탐이 곧 외부 신고가 되지 않도록 human review를 책임 경계로 둠", 0.8, 6.15, 6.4, fill=PPT_TEAL)
    ppt_text(slide, "UC-09 개인정보와 운영 보호 · receipt 기록 도구는 있으나 실제 기관 접수 증거는 아직 없습니다.", 0.8, 6.65, 8.3, 0.25, size=10.2, color=PPT_MUTED)

    # 25. Program list
    slide = slides[24]
    ppt_title(slide, "제품 흐름별 모듈 책임을 분리", 25, section="프로그램 목록")
    ppt_table(slide, [
        ["영역", "핵심 모듈", "역할"],
        ["Web", "page.tsx · hooks", "주 사용자 상태기"],
        ["Risk", "risk-evaluator · risk-guidance", "위험·발화 정책"],
        ["Report", "useAutoReportV2 · report-api", "자동·음성 신고"],
        ["Backend", "reports · storage · read_audit", "journal 저장·검수·상태/반출/조회 감사"],
        ["Navigation", "navigation · tmap_pedestrian", "목적지·보행 경로"],
        ["Admin", "app/admin/page", "검수·필터·상태·CSV"],
        ["Model", "dataset/train/eval scripts", "13-class pipeline"],
        ["Android", "depth · inference", "연구 보조"],
    ], 0.55, 0.98, 8.9, 5.9, font_size=11.3, widths=[1.05, 2.55, 2.4])

    # 26. Table definition
    slide = slides[25]
    ppt_title(slide, "정확히 5개 application table로 원천·감사·공유 제한을 보존", 26, section="테이블 정의서")
    ppt_table(slide, [
        ["테이블", "핵심 컬럼", "책임·제약"],
        ["reports", "UUID·status·class·bbox·Point(4326)·metadata", "신고 원천 · enum/range/SRID CHECK"],
        ["report_status_audits", "report_id·previous/next·actor·note", "논리 참조 · append-only"],
        ["report_export_audits", "audit/requested ID·actor·profile·rows hash", "반출 목적·범위 · append-only"],
        ["report_read_audits", "actor·purpose·resource type/id·details·time", "list/detail/image 조회 · append-only"],
        ["actor_rate_limit_events", "actor digest·rate group·observed_at", "공유 원자 60초 window · non-audit"],
        ["DB 밖 journal/image", "UUID filename·MIME·image_path", "durable journal·startup reconciliation"],
    ], 0.4, 1.0, 9.2, 5.15, font_size=12, widths=[1.8, 3.7, 2.6])
    ppt_label(slide, "3종 감사만 UPDATE/DELETE trigger 거부", 0.65, 6.55, 4.1, fill=PPT_TEAL)

    # 27. Web source
    slide = slides[26]
    ppt_title(slide, "손상 시설은 안내가 아니라 신고로 분기", 27, section="핵심 소스코드 1/2")
    code = source_excerpt("apps/web/app/_walksafe/risk-evaluator.ts", "if (isReportableTactileDamage", 10)
    ppt_rect(slide, 0.7, 1.15, 5.25, 4.9, fill=RGBColor(249, 250, 251), line=PPT_BORDER)
    ppt_text(slide, code, 0.9, 1.36, 4.85, 4.45, size=10.5, color=PPT_INK, font="D2Coding")
    ppt_bullets(slide, ["완결된 실제 if 분기", "damaged tactile은 안내가 아닌 신고로 분리", "reportable=true · alertable=false", "대응 시험: report-only 유지와 기본 사용자 경고 금지"], 6.25, 1.35, 3.0, 4.4, size=14)
    ppt_label(slide, "fail-closed 위험 정책", 6.28, 5.95, 2.35, fill=PPT_TEAL)

    # 28. Backend source
    slide = slides[27]
    ppt_title(slide, "primary ID는 유지하고 내부 중복 후보만 제거", 28, section="핵심 소스코드 2/2")
    code = source_excerpt("backend/app/api/reports.py", "def _field_creation_response", 13)
    ppt_rect(slide, 0.7, 1.15, 5.25, 4.9, fill=RGBColor(249, 250, 251), line=PPT_BORDER)
    ppt_text(slide, code, 0.9, 1.33, 4.85, 4.5, size=9.8, color=PPT_INK, font="D2Coding")
    ppt_bullets(slide, ["완결된 실제 응답 정제 함수", "primary report의 id·duplicate_count는 유지", "다른 candidate ID만 field 응답에서 제거", "대응 시험: ID 유지·내부 후보 비노출"], 6.25, 1.35, 3.0, 4.4, size=14)
    ppt_label(slide, "조직 IdP/RBAC는 운영 전 필요", 6.28, 5.95, 2.9, fill=PPT_CORAL)

    # 29. Environment
    slide = slides[28]
    ppt_title(slide, "소스부터 검증까지 재현 체인으로 기록", 29, section="개발 환경")
    repro_steps = [
        (0.65, "1", "소스·환경", "Git 상태·source hash\nlockfile·실행 설정", PPT_INK),
        (2.95, "2", "빌드", "Web production\nAPK·Office 산출물", PPT_YELLOW),
        (5.25, "3", "산출물", "checkpoint·TFLite\nDOCX·PPTX SHA-256", PPT_CORAL),
        (7.55, "4", "검증", "suite별 격리 범위\nFIELD·RELEASE 분리", PPT_TEAL),
    ]
    for x, number, title, body, color in repro_steps:
        ppt_rect(slide, x, 1.55, 1.8, 3.25, fill=PPT_LIGHT, line=color)
        ppt_label(slide, number, x + 0.15, 1.75, 0.38, fill=color, color=PPT_WHITE if color != PPT_YELLOW else PPT_INK)
        ppt_text(slide, title, x + 0.15, 2.35, 1.5, 0.35, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
        ppt_text(slide, body, x + 0.15, 3.0, 1.5, 1.15, size=11.5, color=PPT_MUTED, align=PP_ALIGN.CENTER)
        if x < 7.5:
            ppt_text(slide, "→", x + 1.8, 2.85, 0.45, 0.4, size=22, bold=True, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_table(slide, [
        ["Web", "Next.js 16.2.6 · React 19.2.6"],
        ["Backend/DB", "FastAPI 0.128.8 · SQLAlchemy 2.0.49 · PostgreSQL/PostGIS"],
        ["Model/Android", "Ultralytics 8.4.48 · img768 · Kotlin 2.2.21 · LiteRT 1.4"],
        ["배포 보호", "DB sslmode=verify-full+gssencmode=disable · shared limiter/nonblocking readiness · VOICE_SERVICE_TOKEN/single runtime · auth OpenAPI"],
    ], 1.0, 5.3, 8.0, 1.4, font_size=10.8, widths=[1.7, 5.4])

    # 30. Software inspection
    slide = slides[29]
    ppt_title(slide, "무엇을 검증했고 무엇은 아직 검증하지 않았는지 분리", 30, section="S/W 기능 실사")
    ppt_text(slide, VERIFICATION_SCOPE_NOTE, 0.55, 0.72, 8.9, 0.2, size=8.2, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    metrics = [
        (0.45, str(VERIFICATION_COUNTS["unit_python"]), "Python unit", PPT_TEAL),
        (1.97, str(VERIFICATION_COUNTS["functional_python"]), "DB functional", PPT_TEAL),
        (3.49, str(VERIFICATION_COUNTS["integration_python"]), "integration", PPT_TEAL),
        (5.01, str(VERIFICATION_COUNTS["voice"]), "Voice", PPT_TEAL),
        (6.53, str(VERIFICATION_COUNTS["android_jvm"]), "Android JVM", PPT_TEAL),
        (8.05, ANDROID_DEVICE_CURRENT_METRIC, "device current", PPT_YELLOW),
    ]
    for x, value, label, color in metrics:
        ppt_metric(slide, value, label, x, 1.02, 1.42, accent=color)
    evidence_panels = [
        (0.55, "확인한 범위", "Web build·headless PWA\n격리 DB·정책·통합 회귀\nAndroid invoke 한정", PPT_TEAL),
        (3.55, "안전 경계", "suite 수치 비합산\nserver-sampled v2\nWebXR display-only\n외부 제출은 관리자 책임", PPT_YELLOW),
        (6.55, "남은 blocker", "실폰·실외 사용자 과업\nsequence-safe 재분할·독립 test\n조직 IAM·복구·기관 접수", PPT_CORAL),
    ]
    for x, title, body, color in evidence_panels:
        ppt_rect(slide, x, 2.55, 2.6, 2.65, fill=PPT_LIGHT, line=color)
        ppt_text(slide, title, x + 0.18, 2.85, 2.24, 0.35, size=17, bold=True, color=color, align=PP_ALIGN.CENTER)
        ppt_text(slide, body, x + 0.18, 3.5, 2.24, 1.25, size=12.3, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_text(slide, f"실행: server-sampled v2 · NFT {VERIFICATION_COUNTS['web_trace_count']} traces/{VERIFICATION_COUNTS['web_trace_files']:,} files · MediaRecorder→faster-whisper→browser speechSynthesis TTS · WebXR 표시 전용·위험판단 미연결", 0.55, 5.5, 8.9, 0.22, size=8.8, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_text(slide, "운영: 안전 API network-only · agency CSV 수동 제출·기관 API 없음 · val truncated 2,331장 제외·독립 test 없음", 0.55, 5.82, 8.9, 0.22, size=8.8, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_text(slide, "Voice: body copy 전 auth·Content-Length·global/actor/IP rate → copy 중 byte 상한 → copy 후 duration·inference queue", 0.55, 6.06, 8.9, 0.2, size=8.4, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_label(slide, "자동검증 PASS ≠ 실사용 안전 PASS", 3.1, 6.36, 3.8, fill=PPT_INK)

    # 31. Hardware
    slide = slides[30]
    ppt_title(slide, "기성 스마트폰 센서와 Android 연구 경로의 현재 수준", 31, section="H/W 기능 실사")
    android_panels = [
        (0.55, "실행 자원", "camera · GPS · 방향/동작\nmicrophone · speaker · vibration\n별도 제작 H/W 없음", PPT_INK),
        (
            3.55,
            "확인한 근거",
            "unified img768 13-class asset\n[1,768,768,3]→[1,300,6]\n"
            f"{ANDROID_DEVICE_HISTORICAL['executed_at']} historical instrumentation "
            f"{ANDROID_DEVICE_HISTORICAL['passed']}/{ANDROID_DEVICE_HISTORICAL['total']}\n"
            f"current: {ANDROID_DEVICE_CURRENT['status']}\nnative 음성 action JVM",
            PPT_TEAL,
        ),
        (6.55, "다음 FIELD", "대화형 camera→inference\nfallback 미발생·FPS\nbbox-depth·N보\nmic·TTS·진동·실외", PPT_CORAL),
    ]
    for x, title, body, color in android_panels:
        ppt_rect(slide, x, 1.35, 2.6, 4.35, fill=PPT_LIGHT, line=color)
        ppt_rect(slide, x, 1.35, 0.1, 4.35, fill=color)
        ppt_text(slide, title, x + 0.22, 1.75, 2.15, 0.4, size=18, bold=True, color=color, align=PP_ALIGN.CENTER)
        ppt_text(slide, body, x + 0.22, 2.55, 2.15, 2.45, size=14, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_label(slide, "Android는 연구 보조 · Web/PWA Release 근거와 분리", 0.75, 6.15, 4.7, fill=PPT_YELLOW, color=PPT_INK)
    ppt_text(slide, "WebXR depth 역시 중앙 거리 표시 전용이며 객체 위험 판단에 연결하지 않습니다.", 0.75, 6.65, 8.5, 0.25, size=10.5, color=PPT_MUTED)

    # 32. Video storyboard
    slide = slides[31]
    ppt_title(slide, "90초 제출 영상 촬영 계획 — 실제 결과 확보 후 완성", 32, section="동영상 촬영 콘티")
    ppt_table(slide, [
        ["구간", "장면", "증거·캡션"],
        ["0–10초", "57.2% 문제와 가치 제안", "공식 출처·측정되지 않은 효과 구분"],
        ["10–35초", "Web/PWA camera·탐지·위험 선별", "detector source·fake/real 명시"],
        ["35–50초", "음성 목적지·길안내", "faster-whisper·TMAP·browser TTS"],
        ["50–70초", "손상 후보 신고", "3 frame·700ms·GPS gate"],
        ["70–82초", "Admin 검수·agency 자료", "fixture·기관 수동 제출"],
        ["82–90초", "자동검증과 다음 실증", "suite별 결과·실폰/실외·독립 test"],
    ], 0.55, 1.02, 8.9, 5.35, font_size=11.7, widths=[1.0, 2.7, 3.3])
    ppt_label(slide, "현재는 촬영 콘티 · 실제 영상·QR·타임코드는 확보 뒤 반영", 0.75, 6.48, 5.8, fill=PPT_CORAL)

    # 33. Project management
    slide = slides[32]
    ppt_title(slide, "문제를 근거로 범위를 조정하고 다음 gate를 남김", 33, section="프로젝트 관리·문제 해결")
    timeline = [
        (0.65, "4–5월", "원안 계획\n실적 추정 안 함", PPT_MUTED),
        (2.85, "6월", "설계 착수\n요구·구조 구체화", PPT_TEAL),
        (5.05, "7월", "model·Web·Backend\n자동검증·invoke", PPT_TEAL),
        (7.25, "8–10월", "재분할·실폰·실외\n운영 실증 계획", PPT_YELLOW),
    ]
    for x, period, body, color in timeline:
        ppt_rect(slide, x, 1.0, 2.05, 1.35, fill=PPT_LIGHT, line=color)
        ppt_text(slide, period, x + 0.15, 1.18, 1.75, 0.3, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
        ppt_text(slide, body, x + 0.15, 1.6, 1.75, 0.55, size=11, color=PPT_MUTED, align=PP_ALIGN.CENTER)
        if x < 7:
            ppt_text(slide, "→", x + 2.02, 1.5, 0.28, 0.35, size=20, bold=True, color=PPT_MUTED, align=PP_ALIGN.CENTER)
    ppt_table(slide, [
        ["문제", "조치", "확보한 결과·남은 gate"],
        ["PWA field-session 503", "actor·token·session secret 분리", "install·update·offline shell 재검증"],
        [
            "Android runtime 불확실",
            "unified img768 asset과 device smoke",
            f"{ANDROID_DEVICE_CURRENT_LABEL} · {ANDROID_DEVICE_HISTORICAL_LABEL} · 전체 FIELD 대기",
        ],
        ["sequence 겹침·weak class", "deployment_eligible=false로 승인 보류", "재분할·독립 test·curb/uneven 개선"],
    ], 0.55, 2.8, 8.9, 3.1, font_size=11.5, widths=[2.1, 3.1, 3.5])
    ppt_label(slide, "실패를 완료로 숨기지 않고 다음 검증 조건으로 전환", 1.0, 6.25, 5.4, fill=PPT_TEAL)

    # 34. End
    slide = slides[33]
    ppt_rect(slide, 0, 0, 10, 7.5, fill=PPT_INK)
    ppt_rect(slide, 0, 0, 0.16, 7.5, fill=PPT_TEAL)
    ppt_text(slide, "현재 위험은 필요한 순간에 선별 안내하고,", 0.75, 0.75, 8.5, 0.62, size=28, bold=True, color=PPT_WHITE, align=PP_ALIGN.CENTER)
    ppt_text(slide, "손상 시설은 사람이 검수할 수 있는 자료로 남깁니다.", 0.65, 1.42, 8.7, 0.62, size=24, bold=True, color=RGBColor(94, 234, 212), align=PP_ALIGN.CENTER)
    ppt_text(slide, "다음 실증 3단계", 3.7, 2.18, 2.6, 0.32, size=16, bold=True, color=PPT_YELLOW, align=PP_ALIGN.CENTER)
    end_cards = [
        (0.8, "통제 실폰 실증", "camera·mic·TTS·진동\n지연·오경고·과업완료율", PPT_TEAL),
        (3.65, "사용자·복지기관 검토", "이해도·중지 성공·실패 대응\n안전 프로토콜과 접근성 의견", PPT_YELLOW),
        (6.5, "시설관리 운영 실증", "검수 승인율·처리시간\nagency 자료·실제 접수 receipt", PPT_CORAL),
    ]
    for x, title, body, color in end_cards:
        ppt_rect(slide, x, 2.65, 2.55, 2.65, fill=RGBColor(34, 50, 68), line=color)
        ppt_rect(slide, x, 2.65, 0.09, 2.65, fill=color)
        ppt_text(slide, title, x + 0.2, 3.02, 2.15, 0.55, size=15.5, bold=True, color=color, align=PP_ALIGN.CENTER)
        ppt_text(slide, body, x + 0.2, 3.82, 2.15, 0.9, size=12.5, color=PPT_WHITE, align=PP_ALIGN.CENTER)
    ppt_text(slide, "다음 milestone: sequence-safe 재평가와 실폰·실외 사용자 과업을 통과한 뒤 운영 실증으로 이동", 0.85, 5.75, 8.3, 0.45, size=14, bold=True, color=PPT_YELLOW, align=PP_ALIGN.CENTER)
    ppt_text(slide, "Web/PWA 주 사용자 앱 · Android 연구 보조 · 외부 제출은 사람의 책임", 1.0, 6.45, 8.0, 0.35, size=12.5, color=PPT_WHITE, align=PP_ALIGN.CENTER)

    with atomic_output_path(PPTX_OUTPUT) as temporary:
        presentation.save(temporary)
        normalize_ooxml(temporary)


def main() -> int:
    with submission_build_lock(REPO_ROOT):
        require_clean_source_revision(REPO_ROOT, ALL_SUBMISSION_GENERATED_PATHS)
        submission_toolchain_attestation(REPO_ROOT)
        require_no_unexpected_files(TEMPLATE_DIR, ALLOWED_TEMPLATE_FILES)
        verify_inputs()
        build_docx()
        build_pptx()
        for source, expected in EXPECTED_HASHES.items():
            if sha256(source) != expected:
                raise SystemExit(f"source template mutated while building: {source.name}")
        write_build_manifest()
        require_exact_file_set(TEMPLATE_DIR, ALLOWED_TEMPLATE_FILES)
    print(f"generated {DOCX_OUTPUT.relative_to(REPO_ROOT)} ({DOCX_OUTPUT.stat().st_size} bytes)")
    print(f"generated {PPTX_OUTPUT.relative_to(REPO_ROOT)} ({PPTX_OUTPUT.stat().st_size} bytes)")
    print(f"generated {FORM_BUILD_MANIFEST.relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

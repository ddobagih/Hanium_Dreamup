#!/usr/bin/env python3
"""Validate the completed development report and production design forms."""

from __future__ import annotations

import sys

if __name__ == "__main__":
    _verified_submission_commit = getattr(
        sys.modules.get("__main__"), "_walksafe_submission_source_commit", None
    )
    _startup_flags = (
        sys.flags.isolated,
        sys.flags.no_site,
        sys.flags.ignore_environment,
        int(getattr(sys.flags, "safe_path", False)),
        sys.flags.dont_write_bytecode,
    )
    if (
        any(value != 1 for value in _startup_flags)
        or "site" in sys.modules
        or not isinstance(_verified_submission_commit, str)
        or len(_verified_submission_commit) != 40
        or any(character not in "0123456789abcdef" for character in _verified_submission_commit)
    ):
        raise SystemExit(
            "WalkSafe submission CLI requires Python -I -S -B through the submission runner"
        )

import argparse
import hashlib
import json
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from PIL import Image

try:
    from docx import Document
    from pptx import Presentation
except ImportError as exc:  # pragma: no cover - actionable environment failure
    raise SystemExit(
        "python-docx and python-pptx are required; run with .venv/bin/python"
    ) from exc

from submission_manifest_policy import (
    ALL_SUBMISSION_GENERATED_PATHS,
    ASSET_DIRECTORY_FILES,
    DESIGN_ASSET_PATHS,
    DESIGN_BUILD_INPUT_PATHS,
    DESIGN_DOCUMENT_NAMES,
    DESIGN_OUTPUT_NAMES,
    FINAL_OUTPUT_NAMES,
    FINAL_SECTION_PATHS,
    FORM_ARTIFACT_PATHS,
    FORM_BUILD_INPUT_PATHS,
    FORM_TEMPLATE_FILES,
    build_tool_provenance,
    record_bundle_sha256,
    require_exact_file_set,
    require_exact_manifest_paths,
    require_review_identity,
    run_attested_submission_tool,
    source_revision,
    submission_toolchain_attestation,
)
from validate_submission_materials_20260710 import (
    validate_android_device_verification,
    validate_core_semantic_facts,
)


REPO_ROOT = Path(__file__).resolve().parents[1]
FACTS_PATH = REPO_ROOT / "docs/submission/form_materials/09_제출_사실_기준.json"

APPROVED_SOURCE_FREEZE_POLICY = (
    "사용자가 승인한 source-freeze commit의 자동검증 snapshot이며 최종 source-freeze "
    "근거로 채택했다. Field 또는 Release 증거로 확대하지 않고 PASS와 FAIL을 함께 "
    "기록하며 서로 다른 계층 수를 합산하지 않는다. current source-freeze Android "
    "device run은 NOT_RUN_CURRENT_SOURCE_FREEZE로 기록하고, 2026-07-13 SM-G981N "
    "2/2는 별도 historical evidence로만 보존한다."
)


def verification_snapshot_boundary_and_policy(snapshot: object) -> str:
    """Independently validate the forms snapshot state and exact policy contract."""

    if not isinstance(snapshot, dict):
        raise ValueError("canonical facts must contain a verification snapshot")
    status = snapshot.get("status")
    executed_at = snapshot.get("executed_at")
    policy = snapshot.get("policy")
    try:
        if (
            not isinstance(executed_at, str)
            or re.fullmatch(r"\d{4}-\d{2}-\d{2}", executed_at) is None
        ):
            raise ValueError
        datetime.strptime(executed_at, "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError("verification snapshot executed_at must be an ISO date") from exc

    if status == f"HISTORICAL_PRE_SOURCE_FREEZE_{executed_at}":
        boundary = "과거 pre-source-freeze snapshot · 현재 hardening 수치 아님"
        expected_policy = (
            f"이 snapshot은 {executed_at} 기준의 과거 계층 증거이며 현재 hardening 수치가 "
            "아니다. PASS와 FAIL을 함께 기록하고 Unit·격리 DB·build를 Field 또는 Release "
            "증거로 확대하지 않으며, source freeze 후 전체 재실행 수치로 교체한다. current "
            "source-freeze에서 Android device run이 없으면 NOT_RUN_CURRENT_SOURCE_FREEZE로 "
            "기록하고, 2026-07-13 SM-G981N 2/2는 별도 historical evidence로만 보존한다."
        )
        label = "historical"
    elif status == f"SOURCE_FREEZE_CANDIDATE_{executed_at}":
        boundary = "source-freeze 후보 검증 snapshot · 최종 승인 source-freeze 아님"
        expected_policy = (
            f"이 snapshot은 {executed_at} hardening 작업 후보의 임시 자동검증 수치이며, 새 "
            "clean source candidate의 전체 재현 대조를 통과해야 채택한다. 사용자가 승인한 "
            "최종 source-freeze commit이나 Field 또는 Release 증거가 아니고 PASS와 FAIL을 "
            "함께 기록하며 서로 다른 계층 수를 합산하지 않는다. current source-freeze "
            "Android device run은 NOT_RUN_CURRENT_SOURCE_FREEZE로 기록하고, 2026-07-13 "
            "SM-G981N 2/2는 별도 historical evidence로만 보존한다."
        )
        label = "candidate"
    elif status == f"SOURCE_FREEZE_APPROVED_{executed_at}":
        boundary = "승인된 source-freeze 검증 snapshot"
        expected_policy = APPROVED_SOURCE_FREEZE_POLICY
        label = "approved"
    else:
        raise ValueError(f"unsupported verification snapshot status: {status}")
    if policy != expected_policy:
        raise ValueError(
            f"verification snapshot must use the exact {label} policy contract"
        )
    return boundary


with FACTS_PATH.open(encoding="utf-8") as handle:
    FACTS = json.load(handle)
if FACTS.get("schema_version") != "walksafe.submission_facts.v2":
    raise SystemExit(f"unexpected facts schema: {FACTS.get('schema_version')!r}")
VERIFICATION = FACTS["verification_snapshot"]
VERIFICATION_COUNTS = VERIFICATION["counts"]
VERIFICATION_STATUS = VERIFICATION.get("status")
VERIFICATION_EXECUTED_AT = VERIFICATION.get("executed_at")
try:
    VERIFICATION_SCOPE_BOUNDARY = verification_snapshot_boundary_and_policy(
        VERIFICATION
    )
except ValueError as exc:
    raise SystemExit(str(exc)) from exc
_device_errors: list[str] = []
_device_spec = validate_android_device_verification(VERIFICATION, _device_errors)
if _device_spec is None or _device_errors:
    raise SystemExit("; ".join(_device_errors))
(
    ANDROID_DEVICE_CURRENT,
    ANDROID_DEVICE_HISTORICAL,
    ANDROID_DEVICE_CURRENT_LABEL,
    ANDROID_DEVICE_HISTORICAL_LABEL,
) = _device_spec
FINAL_OUTPUT_DIR = REPO_ROOT / "docs/submission/final"

ORIGINALS = {
    "docx": (
        REPO_ROOT / "templates/[서식1] 2026 한이음 드림업 개발보고서 양식.docx",
        "b8d52b810760c54c61646566ea3ee78208f43d7b1c0e3b65d53d99dc8cde22e3",
    ),
    "pptx": (
        REPO_ROOT / "templates/[서식2] 2026년_제작설계서_일반(개인정보 기재x).pptx",
        "140628cdd6b160c8d86a2919ce01620c274c54362bd71c9e82a77670a7621336",
    ),
}
OUTPUT_PAIRS = {
    "docx": (
        REPO_ROOT / "templates/개발보고서 양식.docx",
        FINAL_OUTPUT_DIR / "개발보고서 양식.docx",
    ),
    "pptx": (
        REPO_ROOT / "templates/제작설계서_일반.pptx",
        FINAL_OUTPUT_DIR / "제작설계서_일반.pptx",
    ),
}
FINAL_MANIFEST_PATH = FINAL_OUTPUT_DIR / "BUILD_MANIFEST.json"
FINAL_README_PATH = FINAL_OUTPUT_DIR / "README.md"
FORM_BUILD_MANIFEST_PATH = REPO_ROOT / "templates/SUBMISSION_BUILD_MANIFEST.json"
SOURCE_PARITY_PATHS = {
    "requirements": REPO_ROOT / "docs/submission/deliverables/요구사항_정의서.md",
    "use_cases": REPO_ROOT / "docs/submission/deliverables/유스케이스_정의서.md",
    "trace": REPO_ROOT / "docs/submission/drafts/요구사항_추적표.md",
    "report_manuscript": REPO_ROOT / "docs/submission/form_materials/04_개발보고서_본문원고.md",
    "slide_manuscript": REPO_ROOT / "docs/submission/form_materials/05_제작설계서_34장_원고.md",
}

OFFICIAL_REPORT_HEADINGS = (
    "I. 프로젝트 개요",
    "II. 프로젝트 내용",
    "III. 프로젝트 수행 내용",
    "Ⅳ. 기대효과 및 활용분야",
)
OFFICIAL_EVALUATION_BARS = (
    "※ 평가항목 : 기획력 (필요성, 차별성)",
    "※ 평가항목 : 기술력 (기능구체성, 난이도, 완성도)",
    "※ 평가항목 : 수행능력 (문서완성도, 문제해결능력, 수행충실성)",
    "※ 평가항목 : 기획력 (활용가능성)",
)

FORBIDDEN_VISIBLE_TEXT = {
    "English sample text": re.compile(r"\bsamples?\b", re.IGNORECASE),
    "English placeholder text": re.compile(r"\bplaceholders?\b", re.IGNORECASE),
    "Korean sample/template instruction": re.compile(
        r"샘플|본문의\s*예시\s*내용|예시\s*내용을\s*지우고"
    ),
    "asset placeholder": re.compile(r"\[\s*A\d{2,3}\s*\]", re.IGNORECASE),
    "Android-primary claim": re.compile(
        r"(?:Android[^.\n]{0,50}(?:주\s*(?:앱|플랫폼)|1차\s*실행\s*환경|primary\s*(?:app|platform|client))|"
        r"(?:주\s*(?:앱|플랫폼)|primary)\s*(?:은|는|:)?\s*Android)",
        re.IGNORECASE,
    ),
    "Kakao map residue": re.compile(r"(?<![A-Za-z])kakao(?![A-Za-z])|카카오", re.IGNORECASE),
    "stale voice intent count": re.compile(
        r"(?<!\d)11\s*(?:개\s*)?(?:rule\s*)?intents?(?![A-Za-z])", re.IGNORECASE
    ),
    "stale test or PostGIS status": re.compile(
        r"(?:63|82|145)\s+passed|"
        r"통합\s+Python\s*:\s*197\s+passed|"
        r"Root\s+Python[^.\n]{0,30}85\s+passed|"
        r"Android(?:\s+JVM)?\s+101(?:\s+tests?|\s+passed)?|"
        r"49\s+(?:skipped|개)|"
        r"(?:DB|PostGIS)\s+no-skip(?:\s+검증)?\s*(?:필요|남음)?",
        re.IGNORECASE,
    ),
    "stale 2026-07-13 suite counts": re.compile(
        r"(?:selected\s+Python\s+213|213\s+Python\s+unit|"
        r"functional\s+532|532\s+DB\s+functional|"
        r"integration\s+84|84\s+integration|"
        r"Voice\s+125|125\s+Voice|"
        r"Android\s+JVM\s+206|206\s+Android\s+JVM)",
        re.IGNORECASE,
    ),
    "brittle fixed test count": re.compile(
        r"(?<!\d)\d{2,4}\s+(?:tests?|passed)(?![\w-])",
        re.IGNORECASE,
    ),
}

CANONICAL_VERIFICATION_PATTERNS = {
    "Python unit": re.compile(
        rf"(?:Python\s+unit\s+{VERIFICATION_COUNTS['unit_python']}|"
        rf"{VERIFICATION_COUNTS['unit_python']}\s+Python\s+unit)",
        re.IGNORECASE,
    ),
    "DB functional": re.compile(
        rf"(?:functional\s+{VERIFICATION_COUNTS['functional_python']}|"
        rf"{VERIFICATION_COUNTS['functional_python']}\s+DB\s+functional)",
        re.IGNORECASE,
    ),
    "Integration": re.compile(
        rf"(?:integration\s+{VERIFICATION_COUNTS['integration_python']}|"
        rf"{VERIFICATION_COUNTS['integration_python']}\s+integration)",
        re.IGNORECASE,
    ),
    "Voice": re.compile(
        rf"(?:Voice\s+{VERIFICATION_COUNTS['voice']}|"
        rf"{VERIFICATION_COUNTS['voice']}\s+Voice)",
        re.IGNORECASE,
    ),
    "Android JVM": re.compile(
        rf"(?:Android\s+JVM\s+{VERIFICATION_COUNTS['android_jvm']}|"
        rf"{VERIFICATION_COUNTS['android_jvm']}\s+Android\s+JVM)",
        re.IGNORECASE,
    ),
    "Web trace": re.compile(
        rf"NFT\s+{VERIFICATION_COUNTS['web_trace_count']}\s+traces/"
        rf"{VERIFICATION_COUNTS['web_trace_files']:,}\s+(?:unique\s+)?files",
        re.IGNORECASE,
    ),
    "Android device current source-freeze": re.compile(
        re.escape(ANDROID_DEVICE_CURRENT_LABEL),
        re.IGNORECASE,
    ),
    "Android device historical evidence": re.compile(
        re.escape(ANDROID_DEVICE_HISTORICAL_LABEL),
        re.IGNORECASE,
    ),
}
for index, phrase in enumerate(FACTS["forbidden_claims"], start=1):
    FORBIDDEN_VISIBLE_TEXT[f"canonical forbidden claim {index:02d}"] = re.compile(
        re.escape(phrase), re.IGNORECASE
    )

REQUIRED_SCOPE_MARKERS = {
    "verification snapshot status": re.compile(re.escape(VERIFICATION_STATUS)),
    "verification snapshot date": re.compile(re.escape(VERIFICATION_EXECUTED_AT)),
    "verification snapshot time boundary": re.compile(
        re.escape(VERIFICATION_SCOPE_BOUNDARY),
        re.IGNORECASE,
    ),
    "Web/PWA": re.compile(r"Web\s*/\s*PWA", re.IGNORECASE),
    "Web/PWA primary": re.compile(
        r"(?:Web\s*/\s*PWA[^.\n]{0,80}(?:주\s*(?:앱|플랫폼|사용자\s*앱)|primary)|"
        r"(?:주\s*(?:앱|플랫폼|사용자\s*앱)|primary)[^.\n]{0,80}Web\s*/\s*PWA)",
        re.IGNORECASE,
    ),
    "Android research support": re.compile(
        r"(?:Android[^.\n]{0,80}연구\s*보조|연구\s*보조[^.\n]{0,80}Android)",
        re.IGNORECASE,
    ),
    "CSV manual submission": re.compile(
        r"(?:agency|기관용|CSV)[^.\n]{0,120}수동\s*(?:제출|신고)|"
        r"수동\s*(?:제출|신고)[^.\n]{0,120}(?:agency|기관|CSV)",
        re.IGNORECASE,
    ),
    "no agency API": re.compile(
        r"기관\s*API[^.\n]{0,60}(?:없|부재|범위\s*밖|미지원)|"
        r"(?:없|부재|범위\s*밖|미지원)[^.\n]{0,60}기관\s*API",
        re.IGNORECASE,
    ),
    "TMAP only provider": re.compile(r"(?<![A-Z])TMAP(?![A-Z])", re.IGNORECASE),
    "Web non-metric advisory tier": re.compile(r"CAMERA_NON_METRIC_ADVISORY"),
    "Android unsupported non-metric tier": re.compile(r"CAMERA_IMU_NON_METRIC"),
    "Web optional IMU": re.compile(
        r"Web[^.\n]{0,120}CAMERA_NON_METRIC_ADVISORY[^.\n]{0,80}IMU\s*선택",
        re.IGNORECASE,
    ),
    "Android fresh IMU required": re.compile(
        r"Android[^.\n]{0,120}CAMERA_IMU_NON_METRIC[^.\n]{0,100}fresh\s+IMU\s+필수",
        re.IGNORECASE,
    ),
    "non-metric low directions and stability": re.compile(
        r"low\s+좌/중앙/우[^.\n]{0,80}3\s+distinct\s+frame\+700ms",
        re.IGNORECASE,
    ),
    "non-metric forbidden outputs": re.compile(
        r"거리·N보·STOP/high·local\s+steering·안전\s+경로·경로\s+변경·자동\s+신고[^.\n]{0,30}금지",
        re.IGNORECASE,
    ),
    "non-metric report isolation": re.compile(r"reports\s*=\s*false", re.IGNORECASE),
    "existing Web risk and consent report preserved": re.compile(
        r"기존\s+Web\s+alertable\s+risk[^.\n]{0,80}명시\s+동의\s+damaged-report[^.\n]{0,30}유지",
        re.IGNORECASE,
    ),
    "equal severity 3A delivery recheck": re.compile(
        r"3A\s+bounded\s+sequential\s+handoff[^.\n]{0,80}전달\s+직전\s+recheck",
        re.IGNORECASE,
    ),
    "non-metric Field open": re.compile(
        r"Web\s+실폰[^.\n]{0,80}ARCore\s+미지원\s+Android\s+Field[^.\n]{0,40}UNVERIFIED_OPEN",
        re.IGNORECASE,
    ),
    "13 intent schema": re.compile(r"(?<!\d)13\s*(?:개\s*)?(?:rule\s+)?intents?", re.IGNORECASE),
    "speech priority": re.compile(r"위험\s*>\s*상호작용\s*>\s*길안내"),
    "server v2 or sampled boundary": re.compile(r"server[-\s]?(?:v2|sampled)", re.IGNORECASE),
    "browser STT and TTS boundary": re.compile(
        r"MediaRecorder[^.\n]{0,80}faster-whisper[^.\n]{0,80}(?:speechSynthesis|TTS)",
        re.IGNORECASE,
    ),
    "WebXR display only": re.compile(
        r"WebXR[^.\n]{0,100}(?:중앙\s*거리\s*표시|표시\s*전용)", re.IGNORECASE
    ),
    "WebXR risk disconnected": re.compile(
        r"WebXR[^.\n]{0,180}(?:위험\s*(?:판단|평가)[^.\n]{0,40}(?:미연결|사용하지\s*않|연결되지\s*않)|위험판정\s*미연결)",
        re.IGNORECASE,
    ),
    "primary report id contract": re.compile(
        r"primary[^.\n]{0,100}(?:신고\s*)?ID[^.\n]{0,100}duplicate_count",
        re.IGNORECASE,
    ),
    "PWA network only boundary": re.compile(r"network-only", re.IGNORECASE),
    "reports table": re.compile(r"(?<![A-Za-z_])reports(?![A-Za-z_])"),
    "status audit table": re.compile(r"report_status_audits"),
    "export audit table": re.compile(r"report_export_audits"),
    "read audit table": re.compile(r"report_read_audits"),
    "non-audit rate event table": re.compile(r"actor_rate_limit_events"),
    "durable report journal": re.compile(r"durable\s+journal", re.IGNORECASE),
    "four separate privacy scopes": re.compile(
        r"Web\s+frame[^.\n]{0,50}Web\s+report[^.\n]{0,50}Android\s+report[^.\n]{0,50}telemetry",
        re.IGNORECASE,
    ),
    "consent withdrawal cancellation": re.compile(
        r"철회[^.\n]{0,80}in-flight[^.\n]{0,40}취소|in-flight[^.\n]{0,80}철회[^.\n]{0,40}취소",
        re.IGNORECASE,
    ),
    "remote database verify-full": re.compile(r"sslmode\s*=\s*verify-full", re.IGNORECASE),
    "shared limiter and nonblocking readiness": re.compile(
        r"(?:shared|공유)[^.\n]{0,50}limiter[^.\n]{0,80}nonblocking",
        re.IGNORECASE,
    ),
    "dedicated Voice token": re.compile(r"VOICE_SERVICE_TOKEN"),
    "Voice copy-phase limits": re.compile(
        r"(?:pre-copy|body\s+copy\s+전)[^.\n]{0,180}auth[^.\n]{0,80}Content-Length"
        r"[^.\n]{0,120}(?:global[^.\n]{0,40}actor[^.\n]{0,40}IP|global/actor/IP)"
        r"[^.\n]{0,180}(?:post-copy|copy\s+후)[^.\n]{0,80}duration"
        r"[^.\n]{0,60}inference\s+queue",
        re.IGNORECASE,
    ),
    "authenticated OpenAPI": re.compile(
        r"(?:auth(?:enticated)?\s+OpenAPI|OpenAPI[^.\n]{0,30}인증)", re.IGNORECASE
    ),
    "model deployment held": re.compile(r"deployment_eligible\s*=\s*false", re.IGNORECASE),
    "independent test missing": re.compile(r"독립\s*test\s*(?:없|부재)", re.IGNORECASE),
}

CANONICAL_VALUE_MARKERS = {
    "dataset total images": f"{FACTS['model_data']['dataset']['total_images']:,}",
    "dataset total boxes": f"{FACTS['model_data']['dataset']['total_boxes']:,}",
    "model mAP50": f"{FACTS['model_data']['internal_validation']['map50']:.5f}",
    "model mAP50-95": f"{FACTS['model_data']['internal_validation']['map50_95']:.5f}",
    "excluded validation images": f"{FACTS['model_data']['internal_validation']['excluded_truncated_images']:,}",
    "overlapping capture sequences": str(FACTS["model_data"]["sequence_leakage"]["overlapping_sequences"]),
}

EMAIL_RE = re.compile(
    r"(?<![\w.+-])[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}(?![\w.-])",
    re.IGNORECASE,
)
HOME_PATH_RE = re.compile(
    r"(?:file:/+)?/home/[A-Za-z0-9._-]+|[A-Z]:[\\/]Users[\\/][^<>\"\s]+",
    re.IGNORECASE,
)
ONMICROSOFT_RE = re.compile(r"onmicrosoft\.com", re.IGNORECASE)
KNOWN_TEMPLATE_IDENTITIES_RE = re.compile(
    r"이\s*낙선|이\s*상민|lsm@|\bSAMSUNG\b", re.IGNORECASE
)
NEUTRAL_AUTHORS = {
    "",
    "walksafe assist",
    "walksafe assist team",
    "walksafe assist 프로젝트",
    "walksafe",
    "hanium dreamup",
    "한이음 드림업",
    "submission builder",
    "python-docx",
    "python-pptx",
    "libreoffice",
}
CORE_NS = {
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    "dc": "http://purl.org/dc/elements/1.1/",
}


@dataclass(frozen=True)
class PackageReport:
    path: Path
    sha256: str
    member_count: int
    xml_count: int
    media_count: int
    media_bytes: int


def rel(path: Path) -> str:
    try:
        return path.relative_to(REPO_ROOT).as_posix()
    except ValueError:
        return str(path)


def configure_final_output_dir(path: Path) -> None:
    global FINAL_OUTPUT_DIR, FINAL_MANIFEST_PATH, FINAL_README_PATH, OUTPUT_PAIRS

    if path.is_symlink() or not path.is_dir():
        raise ValueError(f"final validation directory is missing or unsafe: {path}")
    resolved = path.resolve()
    try:
        resolved.relative_to(REPO_ROOT.resolve())
    except ValueError as exc:
        raise ValueError("final validation directory must stay inside the repository") from exc
    FINAL_OUTPUT_DIR = resolved
    FINAL_MANIFEST_PATH = resolved / "BUILD_MANIFEST.json"
    FINAL_README_PATH = resolved / "README.md"
    OUTPUT_PAIRS = {
        "docx": (REPO_ROOT / "templates/개발보고서 양식.docx", resolved / "개발보고서 양식.docx"),
        "pptx": (REPO_ROOT / "templates/제작설계서_일반.pptx", resolved / "제작설계서_일반.pptx"),
    }


def final_manifest_target(path_value: str) -> Path:
    prefix = "docs/submission/final/"
    if path_value.startswith(prefix):
        name = path_value.removeprefix(prefix)
        if not name or "/" in name:
            raise ValueError(f"invalid logical final path: {path_value}")
        return FINAL_OUTPUT_DIR / name
    return REPO_ROOT / path_value


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_json_no_duplicates(path: Path) -> dict[str, object]:
    def reject_duplicates(pairs: list[tuple[str, object]]) -> dict[str, object]:
        result: dict[str, object] = {}
        for key, value in pairs:
            if key in result:
                raise ValueError(f"duplicate JSON key: {key}")
            result[key] = value
        return result

    payload = json.loads(
        path.read_text(encoding="utf-8"),
        object_pairs_hook=reject_duplicates,
    )
    if not isinstance(payload, dict):
        raise ValueError("expected JSON object")
    return payload


def normalize_visible_text(text: str) -> str:
    return "\n".join(
        re.sub(r"[\t \u00a0]+", " ", line).strip()
        for line in text.replace("\r", "\n").splitlines()
        if line.strip()
    )


def add_check(
    checks: list[tuple[str, bool, str]], name: str, ok: bool, detail: str = ""
) -> None:
    checks.append((name, ok, detail))


def inspect_package(
    path: Path, checks: list[tuple[str, bool, str]]
) -> PackageReport | None:
    label = rel(path)
    if not path.is_file():
        add_check(checks, f"package exists: {label}", False, "file is missing")
        return None

    add_check(checks, f"package exists: {label}", path.stat().st_size > 0, f"{path.stat().st_size} bytes")
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            duplicates = sorted({name for name in names if names.count(name) > 1})
            add_check(
                checks,
                f"ZIP member names are unique: {label}",
                not duplicates,
                ", ".join(duplicates[:10]),
            )

            bad_member = archive.testzip()
            add_check(
                checks,
                f"ZIP CRC integrity: {label}",
                bad_member is None,
                bad_member or "all members readable",
            )

            xml_names = [
                name
                for name in names
                if name.lower().endswith((".xml", ".rels", ".vml"))
            ]
            malformed: list[str] = []
            for name in xml_names:
                try:
                    ET.fromstring(archive.read(name))
                except (ET.ParseError, UnicodeError, KeyError) as exc:
                    malformed.append(f"{name}: {exc}")
            add_check(
                checks,
                f"OOXML parts parse: {label}",
                not malformed,
                "; ".join(malformed[:5]) if malformed else f"{len(xml_names)} XML parts",
            )

            media_infos = [
                info
                for info in archive.infolist()
                if "/media/" in info.filename.lower() and not info.is_dir()
            ]
            return PackageReport(
                path=path,
                sha256=sha256(path),
                member_count=len(names),
                xml_count=len(xml_names),
                media_count=len(media_infos),
                media_bytes=sum(info.file_size for info in media_infos),
            )
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        add_check(checks, f"ZIP open: {label}", False, str(exc))
        return None


def extract_docx_text(document: Document) -> str:
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    for section in document.sections:
        chunks.extend(paragraph.text for paragraph in section.header.paragraphs)
        chunks.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return normalize_visible_text("\n".join(chunks))


def extract_shape_text(shape: object) -> list[str]:
    chunks: list[str] = []
    if getattr(shape, "has_text_frame", False):
        chunks.append(getattr(shape, "text", ""))
    if getattr(shape, "has_table", False):
        table = getattr(shape, "table")
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
        for child in getattr(shape, "shapes", []):
            chunks.extend(extract_shape_text(child))
    return chunks


def extract_pptx_text(presentation: Presentation) -> str:
    chunks: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            chunks.extend(extract_shape_text(shape))
    return normalize_visible_text("\n".join(chunks))


def check_docx(
    path: Path, checks: list[tuple[str, bool, str]], visible_texts: dict[Path, str]
) -> None:
    label = rel(path)
    if not path.is_file():
        return
    try:
        document = Document(path)
    except Exception as exc:  # python-docx can raise several package exceptions
        add_check(checks, f"python-docx reopen: {label}", False, str(exc))
        return

    add_check(checks, f"python-docx reopen: {label}", True)
    add_check(
        checks,
        f"DOCX table count: {label}",
        len(document.tables) == 30,
        f"found {len(document.tables)}, expected 30",
    )
    text = extract_docx_text(document)
    visible_texts[path] = text
    for heading in OFFICIAL_REPORT_HEADINGS:
        add_check(
            checks,
            f"DOCX heading '{heading}': {label}",
            heading in text,
            "required official heading",
        )
    for label_text in OFFICIAL_EVALUATION_BARS:
        add_check(
            checks,
            f"DOCX evaluation bar '{label_text}': {label}",
            label_text in text,
            "required official evaluation label",
        )
    add_check(
        checks,
        f"DOCX H/W table retained: {label}",
        "별도 제작 H/W" in text and "해당 없음" in text,
        "the official H/W table must remain even when no H/W is built",
    )


def check_pptx(
    path: Path, checks: list[tuple[str, bool, str]], visible_texts: dict[Path, str]
) -> None:
    label = rel(path)
    if not path.is_file():
        return
    try:
        presentation = Presentation(path)
    except Exception as exc:  # python-pptx can raise several package exceptions
        add_check(checks, f"python-pptx reopen: {label}", False, str(exc))
        return

    add_check(checks, f"python-pptx reopen: {label}", True)
    add_check(
        checks,
        f"PPTX slide count: {label}",
        len(presentation.slides) == 34,
        f"found {len(presentation.slides)}, expected 34",
    )
    is_four_by_three = presentation.slide_width * 3 == presentation.slide_height * 4
    add_check(
        checks,
        f"PPTX aspect ratio: {label}",
        is_four_by_three,
        f"{presentation.slide_width} x {presentation.slide_height} EMU; expected 4:3",
    )
    slide_two_tables = [shape.table for shape in presentation.slides[1].shapes if shape.has_table]
    official_matrix_ok = False
    matrix_detail = f"found {len(slide_two_tables)} table(s)"
    if len(slide_two_tables) == 1:
        matrix = slide_two_tables[0]
        expected_cells = {
            (0, 0): "단계",
            (0, 1): "산출물",
            (5, 1): "요구사항 정의서",
            (7, 1): "서비스 구성도(시스템 구성도)",
            (16, 1): "데이터 수집처리 정의서",
            (20, 1): "핵심 소스코드",
        }
        official_matrix_ok = len(matrix.rows) == 21 and len(matrix.columns) == 5 and all(
            matrix.cell(row, column).text.strip() == expected
            for (row, column), expected in expected_cells.items()
        )
        matrix_detail = f"{len(matrix.rows)}x{len(matrix.columns)}"
    add_check(
        checks,
        f"PPTX official slide-2 deliverables matrix preserved: {label}",
        official_matrix_ok,
        matrix_detail,
    )
    text = extract_pptx_text(presentation)
    visible_texts[path] = text
    for item in FACTS["requirements"]:
        add_check(
            checks,
            f"PPTX canonical requirement {item['id']}: {label}",
            item["id"] in text and item["title"] in text,
            item["title"],
        )
    for item in FACTS["use_cases"]:
        add_check(
            checks,
            f"PPTX canonical use case {item['id']}: {label}",
            item["id"] in text and item["title"] in text,
            item["title"],
        )
    add_check(
        checks,
        f"PPTX canonical PARTIAL status: {label}",
        "PARTIAL" in text,
        "REQ/UC overall status must remain visible",
    )


def check_visible_scope(
    path: Path, text: str, checks: list[tuple[str, bool, str]]
) -> None:
    label = rel(path)
    for name, pattern in FORBIDDEN_VISIBLE_TEXT.items():
        match = pattern.search(text)
        excerpt = ""
        if match:
            excerpt = re.sub(r"\s+", " ", text[max(0, match.start() - 40) : match.end() + 60])
        add_check(
            checks,
            f"forbidden visible text ({name}): {label}",
            match is None,
            excerpt,
        )

    for name, pattern in REQUIRED_SCOPE_MARKERS.items():
        add_check(
            checks,
            f"required scope marker ({name}): {label}",
            pattern.search(text) is not None,
            "marker not found" if pattern.search(text) is None else "",
        )

    for name, marker in CANONICAL_VALUE_MARKERS.items():
        add_check(
            checks,
            f"canonical value marker ({name}): {label}",
            marker in text,
            f"expected {marker}",
        )

    for name, pattern in CANONICAL_VERIFICATION_PATTERNS.items():
        found = pattern.search(text) is not None
        add_check(
            checks,
            f"canonical verification marker ({name}): {label}",
            found,
            "" if found else "current verification count not found",
        )

def inspect_xml_privacy(
    path: Path, checks: list[tuple[str, bool, str]], is_pptx: bool
) -> None:
    label = rel(path)
    if not path.is_file():
        return
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            xml_names = [
                name
                for name in names
                if name.lower().endswith((".xml", ".rels", ".vml"))
            ]
            findings: list[str] = []
            for name in xml_names:
                raw = archive.read(name)
                decoded = raw.decode("utf-8", errors="replace")
                for finding_name, pattern in (
                    ("email", EMAIL_RE),
                    ("onmicrosoft domain", ONMICROSOFT_RE),
                    ("local home path", HOME_PATH_RE),
                    ("template author identity", KNOWN_TEMPLATE_IDENTITIES_RE),
                ):
                    match = pattern.search(decoded)
                    if match:
                        findings.append(f"{name}: {finding_name}={match.group(0)!r}")
            add_check(
                checks,
                f"OOXML privacy scan: {label}",
                not findings,
                "; ".join(findings[:10]) if findings else f"scanned {len(xml_names)} XML parts",
            )

            core_authors: list[str] = []
            if "docProps/core.xml" in names:
                core = ET.fromstring(archive.read("docProps/core.xml"))
                for xpath in ("dc:creator", "cp:lastModifiedBy"):
                    element = core.find(xpath, CORE_NS)
                    if element is not None and element.text:
                        core_authors.append(element.text.strip())
            non_neutral = [
                author
                for author in core_authors
                if author.casefold() not in NEUTRAL_AUTHORS
            ]
            add_check(
                checks,
                f"neutral core author metadata: {label}",
                not non_neutral,
                f"authors={core_authors!r}",
            )

            if is_pptx:
                review_parts = sorted(
                    name
                    for name in names
                    if name.lower() == "ppt/authors.xml"
                    or "comment" in name.lower()
                    or name.lower() == "ppt/people.xml"
                )
                add_check(
                    checks,
                    f"PPTX authors/comments removed: {label}",
                    not review_parts,
                    ", ".join(review_parts[:20]),
                )
    except (OSError, zipfile.BadZipFile, ET.ParseError, KeyError) as exc:
        add_check(checks, f"OOXML privacy scan: {label}", False, str(exc))


def check_source_parity(checks: list[tuple[str, bool, str]]) -> None:
    semantic_errors: list[str] = []
    validate_core_semantic_facts(FACTS, semantic_errors)
    add_check(
        checks,
        "canonical core semantic contract",
        not semantic_errors,
        "; ".join(semantic_errors),
    )

    source_texts: dict[str, str] = {}
    for name, path in SOURCE_PARITY_PATHS.items():
        exists = path.is_file()
        add_check(checks, f"canonical source exists: {name}", exists, rel(path))
        if exists:
            source_texts[name] = path.read_text(encoding="utf-8")
    if len(source_texts) != len(SOURCE_PARITY_PATHS):
        return

    capability_source_markers = (
        "CAMERA_NON_METRIC_ADVISORY",
        "CAMERA_IMU_NON_METRIC",
        "TMAP_ONLY",
        "low 좌/중앙/우",
        "3 distinct frame+700ms",
        "fresh IMU",
        "reports=false",
        "기존 Web alertable risk",
        "명시 동의 damaged-report",
        "3A bounded sequential handoff",
        "UNVERIFIED_OPEN",
    )
    semantic_source_markers = {
        "requirements": (
            "durable journal",
            "startup reconciliation",
            "actor·resource·time·purpose",
            "actor_rate_limit_events",
            "Web frame",
            "Android 신고",
            "sslmode=verify-full",
            "gssencmode=disable",
            "PostgreSQL shared",
            "nonblocking",
            "전용 token",
            "OpenAPI",
            "TMAP-only",
            "선택적 IMU",
        ) + capability_source_markers,
        "use_cases": (
            "exact report list/detail/image",
            "read audit",
            "Web frame",
            "Android 신고",
            "telemetry",
            "철회",
            "in-flight",
            "TMAP-only",
            "IMU는 제공될 때",
        ) + capability_source_markers,
        "trace": (
            "durable journal",
            "startup reconciliation",
            "5개 application table",
            "Web frame",
            "Web report",
            "Android report",
            "telemetry",
            "PostgreSQL shared",
            "nonblocking",
            "전용 token",
            "OpenAPI",
            "TMAP-only",
            "IMU 선택",
        ) + capability_source_markers,
        "report_manuscript": (
            "report_read_audits",
            "actor_rate_limit_events",
            "durable journal",
            "startup reconciliation",
            "actor·resource·time·purpose",
            "Web frame",
            "Android 신고",
            "sslmode=verify-full",
            "gssencmode=disable",
            "PostgreSQL shared",
            "VOICE_SERVICE_TOKEN",
            "body copy 전 auth·Content-Length·global/actor/IP rate",
            "copy 중 byte 상한",
            "copy 후 duration 검증·inference queue",
            "OpenAPI",
            "TMAP-only",
            "선택적 IMU",
        ) + capability_source_markers,
        "slide_manuscript": (
            "report_read_audits",
            "actor_rate_limit_events",
            "durable journal",
            "startup reconciliation",
            "actor·resource·time·purpose",
            "telemetry",
            "철회",
            "in-flight",
            "sslmode=verify-full",
            "gssencmode=disable",
            "PostgreSQL shared",
            "전용 token",
            "body copy 전 auth·Content-Length·global/actor/IP rate",
            "copy 중 byte 상한",
            "copy 후 duration·inference queue",
            "OpenAPI",
            "TMAP-only",
            "선택적 IMU",
        ) + capability_source_markers,
    }
    for source_name, markers in semantic_source_markers.items():
        missing = [marker for marker in markers if marker not in source_texts[source_name]]
        add_check(
            checks,
            f"source semantic markers: {source_name}",
            not missing,
            f"missing={missing!r}",
        )

    requirement_targets = (
        source_texts["requirements"],
        source_texts["trace"],
        source_texts["slide_manuscript"],
    )
    for item in FACTS["requirements"]:
        expected = f"{item['id']}"
        ok = all(expected in text and item["title"] in text for text in requirement_targets)
        add_check(
            checks,
            f"source parity requirement {item['id']}",
            ok,
            item["title"],
        )

    use_case_targets = (source_texts["use_cases"], source_texts["slide_manuscript"])
    for item in FACTS["use_cases"]:
        ok = all(item["id"] in text and item["title"] in text for text in use_case_targets)
        add_check(
            checks,
            f"source parity use case {item['id']}",
            ok,
            item["title"],
        )

    all_sources = "\n".join(source_texts.values())
    for phrase in FACTS["forbidden_claims"]:
        add_check(
            checks,
            f"canonical forbidden claim absent from source: {phrase}",
            phrase.casefold() not in all_sources.casefold(),
            "exact unsafe claim found",
        )

    expected_tables = [table["name"] for table in FACTS["database"]["tables"]]
    database_targets = (
        source_texts["report_manuscript"],
        source_texts["slide_manuscript"],
    )
    add_check(
        checks,
        "source parity canonical database tables",
        all(all(table in text for table in expected_tables) for text in database_targets),
        ", ".join(expected_tables),
    )
    add_check(
        checks,
        "source parity WebXR risk boundary",
        "estimateDepthForDetection" in all_sources
        and "위험" in all_sources
        and ("미연결" in all_sources or "사용하지 않는다" in all_sources),
        "WebXR must remain display-only and disconnected from object risk",
    )
    add_check(
        checks,
        "source parity primary report response contract",
        "primary" in all_sources and "duplicate_count" in all_sources,
        "primary ID and duplicate_count must be retained",
    )


def check_form_build_manifest(checks: list[tuple[str, bool, str]]) -> dict[str, object] | None:
    """Verify the builder receipt against current inputs and working artifacts."""
    try:
        require_exact_file_set(REPO_ROOT / "templates", FORM_TEMPLATE_FILES)
    except (FileNotFoundError, ValueError) as exc:
        add_check(checks, "working template file set", False, str(exc))
    else:
        add_check(checks, "working template file set", True, f"{len(FORM_TEMPLATE_FILES)} files")

    if not FORM_BUILD_MANIFEST_PATH.is_file():
        add_check(checks, "form build manifest exists", False, rel(FORM_BUILD_MANIFEST_PATH))
        return None
    try:
        manifest = json.loads(FORM_BUILD_MANIFEST_PATH.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        add_check(checks, "form build manifest parses", False, str(exc))
        return None
    add_check(checks, "form build manifest parses", True)
    add_check(
        checks,
        "form build manifest schema",
        manifest.get("schema_version") == "walksafe.submission-forms-build.v1",
        str(manifest.get("schema_version")),
    )
    add_check(
        checks,
        "form build allowed template files",
        manifest.get("allowed_template_files") == sorted(FORM_TEMPLATE_FILES),
        repr(manifest.get("allowed_template_files")),
    )
    revision = source_revision(REPO_ROOT, ALL_SUBMISSION_GENERATED_PATHS)
    add_check(checks, "current submission source is clean", revision.get("source_dirty") is False)
    add_check(checks, "form build manifest source is clean", manifest.get("source_dirty") is False)
    add_check(
        checks,
        "form build source revision",
        all(manifest.get(key) == value for key, value in revision.items()),
        f"commit={manifest.get('source_commit')}; dirty={manifest.get('source_dirty')}",
    )
    expected_tools = build_tool_provenance(("Pillow", "python-docx", "python-pptx"))
    add_check(
        checks,
        "form build tool provenance",
        manifest.get("build_tools") == expected_tools,
        repr(manifest.get("build_tools")),
    )
    try:
        expected_toolchain = submission_toolchain_attestation(REPO_ROOT)
    except ValueError as exc:
        add_check(checks, "pinned submission toolchain", False, str(exc))
        expected_toolchain = None
    add_check(
        checks,
        "form build submission toolchain",
        expected_toolchain is not None and manifest.get("submission_toolchain") == expected_toolchain,
        repr(manifest.get("submission_toolchain")),
    )

    verified_inputs: list[dict[str, object]] = []
    for section, expected_paths in (
        ("inputs", FORM_BUILD_INPUT_PATHS),
        ("artifacts", FORM_ARTIFACT_PATHS),
    ):
        try:
            entries = require_exact_manifest_paths(section, manifest.get(section), expected_paths)
        except ValueError as exc:
            add_check(checks, f"form build {section} path set", False, str(exc))
            continue
        add_check(checks, f"form build {section} path set", True, f"{len(entries)} paths")
        for entry in entries:
            path_value = str(entry["path"])
            candidate = REPO_ROOT / path_value
            target = candidate.resolve()
            try:
                target.relative_to(REPO_ROOT.resolve())
            except ValueError:
                add_check(checks, f"form build path stays in repository: {path_value}", False)
                continue
            exists = target.is_file() and not candidate.is_symlink()
            actual_hash = sha256(target) if exists else "missing"
            actual_bytes = target.stat().st_size if exists else -1
            matches = actual_hash == entry.get("sha256") and actual_bytes == entry.get("bytes")
            add_check(
                checks,
                f"form build source binding: {path_value}",
                matches,
                f"sha256={actual_hash}; bytes={actual_bytes}",
            )
            if section == "inputs" and matches:
                verified_inputs.append(entry)

    actual_bundle = (
        record_bundle_sha256(verified_inputs)
        if len(verified_inputs) == len(FORM_BUILD_INPUT_PATHS)
        else "incomplete"
    )
    add_check(
        checks,
        "form build input bundle SHA-256",
        actual_bundle == manifest.get("input_bundle_sha256"),
        f"actual={actual_bundle}; expected={manifest.get('input_bundle_sha256')}",
    )
    return manifest


def render_final_office_pages(report: Path, slides_path: Path) -> tuple[int, int]:
    toolchain = submission_toolchain_attestation(REPO_ROOT)
    with tempfile.TemporaryDirectory(prefix="walksafe-final-render-") as temp:
        root = Path(temp)
        output_dir = root / "pdf"
        profile_dir = root / "lo-profile"
        output_dir.mkdir()
        profile_dir.mkdir()
        run_attested_submission_tool(
            toolchain,
            "libreoffice",
            [
                "--headless",
                f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(report),
                str(slides_path),
            ],
            check=True,
            capture_output=True,
            text=True,
        )

        def pages(source: Path) -> int:
            pdf = output_dir / f"{source.stem}.pdf"
            if pdf.is_symlink() or not pdf.is_file():
                raise ValueError(f"missing rendered PDF: {source.name}")
            completed = run_attested_submission_tool(
                toolchain,
                "pdfinfo",
                [str(pdf)],
                check=True,
                capture_output=True,
                text=True,
            )
            match = re.search(r"^Pages:\s*(\d+)\s*$", completed.stdout, re.MULTILINE)
            if match is None or int(match.group(1)) < 1:
                raise ValueError(f"invalid rendered PDF page count: {source.name}")
            return int(match.group(1))

        return pages(report), pages(slides_path)


def check_final_privacy_receipt(
    checks: list[tuple[str, bool, str]],
) -> tuple[int, int] | None:
    receipt_path = FINAL_OUTPUT_DIR / "ASSISTANT_VISUAL_PRIVACY_REVIEW.json"
    if receipt_path.is_symlink() or not receipt_path.is_file():
        add_check(checks, "final human privacy receipt exists safely", False, rel(receipt_path))
        return None
    try:
        receipt = load_json_no_duplicates(receipt_path)
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        add_check(checks, "final human privacy receipt parses uniquely", False, str(exc))
        return None
    add_check(checks, "final human privacy receipt parses uniquely", True)
    expected_receipt_keys = {
        "schema_version",
        "source_commit",
        "reviewed_at",
        "reviewer",
        "reviewer_kind",
        "automatic",
        "manual_assertions",
        "release_ready",
    }
    add_check(
        checks,
        "final human privacy receipt exact shape",
        set(receipt) == expected_receipt_keys,
        repr(sorted(receipt)),
    )
    try:
        require_review_identity(receipt.get("reviewer"), receipt.get("reviewed_at"))
    except ValueError as exc:
        identity_ok = False
        identity_detail = str(exc)
    else:
        identity_ok = True
        identity_detail = "human identity is an external trust boundary, not a cryptographic signature"
    add_check(
        checks,
        "final human privacy receipt identity boundary",
        receipt.get("schema_version") == "walksafe.submission-visual-privacy.v2"
        and receipt.get("source_commit")
        == source_revision(REPO_ROOT, ALL_SUBMISSION_GENERATED_PATHS).get("source_commit")
        and receipt.get("reviewer_kind") == "human"
        and identity_ok
        and receipt.get("release_ready") is True,
        identity_detail,
    )
    manual = receipt.get("manual_assertions")
    required_assertions = {
        "no_visible_faces",
        "no_visible_license_plates",
        "no_personal_addresses_or_accounts",
        "office_pdf_render_reviewed",
    }
    add_check(
        checks,
        "final human privacy assertions exact and true",
        isinstance(manual, dict)
        and set(manual) == required_assertions
        and all(manual.get(key) is True for key in required_assertions),
    )
    automatic = receipt.get("automatic")
    if not isinstance(automatic, dict):
        add_check(checks, "final privacy automatic evidence", False, "missing")
        return None
    expected_automatic_keys = {
        "assets",
        "metadata_violations",
        "face_candidates",
        "automatic_face_detection_limit",
        "documents",
        "final_documents",
        "office_render",
    }
    add_check(
        checks,
        "final privacy automatic evidence exact shape",
        set(automatic) == expected_automatic_keys
        and isinstance(automatic.get("face_candidates"), list)
        and automatic.get("automatic_face_detection_limit")
        == "Haar 후보 검사이며 얼굴·차량번호·OCR 부재를 단독 보장하지 않음",
        repr(sorted(automatic)),
    )
    add_check(
        checks,
        "final privacy metadata violations empty",
        automatic.get("metadata_violations") == [],
    )

    def exact_records(
        records: object,
        paths: dict[str, Path],
        label: str,
        expected_keys: set[str],
    ) -> bool:
        if not isinstance(records, list) or any(not isinstance(item, dict) for item in records):
            add_check(checks, f"final privacy {label} records", False, "invalid list")
            return False
        if any(set(item) != expected_keys for item in records):
            add_check(checks, f"final privacy {label} record shape", False, repr(records))
            return False
        names = [str(item.get("name")) for item in records]
        safe = all(not path.is_symlink() and path.is_file() for path in paths.values())
        recorded = {str(item.get("name")): str(item.get("sha256")) for item in records}
        actual = {name: sha256(path) for name, path in paths.items()} if safe else {}
        ok = len(names) == len(set(names)) and set(names) == set(paths) and safe and recorded == actual
        add_check(checks, f"final privacy exact {label} hashes", ok)
        return ok

    asset_paths = {
        name: REPO_ROOT / "docs/submission/form_materials/assets" / name
        for name in ASSET_DIRECTORY_FILES
        if name.endswith(".png")
    }
    assets_ok = exact_records(
        automatic.get("assets"),
        asset_paths,
        "18 asset",
        {"name", "sha256", "width", "height"},
    )
    if assets_ok:
        metadata_findings: list[str] = []
        dimensions = {
            str(item["name"]): (item["width"], item["height"])
            for item in automatic["assets"]
        }
        for name, path in asset_paths.items():
            try:
                with Image.open(path) as image:
                    image.verify()
                with Image.open(path) as image:
                    if dimensions.get(name) != image.size:
                        metadata_findings.append(f"{name}:dimensions")
                    if image.getexif().get(34853):
                        metadata_findings.append(f"{name}:gps")
                    if any(
                        key.casefold() in {"author", "comment", "description", "location"}
                        for key in image.info
                    ):
                        metadata_findings.append(f"{name}:metadata")
            except OSError:
                metadata_findings.append(f"{name}:decode")
        add_check(checks, "final privacy current asset metadata", not metadata_findings, repr(metadata_findings))
    exact_records(
        automatic.get("documents"),
        {
            name: REPO_ROOT / "docs/submission/design_documents" / name
            for name in DESIGN_DOCUMENT_NAMES
        },
        "8 design document",
        {"name", "sha256"},
    )
    report = FINAL_OUTPUT_DIR / "개발보고서 양식.docx"
    slides_path = FINAL_OUTPUT_DIR / "제작설계서_일반.pptx"
    office_ok = exact_records(
        automatic.get("final_documents"),
        {report.name: report, slides_path.name: slides_path},
        "2 review Office",
        {"name", "sha256"},
    )
    office_render = automatic.get("office_render")
    if not office_ok or not isinstance(office_render, dict):
        add_check(checks, "final privacy Office render evidence", False, "missing")
        return None
    report_render = office_render.get("report")
    slides_render = office_render.get("slides")
    if not isinstance(report_render, dict) or not isinstance(slides_render, dict):
        add_check(checks, "final privacy Office render evidence", False, "invalid")
        return None
    try:
        presentation = Presentation(slides_path)
        actual_slides = len(presentation.slides)
        actual_ratio = presentation.slide_width * 3 == presentation.slide_height * 4
        actual_report_pages, actual_slide_pages = render_final_office_pages(report, slides_path)
    except Exception as exc:
        add_check(checks, "final Office independent render", False, str(exc))
        return None
    render_ok = (
        set(report_render) == {"name", "sha256", "pdf_pages"}
        and report_render.get("name") == report.name
        and report_render.get("sha256") == sha256(report)
        and report_render.get("pdf_pages") == actual_report_pages
        and set(slides_render) == {"name", "sha256", "pdf_pages", "slides", "aspect_ratio"}
        and slides_render.get("name") == slides_path.name
        and slides_render.get("sha256") == sha256(slides_path)
        and slides_render.get("slides") == 34 == actual_slides
        and slides_render.get("pdf_pages") == 34 == actual_slide_pages
        and slides_render.get("aspect_ratio") == "4:3"
        and actual_ratio
    )
    add_check(checks, "final Office independent 34-slide/4:3/page evidence", render_ok)
    return (actual_report_pages, actual_slides) if render_ok else None


def check_final_manifest(
    checks: list[tuple[str, bool, str]],
    form_build_manifest: dict[str, object] | None,
) -> None:
    """Bind promoted Office files and their declared sources to actual bytes."""
    for directory, expected_names, label in (
        (FINAL_OUTPUT_DIR, FINAL_OUTPUT_NAMES, "final output file set"),
        (REPO_ROOT / "docs/submission/design_documents", DESIGN_OUTPUT_NAMES, "design output file set"),
    ):
        try:
            require_exact_file_set(directory, expected_names)
        except (FileNotFoundError, ValueError) as exc:
            add_check(checks, label, False, str(exc))
        else:
            add_check(checks, label, True, f"{len(expected_names)} files")
    if not FINAL_MANIFEST_PATH.is_file():
        add_check(checks, "final build manifest exists", False, rel(FINAL_MANIFEST_PATH))
        return
    office_evidence = check_final_privacy_receipt(checks)
    try:
        manifest = load_json_no_duplicates(FINAL_MANIFEST_PATH)
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        add_check(checks, "final build manifest parses", False, str(exc))
        return

    add_check(checks, "final build manifest parses", True)
    add_check(
        checks,
        "final build manifest schema",
        manifest.get("schema_version") == "walksafe.submission-final.v1",
        str(manifest.get("schema_version")),
    )
    expected_top_level = {
        "schema_version",
        "as_of_date",
        "product_status",
        "candidate_status",
        "product_scope",
        "repository",
        "canonical_facts",
        "official_templates",
        "critical_sources",
        "build_sources",
        "design_pack",
        "form_build",
        "artifacts",
        "supporting_files",
        "validation",
        "known_release_blockers",
    }
    add_check(
        checks,
        "final manifest exact top-level shape",
        set(manifest) == expected_top_level,
        repr(sorted(manifest)),
    )
    expected_product_scope = {
        "primary_app": FACTS["project"]["primary_app"],
        "android_role": FACTS["project"]["android_role"],
        "institution_submission": "named admin review, agency file preparation, manual external-channel submission",
    }
    add_check(
        checks,
        "final manifest canonical date/status/scope/candidate",
        manifest.get("as_of_date") == FACTS["as_of_date"]
        and manifest.get("product_status") == FACTS["project"]["overall_status"]
        and manifest.get("candidate_status") == "AUTOMATED_AND_HUMAN_VISUAL_QA_PASS"
        and manifest.get("product_scope") == expected_product_scope,
        repr(manifest.get("product_scope")),
    )
    add_check(
        checks,
        "final manifest canonical release blockers",
        manifest.get("known_release_blockers") == FACTS["release_gates"],
        repr(manifest.get("known_release_blockers")),
    )

    entries: list[tuple[str, dict[str, object]]] = []
    canonical = manifest.get("canonical_facts")
    if isinstance(canonical, dict):
        entries.append(("canonical facts", canonical))
        add_check(
            checks,
            "final manifest canonical facts exact shape",
            set(canonical) == {"path", "sha256", "bytes", "schema_version"},
            repr(canonical),
        )
        add_check(
            checks,
            "final manifest canonical facts schema",
            canonical.get("schema_version") == "walksafe.submission_facts.v2",
            str(canonical.get("schema_version")),
        )
        add_check(
            checks,
            "final manifest canonical facts path",
            canonical.get("path") == "docs/submission/form_materials/09_제출_사실_기준.json",
            str(canonical.get("path")),
        )
    else:
        add_check(checks, "final manifest canonical facts entry", False, "missing")

    for section in (
        "official_templates",
        "critical_sources",
        "build_sources",
        "artifacts",
        "supporting_files",
    ):
        section_entries = manifest.get(section)
        try:
            exact_entries = require_exact_manifest_paths(
                section,
                section_entries,
                FINAL_SECTION_PATHS[section],
            )
        except ValueError as exc:
            add_check(checks, f"final manifest {section} path set", False, str(exc))
            continue
        add_check(checks, f"final manifest {section} path set", True, f"{len(exact_entries)} paths")
        if section != "artifacts":
            add_check(
                checks,
                f"final manifest {section} exact record shape",
                all(set(entry) == {"path", "sha256", "bytes"} for entry in exact_entries),
                repr(exact_entries),
            )
        entries.extend(
            (f"{section}[{index}]", entry)
            for index, entry in enumerate(exact_entries, start=1)
        )

    artifacts = manifest.get("artifacts")
    artifact_by_path = {
        str(entry.get("path")): entry
        for entry in artifacts
        if isinstance(entry, dict)
    } if isinstance(artifacts, list) else {}
    report_entry = artifact_by_path.get("docs/submission/final/개발보고서 양식.docx")
    slides_entry = artifact_by_path.get("docs/submission/final/제작설계서_일반.pptx")
    artifact_semantics_ok = (
        office_evidence is not None
        and isinstance(report_entry, dict)
        and set(report_entry) == {"path", "sha256", "bytes", "pages"}
        and report_entry.get("pages") == office_evidence[0]
        and isinstance(slides_entry, dict)
        and set(slides_entry) == {"path", "sha256", "bytes", "slides", "aspect_ratio"}
        and slides_entry.get("slides") == office_evidence[1] == 34
        and slides_entry.get("aspect_ratio") == "4:3"
    )
    add_check(checks, "final manifest actual Office page/slide semantics", artifact_semantics_ok)
    if office_evidence is not None:
        counts = VERIFICATION_COUNTS
        report_pages, slides = office_evidence
        expected_validation = {
            "test_layers": (
                f"PASS: Unit Python {counts['unit_python']}; "
                f"Functional Python/PostGIS {counts['functional_python']}; "
                f"Integration Python {counts['integration_python']}"
            ),
            "backend_voice_android": (
                f"PASS: Backend {counts['backend_full']}; Voice {counts['voice']}; "
                f"Android JVM {counts['android_jvm']}/{counts['android_jvm']}"
            ),
            "android_device": (
                f"{ANDROID_DEVICE_CURRENT_LABEL}; {ANDROID_DEVICE_HISTORICAL_LABEL}"
            ),
            "web": (
                "PASS: test/typecheck/lint/build, PWA lifecycle, "
                f"NFT {counts['web_trace_count']} traces/{counts['web_trace_files']} unique files"
            ),
            "submission_materials": (
                f"PASS: canonical facts and generated design/forms, "
                f"{len([name for name in ASSET_DIRECTORY_FILES if name.endswith('.png')])} assets, "
                f"{slides} slides, REQ-001..{len(FACTS['requirements']):03d}, "
                f"UC-01..{len(FACTS['use_cases']):02d}"
            ),
            "office_render": (
                f"PASS_HUMAN: report {report_pages} pages, design {slides} slides; "
                "exact visual/privacy receipt bound"
            ),
        }
        add_check(
            checks,
            "final manifest canonical-derived validation evidence",
            manifest.get("validation") == expected_validation,
            repr(manifest.get("validation")),
        )

    design_pack = manifest.get("design_pack")
    if isinstance(design_pack, dict):
        add_check(
            checks,
            "design pack manifest exact shape and semantics",
            set(design_pack) == {"document_count", "manifest_path", "manifest_sha256"}
            and design_pack.get("document_count") == len(DESIGN_DOCUMENT_NAMES)
            and design_pack.get("manifest_path")
            == "docs/submission/design_documents/BUILD_MANIFEST.json",
            repr(design_pack),
        )
        entries.append(
            (
                "design pack manifest",
                {
                    "path": design_pack.get("manifest_path"),
                    "sha256": design_pack.get("manifest_sha256"),
                },
            )
        )
    else:
        add_check(checks, "final manifest design pack entry", False, "missing")

    form_build = manifest.get("form_build")
    if isinstance(form_build, dict):
        add_check(
            checks,
            "final manifest form build exact shape and semantics",
            set(form_build) == {"schema_version", "manifest_path", "manifest_sha256"}
            and form_build.get("manifest_path") == "templates/SUBMISSION_BUILD_MANIFEST.json"
            and form_build.get("schema_version") == "walksafe.submission-forms-build.v1",
            repr(form_build),
        )
        entries.append(
            (
                "form build manifest",
                {
                    "path": form_build.get("manifest_path"),
                    "sha256": form_build.get("manifest_sha256"),
                },
            )
        )
    else:
        add_check(checks, "final manifest form build entry", False, "missing")

    verified_paths: dict[str, Path] = {}
    for label, entry in entries:
        path_value = entry.get("path")
        expected_hash = entry.get("sha256")
        if not isinstance(path_value, str) or not isinstance(expected_hash, str):
            add_check(checks, f"final manifest entry fields: {label}", False, repr(entry))
            continue
        candidate = final_manifest_target(path_value)
        target = candidate.resolve()
        try:
            target.relative_to(REPO_ROOT.resolve())
        except ValueError:
            add_check(checks, f"final manifest path stays in repository: {label}", False, path_value)
            continue
        exists = not candidate.is_symlink() and target.is_file()
        actual_hash = sha256(target) if exists else "missing"
        add_check(
            checks,
            f"final manifest SHA-256: {label}",
            exists and actual_hash == expected_hash,
            f"{path_value} sha256={actual_hash}; expected={expected_hash}",
        )
        expected_bytes = entry.get("bytes")
        if isinstance(expected_bytes, int):
            actual_bytes = target.stat().st_size if exists else -1
            add_check(
                checks,
                f"final manifest byte size: {label}",
                actual_bytes == expected_bytes,
                f"actual={actual_bytes}; expected={expected_bytes}",
            )
        if exists:
            verified_paths[path_value] = target

    bundle_entries: list[dict[str, object]] = []
    if isinstance(canonical, dict):
        bundle_entries.append(canonical)
    for section in ("critical_sources", "build_sources"):
        section_entries = manifest.get(section)
        if isinstance(section_entries, list):
            bundle_entries.extend(entry for entry in section_entries if isinstance(entry, dict))
    bundle_lines: list[str] = []
    for entry in bundle_entries:
        path_value = entry.get("path")
        if isinstance(path_value, str) and path_value in verified_paths:
            bundle_lines.append(f"{sha256(verified_paths[path_value])}  {path_value}\n")
    expected_entry_count = (
        1
        + len(FINAL_SECTION_PATHS["critical_sources"])
        + len(FINAL_SECTION_PATHS["build_sources"])
    )
    repository = manifest.get("repository")
    expected_bundle_hash = repository.get("source_bundle_sha256") if isinstance(repository, dict) else None
    expected_bundle_algorithm = (
        "SHA-256 of sorted sha256sum lines for canonical facts, "
        f"{len(FINAL_SECTION_PATHS['critical_sources'])} critical sources, and "
        f"{len(FINAL_SECTION_PATHS['build_sources'])} build/validation sources"
    )
    add_check(
        checks,
        "final manifest repository provenance fields",
        isinstance(repository, dict)
        and set(repository)
        == {
            "commit",
            "dirty",
            "source_bundle_sha256",
            "source_bundle_algorithm",
            "source_bundle_entry_count",
        }
        and re.fullmatch(r"[0-9a-f]{40}", str(repository.get("commit"))) is not None
        and repository.get("dirty") is False
        and repository.get("source_bundle_algorithm") == expected_bundle_algorithm
        and repository.get("source_bundle_entry_count") == expected_entry_count,
        repr(repository),
    )
    actual_bundle_hash = hashlib.sha256("".join(sorted(bundle_lines)).encode()).hexdigest()
    add_check(
        checks,
        "final manifest source bundle SHA-256",
        len(bundle_lines) == expected_entry_count and actual_bundle_hash == expected_bundle_hash,
        f"entries={len(bundle_lines)}/{expected_entry_count}; actual={actual_bundle_hash}; expected={expected_bundle_hash}",
    )
    current_revision = source_revision(REPO_ROOT, ALL_SUBMISSION_GENERATED_PATHS)
    add_check(
        checks,
        "final manifest repository matches clean current/form source revision",
        isinstance(repository, dict)
        and isinstance(form_build_manifest, dict)
        and repository.get("commit") == current_revision.get("source_commit")
        and current_revision.get("source_dirty") is False
        and repository.get("commit") == form_build_manifest.get("source_commit")
        and form_build_manifest.get("source_dirty") is False
        and repository.get("dirty") is False,
        f"repository={repository!r}",
    )

    if isinstance(design_pack, dict):
        design_manifest_path = design_pack.get("manifest_path")
        design_target = verified_paths.get(design_manifest_path) if isinstance(design_manifest_path, str) else None
        try:
            design_manifest = load_json_no_duplicates(design_target) if design_target else {}
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError):
            design_manifest = {}
        documents = design_manifest.get("documents") if isinstance(design_manifest, dict) else None
        add_check(
            checks,
            "design manifest schema",
            isinstance(design_manifest, dict)
            and design_manifest.get("schema_version") == "walksafe.design-documents.v2",
            str(design_manifest.get("schema_version")) if isinstance(design_manifest, dict) else "unreadable",
        )
        design_revision = source_revision(REPO_ROOT, ALL_SUBMISSION_GENERATED_PATHS)
        add_check(
            checks,
            "design manifest clean source provenance",
            isinstance(design_manifest, dict)
            and design_manifest.get("source_dirty") is False
            and all(
                design_manifest.get(key) == value
                for key, value in design_revision.items()
            ),
            repr(
                {
                    key: design_manifest.get(key)
                    for key in design_revision
                }
            ) if isinstance(design_manifest, dict) else "unreadable",
        )
        add_check(
            checks,
            "final manifest design document count",
            isinstance(documents, list) and len(documents) == design_pack.get("document_count"),
            f"actual={len(documents) if isinstance(documents, list) else 'unreadable'}; expected={design_pack.get('document_count')}",
        )
        document_names = [
            document.get("file")
            for document in documents
            if isinstance(document, dict)
        ] if isinstance(documents, list) else []
        add_check(
            checks,
            "design manifest document path set",
            len(document_names) == len(set(document_names))
            and set(document_names) == DESIGN_DOCUMENT_NAMES,
            repr(document_names),
        )
        expected_design_tools = build_tool_provenance(("Pillow", "python-docx"))
        add_check(
            checks,
            "design manifest build tool provenance",
            isinstance(design_manifest, dict)
            and design_manifest.get("build_tools") == expected_design_tools,
            repr(design_manifest.get("build_tools")) if isinstance(design_manifest, dict) else "unreadable",
        )
        try:
            expected_submission_toolchain = submission_toolchain_attestation(REPO_ROOT)
        except ValueError as exc:
            add_check(checks, "pinned submission toolchain for design pack", False, str(exc))
            expected_submission_toolchain = None
        add_check(
            checks,
            "design manifest submission toolchain",
            isinstance(design_manifest, dict)
            and expected_submission_toolchain is not None
            and design_manifest.get("submission_toolchain") == expected_submission_toolchain,
            repr(design_manifest.get("submission_toolchain")) if isinstance(design_manifest, dict) else "unreadable",
        )

        def check_design_file(
            label: str,
            path_value: object,
            expected_hash: object,
            expected_bytes: object = None,
            *,
            base: Path = REPO_ROOT,
        ) -> None:
            if not isinstance(path_value, str) or not isinstance(expected_hash, str):
                add_check(checks, f"design manifest entry fields: {label}", False, repr(path_value))
                return
            candidate = base / path_value
            target = candidate.resolve()
            try:
                target.relative_to(REPO_ROOT.resolve())
            except ValueError:
                add_check(checks, f"design manifest path stays in repository: {label}", False, path_value)
                return
            exists = not candidate.is_symlink() and target.is_file()
            actual_hash = sha256(target) if exists else "missing"
            add_check(
                checks,
                f"design manifest SHA-256: {label}",
                exists and actual_hash == expected_hash,
                f"{rel(target)} sha256={actual_hash}; expected={expected_hash}",
            )
            if isinstance(expected_bytes, int):
                actual_bytes = target.stat().st_size if exists else -1
                add_check(
                    checks,
                    f"design manifest byte size: {label}",
                    actual_bytes == expected_bytes,
                    f"actual={actual_bytes}; expected={expected_bytes}",
                )

        if isinstance(design_manifest, dict):
            design_facts = design_manifest.get("facts")
            if isinstance(design_facts, dict):
                check_design_file(
                    "canonical facts",
                    design_facts.get("file"),
                    design_facts.get("sha256"),
                )
            else:
                add_check(checks, "design manifest canonical facts entry", False, "missing")

            for section in (
                "assets",
                "build_inputs",
                "implementation_bindings",
                "supporting_files",
            ):
                records = design_manifest.get(section)
                if not isinstance(records, list) or not records:
                    add_check(checks, f"design manifest {section} entries", False, "missing or empty")
                    continue
                for index, record in enumerate(records, start=1):
                    if not isinstance(record, dict):
                        add_check(checks, f"design manifest {section}[{index}] entry", False, "expected object")
                        continue
                    check_design_file(
                        f"{section}[{index}]",
                        record.get("file"),
                        record.get("sha256"),
                        record.get("bytes"),
                    )
                if section in {"assets", "build_inputs", "supporting_files"} and isinstance(records, list):
                    recorded_paths = [
                        record.get("file") for record in records if isinstance(record, dict)
                    ]
                    expected_paths = {
                        "assets": DESIGN_ASSET_PATHS,
                        "build_inputs": DESIGN_BUILD_INPUT_PATHS,
                        "supporting_files": {
                            "docs/submission/design_documents/README.md"
                        },
                    }[section]
                    add_check(
                        checks,
                        f"design manifest {section} path set",
                        len(recorded_paths) == len(set(recorded_paths))
                        and set(recorded_paths) == expected_paths,
                        repr(recorded_paths),
                    )

            if isinstance(documents, list) and design_target is not None:
                for index, document in enumerate(documents, start=1):
                    if not isinstance(document, dict):
                        add_check(checks, f"design manifest documents[{index}] entry", False, "expected object")
                        continue
                    check_design_file(
                        f"documents[{index}]",
                        document.get("file"),
                        document.get("sha256"),
                        document.get("bytes"),
                        base=design_target.parent,
                    )
                    check_design_file(
                        f"documents[{index}] source",
                        document.get("source"),
                        document.get("source_sha256"),
                    )

    readme_text = FINAL_README_PATH.read_text(encoding="utf-8") if FINAL_README_PATH.is_file() else ""
    artifacts = manifest.get("artifacts")
    artifact_hashes = [
        entry.get("sha256")
        for entry in artifacts
        if isinstance(entry, dict) and isinstance(entry.get("sha256"), str)
    ] if isinstance(artifacts, list) else []
    add_check(
        checks,
        "final README lists promoted artifact hashes",
        bool(readme_text) and bool(artifact_hashes) and all(value in readme_text for value in artifact_hashes),
        rel(FINAL_README_PATH),
    )
    navigation = FACTS["runtime"]["navigation"]
    counts = VERIFICATION_COUNTS
    semantic_markers = (
        str(FACTS["as_of_date"]),
        str(FACTS["project"]["overall_status"]),
        str(FACTS["project"]["primary_app"]),
        str(FACTS["project"]["android_role"]),
        str(navigation["web_supervised_tactile_local_steering"]["status"]),
        str(navigation["android_tactile_local_steering"]["status"]),
        "human identity는 암호학적 서명이 아니라 외부 신뢰 경계",
        (
            f"Unit Python {counts['unit_python']}, Functional Python/PostGIS "
            f"{counts['functional_python']}, Integration Python {counts['integration_python']}, "
            f"Backend full {counts['backend_full']}"
        ),
        (
            f"Voice {counts['voice']}, Android JVM {counts['android_jvm']}/"
            f"{counts['android_jvm']}"
        ),
        ANDROID_DEVICE_CURRENT_LABEL,
        ANDROID_DEVICE_HISTORICAL_LABEL,
        f"Web NFT {counts['web_trace_count']} traces/{counts['web_trace_files']} unique files",
        *(
            f"`{gate['id']}`: {gate['evidence_required']}"
            for gate in FACTS["release_gates"]
        ),
    )
    add_check(
        checks,
        "final README canonical semantics",
        bool(readme_text) and all(marker in readme_text for marker in semantic_markers),
        rel(FINAL_README_PATH),
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--working-only",
        action="store_true",
        help="validate generated files under templates/ before final promotion",
    )
    parser.add_argument(
        "--final-dir",
        type=Path,
        help="validate a staged final tree in place before atomic publication",
    )
    args = parser.parse_args()
    if args.working_only and args.final_dir is not None:
        parser.error("--working-only and --final-dir cannot be combined")
    if args.final_dir is not None:
        try:
            configure_final_output_dir(args.final_dir)
        except ValueError as exc:
            parser.error(str(exc))
    checks: list[tuple[str, bool, str]] = []
    package_reports: list[PackageReport] = []
    visible_texts: dict[Path, str] = {}

    revision = source_revision(REPO_ROOT, ALL_SUBMISSION_GENERATED_PATHS)
    add_check(checks, "current submission source is clean", revision.get("source_dirty") is False)

    add_check(
        checks,
        "facts schema version",
        FACTS.get("schema_version") == "walksafe.submission_facts.v2",
        str(FACTS.get("schema_version")),
    )
    add_check(
        checks,
        "facts canonical counts/status",
        len(FACTS["requirements"]) == 14
        and len(FACTS["use_cases"]) == 9
        and all(item["status"] == "PARTIAL" for item in [*FACTS["requirements"], *FACTS["use_cases"]]),
        "expected 14 REQ, 9 UC, all PARTIAL",
    )
    check_source_parity(checks)
    form_build_manifest = check_form_build_manifest(checks)
    if not args.working_only:
        check_final_manifest(checks, form_build_manifest)

    for kind, (path, expected_hash) in ORIGINALS.items():
        exists = path.is_file()
        actual_hash = sha256(path) if exists else "missing"
        add_check(
            checks,
            f"original {kind} hash unchanged",
            exists and actual_hash == expected_hash,
            f"{rel(path)} sha256={actual_hash}",
        )

    active_output_pairs = {
        kind: (paths[:1] if args.working_only else paths)
        for kind, paths in OUTPUT_PAIRS.items()
    }
    all_paths = [
        ORIGINALS["docx"][0],
        ORIGINALS["pptx"][0],
        *active_output_pairs["docx"],
        *active_output_pairs["pptx"],
    ]
    for path in all_paths:
        report = inspect_package(path, checks)
        if report is not None:
            package_reports.append(report)

    for kind, (working, final) in OUTPUT_PAIRS.items():
        if args.working_only:
            if working.is_file():
                working_hash = sha256(working)
                add_check(
                    checks,
                    f"completed {kind} differs from original template",
                    working_hash != ORIGINALS[kind][1],
                    f"working={working_hash}; original={ORIGINALS[kind][1]}",
                )
            else:
                add_check(checks, f"working {kind} exists", False, rel(working))
            continue
        if working.is_file() and final.is_file():
            working_hash = sha256(working)
            final_hash = sha256(final)
            add_check(
                checks,
                f"working/final {kind} hashes match",
                working_hash == final_hash,
                f"working={working_hash}; final={final_hash}",
            )
            original_hash = ORIGINALS[kind][1]
            add_check(
                checks,
                f"completed {kind} differs from original template",
                working_hash != original_hash,
                f"working={working_hash}; original={original_hash}",
            )
        else:
            add_check(
                checks,
                f"working/final {kind} hashes match",
                False,
                "one or both files are missing",
            )

    for path in active_output_pairs["docx"]:
        check_docx(path, checks, visible_texts)
        inspect_xml_privacy(path, checks, is_pptx=False)
    for path in active_output_pairs["pptx"]:
        check_pptx(path, checks, visible_texts)
        inspect_xml_privacy(path, checks, is_pptx=True)

    for path, text in visible_texts.items():
        check_visible_scope(path, text, checks)

    for report in package_reports:
        print(
            f"[INFO] {rel(report.path)}: sha256={report.sha256}, "
            f"members={report.member_count}, xml={report.xml_count}, "
            f"media={report.media_count} ({report.media_bytes} bytes)"
        )

    failures = 0
    for name, ok, detail in checks:
        status = "PASS" if ok else "FAIL"
        suffix = f" - {detail}" if detail else ""
        print(f"[{status}] {name}{suffix}")
        failures += int(not ok)

    print(
        f"submission forms {'PASS' if failures == 0 else 'FAIL'}: "
        f"{len(checks) - failures}/{len(checks)} checks passed, "
        f"{failures} failed"
    )
    return 0 if failures == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
